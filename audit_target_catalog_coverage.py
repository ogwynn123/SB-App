#!/usr/bin/env python3
"""Audit a target Excel catalog against the local materialized survey repository."""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


CODE_RE = re.compile(r"NPA\s*(\d{4})", re.IGNORECASE)


@dataclass
class TargetRow:
    row_number: int
    code: str
    detail: str
    client: str
    geo: str
    month: str
    year: str


def normalize_code(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    match = CODE_RE.search(text)
    if match:
        return f"NPA{match.group(1)}"
    if text.isdigit() and len(text) == 4:
        return f"NPA{text}"
    return None


def text_or_blank(value: object) -> str:
    return "" if value is None else str(value).strip()


def workbook_stem_slug(path: Path) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "_", path.stem.strip()).strip("_").lower()
    return slug or "target_catalog"


def load_target_rows(workbook_path: Path) -> tuple[list[TargetRow], dict[str, list[int]]]:
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb.active

    rows: list[TargetRow] = []
    row_map: dict[str, list[int]] = defaultdict(list)
    seen_codes: set[str] = set()
    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        code = normalize_code(row[6]) or normalize_code(row[0])
        if not code:
            continue
        row_map[code].append(row_number)
        if code in seen_codes:
            continue
        seen_codes.add(code)
        rows.append(
            TargetRow(
                row_number=row_number,
                code=code,
                detail=text_or_blank(row[1]),
                client=text_or_blank(row[2]),
                geo=text_or_blank(row[3]),
                month=text_or_blank(row[4]),
                year=text_or_blank(row[5]),
            )
        )
    return rows, row_map


def survey_folders(materialized_root: Path) -> dict[str, Path]:
    return {
        path.name: path
        for path in materialized_root.iterdir()
        if path.is_dir() and re.fullmatch(r"NPA\d{4}", path.name)
    }


def manifest_payload(folder: Path) -> dict:
    manifest_path = folder / "materialization.manifest.json"
    if not manifest_path.exists():
        return {}
    try:
        return json.loads(manifest_path.read_text(encoding="utf-8"))
    except Exception:
        return {}


def any_files(paths: Iterable[Path], suffix: str) -> bool:
    suffix = suffix.lower()
    return any(path.is_file() and path.suffix.lower() == suffix for path in paths)


def folder_file_count(folder: Path) -> int:
    return sum(1 for path in folder.rglob("*") if path.is_file())


def classify_folder(folder: Path, code: str) -> dict[str, object]:
    manifest = manifest_payload(folder)
    docx_dir = folder / "docx"
    csv_dir = folder / "csv"
    docx_paths = list(docx_dir.iterdir()) if docx_dir.exists() else []
    csv_paths = list(csv_dir.iterdir()) if csv_dir.exists() else []

    has_manifest = (folder / "materialization.manifest.json").exists()
    has_docx = any_files(docx_paths, ".docx")
    has_csv = any(path.is_file() and path.suffix.lower() == ".csv" for path in folder.rglob("*.csv"))
    has_pdf = (folder / f"{code}.pdf").exists()
    has_raw = (folder / f"{code}RAW.txt").exists()
    has_json = (folder / f"{code}JSON.json").exists()
    has_tables = (folder / f"{code}TABLES.autogen.txt").exists()
    files_total = folder_file_count(folder)
    manifest_copied = manifest.get("copied", [])
    source_dirs = manifest.get("source_dirs", [])
    copied_count = len(manifest_copied) if isinstance(manifest_copied, list) else 0

    if has_docx and has_csv:
        status = "docx+csv"
    elif has_docx:
        status = "docx_only"
    elif has_csv:
        status = "csv_only"
    elif has_manifest and files_total == 1:
        status = "manifest_only"
    else:
        status = "other_files_only"

    return {
        "status": status,
        "has_manifest": has_manifest,
        "manifest_copied_count": copied_count,
        "manifest_source_dir_count": len(source_dirs) if isinstance(source_dirs, list) else 0,
        "manifest_source_dirs": " | ".join(str(value) for value in source_dirs[:3]) if isinstance(source_dirs, list) else "",
        "has_docx": has_docx,
        "has_csv": has_csv,
        "has_pdf": has_pdf,
        "has_raw": has_raw,
        "has_json": has_json,
        "has_tables": has_tables,
        "file_count": files_total,
    }


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    if not rows:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def write_codes_txt(path: Path, codes: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(codes) + ("\n" if codes else ""), encoding="utf-8")


def build_summary_lines(
    *,
    workbook_path: Path,
    materialized_root: Path,
    target_rows: list[TargetRow],
    duplicate_row_map: dict[str, list[int]],
    full_rows: list[dict[str, object]],
    extra_local_codes: list[str],
) -> list[str]:
    total_targets = len(target_rows)
    duplicate_codes = {code: rows for code, rows in duplicate_row_map.items() if len(rows) > 1}
    matched = sum(1 for row in full_rows if row["present_locally"] == "yes")
    missing = total_targets - matched
    status_counts = Counter(str(row["status"]) for row in full_rows)

    by_year: dict[str, Counter[str]] = defaultdict(Counter)
    by_client_missing: Counter[str] = Counter()
    for row in full_rows:
        year = str(row["year"] or "blank")
        by_year[year]["target"] += 1
        if row["present_locally"] == "yes":
            by_year[year]["present"] += 1
            if row["has_docx"] == "yes":
                by_year[year]["docx"] += 1
            if row["has_csv"] == "yes":
                by_year[year]["csv"] += 1
            if row["has_json"] == "yes":
                by_year[year]["json"] += 1
            if row["has_tables"] == "yes":
                by_year[year]["tables"] += 1
        else:
            by_client_missing[str(row["client"] or "(blank)")] += 1

    lines = [
        f"Workbook: {workbook_path}",
        f"Materialized root: {materialized_root}",
        "",
        f"Target jobs (unique codes): {total_targets}",
        f"Duplicate workbook codes: {len(duplicate_codes)}",
        f"Matched local folders: {matched}",
        f"Missing local folders: {missing}",
        f"Extra local folders not in workbook: {len(extra_local_codes)}",
        "",
        "Status counts:",
    ]
    lines.extend(f"  {status}: {count}" for status, count in status_counts.most_common())
    lines.append("")
    lines.append("Coverage by year:")
    for year in sorted(by_year, key=lambda value: int(value) if value.isdigit() else 9999):
        counts = by_year[year]
        present = counts["present"]
        target = counts["target"]
        pct = (present / target * 100.0) if target else 0.0
        lines.append(
            f"  {year}: target={target} present={present} ({pct:.1f}%) "
            f"docx={counts['docx']} csv={counts['csv']} json={counts['json']} tables={counts['tables']}"
        )
    if duplicate_codes:
        lines.append("")
        lines.append("Duplicate codes in workbook:")
        lines.extend(f"  {code}: rows {rows}" for code, rows in sorted(duplicate_codes.items()))
    if by_client_missing:
        lines.append("")
        lines.append("Top missing clients:")
        for client, count in by_client_missing.most_common(20):
            lines.append(f"  {client}: {count}")
    if extra_local_codes:
        lines.append("")
        lines.append("Extra local codes:")
        lines.extend(f"  {code}" for code in extra_local_codes[:50])
    return lines


def main() -> int:
    parser = argparse.ArgumentParser(description="Audit target catalog coverage against materialized survey folders.")
    parser.add_argument("workbook", help="Excel workbook with target job list.")
    parser.add_argument(
        "--materialized-root",
        default="repository/materialized_surveys_csv_priority",
        help="Materialized survey root to compare against.",
    )
    parser.add_argument(
        "--out-dir",
        default=".",
        help="Directory for generated audit artifacts.",
    )
    args = parser.parse_args()

    workbook_path = Path(args.workbook).resolve()
    materialized_root = Path(args.materialized_root).resolve()
    out_dir = Path(args.out_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    target_rows, duplicate_row_map = load_target_rows(workbook_path)
    local_folders = survey_folders(materialized_root)
    target_codes = {row.code for row in target_rows}
    extra_local_codes = sorted(set(local_folders) - target_codes)

    full_rows: list[dict[str, object]] = []
    missing_rows: list[dict[str, object]] = []
    missing_rows_with_detail: list[dict[str, object]] = []
    manifest_only_rows: list[dict[str, object]] = []
    csv_only_rows: list[dict[str, object]] = []

    for row in target_rows:
        folder = local_folders.get(row.code)
        payload = {
            "row_number": row.row_number,
            "code": row.code,
            "detail": row.detail,
            "client": row.client,
            "geo": row.geo,
            "month": row.month,
            "year": row.year,
        }
        if folder is None:
            full_row = {
                **payload,
                "present_locally": "no",
                "status": "missing_folder",
                "has_manifest": "no",
                "manifest_copied_count": 0,
                "manifest_source_dir_count": 0,
                "manifest_source_dirs": "",
                "has_docx": "no",
                "has_csv": "no",
                "has_pdf": "no",
                "has_raw": "no",
                "has_json": "no",
                "has_tables": "no",
                "file_count": 0,
            }
            full_rows.append(full_row)
            missing_rows.append(full_row)
            if str(full_row["detail"]).strip():
                missing_rows_with_detail.append(full_row)
            continue

        status = classify_folder(folder, row.code)
        full_row = {
            **payload,
            "present_locally": "yes",
            "status": status["status"],
            "has_manifest": "yes" if status["has_manifest"] else "no",
            "manifest_copied_count": status["manifest_copied_count"],
            "manifest_source_dir_count": status["manifest_source_dir_count"],
            "manifest_source_dirs": status["manifest_source_dirs"],
            "has_docx": "yes" if status["has_docx"] else "no",
            "has_csv": "yes" if status["has_csv"] else "no",
            "has_pdf": "yes" if status["has_pdf"] else "no",
            "has_raw": "yes" if status["has_raw"] else "no",
            "has_json": "yes" if status["has_json"] else "no",
            "has_tables": "yes" if status["has_tables"] else "no",
            "file_count": status["file_count"],
        }
        full_rows.append(full_row)
        if full_row["status"] == "manifest_only":
            manifest_only_rows.append(full_row)
        elif full_row["status"] == "csv_only":
            csv_only_rows.append(full_row)

    slug = workbook_stem_slug(workbook_path)
    full_csv = out_dir / f"{slug}_coverage_full.csv"
    missing_csv = out_dir / f"{slug}_missing_pull_list.csv"
    missing_txt = out_dir / f"{slug}_missing_codes.txt"
    missing_detail_csv = out_dir / f"{slug}_missing_pull_list_detail_only.csv"
    missing_detail_txt = out_dir / f"{slug}_missing_codes_detail_only.txt"
    manifest_only_csv = out_dir / f"{slug}_manifest_only.csv"
    csv_only_csv = out_dir / f"{slug}_csv_only.csv"
    extra_codes_txt = out_dir / f"{slug}_extra_local_codes.txt"
    summary_txt = out_dir / f"{slug}_coverage_summary.txt"

    write_csv(full_csv, full_rows)
    write_csv(missing_csv, missing_rows)
    write_csv(missing_detail_csv, missing_rows_with_detail)
    write_csv(manifest_only_csv, manifest_only_rows)
    write_csv(csv_only_csv, csv_only_rows)
    write_codes_txt(missing_txt, [str(row["code"]) for row in missing_rows])
    write_codes_txt(missing_detail_txt, [str(row["code"]) for row in missing_rows_with_detail])
    write_codes_txt(extra_codes_txt, extra_local_codes)

    summary_lines = build_summary_lines(
        workbook_path=workbook_path,
        materialized_root=materialized_root,
        target_rows=target_rows,
        duplicate_row_map=duplicate_row_map,
        full_rows=full_rows,
        extra_local_codes=extra_local_codes,
    )
    summary_txt.write_text("\n".join(summary_lines) + "\n", encoding="utf-8")

    print(f"Target jobs (unique codes): {len(target_rows)}")
    print(f"Matched local folders: {sum(1 for row in full_rows if row['present_locally'] == 'yes')}")
    print(f"Missing local folders: {len(missing_rows)}")
    print(f"Missing local folders with detail: {len(missing_rows_with_detail)}")
    print(f"Manifest-only local folders: {len(manifest_only_rows)}")
    print(f"CSV-only local folders: {len(csv_only_rows)}")
    print(f"Full audit CSV: {full_csv}")
    print(f"Missing pull list CSV: {missing_csv}")
    print(f"Missing codes TXT: {missing_txt}")
    print(f"Missing pull list CSV (detail only): {missing_detail_csv}")
    print(f"Missing codes TXT (detail only): {missing_detail_txt}")
    print(f"Manifest-only CSV: {manifest_only_csv}")
    print(f"CSV-only CSV: {csv_only_csv}")
    print(f"Summary TXT: {summary_txt}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
