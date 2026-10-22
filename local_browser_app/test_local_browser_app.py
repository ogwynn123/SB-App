from __future__ import annotations

import csv
import json
from pathlib import Path
import shutil
import threading
import unittest
from urllib.request import Request, urlopen
import uuid

from openpyxl import Workbook

from local_browser_app import services
from local_browser_app.server import SurveyBrowserServer


class ServiceTests(unittest.TestCase):
    def setUp(self) -> None:
        base = services.PLATFORM_ROOT / ".tmp_tests"
        base.mkdir(parents=True, exist_ok=True)
        self.temp_root = base / f"local_browser_{uuid.uuid4().hex}"
        self.temp_root.mkdir()

    def tearDown(self) -> None:
        shutil.rmtree(self.temp_root, ignore_errors=True)

    def test_table_discovery_finds_matching_pair(self) -> None:
        root = self.temp_root
        (root / "NPA9999 Final.docx").write_bytes(b"docx")
        (root / "NPA9999 Final.pdf").write_bytes(b"pdf")
        result = services.table_discover({"root": str(root)})
        self.assertEqual(result["count"], 1)
        self.assertEqual(result["jobs"][0]["stem"], "NPA9999 Final")
        self.assertTrue(result["jobs"][0]["tables_out"].endswith("TABLES.autogen.txt"))

    def test_sheet_merge_preview_and_run_preserve_secondary_extra(self) -> None:
        root = self.temp_root
        primary = root / "primary.csv"
        secondary = root / "delivery.csv"
        output = root / "merged.csv"
        primary.write_text("Survnum,q1\n1,old\n", encoding="utf-8")
        secondary.write_text("Survnum,Q1,extra_note\n2,new,keep me\n", encoding="utf-8")
        payload = {
            "primary_path": str(primary),
            "secondary_path": str(secondary),
            "output_path": str(output),
            "report_enabled": False,
            "archive_secondary": False,
        }
        preview = services.sheet_preview(payload)
        extras = [row for row in preview["mapping"] if row["method"] == "secondary_extra"]
        self.assertEqual(extras[0]["secondary"], "extra_note")

        payload["mapping_text"] = preview["reviewed_mapping_text"]
        result = services.sheet_run(payload, lambda _message: None)
        self.assertEqual(Path(result["output_path"]), output)
        with output.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        self.assertIn("secondary_extra_note", rows[0])
        self.assertEqual(rows[0]["secondary_extra_note"], "keep me")

    def test_quota_review_preserves_sheet_order_and_maps_dynamic_variables(self) -> None:
        quota = self.temp_root / "quotas.xlsx"
        sample = self.temp_root / "sample.csv"
        wb = Workbook()
        ws = wb.active
        ws.title = "Quotas"
        ws["A1"] = "Test Quotas"
        for column, value in enumerate(["Region", "Region Code", "Men", "Women", "Total"], start=1):
            ws.cell(3, column).value = value
        for column, value in enumerate(["North", 1, 2, 2, 4], start=1):
            ws.cell(4, column).value = value
        for column, value in enumerate(["Age", "Minimum", "Maximum", "", "Age", "Minimum", "Maximum"], start=1):
            ws.cell(7, column).value = value
        for column, value in enumerate(["18-34", 2, "", "", "18-34", 0.5, ""], start=1):
            ws.cell(8, column).value = value
        for column, value in enumerate(["35+", "", 2, "", "35+", "", 0.5], start=1):
            ws.cell(9, column).value = value
        for column, value in enumerate(["Education", "Minimum", "Maximum", "", "Education", "Minimum", "Maximum"], start=1):
            ws.cell(12, column).value = value
        for column, value in enumerate(["College grad", "", 2, "", "College grad", "", 0.5], start=1):
            ws.cell(13, column).value = value
        wb.save(quota)
        sample.write_text(
            "regioncode,gendergrp,agegrp,education\n"
            "1,1,1,College grad\n"
            "1,1,1,College grad\n"
            "1,2,2,Non-college\n",
            encoding="utf-8",
        )

        result = services.quota_review({"quota_path": str(quota), "sample_path": str(sample)})
        self.assertEqual([block["title"] for block in result["blocks"]], ["Exact base cells", "Age", "Education"])
        category_columns = {item["category"]: item["column"] for item in result["mapping"]["categories"]}
        self.assertEqual(category_columns, {"Age": "agegrp", "Education": "education"})
        exact = result["blocks"][0]
        women = exact["rows"][0]["cells"]["all::Women"]
        self.assertEqual(women["current"], 1)
        self.assertEqual(women["status"], "Open")
        age_rows = result["blocks"][1]["rows"]
        self.assertEqual(age_rows[0]["match_values"], "1")
        self.assertEqual(age_rows[0]["current"], 2)
        self.assertEqual(age_rows[0]["status"], "OK")

    def test_quota_review_separates_multiple_oversample_codes(self) -> None:
        quota = self.temp_root / "multi_over.xlsx"
        sample = self.temp_root / "multi_over.csv"
        wb = Workbook()
        ws = wb.active
        ws.title = "Quotas"
        ws["A1"] = "Multiple Oversamples"
        for column, value in enumerate(["Region", "Region Code", "Men", "Women", "Total"], start=1):
            ws.cell(3, column).value = value
        for column, value in enumerate(["North", 1, 1, 1, 2], start=1):
            ws.cell(4, column).value = value
        ws["A7"] = "FIRST OVERSAMPLE"
        ws["B9"] = "Men"
        ws["C9"] = "Women"
        ws["D9"] = "Total"
        for column, value in enumerate(["First", 1, 1, 2], start=1):
            ws.cell(10, column).value = value
        ws["A12"] = "SECOND OVERSAMPLE"
        ws["B14"] = "Men"
        ws["C14"] = "Women"
        ws["D14"] = "Total"
        for column, value in enumerate(["Second", 1, 1, 2], start=1):
            ws.cell(15, column).value = value
        wb.save(quota)
        sample.write_text(
            "regioncode,hgender,SampleType\n"
            "1,1,1\n1,2,1\n1,1,2\n1,2,3\n",
            encoding="utf-8",
        )

        result = services.quota_review({"quota_path": str(quota), "sample_path": str(sample)})
        filters = result["mapping"]["sections"]
        self.assertEqual([(row["column"], row["values"], row["ready"]) for row in filters], [
            ("SampleType", "2", True),
            ("SampleType", "3", True),
        ])
        self.assertEqual(result["summary"]["section_counts"], {
            "Base Sample": 2,
            "FIRST OVERSAMPLE": 1,
            "SECOND OVERSAMPLE": 1,
        })

    def test_oe_import_assign_and_export(self) -> None:
        root = self.temp_root
        data = root / "survey.csv"
        session = root / "survey.oe_session.sqlite"
        output = root / "survey.OE-coded.csv"
        data.write_text(
            "Survnum,q1oe,c1_1,c1_2,c1_3\n1,Helpful staff,101,,\n2,Long wait,202,,\n",
            encoding="utf-8",
        )
        imported = services.oe_import(
            {"data_path": str(data), "session_path": str(session), "key_column": "Survnum"}
        )
        question_id = imported["dashboard"]["questions"][0]["id"]
        response_payload = {"session_path": str(session), "question_id": question_id}
        responses = services.oe_responses(response_payload)["responses"]
        services.oe_assign(
            {
                "session_path": str(session),
                "response_id": responses[0]["id"],
                "codes": "500",
                "note": "browser test",
            }
        )
        services.oe_export(
            {"session_path": str(session), "kind": "coded_csv", "output_path": str(output)},
            lambda _message: None,
        )
        with output.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        self.assertEqual(rows[0]["c1_1"], "500")
        self.assertEqual(rows[1]["c1_1"], "202")

    def test_oe_lists_latest_wave_and_keeps_prior_waves_read_only(self) -> None:
        root = self.temp_root
        data = root / "trended.csv"
        session = root / "trended.oe_session.sqlite"
        data.write_text(
            "Survnum,Timeseries,q1oe,c1_1,c1_2,c1_3\n"
            "1,1,Older answer,101,,\n1,2,Newer answer,202,,\n2,2,Another,202,,\n",
            encoding="utf-8",
        )
        imported = services.oe_import({"data_path": str(data), "session_path": str(session)})
        question_id = imported["dashboard"]["questions"][0]["id"]
        base = {"session_path": str(session), "question_id": question_id}

        latest = services.oe_responses(base)["responses"]
        every_wave = services.oe_responses({**base, "all_waves": True})["responses"]

        self.assertEqual([row["record_key"] for row in latest], ["2|1", "2|2"])
        self.assertEqual(len(every_wave), 3)
        prior = next(row for row in every_wave if row["record_key"] == "1|1")
        with self.assertRaisesRegex(ValueError, "prior waves are read-only"):
            services.oe_assign({"session_path": str(session), "response_id": prior["id"], "codes": "500"})
        with self.assertRaisesRegex(ValueError, "prior waves are read-only"):
            services.oe_reset({"session_path": str(session), "response_id": prior["id"]})

    def test_question_search_reads_existing_repository(self) -> None:
        result = services.question_search({"query": "right direction", "page_size": 10})
        self.assertGreater(result["total"], 0)
        self.assertLessEqual(len(result["results"]), 10)
        self.assertIn("question_text", result["results"][0])

    def test_auto_results_empty_input_has_specific_error(self) -> None:
        root = self.temp_root
        marginals = root / "empty.docx"
        template = root / "template.docx"
        marginals.write_bytes(b"")
        template.write_bytes(b"")
        with self.assertRaisesRegex(RuntimeError, "Empty Input File"):
            services.auto_results_run(
                {
                    "marginals_path": str(marginals),
                    "template_path": str(template),
                    "output_path": str(root / "output.docx"),
                },
                lambda _message: None,
            )


class HttpSmokeTests(unittest.TestCase):
    def test_static_health_and_search_api(self) -> None:
        token = "test-token"
        server = SurveyBrowserServer(("127.0.0.1", 0), token)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        host, port = server.server_address[:2]
        try:
            with urlopen(f"http://{host}:{port}/", timeout=5) as response:
                html = response.read().decode("utf-8")
            self.assertIn("Survey Workbench", html)
            self.assertIn(token, html)
            self.assertIn("Quota Checker", html)

            health = Request(
                f"http://{host}:{port}/api/health",
                headers={"X-Survey-App-Token": token},
            )
            with urlopen(health, timeout=5) as response:
                payload = json.loads(response.read().decode("utf-8"))
            self.assertEqual(payload["status"], "ok")

            request = Request(
                f"http://{host}:{port}/api/questions/search",
                method="POST",
                data=json.dumps({"query": "right direction", "page_size": 10}).encode("utf-8"),
                headers={"Content-Type": "application/json", "X-Survey-App-Token": token},
            )
            with urlopen(request, timeout=10) as response:
                payload = json.loads(response.read().decode("utf-8"))
            self.assertTrue(payload["ok"])
            self.assertGreater(payload["result"]["total"], 0)
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=5)


if __name__ == "__main__":
    unittest.main()
