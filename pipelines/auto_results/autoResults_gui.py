#!/usr/bin/env python3
"""
Simple GUI wrapper for AutoResults processing.

Modes:
- no_fix: legacy replacement behavior (single placeholder replacement per run).
- fix: improved replacement behavior (multi-replace per run + split-run placeholder handling).
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
import difflib
import os
from pathlib import Path
import re
import sys
import traceback
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import zipfile
import xml.etree.ElementTree as ET

try:
    import win32com.client  # type: ignore
except Exception:  # pragma: no cover - optional Windows/Word integration
    win32com = None

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except Exception as exc:  # pragma: no cover
    Document = None
    Pt = None
    RGBColor = None
    Table = None
    Paragraph = None
    DOCX_IMPORT_ERROR = exc
else:
    DOCX_IMPORT_ERROR = None


PERCENT_RE = re.compile(r"(\d+)\s*%")
STAR_PERCENT_RE = re.compile(r"\*%\s*")
ZERO_RE = re.compile(r"[^-]-\s*$")
TRAILING_EMPTY_RE = re.compile(r"(?i)(?P<prefix>.*?)(?P<empty>\bempty)$")


@dataclass(frozen=True)
class RunStats:
    extracted_values: int
    used_values: int
    unused_values: int
    unresolved_placeholders: int
    output_path: Path
    alignment_report_path: Path | None = None
    alignment_issue_count: int = 0


@dataclass(frozen=True)
class AlignmentSlot:
    seq: int
    location: str
    question: str
    label: str
    value: str
    context: str
    placeholder_index: int = 1
    placeholder_count: int = 1


@dataclass(frozen=True)
class XlsxSlot:
    seq: int
    sheet: str
    table: str
    label: str
    value: str


def xlsx_slots_to_alignment(slots: list[XlsxSlot]) -> list[AlignmentSlot]:
    return [
        AlignmentSlot(
            seq=slot.seq,
            location=f"xlsx:{slot.sheet}:{slot.seq}",
            question=slot.table or slot.sheet,
            label=slot.label,
            value=slot.value,
            context=f"{slot.sheet} | {slot.table} | {slot.label} | {slot.value}",
        )
        for slot in slots
    ]


def _configure_tcl_env() -> None:
    if os.environ.get("TCL_LIBRARY") and os.environ.get("TK_LIBRARY"):
        return

    frozen_root = Path(getattr(sys, "_MEIPASS", "")) if getattr(sys, "_MEIPASS", None) else None
    py_dir = Path(sys.executable).resolve().parent
    candidate_pairs: list[tuple[Path, Path]] = []
    if frozen_root:
        candidate_pairs.append((frozen_root / "_tcl_data", frozen_root / "_tk_data"))
        candidate_pairs.append((frozen_root / "tcl" / "tcl8.6", frozen_root / "tcl" / "tk8.6"))
    roots = [
        py_dir / "tcl",
        py_dir.parent / "tcl",
        Path.home() / "tcl",
        Path.cwd() / "tcl",
    ]
    for root in roots:
        candidate_pairs.append((root / "tcl8.6", root / "tk8.6"))

    for tcl_dir, tk_dir in candidate_pairs:
        if not (tcl_dir / "init.tcl").exists():
            continue
        if not (tk_dir / "tk.tcl").exists():
            continue
        os.environ.setdefault("TCL_LIBRARY", str(tcl_dir))
        os.environ.setdefault("TK_LIBRARY", str(tk_dir))
        return


def _tcl_runtime_hint(script_name: str, exc: BaseException) -> str:
    return (
        "Failed to initialize Tk (Tcl/Tk runtime not found).\n"
        f"Error: {exc}\n\n"
        "Try launching with the bundled .bat file in this folder, or set "
        "TCL_LIBRARY and TK_LIBRARY to a Python installation that includes tkinter.\n"
        f"Source fallback:\npy -3 .\\{script_name}\n"
    )


def extract_values_from_marginals(document) -> list[str]:
    values: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text
        match_percent = PERCENT_RE.search(text)
        match_star_percent = STAR_PERCENT_RE.search(text)
        match_zero = ZERO_RE.search(text)

        if match_percent:
            values.append(match_percent.group(1))
        elif match_star_percent:
            values.append("<1")
        elif match_zero:
            values.append("0")
    return values


def apply_font_settings(run) -> None:
    run.font.size = Pt(9)
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor(0, 0, 0)


def _process_paragraph_no_fix(paragraph, values: list[str], placeholder: str) -> None:
    for run in paragraph.runs:
        if placeholder in run.text:
            value = values.pop(0) if values else ""
            run.text = run.text.replace(placeholder, value, 1)
            apply_font_settings(run)


def _process_paragraph_fix(paragraph, values: list[str], placeholder: str) -> None:
    # Walk runs left-to-right so split placeholders are filled in visual order.
    if len(placeholder) < 2:
        for run in paragraph.runs:
            while placeholder in run.text:
                value = values.pop(0) if values else ""
                run.text = run.text.replace(placeholder, value, 1)
                apply_font_settings(run)
        return

    runs = paragraph.runs
    i = 0
    while i < len(runs):
        run = runs[i]
        full_idx = run.text.find(placeholder)
        if full_idx >= 0:
            value = values.pop(0) if values else ""
            run.text = run.text[:full_idx] + value + run.text[full_idx + len(placeholder) :]
            apply_font_settings(run)
            continue

        split_replaced = False
        if i < len(runs) - 1:
            right = runs[i + 1]
            for split_idx in range(len(placeholder) - 1, 0, -1):
                prefix = placeholder[:split_idx]
                suffix = placeholder[split_idx:]
                if run.text.endswith(prefix) and right.text.startswith(suffix):
                    value = values.pop(0) if values else ""
                    run.text = run.text[: -len(prefix)] + value
                    right.text = right.text[len(suffix) :]
                    apply_font_settings(run)
                    split_replaced = True
                    break

        if not split_replaced:
            i += 1


def _process_document(document, values: list[str], placeholder: str, mode: str) -> None:
    processor = _process_paragraph_fix if mode == "fix" else _process_paragraph_no_fix

    for content in document.element.body:
        if content.tag.endswith("p"):
            paragraph = Paragraph(content, document)
            processor(paragraph, values, placeholder)
        elif content.tag.endswith("tbl"):
            table = Table(content, document)
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        processor(paragraph, values, placeholder)
        else:
            break


def count_unresolved_placeholders(document, placeholder: str) -> int:
    count = 0
    for content in document.element.body:
        if content.tag.endswith("p"):
            paragraph = Paragraph(content, document)
            for run in paragraph.runs:
                count += run.text.count(placeholder)
        elif content.tag.endswith("tbl"):
            table = Table(content, document)
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            count += run.text.count(placeholder)
    return count


def _iter_document_blocks(document):
    for block_idx, content in enumerate(document.element.body):
        if content.tag.endswith("p"):
            yield "p", block_idx, Paragraph(content, document)
        elif content.tag.endswith("tbl"):
            yield "tbl", block_idx, Table(content, document)


def _clean_text(text: str) -> str:
    text = text.replace("\u200b", "")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _norm_text(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", _clean_text(text).lower()).strip()


def _similarity(left: str, right: str) -> float:
    left_norm = _norm_text(left)
    right_norm = _norm_text(right)
    if not left_norm or not right_norm:
        return 0.0
    if left_norm == right_norm:
        return 1.0
    if left_norm in right_norm or right_norm in left_norm:
        return 0.92
    return difflib.SequenceMatcher(None, left_norm, right_norm).ratio()


def _question_id(text: str) -> str:
    match = re.search(r"^\s*q\s*(\d+[a-z]?)\b", _clean_text(text), flags=re.IGNORECASE)
    return match.group(1).lower() if match else ""


def _display_value_from_text(text: str) -> tuple[str | None, re.Match[str] | None]:
    match_percent = PERCENT_RE.search(text)
    if match_percent:
        return match_percent.group(1), match_percent
    match_star_percent = STAR_PERCENT_RE.search(text)
    if match_star_percent:
        return "<1", match_star_percent
    match_zero = ZERO_RE.search(text)
    if match_zero:
        return "0", match_zero
    return None, None


def _canonical_value(value: str) -> str:
    cleaned = _clean_text(str(value)).replace("%", "")
    if not cleaned:
        return ""
    if cleaned in {"-", "<1", "*"}:
        return "0"
    try:
        return str(round(float(cleaned) * 100)) if 0 < float(cleaned) < 1 else str(round(float(cleaned)))
    except Exception:
        return cleaned.lower()


def _format_xlsx_value(value: str) -> str:
    cleaned = _clean_text(str(value))
    if cleaned in {"", "-", "<1", "*"}:
        return cleaned
    try:
        numeric = float(cleaned)
    except Exception:
        return cleaned
    if 0 < numeric < 1:
        return str(round(numeric * 100))
    return str(round(numeric))


def _looks_like_heading(text: str) -> bool:
    cleaned = _clean_text(text)
    if len(cleaned) < 8:
        return False
    if "!!" in cleaned:
        return False
    if _display_value_from_text(cleaned)[0] is not None:
        return False
    return True


def _label_before_match(text: str, match: re.Match[str] | None) -> str:
    if match is None:
        return _clean_text(text)
    label = text[: match.start()]
    label = label.replace("\t", " ")
    return _clean_text(label)


def _label_before_placeholder(text: str, placeholder: str) -> str:
    idx = text.find(placeholder)
    if idx < 0:
        return _clean_text(text)
    label = text[:idx].replace("\t", " ")
    return _clean_text(label)


def extract_marginal_slots(document) -> list[AlignmentSlot]:
    slots: list[AlignmentSlot] = []
    current_question = ""

    def add_from_text(text: str, location: str) -> None:
        nonlocal current_question
        cleaned = _clean_text(text)
        value, match = _display_value_from_text(cleaned)
        if value is not None:
            slots.append(
                AlignmentSlot(
                    seq=len(slots) + 1,
                    location=location,
                    question=current_question,
                    label=_label_before_match(cleaned, match),
                    value=value,
                    context=cleaned,
                )
            )
        elif _looks_like_heading(cleaned):
            current_question = cleaned

    for kind, block_idx, block in _iter_document_blocks(document):
        if kind == "p":
            add_from_text(block.text, f"p{block_idx}")
            continue
        for row_idx, row in enumerate(block.rows, start=1):
            row_text = " | ".join(_clean_text(cell.text) for cell in row.cells)
            if _looks_like_heading(row_text):
                current_question = row_text
            for col_idx, cell in enumerate(row.cells, start=1):
                for para_idx, paragraph in enumerate(cell.paragraphs, start=1):
                    add_from_text(
                        paragraph.text,
                        f"tbl{block_idx} r{row_idx} c{col_idx} p{para_idx}",
                    )
    return slots


def extract_placeholder_slots(document, placeholder: str) -> list[AlignmentSlot]:
    slots: list[AlignmentSlot] = []
    current_question = ""

    def add_from_text(
        text: str,
        location: str,
        fallback_context: str = "",
        placeholder_count_override: int | None = None,
        placeholder_index_offset: int = 0,
    ) -> int:
        nonlocal current_question
        cleaned = _clean_text(text)
        context = _clean_text(fallback_context) or cleaned
        if placeholder not in cleaned:
            if _looks_like_heading(cleaned):
                current_question = cleaned
            return 0
        count = cleaned.count(placeholder)
        placeholder_count = placeholder_count_override or count
        label = _label_before_placeholder(cleaned, placeholder)
        if not label and context:
            label = _label_before_placeholder(context, placeholder)
        for idx in range(1, count + 1):
            slots.append(
                AlignmentSlot(
                    seq=len(slots) + 1,
                    location=location,
                    question=current_question,
                    label=label,
                    value="",
                    context=context,
                    placeholder_index=placeholder_index_offset + idx,
                    placeholder_count=placeholder_count,
                )
            )
        return count

    for kind, block_idx, block in _iter_document_blocks(document):
        if kind == "p":
            add_from_text(block.text, f"p{block_idx}")
            continue
        for row_idx, row in enumerate(block.rows, start=1):
            row_text = " | ".join(_clean_text(cell.text) for cell in row.cells)
            row_placeholder_count = row_text.count(placeholder)
            row_placeholder_seen = 0
            if placeholder not in row_text and _looks_like_heading(row_text):
                current_question = row_text
            for col_idx, cell in enumerate(row.cells, start=1):
                for para_idx, paragraph in enumerate(cell.paragraphs, start=1):
                    added = add_from_text(
                        paragraph.text,
                        f"tbl{block_idx} r{row_idx} c{col_idx} p{para_idx}",
                        fallback_context=row_text,
                        placeholder_count_override=row_placeholder_count or None,
                        placeholder_index_offset=row_placeholder_seen,
                    )
                    row_placeholder_seen += added
    return slots


def extract_xlsx_slots(path: Path) -> list[XlsxSlot]:
    slots: list[XlsxSlot] = []
    with zipfile.ZipFile(path) as archive:
        shared_strings = _xlsx_shared_strings(archive)
        for sheet_name, sheet_path in _xlsx_sheet_paths(archive):
            if sheet_name.lower() == "table of contents":
                continue
            table_title = ""
            for row in _xlsx_rows(archive, sheet_path, shared_strings):
                tag = row[1] if len(row) > 1 else None
                label = row[2] if len(row) > 2 else None
                value = row[3] if len(row) > 3 else None
                if tag == "Table Title 1" and label:
                    table_title = str(label)
                if tag == "Vertical % Row" and label is not None and value is not None:
                    slots.append(
                        XlsxSlot(
                            seq=len(slots) + 1,
                            sheet=sheet_name,
                            table=table_title,
                            label=_clean_text(str(label)),
                            value=_format_xlsx_value(str(value)),
                        )
                    )
    return slots


def _xml_ns(tag: str) -> str:
    return tag.split("}", 1)[0].strip("{") if tag.startswith("{") else ""


def _xlsx_shared_strings(archive: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in archive.namelist():
        return []
    root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
    ns = _xml_ns(root.tag)
    strings: list[str] = []
    for si in root.findall(f"{{{ns}}}si"):
        parts = [node.text or "" for node in si.iter(f"{{{ns}}}t")]
        strings.append("".join(parts))
    return strings


def _xlsx_sheet_paths(archive: zipfile.ZipFile) -> list[tuple[str, str]]:
    workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
    workbook_ns = _xml_ns(workbook_root.tag)
    rels_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rels = {
        rel.attrib.get("Id", ""): rel.attrib.get("Target", "")
        for rel in rels_root
    }
    sheets: list[tuple[str, str]] = []
    for sheet in workbook_root.findall(f".//{{{workbook_ns}}}sheet"):
        name = sheet.attrib.get("name", "")
        rel_id = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", "")
        target = rels.get(rel_id, "")
        if not target:
            continue
        path = target.lstrip("/")
        if not path.startswith("xl/"):
            path = f"xl/{path}"
        sheets.append((name, path.replace("\\", "/")))
    return sheets


def _xlsx_col_index(cell_ref: str) -> int:
    letters = "".join(ch for ch in cell_ref if ch.isalpha()).upper()
    value = 0
    for ch in letters:
        value = value * 26 + (ord(ch) - ord("A") + 1)
    return max(0, value - 1)


def _xlsx_cell_value(cell, shared_strings: list[str], ns: str) -> str:
    cell_type = cell.attrib.get("t", "")
    if cell_type == "inlineStr":
        return "".join(node.text or "" for node in cell.iter(f"{{{ns}}}t"))
    value_node = cell.find(f"{{{ns}}}v")
    if value_node is None or value_node.text is None:
        return ""
    raw = value_node.text
    if cell_type == "s":
        try:
            return shared_strings[int(raw)]
        except Exception:
            return raw
    return raw


def _xlsx_rows(archive: zipfile.ZipFile, sheet_path: str, shared_strings: list[str]) -> list[list[str]]:
    root = ET.fromstring(archive.read(sheet_path))
    ns = _xml_ns(root.tag)
    rows: list[list[str]] = []
    for row in root.findall(f".//{{{ns}}}row"):
        values: list[str] = []
        for cell in row.findall(f"{{{ns}}}c"):
            idx = _xlsx_col_index(cell.attrib.get("r", ""))
            while len(values) <= idx:
                values.append("")
            values[idx] = _xlsx_cell_value(cell, shared_strings, ns)
        rows.append(values)
    return rows


def _short_label(text: str) -> bool:
    cleaned = _clean_text(text)
    return 0 < len(cleaned) <= 80


def _is_summary_label(text: str) -> bool:
    label = _norm_text(text)
    return label.startswith("total ") or label in {"net", "minus"}


def _is_expected_unused_region_slot(source: AlignmentSlot | None, xlsx_slot: XlsxSlot | None) -> bool:
    if xlsx_slot is not None:
        sheet = _norm_text(xlsx_slot.sheet)
        table = _norm_text(xlsx_slot.table)
        if sheet == "region" or table == "region":
            return True
    if source is not None:
        question = _norm_text(source.question)
        location = _norm_text(source.location)
        if question == "region" or location.startswith("xlsx region"):
            return True
    return False


def _response_key(label: str) -> str:
    norm = _norm_text(label)
    mapping = {
        "much more likely": "much_more",
        "somewhat more likely": "somewhat_more",
        "smwt more likely": "somewhat_more",
        "no difference": "no_difference",
        "less likely": "less",
        "not true": "not_true",
        "don t know": "dont_know",
        "dont know": "dont_know",
        "dk": "dont_know",
        "total more likely": "total_more",
        "much less likely": "much_less",
        "somewhat less likely": "somewhat_less",
        "smwt less likely": "somewhat_less",
        "more likely": "more",
        "total less likely": "total_less",
    }
    return mapping.get(norm, norm)


def _expected_summary_index(label: str, placeholder_count: int) -> int | None:
    key = _response_key(label)
    if placeholder_count == 7 and key in {"total_more", "total_less"}:
        return 4
    return None


def _reorder_source_group_for_destination(
    source_group: list[AlignmentSlot],
    destination_group: list[AlignmentSlot],
) -> tuple[list[AlignmentSlot], bool]:
    if not source_group or len(source_group) != len(destination_group):
        return source_group, False
    if len(source_group) != 7 or destination_group[0].placeholder_count != 7:
        return source_group, False

    by_key: dict[str, AlignmentSlot] = {}
    for slot in source_group:
        key = _response_key(slot.label)
        if key in by_key:
            return source_group, False
        by_key[key] = slot

    orders = [
        ["much_more", "somewhat_more", "no_difference", "total_more", "not_true", "dont_know", "less"],
        ["much_less", "somewhat_less", "no_difference", "total_less", "not_true", "dont_know", "more"],
    ]
    for order in orders:
        if all(key in by_key for key in order):
            planned = [by_key[key] for key in order]
            return planned, planned != source_group
    return source_group, False


def build_fill_plan(
    source_slots: list[AlignmentSlot],
    destination_slots: list[AlignmentSlot],
) -> tuple[list[AlignmentSlot], int]:
    planned: list[AlignmentSlot] = []
    source_idx = 0
    dest_idx = 0
    reordered_groups = 0
    while dest_idx < len(destination_slots) and source_idx < len(source_slots):
        group_size = max(1, destination_slots[dest_idx].placeholder_count)
        group_size = min(group_size, len(destination_slots) - dest_idx)
        source_group = source_slots[source_idx : source_idx + group_size]
        destination_group = destination_slots[dest_idx : dest_idx + group_size]
        if len(source_group) < group_size:
            planned.extend(source_group)
            source_idx += len(source_group)
            dest_idx += group_size
            continue
        reordered, changed = _reorder_source_group_for_destination(source_group, destination_group)
        planned.extend(reordered)
        if changed:
            reordered_groups += 1
        source_idx += group_size
        dest_idx += group_size
    if source_idx < len(source_slots):
        planned.extend(source_slots[source_idx:])
    return planned, reordered_groups


def _alignment_issue_types(
    source: AlignmentSlot | None,
    destination: AlignmentSlot | None,
    xlsx_slot: XlsxSlot | None,
) -> list[str]:
    issues: list[str] = []
    if source is None and destination is None:
        if xlsx_slot is not None:
            issues.append("EXTRA_XLSX_ROW")
        return issues
    if source is None:
        issues.append("EXTRA_PLACEHOLDER")
        return issues
    if destination is None:
        issues.append("EXTRA_SOURCE_VALUE")
        return issues

    question_similarity = _similarity(source.question, destination.question)
    label_similarity = _similarity(source.label, destination.label)
    source_qid = _question_id(source.question)
    dest_qid = _question_id(destination.question)
    if source_qid and dest_qid and source_qid != dest_qid:
        issues.append("QUESTION_MISMATCH")
    elif (
        source.question
        and destination.question
        and question_similarity < 0.20
        and label_similarity < 0.45
        and destination.placeholder_count == 1
    ):
        issues.append("QUESTION_MISMATCH")

    if (
        "QUESTION_MISMATCH" not in issues
        and destination.placeholder_count == 1
        and _short_label(source.label)
        and _short_label(destination.label)
        and label_similarity < 0.45
    ):
        issues.append("ORDER_MISMATCH")
    expected_summary_index = _expected_summary_index(source.label, destination.placeholder_count)
    if expected_summary_index is not None and destination.placeholder_index != expected_summary_index:
        issues.append("MULTI_VALUE_SUMMARY_ORDER_REVIEW")

    if xlsx_slot is not None:
        if _canonical_value(source.value) != _canonical_value(xlsx_slot.value):
            issues.append("XLSX_VALUE_MISMATCH")
        if _short_label(source.label) and _short_label(xlsx_slot.label) and _similarity(source.label, xlsx_slot.label) < 0.45:
            issues.append("XLSX_LABEL_REVIEW")
    return issues


def _alignment_review_notes(source: AlignmentSlot | None, destination: AlignmentSlot | None) -> list[str]:
    notes: list[str] = []
    if source is not None and destination is not None and destination.placeholder_count > 1 and _short_label(source.label):
        notes.append("MULTI_VALUE_ROW")
    return notes


def write_alignment_report(
    *,
    marginals_path: Path | None,
    template_path: Path,
    report_path: Path,
    placeholder: str,
    xlsx_path: Path | None = None,
    source_mode: str = "docx",
) -> tuple[Path, int]:
    template_doc = Document(str(template_path))
    docx_slots: list[AlignmentSlot] = []
    if marginals_path is not None and marginals_path.exists():
        marginals_doc = Document(str(marginals_path))
        docx_slots = extract_marginal_slots(marginals_doc)
    destination_slots = extract_placeholder_slots(template_doc, placeholder)
    xlsx_slots = extract_xlsx_slots(xlsx_path) if xlsx_path else []
    if source_mode == "xlsx":
        if not xlsx_slots:
            raise ValueError("No values found in marginal XLSX file.")
        source_slots = xlsx_slots_to_alignment(xlsx_slots)
    elif source_mode == "docx":
        if not docx_slots:
            raise ValueError("No values found in marginals document.")
        source_slots = docx_slots
    else:
        raise ValueError(f"Unsupported source mode: {source_mode}")
    docx_audit_available = bool(docx_slots)
    fill_plan, _reordered_groups = build_fill_plan(source_slots, destination_slots)

    max_len = max(len(fill_plan), len(destination_slots), len(xlsx_slots) if xlsx_slots else 0)
    fieldnames = [
        "seq",
        "issue_types",
        "source_type",
        "source_seq",
        "source_value",
        "source_question",
        "source_label",
        "source_location",
        "source_context",
        "destination_question",
        "destination_label",
        "destination_location",
        "destination_placeholder_index",
        "destination_placeholder_count",
        "destination_context",
        "xlsx_sheet",
        "xlsx_table",
        "xlsx_label",
        "xlsx_value",
        "docx_marg_label",
        "docx_marg_value",
        "docx_marg_context",
        "question_similarity",
        "label_similarity",
        "review_notes",
        "fill_plan_notes",
    ]

    issue_count = 0
    report_path.parent.mkdir(parents=True, exist_ok=True)
    with report_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for idx in range(max_len):
            source = fill_plan[idx] if idx < len(fill_plan) else None
            destination = destination_slots[idx] if idx < len(destination_slots) else None
            if source_mode == "xlsx" and source is not None and 0 < source.seq <= len(xlsx_slots):
                xlsx_slot = xlsx_slots[source.seq - 1]
            elif source is not None and 0 < source.seq <= len(xlsx_slots):
                xlsx_slot = xlsx_slots[source.seq - 1]
            else:
                xlsx_slot = xlsx_slots[idx] if idx < len(xlsx_slots) else None
            docx_slot = docx_slots[source.seq - 1] if source is not None and 0 < source.seq <= len(docx_slots) else None
            expected_unused = destination is None and _is_expected_unused_region_slot(source, xlsx_slot)
            issues = [] if expected_unused else _alignment_issue_types(source, destination, xlsx_slot)
            if not expected_unused and source_mode == "xlsx" and source is not None and docx_audit_available:
                if docx_slot is None:
                    issues.append("MISSING_DOCX_MARG_ROW")
                else:
                    if _canonical_value(source.value) != _canonical_value(docx_slot.value):
                        issues.append("DOCX_MARG_VALUE_MISMATCH")
                    if _short_label(source.label) and _short_label(docx_slot.label) and _similarity(source.label, docx_slot.label) < 0.45:
                        issues.append("DOCX_MARG_LABEL_REVIEW")
            notes = _alignment_review_notes(source, destination)
            if expected_unused:
                notes.append("EXPECTED_UNUSED_REGION_ROW")
            fill_notes: list[str] = []
            if source is not None and source.seq != idx + 1:
                fill_notes.append("REORDERED_SOURCE_SLOT")
            if issues:
                issue_count += 1
            writer.writerow(
                {
                    "seq": idx + 1,
                    "issue_types": ";".join(issues),
                    "source_type": source_mode,
                    "source_seq": source.seq if source else "",
                    "source_value": source.value if source else "",
                    "source_question": source.question if source else "",
                    "source_label": source.label if source else "",
                    "source_location": source.location if source else "",
                    "source_context": source.context if source else "",
                    "destination_question": destination.question if destination else "",
                    "destination_label": destination.label if destination else "",
                    "destination_location": destination.location if destination else "",
                    "destination_placeholder_index": destination.placeholder_index if destination else "",
                    "destination_placeholder_count": destination.placeholder_count if destination else "",
                    "destination_context": destination.context if destination else "",
                    "xlsx_sheet": xlsx_slot.sheet if xlsx_slot else "",
                    "xlsx_table": xlsx_slot.table if xlsx_slot else "",
                    "xlsx_label": xlsx_slot.label if xlsx_slot else "",
                    "xlsx_value": xlsx_slot.value if xlsx_slot else "",
                    "docx_marg_label": docx_slot.label if docx_slot else "",
                    "docx_marg_value": docx_slot.value if docx_slot else "",
                    "docx_marg_context": docx_slot.context if docx_slot else "",
                    "question_similarity": f"{_similarity(source.question, destination.question):.3f}" if source and destination else "",
                    "label_similarity": f"{_similarity(source.label, destination.label):.3f}" if source and destination else "",
                    "review_notes": ";".join(notes),
                    "fill_plan_notes": ";".join(fill_notes),
                }
            )
    return report_path, issue_count


def run_autoresults(
    marginals_path: Path | None,
    template_path: Path,
    output_path: Path,
    placeholder: str = "!!",
    mode: str = "fix",
    write_report: bool = False,
    xlsx_path: Path | None = None,
    source_mode: str = "docx",
) -> RunStats:
    if Document is None:
        raise RuntimeError(
            f"python-docx is required. Import error: {DOCX_IMPORT_ERROR}"
        )
    if mode not in {"fix", "no_fix"}:
        raise ValueError(f"Unsupported mode: {mode}")
    if not placeholder:
        raise ValueError("Placeholder cannot be empty.")
    if source_mode not in {"docx", "xlsx"}:
        raise ValueError(f"Unsupported source mode: {source_mode}")
    if source_mode == "docx" and (marginals_path is None or not marginals_path.exists()):
        raise FileNotFoundError(f"Marginals file not found: {marginals_path}")
    if not template_path.exists():
        raise FileNotFoundError(f"Template file not found: {template_path}")
    if xlsx_path is not None and not xlsx_path.exists():
        raise FileNotFoundError(f"Marginal XLSX file not found: {xlsx_path}")
    if source_mode == "xlsx" and xlsx_path is None:
        raise ValueError("Marginal XLSX is required when using XLSX as the primary value source.")

    docx_slots: list[AlignmentSlot] = []
    if marginals_path is not None and marginals_path.exists():
        marginals_doc = Document(str(marginals_path))
        docx_slots = extract_marginal_slots(marginals_doc)
    xlsx_slots = extract_xlsx_slots(xlsx_path) if xlsx_path else []
    source_slots = xlsx_slots_to_alignment(xlsx_slots) if source_mode == "xlsx" else docx_slots
    if not source_slots:
        raise ValueError("No values found in selected marginal source.")

    results_doc = Document(str(template_path))
    destination_slots = extract_placeholder_slots(results_doc, placeholder)
    fill_plan, _reordered_groups = build_fill_plan(source_slots, destination_slots)
    values = [slot.value for slot in fill_plan]
    queue = list(values)
    _process_document(results_doc, queue, placeholder, mode)

    unresolved = count_unresolved_placeholders(results_doc, placeholder)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    results_doc.save(str(output_path))
    report_path = None
    report_issue_count = 0
    if write_report:
        report_path, report_issue_count = write_alignment_report(
            marginals_path=marginals_path,
            template_path=template_path,
            report_path=output_path.with_name(f"{output_path.stem}.alignment.csv"),
            placeholder=placeholder,
            xlsx_path=xlsx_path,
            source_mode=source_mode,
        )

    return RunStats(
        extracted_values=len(source_slots),
        used_values=len(values) - len(queue),
        unused_values=len(queue),
        unresolved_placeholders=unresolved,
        output_path=output_path,
        alignment_report_path=report_path,
        alignment_issue_count=report_issue_count,
    )


def _open_path(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    if sys.platform.startswith("win"):
        os.startfile(str(path))  # type: ignore[attr-defined]
    elif sys.platform == "darwin":
        import subprocess

        subprocess.Popen(["open", str(path)])
    else:
        import subprocess

        subprocess.Popen(["xdg-open", str(path)])


def open_review_pair(marginals_path: Path, output_path: Path) -> str:
    if not marginals_path.exists():
        raise FileNotFoundError(f"Marginals file not found: {marginals_path}")
    if not output_path.exists():
        raise FileNotFoundError(f"Filled results file not found: {output_path}")

    if sys.platform.startswith("win") and win32com is not None:
        try:
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = True
            word.Documents.Open(str(marginals_path.resolve()))
            word.Documents.Open(str(output_path.resolve()))
            try:
                word.Windows.Arrange(0)
            except Exception:
                pass
            return "Review files opened in Word."
        except Exception as exc:
            _open_path(marginals_path)
            _open_path(output_path)
            return f"Review files opened. Word side-by-side arrange failed: {exc}"

    _open_path(marginals_path)
    _open_path(output_path)
    return "Review files opened."


class AutoResultsGUI:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("AutoResults Runner")
        self.root.geometry("960x540")

        self.marginals_var = tk.StringVar()
        self.xlsx_var = tk.StringVar()
        self.template_var = tk.StringVar()
        self.output_var = tk.StringVar()
        self.placeholder_var = tk.StringVar(value="!!")
        self.mode_var = tk.StringVar(value="fix")
        self.open_review_var = tk.BooleanVar(value=False)
        self.write_report_var = tk.BooleanVar(value=True)
        self.use_xlsx_primary_var = tk.BooleanVar(value=False)
        self.status_var = tk.StringVar(value="Select files and click Run.")
        if DOCX_IMPORT_ERROR is not None:
            self.status_var.set("python-docx missing. Install: py -3 -m pip install python-docx")

        self._build_ui()

    def _build_ui(self) -> None:
        outer = ttk.Frame(self.root, padding=10)
        outer.pack(fill=tk.BOTH, expand=True)

        files = ttk.LabelFrame(outer, text="Files")
        files.pack(fill=tk.X, pady=(0, 8))
        self._file_row(files, 0, "Marginals DOCX (optional for XLSX-primary)", self.marginals_var, self._browse_input)
        self._file_row(files, 1, "Marginal XLSX (optional)", self.xlsx_var, self._browse_xlsx)
        self._file_row(files, 2, "Template (.docx)", self.template_var, self._browse_input)
        self._file_row(files, 3, "Output (.docx)", self.output_var, self._browse_output)

        opts = ttk.LabelFrame(outer, text="Options")
        opts.pack(fill=tk.X, pady=(0, 8))
        ttk.Label(opts, text="Placeholder").grid(row=0, column=0, padx=6, pady=6, sticky="w")
        ttk.Entry(opts, textvariable=self.placeholder_var, width=12).grid(
            row=0, column=1, padx=6, pady=6, sticky="w"
        )
        ttk.Label(opts, text="Mode").grid(row=0, column=2, padx=6, pady=6, sticky="w")
        ttk.Combobox(
            opts,
            textvariable=self.mode_var,
            values=["fix", "no_fix"],
            state="readonly",
            width=10,
        ).grid(row=0, column=3, padx=6, pady=6, sticky="w")
        ttk.Checkbutton(
            opts,
            text="Open marginals + filled result after run",
            variable=self.open_review_var,
        ).grid(row=1, column=0, columnspan=4, padx=6, pady=(0, 6), sticky="w")
        ttk.Checkbutton(
            opts,
            text="Write alignment report",
            variable=self.write_report_var,
        ).grid(row=2, column=0, columnspan=4, padx=6, pady=(0, 6), sticky="w")
        ttk.Checkbutton(
            opts,
            text="Use XLSX as primary value source",
            variable=self.use_xlsx_primary_var,
        ).grid(row=3, column=0, columnspan=4, padx=6, pady=(0, 6), sticky="w")

        actions = ttk.Frame(outer)
        actions.pack(fill=tk.X, pady=(0, 8))
        ttk.Button(actions, text="Run", command=self.run_click).pack(side=tk.LEFT, padx=4)
        ttk.Button(actions, text="Close", command=self.root.destroy).pack(side=tk.LEFT, padx=4)

        help_text = (
            "Mode 'fix' handles placeholders split across adjacent runs.\n"
            "Mode 'no_fix' uses the older single-run replacement behavior.\n"
            "XLSX-primary uses structured rows for values; the DOCX marginal remains a review check."
        )
        ttk.Label(outer, text=help_text, justify=tk.LEFT).pack(anchor="w", padx=4, pady=(0, 6))
        ttk.Label(outer, textvariable=self.status_var).pack(anchor="w", padx=4)

    def _file_row(self, parent: ttk.LabelFrame, row: int, label: str, var: tk.StringVar, browse_cb) -> None:
        ttk.Label(parent, text=label).grid(row=row, column=0, padx=6, pady=6, sticky="w")
        ttk.Entry(parent, textvariable=var, width=88).grid(row=row, column=1, padx=6, pady=6, sticky="we")
        parent.grid_columnconfigure(1, weight=1)
        ttk.Button(parent, text="Browse", command=lambda: browse_cb(var)).grid(
            row=row, column=2, padx=6, pady=6
        )

    def _browse_input(self, var: tk.StringVar) -> None:
        path = filedialog.askopenfilename(
            filetypes=[("Word Document", "*.docx"), ("All Files", "*.*")]
        )
        if path:
            var.set(path)
            self._maybe_seed_output()

    def _browse_xlsx(self, var: tk.StringVar) -> None:
        path = filedialog.askopenfilename(
            filetypes=[("Excel Workbook", "*.xlsx"), ("All Files", "*.*")]
        )
        if path:
            var.set(path)

    def _browse_output(self, var: tk.StringVar) -> None:
        path = filedialog.asksaveasfilename(
            defaultextension=".docx",
            filetypes=[("Word Document", "*.docx"), ("All Files", "*.*")],
        )
        if path:
            var.set(path)

    def _maybe_seed_output(self) -> None:
        if self.output_var.get().strip():
            return
        template = self.template_var.get().strip()
        if not template:
            return
        p = Path(template)
        stem = p.stem
        match = TRAILING_EMPTY_RE.match(stem.strip())
        if match:
            stem = f"{match.group('prefix').rstrip()} FILLED"
        else:
            stem = f"{stem} FILLED"
        self.output_var.set(str(p.with_name(f"{stem}{p.suffix}")))

    def run_click(self) -> None:
        try:
            marginals_raw = self.marginals_var.get().strip()
            marginals = Path(marginals_raw) if marginals_raw else None
            template = Path(self.template_var.get().strip())
            output_raw = self.output_var.get().strip()
            if not output_raw:
                self._maybe_seed_output()
                output_raw = self.output_var.get().strip()
            output = Path(output_raw)
            placeholder = self.placeholder_var.get()
            mode = self.mode_var.get().strip()
            xlsx_raw = self.xlsx_var.get().strip()
            xlsx = Path(xlsx_raw) if xlsx_raw else None
            source_mode = "xlsx" if self.use_xlsx_primary_var.get() else "docx"

            stats = run_autoresults(
                marginals_path=marginals,
                template_path=template,
                output_path=output,
                placeholder=placeholder,
                mode=mode,
                write_report=self.write_report_var.get(),
                xlsx_path=xlsx,
                source_mode=source_mode,
            )
            self.status_var.set(
                f"Done: source={source_mode}, used={stats.used_values}, unused={stats.unused_values}, unresolved={stats.unresolved_placeholders}, alignment flags={stats.alignment_issue_count}"
            )
            review_message = ""
            if self.open_review_var.get():
                if marginals is not None and marginals.exists():
                    review_message = open_review_pair(marginals, stats.output_path)
                else:
                    _open_path(stats.output_path)
                    review_message = "Filled result opened. Marginals DOCX was not available."
            messagebox.showinfo(
                "AutoResults Complete",
                "\n".join(
                    [
                        f"Output: {stats.output_path}",
                        f"Source: {source_mode}",
                        f"Extracted values: {stats.extracted_values}",
                        f"Used values: {stats.used_values}",
                        f"Unused values: {stats.unused_values}",
                        f"Unresolved placeholders: {stats.unresolved_placeholders}",
                        f"Alignment flags: {stats.alignment_issue_count}",
                        *( [f"Alignment report: {stats.alignment_report_path}"] if stats.alignment_report_path else [] ),
                        *(["", review_message] if review_message else []),
                    ]
                ),
            )
        except Exception as exc:
            self.status_var.set(f"Error: {exc}")
            detail = traceback.format_exc()
            messagebox.showerror("Run Failed", f"{exc}\n\n{detail}")


def main() -> None:
    _configure_tcl_env()
    try:
        root = tk.Tk()
    except tk.TclError as exc:
        msg = _tcl_runtime_hint("autoResults_gui.py", exc)
        print(msg, file=sys.stderr)
        raise SystemExit(2)
    AutoResultsGUI(root)
    root.mainloop()


if __name__ == "__main__":
    main()
