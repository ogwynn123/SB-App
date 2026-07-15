#!/usr/bin/env python3
"""
Generate an EMPTY results DOCX by replacing response-value numerics with
placeholder tokens:

- numeric slots -> !!
- percent slots -> !!%

Supports two operations:
1) Replace value slots already present in the input DOCX.
2) Insert missing output rows inferred from TABLES.autogen.txt / survey JSON.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from copy import deepcopy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Sequence

try:
    from docx import Document
    from docx.text.paragraph import Paragraph
except Exception as exc:  # pragma: no cover - environment guard
    Document = None
    Paragraph = None
    DOCX_IMPORT_ERROR = exc
else:
    DOCX_IMPORT_ERROR = None


TABLE_ID_LINE_RE = re.compile(r"^[A-Za-z0-9/_-]+\^\d+\s*$")
QUESTION_QID_RE = re.compile(r"^\d+[a-z]?$", re.IGNORECASE)
PERCENT_END_RE = re.compile(
    r"^(?P<head>.*?)(?P<value>(?:<\s*1|\*|\d+(?:\.\d+)?)\s*%)(?P<trail>\s*)$",
    re.IGNORECASE,
)
NUMBER_END_RE = re.compile(r"^(?P<head>.*?)(?P<value>-?\d+(?:\.\d+)?)(?P<trail>\s*)$")
VALUE_ONLY_PERCENT_RE = re.compile(
    r"^\s*(?:<\s*1|\*|\d+(?:\.\d+)?)\s*%\s*$",
    re.IGNORECASE,
)
VALUE_ONLY_NUMBER_RE = re.compile(r"^\s*-?\d+(?:\.\d+)?\s*$")

NO_PERCENT_TOTALS = {
    "total no degree",
    "total coll",
    "total college grad",
    "total liberal",
    "total conservative",
    "total important",
    "total not important",
    "total 18-44",
    "total 45-59",
    "total 60+",
}
SK_ALLOWED_HINTS = (
    "don't know",
    "dont know",
    "undecided",
    "other",
    "can't rate",
    "never heard",
    "refused",
)


@dataclass
class TableRowSpec:
    label: str
    selector: str
    flags: str
    style: str  # "percent" | "numeric" | "none"
    norm_label: str


@dataclass
class TableBlock:
    table_id: str
    qid: str
    banner: str = ""
    header: str = ""
    rows: list[TableRowSpec] = field(default_factory=list)


@dataclass
class GenerationStats:
    output_path: Path
    labels_loaded: int
    matched_label_lines: int
    replaced_percent_slots: int
    replaced_numeric_slots: int
    inserted_rows: int


def _normalize_text(text: str) -> str:
    text = text.replace("\u2019", "'").replace("\u2018", "'")
    text = text.replace("\u2013", "-").replace("\u2014", "-")
    text = re.sub(r"\s+", " ", text.strip())
    return text.casefold()


def _resolve_path(raw: str | os.PathLike[str]) -> Path:
    return Path(os.path.expandvars(os.path.expanduser(str(raw))))


def _default_output_path(docx_path: Path) -> Path:
    stem = docx_path.stem
    upper = stem.upper()
    if upper.endswith(" FILLED"):
        stem = stem[:-7].rstrip()
    if not stem.upper().endswith(" EMPTY"):
        stem = f"{stem} EMPTY"
    return docx_path.with_name(f"{stem}{docx_path.suffix}")


def _default_tables_path(docx_path: Path) -> Path | None:
    candidates = [
        docx_path.with_name(f"{docx_path.stem}TABLES.autogen.txt"),
        docx_path.with_name("TABLES.autogen.txt"),
        docx_path.with_name(f"{docx_path.stem}TABLES.txt"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def _default_json_path(docx_path: Path) -> Path | None:
    candidates = [
        docx_path.with_name(f"{docx_path.stem}JSON.json"),
        docx_path.with_name("JSON.json"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def _table_qid(table_id: str) -> str:
    left = table_id.split("^", 1)[0]
    if left.upper().startswith("TQ"):
        return left[2:]
    return left


def _parse_row_triplet(line: str) -> tuple[str, str, str] | None:
    if "^" not in line:
        return None
    parts = line.split("^")
    if len(parts) < 3:
        return None
    label = parts[0].strip()
    selector = parts[1].strip()
    flags = parts[2].strip()
    if not label:
        return None
    if label.casefold() == "hidden":
        return None
    return label, selector, flags


def _row_style(block: TableBlock, data_index: int, label: str, selector: str, flags: str) -> str:
    norm_label = _normalize_text(label)
    norm_header = _normalize_text(block.header)
    upper_flags = f",{flags.upper()},"

    if not QUESTION_QID_RE.fullmatch(block.qid):
        return "none"
    if norm_label in {"total answering", "total", "^"}:
        return "none"
    if ",HR," in upper_flags:
        return "none"
    if not selector:
        return "none"
    if re.search(r"\b(AND|OR)\b", selector, re.IGNORECASE):
        return "none"

    if "how likely would you say you are to vote" in norm_header and "already voted" in norm_label:
        return "percent"

    if norm_label.startswith("total "):
        if norm_label in NO_PERCENT_TOTALS:
            return "numeric"
        return "percent"

    if ",SK," in upper_flags and not any(token in norm_label for token in SK_ALLOWED_HINTS):
        # Keep terminal rows like "Probably not vote" untouched.
        return "none"

    if "which type of candidate for congress" in norm_header:
        return "percent"

    if data_index == 0:
        return "percent"
    return "numeric"


def _make_row(block: TableBlock, data_index: int, label: str, selector: str, flags: str) -> TableRowSpec:
    style = _row_style(block, data_index, label, selector, flags)
    return TableRowSpec(
        label=label,
        selector=selector,
        flags=flags,
        style=style,
        norm_label=_normalize_text(label),
    )


def _load_blocks_from_tables(tables_path: Path) -> list[TableBlock]:
    blocks: list[TableBlock] = []
    current: TableBlock | None = None
    state = "seek_id"
    data_index = 0

    with tables_path.open("r", encoding="utf-8", errors="replace") as f:
        for raw in f:
            line = raw.rstrip("\n")
            stripped = line.strip()
            if not stripped or stripped in {"```", "[TABLES]"}:
                continue

            if not line[:1].isspace() and TABLE_ID_LINE_RE.fullmatch(stripped):
                if current is not None:
                    blocks.append(current)
                current = TableBlock(table_id=stripped, qid=_table_qid(stripped))
                state = "need_banner"
                data_index = 0
                continue

            if current is None:
                continue

            if state == "need_banner":
                current.banner = stripped
                state = "need_header"
                continue
            if state == "need_header":
                current.header = stripped
                state = "rows"
                continue

            row = _parse_row_triplet(line)
            if row is None:
                continue
            label, selector, flags = row
            row_spec = _make_row(current, data_index, label, selector, flags)
            current.rows.append(row_spec)
            if row_spec.style != "none":
                data_index += 1

    if current is not None:
        blocks.append(current)
    return blocks


def _import_wincross_module():
    here = Path(__file__).resolve().parent
    wc_dir = (here.parent / "wincross_table_generator").resolve()
    if str(wc_dir) not in sys.path:
        sys.path.insert(0, str(wc_dir))
    import wincrossTableCompile as wc  # type: ignore

    return wc


def _default_wincross_rules(wc) -> list:
    return [
        wc.rule_favorability_0_10,
        wc.rule_approve_disapprove_4pt,
        wc.rule_favor_oppose_4pt,
        wc.rule_describes_well_5pt,
        wc.rule_agree_disagree_4pt,
        wc.rule_ideology_5pt,
        wc.rule_race_on_file,
        wc.rule_politics_identity_importance,
        wc.rule_party_on_file,
        wc.rule_motivation_to_vote,
        wc.rule_likelihood_to_vote,
        wc.rule_right_direction_location,
        wc.rule_right_direction_us,
        wc.rule_country_mood_summary,
        wc.rule_candidate_battery_header,
        wc.rule_vote_choice_with_leans,
        wc.rule_age_q102,
        wc.rule_education_q101,
        wc.rule_persuasion_summary,
    ]


def _load_blocks_from_json(json_path: Path) -> list[TableBlock]:
    wc = _import_wincross_module()
    with json_path.open("r", encoding="utf-8") as f:
        survey = json.load(f)

    questions = wc.parse_survey(survey)
    wc.apply_battery_suffixes(questions)
    wc.mark_common_stems(questions)
    wc_tables = wc.build_tables(questions, _default_wincross_rules(wc))

    blocks: list[TableBlock] = []
    for table in wc_tables:
        block = TableBlock(
            table_id=str(table.table_id),
            qid=_table_qid(str(table.table_id)),
            banner=str(table.banner),
            header=str(table.header),
        )
        data_index = 0
        for line in table.lines:
            row = _parse_row_triplet(line)
            if row is None:
                continue
            label, selector, flags = row
            row_spec = _make_row(block, data_index, label, selector, flags)
            block.rows.append(row_spec)
            if row_spec.style != "none":
                data_index += 1
        blocks.append(block)
    return blocks


def _merge_blocks(primary: list[TableBlock], secondary: list[TableBlock]) -> list[TableBlock]:
    if not primary:
        return secondary
    by_id = {b.table_id: b for b in primary}
    for block in secondary:
        if block.table_id not in by_id:
            primary.append(block)
            by_id[block.table_id] = block
    return primary


def _sorted_labels(labels: Iterable[str]) -> list[str]:
    deduped = sorted(set(labels))
    deduped.sort(key=lambda s: (-len(s), s))
    return deduped


def _build_label_style_map(blocks: Sequence[TableBlock]) -> tuple[dict[str, str], list[str]]:
    votes: dict[str, dict[str, int]] = {}
    for block in blocks:
        for row in block.rows:
            if row.style not in {"percent", "numeric"}:
                continue
            bucket = votes.setdefault(row.norm_label, {"percent": 0, "numeric": 0})
            bucket[row.style] += 1

    label_styles: dict[str, str] = {}
    for norm_label, counts in votes.items():
        if counts["percent"] > counts["numeric"]:
            label_styles[norm_label] = "percent"
        else:
            label_styles[norm_label] = "numeric"
    return label_styles, _sorted_labels(label_styles.keys())


def _match_label_prefix(text: str, labels: Sequence[str]) -> str | None:
    normalized = _normalize_text(text)
    if not normalized:
        return None
    for label in labels:
        if not normalized.startswith(label):
            continue
        if len(normalized) == len(label):
            return label
        next_char = normalized[len(label)]
        if not next_char.isalnum():
            return label
    return None


def _replace_label_aware_slot(
    text: str,
    matched_label: str,
    style: str,
) -> tuple[str, str | None]:
    if style not in {"percent", "numeric"}:
        return text, None

    normalized = _normalize_text(text)
    if normalized == matched_label:
        return text, None

    token = "!!%" if style == "percent" else "!!"

    m_pct = PERCENT_END_RE.match(text)
    if m_pct:
        new_text = f"{m_pct.group('head')}{token}{m_pct.group('trail')}"
        if new_text != text:
            return new_text, style

    m_num = NUMBER_END_RE.match(text)
    if m_num:
        new_text = f"{m_num.group('head')}{token}{m_num.group('trail')}"
        if new_text != text:
            return new_text, style

    m_placeholder = re.match(r"^(?P<head>.*?)(?P<value>!!%?)(?P<trail>\s*)$", text)
    if m_placeholder:
        new_text = f"{m_placeholder.group('head')}{token}{m_placeholder.group('trail')}"
        if new_text != text:
            return new_text, style

    return text, None


def _replace_value_only_slot(text: str) -> tuple[str, str | None]:
    if VALUE_ONLY_PERCENT_RE.fullmatch(text):
        leading = text[: len(text) - len(text.lstrip())]
        trailing = text[len(text.rstrip()) :]
        return f"{leading}!!%{trailing}", "percent"
    if VALUE_ONLY_NUMBER_RE.fullmatch(text):
        leading = text[: len(text) - len(text.lstrip())]
        trailing = text[len(text.rstrip()) :]
        return f"{leading}!!{trailing}", "numeric"
    return text, None


def _set_paragraph_text(paragraph, new_text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(new_text)
        return
    runs[0].text = new_text
    for run in runs[1:]:
        run.text = ""


def _record_slot(stats: GenerationStats, kind: str) -> None:
    if kind == "percent":
        stats.replaced_percent_slots += 1
    elif kind == "numeric":
        stats.replaced_numeric_slots += 1


def _display_label(label: str) -> str:
    norm = _normalize_text(label)
    label_map = {
        "total college grad": "Total Coll",
    }
    if norm in label_map:
        return label_map[norm]
    t = label.strip()
    if t and t == t.upper():
        return t.title()
    return t


def _looks_like_scale_line(text: str) -> bool:
    return bool(
        re.search(r"\b0\s+1\s+2\s+3\s+4\b", text)
        and re.search(r"\b10\b", text)
        and re.search(r"\b11\b", text)
        and re.search(r"\b12\b", text)
    )


def _favorability_item_name(block: TableBlock) -> str:
    header = block.header
    if " - " in header:
        name = header.split(" - ", 1)[1].strip()
    else:
        name = header.strip()
    name = re.sub(r"\[[^\]]+\]", "", name).strip()
    if name == name.upper():
        name = name.title()
    return name


def _favorability_letter(block: TableBlock, idx: int) -> str:
    qid = block.qid
    if qid and qid[-1].isalpha():
        return qid[-1].lower()
    return chr(ord("a") + idx)


def _favorability_blocks(blocks: Sequence[TableBlock]) -> list[TableBlock]:
    out: list[TableBlock] = []
    for block in blocks:
        if not QUESTION_QID_RE.fullmatch(block.qid):
            continue
        if "favorability rating" not in _normalize_text(block.header):
            continue
        if not any(row.style in {"percent", "numeric"} for row in block.rows):
            continue
        out.append(block)

    def _sort_key(block: TableBlock) -> tuple[int, str]:
        m = re.match(r"^(?P<num>\d+)(?P<sfx>[a-z]?)$", block.qid, re.IGNORECASE)
        if not m:
            return (10_000, block.qid)
        return (int(m.group("num")), m.group("sfx").lower())

    out.sort(key=_sort_key)
    return out


def _apply_favorability_matrix_transform(document, blocks: Sequence[TableBlock], stats: GenerationStats) -> None:
    fav_blocks = _favorability_blocks(blocks)
    if not fav_blocks:
        return

    scale_paragraphs = [p for p in document.paragraphs if _looks_like_scale_line(p.text)]
    if len(scale_paragraphs) < len(fav_blocks):
        return
    targets = scale_paragraphs[: len(fav_blocks)]

    first_target = targets[0]
    prev_xml = first_target._p.getprevious()
    if prev_xml is not None and Paragraph is not None:
        header_anchor = Paragraph(prev_xml, first_target._parent)
        for line in [
            "Very",
            "Fav",
            "Smwt",
            "Fav",
            "Ntrl",
            "Smwt",
            "Unfav",
            "Very",
            "Unfav",
            "CR",
            "NH",
            "Total",
            "Fav",
            "Total",
            "Unfav",
            "Name",
            "ID",
        ]:
            header_anchor = _insert_paragraph_after(header_anchor, line)

    for idx, (block, paragraph) in enumerate(zip(fav_blocks, targets)):
        label = f"{_favorability_letter(block, idx)}. {_favorability_item_name(block)}"
        _set_paragraph_text(paragraph, label)

        anchor = paragraph
        for row in block.rows:
            if row.style not in {"percent", "numeric"}:
                continue
            token = "!!%" if row.style == "percent" else "!!"
            anchor = _insert_paragraph_after(anchor, token)
            stats.inserted_rows += 1
            _record_slot(stats, row.style)


def _apply_replacement(
    paragraph,
    label_candidates: Sequence[str],
    label_styles: dict[str, str],
    allow_value_only: bool,
    stats: GenerationStats,
) -> None:
    original = paragraph.text
    if not original or not original.strip():
        return
    if "!!" in original:
        return

    matched_label = _match_label_prefix(original, label_candidates)
    if matched_label is not None:
        stats.matched_label_lines += 1
        style = label_styles.get(matched_label, "numeric")
        new_text, kind = _replace_label_aware_slot(original, matched_label, style)
        if kind is not None:
            _set_paragraph_text(paragraph, new_text)
            _record_slot(stats, kind)
            return

    if not allow_value_only:
        return
    new_text, kind = _replace_value_only_slot(original)
    if kind is None:
        return
    _set_paragraph_text(paragraph, new_text)
    _record_slot(stats, kind)


def _row_has_known_label(row, label_candidates: Sequence[str]) -> bool:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            if _match_label_prefix(paragraph.text, label_candidates) is not None:
                return True
    return False


def _process_table(
    table,
    label_candidates: Sequence[str],
    label_styles: dict[str, str],
    stats: GenerationStats,
    seen: set[int],
) -> None:
    for row in table.rows:
        row_target = _row_has_known_label(row, label_candidates)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                key = id(paragraph._p)
                if key in seen:
                    continue
                seen.add(key)
                _apply_replacement(
                    paragraph=paragraph,
                    label_candidates=label_candidates,
                    label_styles=label_styles,
                    allow_value_only=row_target,
                    stats=stats,
                )
            for nested in cell.tables:
                _process_table(nested, label_candidates, label_styles, stats, seen)


def _paragraph_matches_label(paragraph, norm_label: str) -> bool:
    return _match_label_prefix(paragraph.text, [norm_label]) is not None


def _insert_paragraph_after(anchor, text: str):
    if Paragraph is None:
        raise RuntimeError(
            "python-docx paragraph insertion unavailable. "
            f"Import error: {DOCX_IMPORT_ERROR}"
        )
    new_p = deepcopy(anchor._p)
    anchor._p.addnext(new_p)
    new_para = Paragraph(new_p, anchor._parent)
    _set_paragraph_text(new_para, text)
    return new_para


def _apply_missing_rows(document, blocks: Sequence[TableBlock], stats: GenerationStats) -> None:
    cursor = 0
    for block in blocks:
        insertable = [row for row in block.rows if row.style in {"percent", "numeric"}]
        if not insertable:
            continue

        paragraphs = list(document.paragraphs)
        matches: list[tuple[TableRowSpec, object | None, int | None]] = []
        search_start = min(cursor, len(paragraphs))
        found_any = False

        for row in block.rows:
            if row.style not in {"percent", "numeric"}:
                matches.append((row, None, None))
                continue

            found_para = None
            found_idx = None
            for idx in range(search_start, len(paragraphs)):
                if _paragraph_matches_label(paragraphs[idx], row.norm_label):
                    found_para = paragraphs[idx]
                    found_idx = idx
                    search_start = idx + 1
                    found_any = True
                    break
            matches.append((row, found_para, found_idx))

        if not found_any:
            continue

        found_indexes = [idx for _, _, idx in matches if idx is not None]
        if found_indexes:
            cursor = max(found_indexes) + 1

        anchor = None
        for row, paragraph, _ in matches:
            if row.style not in {"percent", "numeric"}:
                continue

            if paragraph is not None:
                new_text, kind = _replace_label_aware_slot(paragraph.text, row.norm_label, row.style)
                if kind is not None:
                    _set_paragraph_text(paragraph, new_text)
                    _record_slot(stats, kind)
                anchor = paragraph
                continue

            if anchor is None:
                continue

            token = "!!%" if row.style == "percent" else "!!"
            new_text = f"{_display_label(row.label)}{token}"
            anchor = _insert_paragraph_after(anchor, new_text)
            stats.inserted_rows += 1
            _record_slot(stats, row.style)


def generate_empty_results(
    *,
    docx_path: Path,
    output_path: Path,
    blocks: Sequence[TableBlock],
    insert_missing: bool,
) -> GenerationStats:
    if Document is None:
        raise RuntimeError(
            "python-docx is required. "
            f"Import error: {DOCX_IMPORT_ERROR}"
        )

    label_styles, label_candidates = _build_label_style_map(blocks)
    if not label_candidates:
        raise ValueError("No schema labels were loaded. Provide --tables and/or --json.")

    stats = GenerationStats(
        output_path=output_path,
        labels_loaded=len(label_candidates),
        matched_label_lines=0,
        replaced_percent_slots=0,
        replaced_numeric_slots=0,
        inserted_rows=0,
    )

    document = Document(str(docx_path))

    if insert_missing:
        _apply_favorability_matrix_transform(document, blocks, stats)
        _apply_missing_rows(document, blocks, stats)

    for paragraph in document.paragraphs:
        _apply_replacement(
            paragraph=paragraph,
            label_candidates=label_candidates,
            label_styles=label_styles,
            allow_value_only=False,
            stats=stats,
        )

    seen: set[int] = set()
    for table in document.tables:
        _process_table(table, label_candidates, label_styles, stats, seen)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return stats


def _build_parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(
        description=(
            "Generate EMPTY results DOCX by replacing response-value numerics "
            "with !! / !!% placeholders."
        )
    )
    ap.add_argument(
        "--docx",
        required=True,
        help="Input results DOCX file to transform.",
    )
    ap.add_argument(
        "--tables",
        help=(
            "Optional TABLES.autogen.txt path. If omitted, defaults are searched "
            "next to --docx."
        ),
    )
    ap.add_argument(
        "--json",
        help=(
            "Optional survey JSON path (for schema generation via wincross rules). "
            "Used as fallback or combined with --tables."
        ),
    )
    ap.add_argument(
        "--output",
        help="Optional output DOCX path. Defaults to '<input stem> EMPTY.docx'.",
    )
    ap.add_argument(
        "--no-insert-missing",
        action="store_true",
        help="Disable insertion of missing rows from schema.",
    )
    return ap


def main(argv: Sequence[str] | None = None) -> int:
    parser = _build_parser()
    args = parser.parse_args(argv)

    docx_path = _resolve_path(args.docx)
    if not docx_path.exists():
        raise SystemExit(f"DOCX not found: {docx_path}")

    tables_path = _resolve_path(args.tables) if args.tables else _default_tables_path(docx_path)
    json_path = _resolve_path(args.json) if args.json else _default_json_path(docx_path)

    blocks: list[TableBlock] = []
    sources: list[str] = []

    if tables_path and tables_path.exists():
        from_tables = _load_blocks_from_tables(tables_path)
        if from_tables:
            blocks = _merge_blocks(blocks, from_tables)
            sources.append(f"tables={tables_path}")
    elif args.tables:
        raise SystemExit(f"TABLES file not found: {tables_path}")

    if json_path and json_path.exists():
        try:
            from_json = _load_blocks_from_json(json_path)
        except Exception as exc:
            if not blocks:
                raise SystemExit(f"Failed to build schema from JSON: {exc}") from exc
            print(
                f"[warn] JSON schema generation failed ({json_path}): {exc}. "
                "Continuing with TABLES schema.",
                file=sys.stderr,
            )
        else:
            if from_json:
                blocks = _merge_blocks(blocks, from_json)
                sources.append(f"json={json_path}")
    elif args.json:
        raise SystemExit(f"JSON file not found: {json_path}")

    if not blocks:
        raise SystemExit(
            "No schema rows available. Provide --tables and/or --json "
            "(or keep defaults next to --docx)."
        )

    output_path = _resolve_path(args.output) if args.output else _default_output_path(docx_path)
    insert_missing = not args.no_insert_missing

    stats = generate_empty_results(
        docx_path=docx_path,
        output_path=output_path,
        blocks=blocks,
        insert_missing=insert_missing,
    )

    print("EMPTY Results Generation Complete")
    print(f"Input DOCX: {docx_path}")
    print(f"Schema source: {', '.join(sources) if sources else '(inferred)'}")
    print(f"Labels loaded: {stats.labels_loaded}")
    print(f"Matched label lines: {stats.matched_label_lines}")
    print(f"Inserted rows: {stats.inserted_rows}")
    print(f"Percent slots written: {stats.replaced_percent_slots}")
    print(f"Numeric slots written: {stats.replaced_numeric_slots}")
    print(f"Output DOCX: {stats.output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
