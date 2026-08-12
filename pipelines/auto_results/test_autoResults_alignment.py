from __future__ import annotations

import csv
import shutil
import sys
import unittest
import uuid
from pathlib import Path

from docx import Document
from openpyxl import Workbook


APP_ROOT = Path(__file__).resolve().parent
if str(APP_ROOT) not in sys.path:
    sys.path.insert(0, str(APP_ROOT))

from autoResults_gui import _dialog_initial_directory, run_autoresults


class AutoResultsAlignmentTests(unittest.TestCase):
    def _tmpdir(self) -> Path:
        root = APP_ROOT / ".tmp_tests" / "alignment"
        root.mkdir(parents=True, exist_ok=True)
        case_dir = root / f"case_{uuid.uuid4().hex}"
        case_dir.mkdir()
        self.addCleanup(lambda path=case_dir: shutil.rmtree(path, ignore_errors=True))
        return case_dir

    def _docx(self, path: Path, lines: list[str]) -> Path:
        doc = Document()
        for line in lines:
            doc.add_paragraph(line)
        doc.save(str(path))
        return path

    def _issues(self, report: Path) -> dict[int, str]:
        return {int(row["seq"]): row["issue_types"] for row in self._rows(report)}

    def _rows(self, report: Path) -> list[dict[str, str]]:
        with report.open("r", encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))

    def test_launcher_work_directory_is_used_as_dialog_fallback(self) -> None:
        fallback = self._tmpdir()
        self.assertEqual(_dialog_initial_directory("", fallback), str(fallback.resolve()))

    def test_alignment_report_has_no_flags_for_matching_order(self) -> None:
        case = self._tmpdir()
        marginals = self._docx(
            case / "marg.docx",
            ["Q1. Favorability", "Very favorable 60%", "Somewhat favorable 40%"],
        )
        template = self._docx(
            case / "empty.docx",
            ["Q1. Favorability", "Very favorable !!%", "Somewhat favorable !!%"],
        )

        stats = run_autoresults(
            marginals,
            template,
            case / "filled.docx",
            write_report=True,
        )

        self.assertEqual(stats.alignment_issue_count, 0)
        self.assertIsNotNone(stats.alignment_report_path)
        self.assertTrue(stats.alignment_report_path.exists())

    def test_alignment_report_flags_wrong_order_within_question(self) -> None:
        case = self._tmpdir()
        marginals = self._docx(
            case / "marg.docx",
            ["Q1. Vote choice", "Red 60%", "Blue 40%"],
        )
        template = self._docx(
            case / "empty.docx",
            ["Q1. Vote choice", "Blue !!%", "Red !!%"],
        )

        stats = run_autoresults(
            marginals,
            template,
            case / "filled.docx",
            write_report=True,
        )

        issues = self._issues(stats.alignment_report_path)
        self.assertIn("ORDER_MISMATCH", issues[1])
        self.assertIn("ORDER_MISMATCH", issues[2])

    def test_alignment_report_flags_question_boundary_shift(self) -> None:
        case = self._tmpdir()
        marginals = self._docx(
            case / "marg.docx",
            ["Q1. Vote choice", "Red 60%", "Blue 40%"],
        )
        template = self._docx(
            case / "empty.docx",
            ["Q1. Vote choice", "Red !!%", "Q2. Mood", "Blue !!%"],
        )

        stats = run_autoresults(
            marginals,
            template,
            case / "filled.docx",
            write_report=True,
        )

        issues = self._issues(stats.alignment_report_path)
        self.assertIn("QUESTION_MISMATCH", issues[2])

    def test_multi_value_more_likely_summary_fills_visual_summary_slot(self) -> None:
        case = self._tmpdir()
        marginals = self._docx(
            case / "marg.docx",
            [
                "Q16b. More Likely",
                "Much more likely 34%",
                "Somewhat more likely 26%",
                "No difference 26%",
                "Less likely 4%",
                "Not true 6%",
                "(Don't know) 4%",
                "TOTAL MORE LIKELY 60%",
            ],
        )
        template = self._docx(
            case / "empty.docx",
            ["[COSTS] Statement !!% !! !! !! !! !! !!%"],
        )

        stats = run_autoresults(
            marginals,
            template,
            case / "filled.docx",
            write_report=True,
        )

        filled_text = "\n".join(p.text for p in Document(str(stats.output_path)).paragraphs)
        self.assertIn("[COSTS] Statement 34% 26 26 60 6 4 4%", filled_text)

        rows = self._rows(stats.alignment_report_path)
        self.assertEqual(rows[3]["source_label"], "TOTAL MORE LIKELY")
        self.assertEqual(rows[3]["source_value"], "60")
        self.assertIn("REORDERED_SOURCE_SLOT", rows[3]["fill_plan_notes"])
        self.assertEqual(rows[3]["issue_types"], "")
        self.assertEqual(rows[6]["source_label"], "Less likely")
        self.assertEqual(rows[6]["source_value"], "4")

    def test_split_placeholder_is_filled_in_text_order(self) -> None:
        case = self._tmpdir()
        marginals = self._docx(
            case / "marg.docx",
            [
                "Q16b. More Likely",
                "Much more likely 34%",
                "Somewhat more likely 26%",
                "No difference 26%",
                "Less likely 4%",
                "Not true 6%",
                "(Don't know) 4%",
                "TOTAL MORE LIKELY 60%",
            ],
        )
        template = case / "empty.docx"
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("[COSTS] Statement\t")
        paragraph.add_run("!!%")
        paragraph.add_run("\t!!")
        paragraph.add_run("\t!")
        paragraph.add_run("!")
        paragraph.add_run("\t!!")
        paragraph.add_run("\t!!")
        paragraph.add_run("\t!!")
        paragraph.add_run("\t!!%")
        doc.save(str(template))

        stats = run_autoresults(
            marginals,
            template,
            case / "filled.docx",
            write_report=True,
        )

        filled_text = "\n".join(p.text for p in Document(str(stats.output_path)).paragraphs)
        self.assertIn("[COSTS] Statement\t34%\t26\t26\t60\t6\t4\t4%", filled_text)

    def test_alignment_report_uses_optional_xlsx_context(self) -> None:
        case = self._tmpdir()
        marginals = self._docx(
            case / "marg.docx",
            ["Q1. Vote choice", "Red 60%", "Blue 40%"],
        )
        template = self._docx(
            case / "empty.docx",
            ["Q1. Vote choice", "Red !!%", "Blue !!%"],
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Q1"
        sheet.append([None, "Table Title 1", "Q1. Vote choice", None])
        sheet.append([None, "Vertical % Row", "Red", 0.6])
        sheet.append([None, "Vertical % Row", "Green", 0.4])
        xlsx = case / "marg.xlsx"
        workbook.save(str(xlsx))

        stats = run_autoresults(
            marginals,
            template,
            case / "filled.docx",
            write_report=True,
            xlsx_path=xlsx,
        )

        issues = self._issues(stats.alignment_report_path)
        self.assertIn("XLSX_LABEL_REVIEW", issues[2])

    def test_xlsx_primary_uses_xlsx_values_and_audits_docx_marg(self) -> None:
        case = self._tmpdir()
        marginals = self._docx(
            case / "marg.docx",
            ["Q1. Vote choice", "Red 60%", "Blue 40%"],
        )
        template = self._docx(
            case / "empty.docx",
            ["Q1. Vote choice", "Red !!%", "Blue !!%"],
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Q1"
        sheet.append([None, "Table Title 1", "Q1. Vote choice", None])
        sheet.append([None, "Vertical % Row", "Red", 0.7])
        sheet.append([None, "Vertical % Row", "Blue", 0.3])
        xlsx = case / "marg.xlsx"
        workbook.save(str(xlsx))

        stats = run_autoresults(
            marginals,
            template,
            case / "filled.docx",
            write_report=True,
            xlsx_path=xlsx,
            source_mode="xlsx",
        )

        filled_text = "\n".join(p.text for p in Document(str(stats.output_path)).paragraphs)
        self.assertIn("Red 70%", filled_text)
        self.assertIn("Blue 30%", filled_text)
        rows = self._rows(stats.alignment_report_path)
        self.assertEqual(rows[0]["source_type"], "xlsx")
        self.assertEqual(rows[0]["source_value"], "70")
        self.assertEqual(rows[0]["docx_marg_value"], "60")
        self.assertIn("DOCX_MARG_VALUE_MISMATCH", rows[0]["issue_types"])

    def test_xlsx_primary_expected_unused_region_rows_are_notes(self) -> None:
        case = self._tmpdir()
        marginals = self._docx(
            case / "marg.docx",
            ["Region", "Boston 9%", "Middlesex 27%"],
        )
        template = self._docx(
            case / "empty.docx",
            ["Region", "Boston !!%", "Middlesex !!%"],
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "REGION"
        sheet.append([None, "Table Title 1", "Region", None])
        sheet.append([None, "Vertical % Row", "Boston", 0.09])
        sheet.append([None, "Vertical % Row", "Middlesex", 0.27])
        sheet.append([None, "Vertical % Row", "Springfield", 0.10])
        xlsx = case / "marg.xlsx"
        workbook.save(str(xlsx))

        stats = run_autoresults(
            marginals,
            template,
            case / "filled.docx",
            write_report=True,
            xlsx_path=xlsx,
            source_mode="xlsx",
        )

        self.assertEqual(stats.used_values, 2)
        self.assertEqual(stats.unused_values, 1)
        self.assertEqual(stats.alignment_issue_count, 0)
        rows = self._rows(stats.alignment_report_path)
        self.assertEqual(rows[2]["issue_types"], "")
        self.assertIn("EXPECTED_UNUSED_REGION_ROW", rows[2]["review_notes"])

    def test_xlsx_primary_can_run_without_docx_marginal(self) -> None:
        case = self._tmpdir()
        template = self._docx(
            case / "empty.docx",
            ["Q1. Vote choice", "Red !!%", "Blue !!%"],
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Q1"
        sheet.append([None, "Table Title 1", "Q1. Vote choice", None])
        sheet.append([None, "Vertical % Row", "Red", 0.7])
        sheet.append([None, "Vertical % Row", "Blue", 0.3])
        xlsx = case / "marg.xlsx"
        workbook.save(str(xlsx))

        stats = run_autoresults(
            None,
            template,
            case / "filled.docx",
            write_report=True,
            xlsx_path=xlsx,
            source_mode="xlsx",
        )

        filled_text = "\n".join(p.text for p in Document(str(stats.output_path)).paragraphs)
        self.assertIn("Red 70%", filled_text)
        self.assertIn("Blue 30%", filled_text)
        rows = self._rows(stats.alignment_report_path)
        self.assertEqual(rows[0]["source_type"], "xlsx")
        self.assertEqual(rows[0]["docx_marg_value"], "")


if __name__ == "__main__":
    unittest.main()
