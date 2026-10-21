#!/usr/bin/env python3
"""
Simple GUI for repository survey status, categorization editing, and question browsing.
"""

from __future__ import annotations

import json
import importlib.util
import os
import re
import shlex
import shutil
import sqlite3
import subprocess
import sys
import csv
import inspect
import threading
import unicodedata
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path


def _configure_tcl_env_early() -> None:
    if os.environ.get("TCL_LIBRARY") and os.environ.get("TK_LIBRARY"):
        return

    py_dir = Path(sys.executable).resolve().parent
    candidates = [
        py_dir / "tcl",
        py_dir.parent / "tcl",
        Path.home() / "tcl",
        Path.cwd() / "tcl",
    ]
    for root in candidates:
        tcl_dir = root / "tcl8.6"
        tk_dir = root / "tk8.6"
        if not (tcl_dir / "init.tcl").exists():
            continue
        if not (tk_dir / "tk.tcl").exists():
            continue
        os.environ.setdefault("TCL_LIBRARY", str(tcl_dir))
        os.environ.setdefault("TK_LIBRARY", str(tk_dir))
        return


_configure_tcl_env_early()

import tkinter as tk
from tkinter import filedialog, messagebox, ttk


BASE_DIR = Path(__file__).resolve().parents[2]
DEFAULT_REPO_ROOT = BASE_DIR / "repository"
CONFIG_FILENAME = "survey_tool_config.json"
REPO_INDEX_FILENAME = "index.json"
REPO_SURVEYS_DIRNAME = "surveys"
REPO_META_FILENAME = "survey.meta.json"
REPO_QUESTIONS_DB_FILENAME = "questions.sqlite"
WORKSPACE_DIRNAME = "workspace"
WORKFLOW_STATUS_FILENAME = "workflow_status.json"
WORKSPACE_MANIFEST_FILENAME = "workspace_manifest.json"
SURVEY_STATE_FILENAME = "survey_state.json"
DASHBOARD_CONFIG_KEY = "dashboard"
USER_CODE_START = 500
SURVEY_CODE_RE = re.compile(r"(?i)npa(?P<code>\d{4,})")
REPO_PLATFORM_STAGE_IDS = ("wincross", "catalog", "merge", "oe", "tomas", "results")
SURVEY_RUN_STAGE_IDS = ("wincross", "merge")
REPO_DISCOVERY_MODES = ("standard", "archive")
PHASE_VALUES = ("Pre-fielding", "In-field Daily", "Final Results", "Archive/Search")
QUESTION_MATCH_MODE_LABELS = ("All words", "Exact phrase", "Any word")
QUESTION_GROUP_PRESET_VALUES = (
    "",
    "approval",
    "ballot",
    "comparison",
    "demographics",
    "directional_mood",
    "evaluation",
    "favorability",
    "issue_position",
    "job_approval",
    "job_performance",
    "likelihood",
    "opinion",
    "party",
    "position",
    "probability",
    "salience",
    "voting",
)
WORKFLOW_STEP_DEFS: dict[str, tuple[tuple[str, str], ...]] = {
    "Pre-fielding": (
        ("final_questionnaire_selected", "Final questionnaire selected"),
        ("wincross_generated", "WinCross tables generated"),
        ("wincross_reviewed", "WinCross tables reviewed"),
        ("tomas_generated", "TOMAS generated, if Spanish"),
        ("tomas_reviewed", "TOMAS reviewed, if Spanish"),
    ),
    "In-field Daily": (
        ("daily_data_current", "Current data delivery tracked"),
        ("daily_data_headers_checked", "Data headers checked"),
        ("daily_codelist_tracked", "Daily codelist tracked, if delivered"),
        ("merge_generated", "Merge generated, if trended"),
        ("oe_codebook_generated", "OE codebook generated, if needed"),
        ("oe_wincross_rendered", "OE WinCross block rendered, if needed"),
        ("oe_reviewed", "OE output reviewed, if needed"),
    ),
    "Final Results": (
        ("final_data_confirmed", "Final data confirmed"),
        ("empty_results_generated", "Empty results generated"),
        ("results_reviewed", "Results reviewed"),
    ),
    "Archive/Search": (
        ("repo_scan_reviewed", "Repo scan reviewed"),
        ("question_index_rebuilt", "Question index rebuilt"),
    ),
}
WORKFLOW_STEP_STATUS_VALUES = ("pending", "done", "reviewed", "skipped")
WORKSPACE_FILE_ROLES: tuple[tuple[str, str, str, tuple[str, ...]], ...] = (
    ("final_questionnaire_docx", "Final Questionnaire DOCX", "inputs/questionnaire/final", (".docx",)),
    ("final_questionnaire_pdf", "Final Questionnaire PDF", "inputs/questionnaire/final", (".pdf",)),
    ("phone_docx", "Phone DOCX", "inputs/questionnaire/phone", (".docx",)),
    ("phone_pdf", "Phone PDF", "inputs/questionnaire/phone", (".pdf",)),
    ("sms_docx", "SMS DOCX", "inputs/questionnaire/sms", (".docx",)),
    ("sms_pdf", "SMS PDF", "inputs/questionnaire/sms", (".pdf",)),
    ("current_data", "Current Data CSV/XLSX", "inputs/data/current", (".csv", ".xlsx", ".xlsm")),
    ("timeseries_current", "Current Timeseries CSV/XLSX", "inputs/data/timeseries_current", (".csv", ".xlsx", ".xlsm")),
    ("daily_codelist", "Daily Codelist CSV/XLSX", "inputs/data/codelists", (".csv", ".xlsx", ".xlsm")),
    ("reviewed_tables", "Reviewed TABLES TXT", "inputs/tables/reviewed", (".txt",)),
    ("results_template_docx", "Results Template DOCX", "inputs/results/templates", (".docx",)),
    ("codebook", "Codebook CSV/XLSX", "inputs/codebooks", (".csv", ".xlsx", ".xlsm")),
    ("wincross_raw", "WinCross RAW TXT", "outputs/wincross/raw", (".txt",)),
    ("wincross_json", "WinCross JSON", "outputs/wincross/json", (".json",)),
    ("wincross_tables", "WinCross TABLES TXT", "outputs/wincross/tables", (".txt",)),
    ("merge_output", "Merge Output CSV/XLSX", "outputs/merge", (".csv", ".xlsx", ".xlsm")),
    ("merge_map", "Merge Map TXT", "inputs/merge/maps", (".txt",)),
    ("oe_codebook", "OE Codebook CSV", "outputs/oe/codebooks", (".csv",)),
    ("oe_wincross_block", "OE WinCross Block TXT", "outputs/oe/wincross_blocks", (".txt",)),
    ("results_output_docx", "Results Output DOCX", "outputs/results", (".docx",)),
    ("other", "Other Workspace File", "inputs/other", ()),
)
WORKSPACE_FILE_ROLE_BY_ID = {role_id: (label, rel_dir, extensions) for role_id, label, rel_dir, extensions in WORKSPACE_FILE_ROLES}
WORKSPACE_FILE_ROLE_LABEL_TO_ID = {label: role_id for role_id, label, _rel_dir, _extensions in WORKSPACE_FILE_ROLES}
OE_QUESTION_COL_RE = re.compile(r"^(?:(?P<prefix>\d{4,})[_-])?q(?P<num>\d{1,3})(?:[_-]?oe)?$", re.IGNORECASE)


@dataclass
class DataIntakeTable:
    headers: list[str]
    rows: list[list[object]]
    pairs: set[tuple[str, str]]
    duplicate_count: int
    duplicate_examples: list[str]
    timeseries_added: bool
    blank_timeseries_filled: int
    timeseries_values: list[str]

    @property
    def row_count(self) -> int:
        return len(self.rows)

    @property
    def pair_count(self) -> int:
        return len(self.pairs)


@dataclass(frozen=True)
class CodelistRecord:
    question_col: str
    code: str
    label: str
    source: str = "delivered"


@dataclass(frozen=True)
class CodelistReconciliationResult:
    current_path: Path
    delivered_path: Path
    user_codes_path: Path
    report_path: Path
    delivered_count: int
    added_count: int
    removed_count: int
    changed_count: int
    manual_review_count: int
    warning_count: int
    next_user_code: int


def _configure_tcl_env() -> None:
    _configure_tcl_env_early()


def _utc_now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def _read_json_dict(path: Path) -> dict:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"Expected object JSON: {path}")
    return payload


def _write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _normalize_categorization(value: object) -> dict[str, str]:
    if not isinstance(value, dict):
        return {"region": "", "client": ""}
    region = str(value.get("region", "")).strip()
    client = str(value.get("client", "")).strip()
    return {"region": region, "client": client}


def _normalize_survey_code(value: str | int | None) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    m = SURVEY_CODE_RE.search(text)
    if m:
        return str(int(m.group("code")))
    if text.isdigit() and len(text) >= 4:
        return str(int(text))
    return None


def _survey_id_desc_sort_key(value: object) -> tuple[int, int, str]:
    text = str(value or "").strip()
    code = _normalize_survey_code(text)
    if code is not None:
        return (0, -int(code), text.casefold())
    return (1, 0, text.casefold())


def _survey_index_row_desc_sort_key(row: dict) -> tuple[int, int, str]:
    survey_id = str(row.get("survey_id", "") or "").strip()
    survey_code = str(row.get("survey_code", "") or "").strip()
    code = _normalize_survey_code(survey_id) or _normalize_survey_code(survey_code)
    label = survey_id or (f"NPA{survey_code}" if survey_code else "")
    if code is not None:
        return (0, -int(code), label.casefold())
    return (1, 0, label.casefold())


def _normalize_question_match_mode(value: object) -> str:
    text = str(value or "").strip().casefold().replace("-", "_").replace(" ", "_")
    aliases = {
        "": "all_words",
        "all": "all_words",
        "all_word": "all_words",
        "all_words": "all_words",
        "all_terms": "all_words",
        "words": "all_words",
        "exact": "exact",
        "exact_phrase": "exact",
        "phrase": "exact",
        "any": "any_word",
        "any_word": "any_word",
        "any_words": "any_word",
        "any_terms": "any_word",
    }
    out = aliases.get(text)
    if out:
        return out
    raise ValueError("Unsupported match mode. Expected All words, Exact phrase, or Any word.")


def _question_query_tokens(query_text: str) -> list[str]:
    tokens: list[str] = []
    seen: set[str] = set()
    for token in re.findall(r"[A-Za-z0-9']+", str(query_text or "").casefold()):
        token = token.strip("'")
        if not token or token in seen:
            continue
        tokens.append(token)
        seen.add(token)
    return tokens


def _question_frequency_ballot_options_key(raw: object) -> str:
    try:
        payload = json.loads(str(raw or ""))
    except Exception:
        payload = []
    if not isinstance(payload, list):
        payload = [str(raw or "")]

    normalized: list[str] = []
    for item in payload:
        text = re.sub(r"[_]+", " ", str(item or "").casefold())
        text = re.sub(r"[^a-z0-9]+", " ", text).strip()
        text = re.sub(r"\s+", " ", text)
        if not text:
            continue
        if text == "other" or text == "none of the above" or text.startswith("other specify"):
            continue
        normalized.append(text)
    return "|".join(sorted(set(normalized)))


_SPANISH_BROWSER_PHRASE_CUES = (
    "ahora le voy",
    "ahora voy",
    "voy a leer",
    "voy a leerle",
    "por favor lea",
    "por favor seleccione",
    "me gustaria",
    "que me dijera",
    "si tiene una opinion",
    "opinion favorable o desfavorable",
    "figuras publicas",
    "personajes publicos",
    "cada persona",
    "cada uno",
    "una escala",
    "escala del",
    "significa que",
    "muy baja",
    "muy alta",
    "no sabe",
    "no conoce",
    "suficiente para calificar",
    "trabajadores domesticos",
    "cuidadores",
)

_SPANISH_BROWSER_TOKEN_CUES = {
    "ahora",
    "leer",
    "leerle",
    "lea",
    "gustaria",
    "dijera",
    "tiene",
    "opinion",
    "desfavorable",
    "figuras",
    "publicas",
    "personajes",
    "publicos",
    "seleccione",
    "cada",
    "escala",
    "significa",
    "baja",
    "alta",
    "nombre",
    "calificar",
    "suficiente",
    "trabajadores",
    "domesticos",
    "cuidadores",
}

_ENGLISH_BROWSER_PHRASE_CUES = (
    "now i'm going",
    "now i am going",
    "i'm going to read",
    "i am going to read",
    "please tell me",
    "favorable or unfavorable",
    "very low or very unfavorable",
    "very high or very favorable",
    "using a scale",
    "don't know enough to rate",
    "do not know enough to rate",
)


def _ascii_fold_text(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.replace("\u2019", "'").replace("\u2018", "'")
    text = re.sub(r"[^a-zA-Z0-9']+", " ", text.casefold())
    return re.sub(r"\s+", " ", text).strip()


def _question_browser_is_english(
    question_text: object,
    item_text: object = "",
    source_path: object = "",
) -> int:
    """Return 1 for rows the Question Browser should show.

    The repository index may contain Spanish-language companion instruments.
    The browser is intentionally English-only, so exclude rows with clear
    Spanish questionnaire phrasing while preserving English questions that
    merely mention Hispanic/Latino/Spanish identity.
    """
    source = _ascii_fold_text(source_path)
    if re.search(r"\b(spanish|espanol|espanol|castellano)\b", source):
        return 0

    blob = _ascii_fold_text(f"{question_text or ''} {item_text or ''}")
    if not blob:
        return 1

    spanish_phrase_score = sum(1 for cue in _SPANISH_BROWSER_PHRASE_CUES if cue in blob)
    english_phrase_score = sum(1 for cue in _ENGLISH_BROWSER_PHRASE_CUES if cue in blob)
    tokens = set(blob.split())
    spanish_token_score = len(tokens & _SPANISH_BROWSER_TOKEN_CUES)
    spanish_score = spanish_phrase_score * 2 + spanish_token_score

    if spanish_phrase_score >= 2 and spanish_score > english_phrase_score:
        return 0
    if blob.startswith(("ahora ", "como ya ", "pensandolo ")) and spanish_score >= 3:
        return 0
    if spanish_score >= 8 and english_phrase_score == 0:
        return 0
    return 1


def _question_search_columns(*, attrs_enabled: bool) -> list[str]:
    columns = [
        "q.question_text",
        "q.item_text",
        "q.option_texts_json",
    ]
    if attrs_enabled:
        columns.extend(
            [
                "q.question_family",
                "q.question_group",
                "q.classification_groups_json",
                "q.question_topics_json",
                "q.structure_tags_json",
                "q.scale_tags_json",
                "q.question_target",
            ]
        )
    return columns


def _question_text_match_sql(
    query_text: str,
    match_mode: str,
    *,
    attrs_enabled: bool,
) -> tuple[str, list[object]]:
    query_text = str(query_text or "").strip()
    if not query_text:
        return "", []

    match_mode = _normalize_question_match_mode(match_mode)
    columns = _question_search_columns(attrs_enabled=attrs_enabled)
    column_clause = " OR ".join(f"LOWER({column}) LIKE ?" for column in columns)

    if match_mode == "exact":
        pattern = f"%{query_text.casefold()}%"
        return f" AND ({column_clause})", [pattern] * len(columns)

    tokens = _question_query_tokens(query_text)
    if not tokens:
        return " AND 0=1", []

    token_clauses: list[str] = []
    params: list[object] = []
    for token in tokens:
        token_clauses.append(f"({column_clause})")
        params.extend([f"%{token}%"] * len(columns))

    joiner = " OR " if match_mode == "any_word" else " AND "
    return " AND (" + joiner.join(token_clauses) + ")", params


def _shorten(text: str, max_len: int) -> str:
    clean = str(text or "").strip()
    if len(clean) <= max_len:
        return clean
    return clean[: max_len - 3] + "..."


def _package_config_path() -> Path:
    return BASE_DIR / CONFIG_FILENAME


def _tomas_feature_enabled() -> bool:
    """TOMAS is off unless features.tomas is true in the config (or SURVEYPLATFORM_TOMAS=1)."""
    env = os.environ.get("SURVEYPLATFORM_TOMAS", "").strip().casefold()
    if env in {"1", "true", "yes", "on"}:
        return True
    if env in {"0", "false", "no", "off"}:
        return False
    features = _read_package_config().get("features")
    return isinstance(features, dict) and features.get("tomas") is True


def _read_package_config() -> dict:
    config_path = _package_config_path()
    if not config_path.exists():
        return {}
    try:
        return _read_json_dict(config_path)
    except Exception:
        return {}


def _package_repo_defaults() -> tuple[str, str]:
    source_root = ""
    discovery_mode = "standard"
    payload = _read_package_config()
    repo_cfg = payload.get("repo", {})
    if not isinstance(repo_cfg, dict):
        return source_root, discovery_mode
    source_root = str(repo_cfg.get("source_root", "")).strip()
    discovery_mode = str(repo_cfg.get("discovery_mode", "standard")).strip().casefold() or "standard"
    if discovery_mode not in REPO_DISCOVERY_MODES:
        discovery_mode = "standard"
    return source_root, discovery_mode


def screen_safe_geometry(root: tk.Tk, width: int, height: int, *, min_width: int = 1100, min_height: int = 700) -> str:
    root.update_idletasks()
    screen_width = max(root.winfo_screenwidth(), 1)
    screen_height = max(root.winfo_screenheight(), 1)
    padding = 80
    width = min(max(width, min_width), max(screen_width - padding, min_width))
    height = min(max(height, min_height), max(screen_height - padding, min_height))
    x = max((screen_width - width) // 2, 0)
    y = max((screen_height - height) // 2, 0)
    return f"{width}x{height}+{x}+{y}"


class RepositoryDashboardGUI:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("Survey Platform Dashboard")
        self.root.geometry(screen_safe_geometry(self.root, 1720, 940))

        default_source_root, default_discovery_mode = _package_repo_defaults()
        self.repo_root_var = tk.StringVar(value=str(DEFAULT_REPO_ROOT))
        self.source_root_var = tk.StringVar(value=default_source_root)
        self.discovery_mode_var = tk.StringVar(value=default_discovery_mode)
        self.summary_var = tk.StringVar(value="Load a repository index to begin.")
        self.status_var = tk.StringVar(value="")
        self.selected_survey_id_var = tk.StringVar(value="")
        self.selected_stage_var = tk.StringVar(value="")
        self.selected_last_run_var = tk.StringVar(value="")
        self.selected_platform_status_var = tk.StringVar(value="")
        self.selected_available_actions_var = tk.StringVar(value="")
        self.phase_var = tk.StringVar(value=PHASE_VALUES[0])
        self.workflow_status_summary_var = tk.StringVar(value="Select and confirm a survey to load workflow status.")
        self.workflow_workspace_var = tk.StringVar(value="")
        self.workflow_detail_var = tk.StringVar(value="")
        self.phase_primary_action_var = tk.StringVar(value="Run Pre-fielding Action")
        self.show_advanced_details_var = tk.BooleanVar(value=False)
        self.survey_selector_var = tk.StringVar(value="")
        self.main_survey_id_var = tk.StringVar(value="")
        self.main_survey_source_var = tk.StringVar(value="")
        self.region_var = tk.StringVar(value="")
        self.client_var = tk.StringVar(value="")
        self.run_stage_var = tk.StringVar(value=SURVEY_RUN_STAGE_IDS[0])
        self.run_args_var = tk.StringVar(value="")

        self.wincross_docx_var = tk.StringVar(value="")
        self.wincross_pdf_var = tk.StringVar(value="")
        self.wincross_raw_input_var = tk.StringVar(value="")
        self.wincross_json_out_var = tk.StringVar(value="")
        self.wincross_tables_out_var = tk.StringVar(value="")
        self.wincross_translate_var = tk.BooleanVar(value=False)
        self.wincross_skip_tables_var = tk.BooleanVar(value=False)
        self.wincross_status_var = tk.StringVar(value="")

        self.tomas_phone_docx_var = tk.StringVar(value="")
        self.tomas_phone_pdf_var = tk.StringVar(value="")
        self.tomas_sms_docx_var = tk.StringVar(value="")
        self.tomas_sms_pdf_var = tk.StringVar(value="")
        self.tomas_out_dir_var = tk.StringVar(value="")
        self.tomas_build_formatted_var = tk.BooleanVar(value=True)
        self.tomas_status_var = tk.StringVar(value="")
        self.tomas_enabled = _tomas_feature_enabled()

        self.merge_primary_survey_var = tk.StringVar(value="")
        self.merge_primary_csv_var = tk.StringVar(value="")
        self.merge_secondary_path_var = tk.StringVar(value="")
        self.merge_output_mode_var = tk.StringVar(value="new_file")
        self.merge_output_path_var = tk.StringVar(value="")
        self.merge_mapping_mode_var = tk.StringVar(value="auto")
        self.merge_threshold_var = tk.StringVar(value="85")
        self.merge_similarity_policy_var = tk.StringVar(value="drop")
        self.merge_status_var = tk.StringVar(value="")

        self.data_delivery_path_var = tk.StringVar(value="")
        self.data_delivery_codelist_path_var = tk.StringVar(value="")
        self.data_delivery_day_var = tk.StringVar(value="1")
        self.data_delivery_timeseries_var = tk.BooleanVar(value=False)
        self.data_delivery_status_var = tk.StringVar(value="")

        self.oe_input_path_var = tk.StringVar(value="")
        self.oe_question_col_var = tk.StringVar(value="")
        self.oe_defaults_path_var = tk.StringVar(value="")
        self.oe_existing_path_var = tk.StringVar(value="")
        self.oe_extract_out_var = tk.StringVar(value="")
        self.oe_include_blank_var = tk.BooleanVar(value=False)
        self.oe_render_codebook_var = tk.StringVar(value="")
        self.oe_render_selector_var = tk.StringVar(value="")
        self.oe_render_title_var = tk.StringVar(value="ABLE TO MENTION")
        self.oe_render_item_label_var = tk.StringVar(value="short")
        self.oe_render_out_var = tk.StringVar(value="")
        self.oe_status_var = tk.StringVar(value="")

        self.results_docx_var = tk.StringVar(value="")
        self.results_tables_var = tk.StringVar(value="")
        self.results_json_var = tk.StringVar(value="")
        self.results_output_var = tk.StringVar(value="")
        self.results_insert_missing_var = tk.BooleanVar(value=True)
        self.results_status_var = tk.StringVar(value="")

        self.question_query_var = tk.StringVar(value="")
        self.question_match_mode_var = tk.StringVar(value=QUESTION_MATCH_MODE_LABELS[0])
        self.question_state_var = tk.StringVar(value="")
        self.question_region_var = tk.StringVar(value="")
        self.question_client_var = tk.StringVar(value="")
        self.question_type_filter_var = tk.StringVar(value="")
        self.question_family_filter_var = tk.StringVar(value="")
        self.question_group_filter_var = tk.StringVar(value="")
        self.question_group_preset_var = tk.StringVar(value="")
        self.question_target_filter_var = tk.StringVar(value="")
        self.question_survey_var = tk.StringVar(value="")
        self.question_limit_var = tk.StringVar(value="30")
        self.question_page_var = tk.StringVar(value="1")
        self.question_page_summary_var = tk.StringVar(value="Page 0 of 0")
        self.question_summary_var = tk.StringVar(value="")
        self.question_status_var = tk.StringVar(value="")

        self.rows_by_id: dict[str, dict] = {}
        self.question_result_rows: dict[str, dict] = {}
        self.question_last_mode: str = ""
        self.question_last_filters: dict[str, object] | None = None
        self.question_last_total_count = 0
        self.tree: ttk.Treeview | None = None
        self.question_tree: ttk.Treeview | None = None
        self.oe_question_combo: ttk.Combobox | None = None
        self.survey_selector_combo: ttk.Combobox | None = None
        self.notebook: ttk.Notebook | None = None
        self.workflow_book: ttk.Notebook | None = None
        self.workflow_status_tree: ttk.Treeview | None = None
        self.phase_action_button: ttk.Button | None = None
        self.advanced_detail_frames: list[tuple[ttk.Widget, dict[str, object]]] = []
        self.question_prev_button: ttk.Button | None = None
        self.question_next_button: ttk.Button | None = None

        self._build_ui()

    def _build_ui(self) -> None:
        style = ttk.Style()
        style.theme_use("clam")

        outer = ttk.Frame(self.root, padding=10)
        outer.pack(fill=tk.BOTH, expand=True)

        controls = ttk.LabelFrame(outer, text="Platform")
        controls.pack(fill=tk.X, pady=(0, 6))

        ttk.Label(controls, text="Repository root").grid(row=0, column=0, sticky="w", padx=6, pady=6)
        ttk.Entry(controls, textvariable=self.repo_root_var, width=100).grid(
            row=0,
            column=1,
            sticky="we",
            padx=6,
            pady=6,
        )
        ttk.Button(controls, text="Browse", command=self._browse_repo_root).grid(
            row=0,
            column=2,
            sticky="w",
            padx=6,
            pady=6,
        )
        ttk.Button(controls, text="Load / Refresh", command=self.load_index).grid(
            row=0,
            column=3,
            sticky="w",
            padx=6,
            pady=6,
        )
        ttk.Button(controls, text="Repo Scan", command=self.scan_repository).grid(
            row=0,
            column=4,
            sticky="w",
            padx=6,
            pady=6,
        )
        ttk.Button(controls, text="Rebuild Question Index", command=self.rebuild_question_index).grid(
            row=0,
            column=5,
            sticky="w",
            padx=6,
            pady=6,
        )
        ttk.Label(controls, text="Source root").grid(row=1, column=0, sticky="w", padx=6, pady=6)
        ttk.Entry(controls, textvariable=self.source_root_var, width=100).grid(
            row=1,
            column=1,
            sticky="we",
            padx=6,
            pady=6,
        )
        ttk.Button(controls, text="Browse", command=self._browse_source_root).grid(
            row=1,
            column=2,
            sticky="w",
            padx=6,
            pady=6,
        )
        ttk.Label(controls, text="Discovery").grid(row=1, column=3, sticky="e", padx=6, pady=6)
        ttk.Combobox(
            controls,
            textvariable=self.discovery_mode_var,
            values=list(REPO_DISCOVERY_MODES),
            state="readonly",
            width=14,
        ).grid(
            row=1,
            column=4,
            sticky="w",
            padx=6,
            pady=6,
        )
        controls.grid_columnconfigure(1, weight=1)

        ttk.Label(outer, textvariable=self.summary_var).pack(anchor="w", padx=2, pady=(0, 4))

        notebook = ttk.Notebook(outer)
        self.notebook = notebook
        notebook.pack(fill=tk.BOTH, expand=True)

        start_tab = ttk.Frame(notebook, padding=4)
        surveys_tab = ttk.Frame(notebook, padding=4)
        workflows_tab = ttk.Frame(notebook, padding=4)
        questions_tab = ttk.Frame(notebook, padding=4)
        notebook.add(start_tab, text="Start")
        notebook.add(surveys_tab, text="Survey Workspaces")
        notebook.add(workflows_tab, text="Phase Workflows")
        notebook.add(questions_tab, text="Question Browser")

        self._build_start_tab(start_tab)
        self._build_surveys_tab(surveys_tab)
        self._build_workflows_tab(workflows_tab)
        self._build_questions_tab(questions_tab)

    def _register_advanced_detail_frame(self, frame: ttk.Widget, **pack_kwargs: object) -> None:
        self.advanced_detail_frames.append((frame, dict(pack_kwargs)))
        if self.show_advanced_details_var.get():
            frame.pack(**pack_kwargs)
        else:
            frame.pack_forget()

    def toggle_advanced_details(self) -> None:
        show = bool(self.show_advanced_details_var.get())
        for frame, pack_kwargs in self.advanced_detail_frames:
            if show:
                frame.pack(**pack_kwargs)
            else:
                frame.pack_forget()

    def _build_start_tab(self, tab: ttk.Frame) -> None:
        selector = ttk.LabelFrame(tab, text="Survey and Phase")
        selector.pack(fill=tk.X, pady=(0, 8))

        ttk.Label(selector, text="Survey").grid(row=0, column=0, sticky="w", padx=8, pady=8)
        self.survey_selector_combo = ttk.Combobox(
            selector,
            textvariable=self.survey_selector_var,
            width=28,
            state="readonly",
        )
        self.survey_selector_combo.grid(row=0, column=1, sticky="w", padx=6, pady=8)
        self.survey_selector_combo.bind("<<ComboboxSelected>>", self.select_survey_from_selector)

        ttk.Label(selector, text="Phase").grid(row=0, column=2, sticky="w", padx=(18, 6), pady=8)
        ttk.Combobox(
            selector,
            textvariable=self.phase_var,
            values=PHASE_VALUES,
            state="readonly",
            width=18,
        ).grid(row=0, column=3, sticky="w", padx=6, pady=8)
        ttk.Button(selector, text="Confirm Survey", command=self.confirm_survey_selection).grid(
            row=0,
            column=4,
            sticky="w",
            padx=6,
            pady=8,
        )
        ttk.Button(selector, text="Set Phase", command=self.save_current_phase).grid(
            row=0,
            column=5,
            sticky="w",
            padx=6,
            pady=8,
        )
        ttk.Button(selector, text="Open Phase", command=self.go_to_selected_phase).grid(
            row=0,
            column=6,
            sticky="w",
            padx=6,
            pady=8,
        )

        ttk.Label(selector, text="Main survey selected").grid(row=1, column=0, sticky="w", padx=8, pady=(0, 8))
        ttk.Entry(selector, textvariable=self.main_survey_id_var, width=18, state="readonly").grid(
            row=1,
            column=1,
            sticky="w",
            padx=6,
            pady=(0, 8),
        )
        ttk.Entry(selector, textvariable=self.main_survey_source_var, width=100, state="readonly").grid(
            row=1,
            column=2,
            columnspan=5,
            sticky="we",
            padx=6,
            pady=(0, 8),
        )
        selector.grid_columnconfigure(6, weight=1)

        home = ttk.LabelFrame(tab, text="Home Base")
        home.pack(fill=tk.X, pady=(0, 8))
        ttk.Label(home, textvariable=self.workflow_status_summary_var).grid(
            row=0,
            column=0,
            columnspan=6,
            sticky="w",
            padx=8,
            pady=(8, 4),
        )
        ttk.Label(home, text="Workspace").grid(row=1, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(home, textvariable=self.workflow_workspace_var, width=110, state="readonly").grid(
            row=1,
            column=1,
            columnspan=4,
            sticky="we",
            padx=6,
            pady=4,
        )
        ttk.Button(home, text="Open Workspace", command=self.open_selected_survey_files).grid(
            row=1,
            column=5,
            sticky="w",
            padx=6,
            pady=4,
        )
        self.phase_action_button = ttk.Button(
            home,
            textvariable=self.phase_primary_action_var,
            command=self.run_phase_primary_action,
        )
        self.phase_action_button.grid(row=2, column=0, sticky="w", padx=8, pady=(6, 8))
        ttk.Button(home, text="Add File to Workspace", command=self.add_file_to_workspace_dialog).grid(
            row=2,
            column=1,
            sticky="w",
            padx=6,
            pady=(6, 8),
        )
        ttk.Button(home, text="Open Phase Details", command=self.go_to_selected_phase).grid(
            row=2,
            column=2,
            sticky="w",
            padx=6,
            pady=(6, 8),
        )
        if self.tomas_enabled:
            ttk.Button(home, text="Run Optional TOMAS", command=self.run_guided_tomas).grid(
                row=2,
                column=3,
                sticky="w",
                padx=6,
                pady=(6, 8),
            )
        ttk.Button(home, text="Open Question Browser", command=lambda: self._select_top_tab("Question Browser")).grid(
            row=2,
            column=4,
            sticky="w",
            padx=6,
            pady=(6, 8),
        )
        ttk.Label(home, textvariable=self.workflow_detail_var).grid(
            row=2,
            column=5,
            columnspan=1,
            sticky="w",
            padx=6,
            pady=(6, 8),
        )
        home.grid_columnconfigure(4, weight=1)

        phases = ttk.LabelFrame(tab, text="Lifecycle")
        phases.pack(fill=tk.X, pady=(0, 8))
        for idx, phase in enumerate(PHASE_VALUES):
            ttk.Radiobutton(
                phases,
                text=phase,
                value=phase,
                variable=self.phase_var,
                command=self.go_to_selected_phase,
            ).grid(row=0, column=idx, sticky="w", padx=8, pady=8)

        checklist = ttk.LabelFrame(tab, text="Workflow Checklist")
        checklist.pack(fill=tk.BOTH, expand=True)
        self.workflow_status_tree = ttk.Treeview(
            checklist,
            columns=("phase", "item", "status", "updated"),
            show="headings",
            selectmode="browse",
            height=15,
        )
        self.workflow_status_tree.heading("phase", text="Phase")
        self.workflow_status_tree.heading("item", text="Item")
        self.workflow_status_tree.heading("status", text="Status")
        self.workflow_status_tree.heading("updated", text="Updated")
        self.workflow_status_tree.column("phase", width=150, anchor="w")
        self.workflow_status_tree.column("item", width=360, anchor="w")
        self.workflow_status_tree.column("status", width=110, anchor="w")
        self.workflow_status_tree.column("updated", width=190, anchor="w")
        self.workflow_status_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(6, 0), pady=6)

        checklist_scroll = ttk.Scrollbar(checklist, orient=tk.VERTICAL, command=self.workflow_status_tree.yview)
        checklist_scroll.pack(side=tk.RIGHT, fill=tk.Y, padx=(0, 6), pady=6)
        self.workflow_status_tree.configure(yscrollcommand=checklist_scroll.set)

        checklist_buttons = ttk.Frame(tab)
        checklist_buttons.pack(fill=tk.X, pady=(6, 0))
        ttk.Button(checklist_buttons, text="Mark Done", command=lambda: self.mark_selected_workflow_step("done")).pack(
            side=tk.LEFT,
            padx=(0, 6),
        )
        ttk.Button(
            checklist_buttons,
            text="Mark Reviewed",
            command=lambda: self.mark_selected_workflow_step("reviewed"),
        ).pack(side=tk.LEFT, padx=6)
        ttk.Button(
            checklist_buttons,
            text="Mark Skipped",
            command=lambda: self.mark_selected_workflow_step("skipped"),
        ).pack(side=tk.LEFT, padx=6)
        ttk.Button(
            checklist_buttons,
            text="Reset Item",
            command=lambda: self.mark_selected_workflow_step("pending"),
        ).pack(side=tk.LEFT, padx=6)
        self._refresh_workflow_home()

    def _build_surveys_tab(self, tab: ttk.Frame) -> None:
        table_frame = ttk.LabelFrame(tab, text="Surveys")
        table_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 6))
        self.tree = ttk.Treeview(
            table_frame,
            columns=(
                "survey_id",
                "stage",
                "wincross",
                "catalog",
                "merge",
                "oe",
                "tomas",
                "results",
                "region",
                "client",
                "last_run",
            ),
            show="headings",
            selectmode="browse",
            height=18,
        )
        self.tree.heading("survey_id", text="Survey")
        self.tree.heading("stage", text="Legacy")
        self.tree.heading("wincross", text="WinCross")
        self.tree.heading("catalog", text="Catalog")
        self.tree.heading("merge", text="Merge")
        self.tree.heading("oe", text="OE")
        self.tree.heading("tomas", text="TOMAS")
        if not self.tomas_enabled:
            self.tree.configure(displaycolumns=[c for c in self.tree["columns"] if c != "tomas"])
        self.tree.heading("results", text="Results")
        self.tree.heading("region", text="Region")
        self.tree.heading("client", text="Client")
        self.tree.heading("last_run", text="Last Run")
        self.tree.column("survey_id", width=120, anchor="w")
        self.tree.column("stage", width=110, anchor="w")
        self.tree.column("wincross", width=95, anchor="w")
        self.tree.column("catalog", width=95, anchor="w")
        self.tree.column("merge", width=95, anchor="w")
        self.tree.column("oe", width=90, anchor="w")
        self.tree.column("tomas", width=95, anchor="w")
        self.tree.column("results", width=95, anchor="w")
        self.tree.column("region", width=200, anchor="w")
        self.tree.column("client", width=220, anchor="w")
        self.tree.column("last_run", width=190, anchor="w")
        self.tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(6, 0), pady=6)
        self.tree.bind("<<TreeviewSelect>>", self._on_select_row)
        self.tree.bind("<Double-1>", self.show_selected_survey)

        scroll = ttk.Scrollbar(table_frame, orient=tk.VERTICAL, command=self.tree.yview)
        scroll.pack(side=tk.RIGHT, fill=tk.Y, padx=(0, 6), pady=6)
        self.tree.configure(yscrollcommand=scroll.set)
        x_scroll = ttk.Scrollbar(tab, orient=tk.HORIZONTAL, command=self.tree.xview)
        x_scroll.pack(fill=tk.X, padx=(6, 6), pady=(0, 6))
        self.tree.configure(xscrollcommand=x_scroll.set)

        details = ttk.LabelFrame(tab, text="Selected Survey Workspace")
        details.pack(fill=tk.X)

        ttk.Label(details, text="Survey").grid(row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(details, textvariable=self.selected_survey_id_var, width=22, state="readonly").grid(
            row=0,
            column=1,
            sticky="w",
            padx=6,
            pady=4,
        )
        ttk.Label(details, text="Stage").grid(row=0, column=2, sticky="w", padx=6, pady=4)
        ttk.Entry(details, textvariable=self.selected_stage_var, width=18, state="readonly").grid(
            row=0,
            column=3,
            sticky="w",
            padx=6,
            pady=4,
        )
        ttk.Label(details, text="Last run").grid(row=0, column=4, sticky="w", padx=6, pady=4)
        ttk.Entry(details, textvariable=self.selected_last_run_var, width=26, state="readonly").grid(
            row=0,
            column=5,
            sticky="w",
            padx=6,
            pady=4,
        )

        ttk.Label(details, text="Region").grid(row=1, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(details, textvariable=self.region_var, width=36).grid(
            row=1,
            column=1,
            columnspan=2,
            sticky="we",
            padx=6,
            pady=4,
        )
        ttk.Label(details, text="Client").grid(row=1, column=3, sticky="w", padx=6, pady=4)
        ttk.Entry(details, textvariable=self.client_var, width=40).grid(
            row=1,
            column=4,
            columnspan=2,
            sticky="we",
            padx=6,
            pady=4,
        )

        ttk.Label(details, text="Platform").grid(row=2, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(details, textvariable=self.selected_platform_status_var, width=112, state="readonly").grid(
            row=2,
            column=1,
            columnspan=5,
            sticky="we",
            padx=6,
            pady=4,
        )

        ttk.Label(details, text="Available").grid(row=3, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(details, textvariable=self.selected_available_actions_var, width=112, state="readonly").grid(
            row=3,
            column=1,
            columnspan=5,
            sticky="we",
            padx=6,
            pady=4,
        )

        ttk.Label(details, text="Main survey selected").grid(row=4, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(details, textvariable=self.main_survey_id_var, width=22, state="readonly").grid(
            row=4,
            column=1,
            sticky="w",
            padx=6,
            pady=4,
        )
        ttk.Entry(details, textvariable=self.main_survey_source_var, width=84, state="readonly").grid(
            row=4,
            column=2,
            columnspan=3,
            sticky="we",
            padx=6,
            pady=4,
        )
        ttk.Button(details, text="Set From Selected", command=self.set_main_survey_from_selection).grid(
            row=4,
            column=5,
            sticky="w",
            padx=6,
            pady=4,
        )

        buttons = ttk.Frame(details)
        buttons.grid(row=5, column=0, columnspan=6, sticky="w", padx=6, pady=(2, 6))
        ttk.Button(buttons, text="Save Categorization", command=self.save_selected).pack(
            side=tk.LEFT,
            padx=(0, 6),
        )
        ttk.Button(buttons, text="Reload", command=self.load_index).pack(side=tk.LEFT, padx=6)
        ttk.Button(buttons, text="Show", command=self.show_selected_survey).pack(side=tk.LEFT, padx=6)
        ttk.Button(buttons, text="Detect", command=self.detect_selected_survey).pack(side=tk.LEFT, padx=6)
        ttk.Button(buttons, text="History", command=self.show_selected_history).pack(side=tk.LEFT, padx=6)
        ttk.Button(buttons, text="Open Files", command=self.open_selected_survey_files).pack(side=tk.LEFT, padx=6)
        ttk.Button(buttons, text="Add File to Workspace", command=self.add_file_to_workspace_dialog).pack(side=tk.LEFT, padx=6)
        ttk.Button(buttons, text="Open Metadata", command=self.open_selected_survey_metadata).pack(side=tk.LEFT, padx=6)
        ttk.Button(buttons, text="Materialize Selected", command=self.materialize_selected_survey).pack(
            side=tk.LEFT,
            padx=6,
        )
        ttk.Button(buttons, text="Materialize All", command=self.materialize_all_surveys).pack(
            side=tk.LEFT,
            padx=6,
        )
        ttk.Button(buttons, text="Export PDF", command=self.export_selected_survey_pdf).pack(side=tk.LEFT, padx=6)
        ttk.Button(buttons, text="JSONify Files", command=self.jsonify_selected_survey_files).pack(side=tk.LEFT, padx=6)
        ttk.Button(buttons, text="Export All PDFs", command=self.export_all_materialized_pdfs).pack(side=tk.LEFT, padx=6)
        ttk.Button(buttons, text="JSONify All", command=self.jsonify_all_materialized_files).pack(side=tk.LEFT, padx=6)
        ttk.Label(buttons, text="Run").pack(side=tk.LEFT, padx=(18, 6))
        ttk.Combobox(
            buttons,
            textvariable=self.run_stage_var,
            values=SURVEY_RUN_STAGE_IDS,
            state="readonly",
            width=10,
        ).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Label(buttons, text="Args").pack(side=tk.LEFT, padx=(6, 6))
        ttk.Entry(buttons, textvariable=self.run_args_var, width=46).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(buttons, text="Run Stage", command=self.run_selected_stage).pack(side=tk.LEFT, padx=6)

        for col in (1, 2, 4, 5):
            details.grid_columnconfigure(col, weight=1)

        ttk.Label(tab, textvariable=self.status_var).pack(anchor="w", pady=(6, 0))

    def _build_workflows_tab(self, tab: ttk.Frame) -> None:
        intro = ttk.LabelFrame(tab, text="Guided Workflows")
        intro.pack(fill=tk.X, pady=(0, 6))
        ttk.Label(
            intro,
            text=(
                "Choose or set one Main Survey on the Survey Workspaces page. These workflow pages "
                "use that survey to preload likely input and output files, while the Question Browser "
                "stays unfiltered by default."
            ),
            wraplength=1250,
            justify=tk.LEFT,
        ).grid(row=0, column=0, columnspan=4, sticky="w", padx=8, pady=(8, 4))
        ttk.Label(intro, text="Main survey selected").grid(row=1, column=0, sticky="w", padx=8, pady=(0, 8))
        ttk.Entry(intro, textvariable=self.main_survey_id_var, width=18, state="readonly").grid(
            row=1,
            column=1,
            sticky="w",
            padx=6,
            pady=(0, 8),
        )
        ttk.Entry(intro, textvariable=self.main_survey_source_var, width=100, state="readonly").grid(
            row=1,
            column=2,
            sticky="we",
            padx=6,
            pady=(0, 8),
        )
        ttk.Button(intro, text="Use Selected Survey", command=self.set_main_survey_from_selection).grid(
            row=1,
            column=3,
            sticky="w",
            padx=6,
            pady=(0, 8),
        )
        ttk.Checkbutton(
            intro,
            text="Show Advanced Details",
            variable=self.show_advanced_details_var,
            command=self.toggle_advanced_details,
        ).grid(row=2, column=0, sticky="w", padx=8, pady=(0, 8))
        intro.grid_columnconfigure(2, weight=1)

        workflow_book = ttk.Notebook(tab)
        self.workflow_book = workflow_book
        workflow_book.pack(fill=tk.BOTH, expand=True)

        pre_field_tab = ttk.Frame(workflow_book, padding=6)
        daily_tab = ttk.Frame(workflow_book, padding=6)
        final_tab = ttk.Frame(workflow_book, padding=6)
        archive_tab = ttk.Frame(workflow_book, padding=6)
        workflow_book.add(pre_field_tab, text="Pre-fielding")
        workflow_book.add(daily_tab, text="In-field Daily")
        workflow_book.add(final_tab, text="Final Results")
        workflow_book.add(archive_tab, text="Archive/Search")

        pre_primary = ttk.LabelFrame(pre_field_tab, text="Pre-fielding Actions")
        pre_primary.pack(fill=tk.X, pady=(0, 8))
        ttk.Button(pre_primary, text="Generate WinCross Tables", command=self.run_guided_wincross).pack(
            side=tk.LEFT,
            padx=6,
            pady=8,
        )
        if self.tomas_enabled:
            ttk.Button(pre_primary, text="Run Optional TOMAS", command=self.run_guided_tomas).pack(
                side=tk.LEFT,
                padx=6,
                pady=8,
            )
            ttk.Button(pre_primary, text="Open Advanced TOMAS", command=self.open_advanced_tomas_editor).pack(
                side=tk.LEFT,
                padx=6,
                pady=8,
            )
        ttk.Label(pre_primary, textvariable=self.wincross_status_var).pack(side=tk.LEFT, padx=12, pady=8)
        if self.tomas_enabled:
            ttk.Label(pre_primary, textvariable=self.tomas_status_var).pack(side=tk.LEFT, padx=12, pady=8)

        wincross = ttk.LabelFrame(pre_field_tab, text="Run WinCross Pipeline")
        self._register_advanced_detail_frame(wincross, fill=tk.X, pady=(0, 8))
        ttk.Label(wincross, text="DOCX").grid(row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(wincross, textvariable=self.wincross_docx_var, width=96).grid(
            row=0, column=1, columnspan=3, sticky="we", padx=6, pady=4
        )
        ttk.Button(wincross, text="Browse", command=self.browse_wincross_docx).grid(
            row=0, column=4, sticky="w", padx=6, pady=4
        )
        ttk.Label(wincross, text="PDF").grid(row=1, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(wincross, textvariable=self.wincross_pdf_var, width=96).grid(
            row=1, column=1, columnspan=3, sticky="we", padx=6, pady=4
        )
        ttk.Button(wincross, text="Browse", command=self.browse_wincross_pdf).grid(
            row=1, column=4, sticky="w", padx=6, pady=4
        )
        ttk.Label(wincross, text="RAW input").grid(row=2, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(wincross, textvariable=self.wincross_raw_input_var, width=96).grid(
            row=2, column=1, columnspan=3, sticky="we", padx=6, pady=4
        )
        ttk.Button(wincross, text="Browse", command=self.browse_wincross_raw_input).grid(
            row=2, column=4, sticky="w", padx=6, pady=4
        )
        ttk.Label(wincross, text="JSON out").grid(row=3, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(wincross, textvariable=self.wincross_json_out_var, width=96).grid(
            row=3, column=1, columnspan=3, sticky="we", padx=6, pady=4
        )
        ttk.Button(wincross, text="Browse", command=self.browse_wincross_json_out).grid(
            row=3, column=4, sticky="w", padx=6, pady=4
        )
        ttk.Label(wincross, text="TABLES out").grid(row=4, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(wincross, textvariable=self.wincross_tables_out_var, width=96).grid(
            row=4, column=1, columnspan=3, sticky="we", padx=6, pady=4
        )
        ttk.Button(wincross, text="Browse", command=self.browse_wincross_tables_out).grid(
            row=4, column=4, sticky="w", padx=6, pady=4
        )
        ttk.Checkbutton(wincross, text="Translate Spanish", variable=self.wincross_translate_var).grid(
            row=5, column=0, sticky="w", padx=6, pady=4
        )
        ttk.Checkbutton(wincross, text="Skip tables", variable=self.wincross_skip_tables_var).grid(
            row=5, column=1, sticky="w", padx=6, pady=4
        )
        win_buttons = ttk.Frame(wincross)
        win_buttons.grid(row=6, column=0, columnspan=5, sticky="w", padx=6, pady=(6, 8))
        ttk.Button(win_buttons, text="Run WinCross", command=self.run_guided_wincross).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(win_buttons, text="Reset WinCross Defaults", command=self.reset_wincross_form).pack(side=tk.LEFT, padx=6)
        ttk.Label(pre_field_tab, textvariable=self.wincross_status_var).pack(anchor="w", pady=(8, 0))
        for col in (1, 2, 3):
            wincross.grid_columnconfigure(col, weight=1)

        if self.tomas_enabled:
            tomas = ttk.LabelFrame(pre_field_tab, text="Optional TOMAS Spanish Pipeline")
            self._register_advanced_detail_frame(tomas, fill=tk.X, pady=(8, 0))
            ttk.Label(tomas, text="Phone DOCX").grid(row=0, column=0, sticky="w", padx=6, pady=4)
            ttk.Entry(tomas, textvariable=self.tomas_phone_docx_var, width=92).grid(
                row=0, column=1, columnspan=3, sticky="we", padx=6, pady=4
            )
            ttk.Button(tomas, text="Browse", command=lambda: self.browse_tomas_input(self.tomas_phone_docx_var, "DOCX")).grid(
                row=0, column=4, sticky="w", padx=6, pady=4
            )
            ttk.Label(tomas, text="Phone PDF").grid(row=1, column=0, sticky="w", padx=6, pady=4)
            ttk.Entry(tomas, textvariable=self.tomas_phone_pdf_var, width=92).grid(
                row=1, column=1, columnspan=3, sticky="we", padx=6, pady=4
            )
            ttk.Button(tomas, text="Browse", command=lambda: self.browse_tomas_input(self.tomas_phone_pdf_var, "PDF")).grid(
                row=1, column=4, sticky="w", padx=6, pady=4
            )
            ttk.Label(tomas, text="SMS DOCX").grid(row=2, column=0, sticky="w", padx=6, pady=4)
            ttk.Entry(tomas, textvariable=self.tomas_sms_docx_var, width=92).grid(
                row=2, column=1, columnspan=3, sticky="we", padx=6, pady=4
            )
            ttk.Button(tomas, text="Browse", command=lambda: self.browse_tomas_input(self.tomas_sms_docx_var, "DOCX")).grid(
                row=2, column=4, sticky="w", padx=6, pady=4
            )
            ttk.Label(tomas, text="SMS PDF").grid(row=3, column=0, sticky="w", padx=6, pady=4)
            ttk.Entry(tomas, textvariable=self.tomas_sms_pdf_var, width=92).grid(
                row=3, column=1, columnspan=3, sticky="we", padx=6, pady=4
            )
            ttk.Button(tomas, text="Browse", command=lambda: self.browse_tomas_input(self.tomas_sms_pdf_var, "PDF")).grid(
                row=3, column=4, sticky="w", padx=6, pady=4
            )
            ttk.Label(tomas, text="Output folder").grid(row=4, column=0, sticky="w", padx=6, pady=4)
            ttk.Entry(tomas, textvariable=self.tomas_out_dir_var, width=92).grid(
                row=4, column=1, columnspan=3, sticky="we", padx=6, pady=4
            )
            ttk.Button(tomas, text="Browse", command=self.browse_tomas_out_dir).grid(
                row=4, column=4, sticky="w", padx=6, pady=4
            )
            ttk.Checkbutton(tomas, text="Build formatted DOCX", variable=self.tomas_build_formatted_var).grid(
                row=5, column=0, sticky="w", padx=6, pady=4
            )
            tomas_buttons = ttk.Frame(tomas)
            tomas_buttons.grid(row=6, column=0, columnspan=5, sticky="w", padx=6, pady=(6, 8))
            ttk.Button(tomas_buttons, text="Run TOMAS", command=self.run_guided_tomas).pack(side=tk.LEFT, padx=(0, 6))
            ttk.Button(tomas_buttons, text="Open Advanced TOMAS", command=self.open_advanced_tomas_editor).pack(
                side=tk.LEFT,
                padx=6,
            )
            ttk.Label(tomas, textvariable=self.tomas_status_var).grid(
                row=7,
                column=0,
                columnspan=5,
                sticky="w",
                padx=6,
                pady=(0, 8),
            )
            for col in (1, 2, 3):
                tomas.grid_columnconfigure(col, weight=1)

        self._build_data_delivery_section(daily_tab)

        merge_quick = ttk.LabelFrame(daily_tab, text="Daily Merge")
        merge_quick.pack(fill=tk.X, pady=(8, 0))
        ttk.Label(merge_quick, text="Secondary file").grid(row=0, column=0, sticky="w", padx=6, pady=6)
        ttk.Entry(merge_quick, textvariable=self.merge_secondary_path_var, width=92).grid(
            row=0, column=1, sticky="we", padx=6, pady=6
        )
        ttk.Button(merge_quick, text="Browse", command=self.browse_merge_secondary).grid(
            row=0, column=2, sticky="w", padx=6, pady=6
        )
        merge_quick_buttons = ttk.Frame(merge_quick)
        merge_quick_buttons.grid(row=1, column=0, columnspan=3, sticky="w", padx=6, pady=(0, 6))
        ttk.Button(
            merge_quick_buttons,
            text="Run Merge with Saved Map",
            command=self.run_guided_merge_with_saved_map,
        ).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(
            merge_quick_buttons,
            text="Open Advanced Merge Editor",
            command=self.open_advanced_merge_editor,
        ).pack(side=tk.LEFT, padx=6)
        ttk.Button(merge_quick_buttons, text="Reset Merge Defaults", command=self.clear_merge_form).pack(
            side=tk.LEFT, padx=6
        )
        ttk.Label(merge_quick, textvariable=self.merge_status_var).grid(
            row=2, column=0, columnspan=3, sticky="w", padx=6, pady=(0, 8)
        )
        merge_quick.grid_columnconfigure(1, weight=1)

        oe_quick = ttk.LabelFrame(daily_tab, text="Open Ends")
        oe_quick.pack(fill=tk.X, pady=(8, 0))
        ttk.Button(oe_quick, text="Open Advanced OE Editor", command=self.open_advanced_oe_editor).grid(
            row=0, column=0, sticky="w", padx=6, pady=8
        )
        ttk.Label(oe_quick, textvariable=self.oe_status_var).grid(
            row=0, column=1, sticky="w", padx=6, pady=8
        )
        oe_quick.grid_columnconfigure(1, weight=1)

        final_primary = ttk.LabelFrame(final_tab, text="Final Results Actions")
        final_primary.pack(fill=tk.X, pady=(0, 8))
        ttk.Button(final_primary, text="Generate Empty Results", command=self.run_empty_results).pack(
            side=tk.LEFT,
            padx=6,
            pady=8,
        )
        ttk.Button(final_primary, text="Mark Final Data Confirmed", command=self.mark_final_data_confirmed).pack(
            side=tk.LEFT,
            padx=6,
            pady=8,
        )
        ttk.Label(final_primary, textvariable=self.results_status_var).pack(side=tk.LEFT, padx=12, pady=8)

        results = ttk.LabelFrame(final_tab, text="Generate Empty Results")
        self._register_advanced_detail_frame(results, fill=tk.X)
        ttk.Label(results, text="Template DOCX").grid(row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(results, textvariable=self.results_docx_var, width=96).grid(
            row=0, column=1, columnspan=3, sticky="we", padx=6, pady=4
        )
        ttk.Button(results, text="Browse", command=self.browse_results_docx).grid(
            row=0, column=4, sticky="w", padx=6, pady=4
        )
        ttk.Label(results, text="TABLES").grid(row=1, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(results, textvariable=self.results_tables_var, width=96).grid(
            row=1, column=1, columnspan=3, sticky="we", padx=6, pady=4
        )
        ttk.Button(results, text="Browse", command=self.browse_results_tables).grid(
            row=1, column=4, sticky="w", padx=6, pady=4
        )
        ttk.Label(results, text="Survey JSON").grid(row=2, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(results, textvariable=self.results_json_var, width=96).grid(
            row=2, column=1, columnspan=3, sticky="we", padx=6, pady=4
        )
        ttk.Button(results, text="Browse", command=self.browse_results_json).grid(
            row=2, column=4, sticky="w", padx=6, pady=4
        )
        ttk.Label(results, text="Output DOCX").grid(row=3, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(results, textvariable=self.results_output_var, width=96).grid(
            row=3, column=1, columnspan=3, sticky="we", padx=6, pady=4
        )
        ttk.Button(results, text="Browse", command=self.browse_results_output).grid(
            row=3, column=4, sticky="w", padx=6, pady=4
        )
        ttk.Checkbutton(results, text="Insert missing rows", variable=self.results_insert_missing_var).grid(
            row=4, column=0, sticky="w", padx=6, pady=4
        )
        result_buttons = ttk.Frame(results)
        result_buttons.grid(row=5, column=0, columnspan=5, sticky="w", padx=6, pady=(6, 8))
        ttk.Button(result_buttons, text="Generate Empty Results", command=self.run_empty_results).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(result_buttons, text="Use Main Survey Defaults", command=self.sync_results_from_main).pack(side=tk.LEFT, padx=6)
        ttk.Label(final_tab, textvariable=self.results_status_var).pack(anchor="w", pady=(8, 0))
        for col in (1, 2, 3):
            results.grid_columnconfigure(col, weight=1)

        archive = ttk.LabelFrame(archive_tab, text="Repository and Search")
        archive.pack(fill=tk.X)
        ttk.Button(archive, text="Repo Scan", command=self.scan_repository).grid(row=0, column=0, sticky="w", padx=6, pady=8)
        ttk.Button(archive, text="Rebuild Question Index", command=self.rebuild_question_index).grid(
            row=0,
            column=1,
            sticky="w",
            padx=6,
            pady=8,
        )
        ttk.Button(archive, text="Open Question Browser", command=lambda: self._select_top_tab("Question Browser")).grid(
            row=0,
            column=2,
            sticky="w",
            padx=6,
            pady=8,
        )

    def _build_data_delivery_section(self, tab: ttk.Frame) -> None:
        delivery = ttk.LabelFrame(tab, text="Track Data Delivery")
        delivery.pack(fill=tk.X, pady=(0, 12))

        ttk.Label(delivery, text="Survey").grid(row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(delivery, textvariable=self.main_survey_id_var, width=18, state="readonly").grid(
            row=0,
            column=1,
            sticky="w",
            padx=6,
            pady=4,
        )
        ttk.Label(delivery, text="Day").grid(row=0, column=2, sticky="w", padx=6, pady=4)
        ttk.Spinbox(delivery, textvariable=self.data_delivery_day_var, from_=1, to=30, width=6).grid(
            row=0,
            column=3,
            sticky="w",
            padx=6,
            pady=4,
        )
        ttk.Checkbutton(delivery, text="Timeseries", variable=self.data_delivery_timeseries_var).grid(
            row=0,
            column=4,
            sticky="w",
            padx=6,
            pady=4,
        )

        ttk.Label(delivery, text="Data CSV/XLSX").grid(row=1, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(delivery, textvariable=self.data_delivery_path_var, width=104).grid(
            row=1,
            column=1,
            columnspan=4,
            sticky="we",
            padx=6,
            pady=4,
        )
        ttk.Button(delivery, text="Browse", command=self.browse_data_delivery).grid(
            row=1,
            column=5,
            sticky="w",
            padx=6,
            pady=4,
        )
        ttk.Label(delivery, text="Codelist").grid(row=2, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(delivery, textvariable=self.data_delivery_codelist_path_var, width=104).grid(
            row=2,
            column=1,
            columnspan=4,
            sticky="we",
            padx=6,
            pady=4,
        )
        ttk.Button(delivery, text="Browse", command=self.browse_data_delivery_codelist).grid(
            row=2,
            column=5,
            sticky="w",
            padx=6,
            pady=4,
        )
        ttk.Button(delivery, text="Track Delivery", command=self.track_data_delivery).grid(
            row=3,
            column=0,
            sticky="w",
            padx=6,
            pady=(6, 8),
        )
        ttk.Label(delivery, textvariable=self.data_delivery_status_var).grid(
            row=3,
            column=1,
            columnspan=5,
            sticky="w",
            padx=6,
            pady=(6, 8),
        )
        for col in (1, 2, 3, 4):
            delivery.grid_columnconfigure(col, weight=1)

    def _build_questions_tab(self, tab: ttk.Frame) -> None:
        filters = ttk.LabelFrame(tab, text="Filters")
        filters.pack(fill=tk.X, pady=(0, 6))

        ttk.Label(filters, text="Query").grid(row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(filters, textvariable=self.question_query_var, width=44).grid(
            row=0,
            column=1,
            sticky="we",
            padx=6,
            pady=4,
        )
        ttk.Button(filters, text="Search", command=self.run_question_search).grid(
            row=0,
            column=2,
            sticky="w",
            padx=6,
            pady=4,
        )
        ttk.Button(filters, text="Frequency", command=self.run_question_frequency).grid(
            row=0,
            column=3,
            sticky="w",
            padx=6,
            pady=4,
        )
        ttk.Button(filters, text="Preview Selected Row", command=self.preview_selected_question_row).grid(
            row=0,
            column=4,
            sticky="w",
            padx=6,
            pady=4,
        )
        ttk.Button(filters, text="Export CSV", command=self.export_question_results_csv).grid(
            row=0,
            column=5,
            sticky="w",
            padx=6,
            pady=4,
        )
        ttk.Button(filters, text="Audit Index", command=self.run_question_audit).grid(
            row=0,
            column=6,
            sticky="w",
            padx=6,
            pady=4,
        )
        ttk.Button(filters, text="Classification Help", command=self.show_question_classification_help).grid(
            row=0,
            column=7,
            sticky="w",
            padx=6,
            pady=4,
        )
        ttk.Button(filters, text="Clear Results", command=self.clear_question_results).grid(
            row=0,
            column=8,
            sticky="w",
            padx=6,
            pady=4,
        )

        ttk.Label(filters, text="State").grid(row=1, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(filters, textvariable=self.question_state_var, width=8).grid(
            row=1,
            column=1,
            sticky="w",
            padx=6,
            pady=4,
        )
        ttk.Label(filters, text="Region").grid(row=1, column=2, sticky="e", padx=6, pady=4)
        ttk.Entry(filters, textvariable=self.question_region_var, width=24).grid(
            row=1,
            column=3,
            sticky="we",
            padx=6,
            pady=4,
        )
        ttk.Label(filters, text="Client").grid(row=1, column=4, sticky="e", padx=6, pady=4)
        ttk.Entry(filters, textvariable=self.question_client_var, width=24).grid(
            row=1,
            column=5,
            sticky="we",
            padx=6,
            pady=4,
        )
        ttk.Label(filters, text="Survey(s)").grid(row=1, column=6, sticky="e", padx=6, pady=4)
        ttk.Entry(filters, textvariable=self.question_survey_var, width=22).grid(
            row=1,
            column=7,
            sticky="we",
            padx=6,
            pady=4,
        )
        ttk.Label(filters, text="Per page").grid(row=1, column=8, sticky="e", padx=6, pady=4)
        ttk.Entry(filters, textvariable=self.question_limit_var, width=8).grid(
            row=1,
            column=9,
            sticky="w",
            padx=6,
            pady=4,
        )
        ttk.Button(filters, text="Rebuild Index", command=self.rebuild_question_index).grid(
            row=1,
            column=10,
            sticky="w",
            padx=6,
            pady=4,
        )
        self.question_prev_button = ttk.Button(filters, text="Prev", command=self.show_previous_question_page)
        self.question_prev_button.grid(row=1, column=11, sticky="w", padx=6, pady=4)
        self.question_next_button = ttk.Button(filters, text="Next", command=self.show_next_question_page)
        self.question_next_button.grid(row=1, column=12, sticky="w", padx=6, pady=4)
        ttk.Label(filters, textvariable=self.question_page_summary_var).grid(
            row=1,
            column=13,
            sticky="w",
            padx=6,
            pady=4,
        )
        ttk.Label(filters, text="Type").grid(row=2, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(filters, textvariable=self.question_type_filter_var, width=24).grid(
            row=2,
            column=1,
            sticky="we",
            padx=6,
            pady=4,
        )
        ttk.Label(filters, text="Family").grid(row=2, column=2, sticky="e", padx=6, pady=4)
        ttk.Entry(filters, textvariable=self.question_family_filter_var, width=24).grid(
            row=2,
            column=3,
            sticky="we",
            padx=6,
            pady=4,
        )
        ttk.Label(filters, text="Group/Tag").grid(row=2, column=4, sticky="e", padx=6, pady=4)
        ttk.Entry(filters, textvariable=self.question_group_filter_var, width=24).grid(
            row=2,
            column=5,
            sticky="we",
            padx=6,
            pady=4,
        )
        group_preset_combo = ttk.Combobox(
            filters,
            textvariable=self.question_group_preset_var,
            values=QUESTION_GROUP_PRESET_VALUES,
            state="readonly",
            width=20,
        )
        group_preset_combo.grid(
            row=2,
            column=6,
            sticky="we",
            padx=6,
            pady=4,
        )
        group_preset_combo.bind("<<ComboboxSelected>>", self._on_question_group_preset_selected)
        ttk.Label(filters, text="Target").grid(row=2, column=7, sticky="e", padx=6, pady=4)
        ttk.Entry(filters, textvariable=self.question_target_filter_var, width=44).grid(
            row=2,
            column=8,
            columnspan=3,
            sticky="we",
            padx=6,
            pady=4,
        )
        ttk.Label(filters, text="Match").grid(row=3, column=0, sticky="w", padx=6, pady=4)
        ttk.Combobox(
            filters,
            textvariable=self.question_match_mode_var,
            values=QUESTION_MATCH_MODE_LABELS,
            state="readonly",
            width=16,
        ).grid(
            row=3,
            column=1,
            sticky="w",
            padx=6,
            pady=4,
        )

        filters.grid_columnconfigure(1, weight=1)
        filters.grid_columnconfigure(3, weight=1)
        filters.grid_columnconfigure(5, weight=1)
        filters.grid_columnconfigure(7, weight=1)
        filters.grid_columnconfigure(8, weight=1)
        filters.grid_columnconfigure(13, weight=1)

        ttk.Label(tab, textvariable=self.question_summary_var).pack(anchor="w", padx=2, pady=(0, 4))

        result_frame = ttk.LabelFrame(tab, text="Results")
        result_frame.pack(fill=tk.BOTH, expand=True)

        self.question_tree = ttk.Treeview(
            result_frame,
            columns=(
                "mode",
                "survey",
                "qid",
                "qtype",
                "qfamily",
                "qgroup",
                "qtarget",
                "state",
                "region",
                "client",
                "surveys",
                "occurrences",
                "question",
                "item",
            ),
            show="headings",
            selectmode="browse",
            height=20,
        )
        self.question_tree.heading("mode", text="Mode")
        self.question_tree.heading("survey", text="Survey")
        self.question_tree.heading("qid", text="QID")
        self.question_tree.heading("qtype", text="Type")
        self.question_tree.heading("qfamily", text="Family")
        self.question_tree.heading("qgroup", text="Group")
        self.question_tree.heading("qtarget", text="Target")
        self.question_tree.heading("state", text="State")
        self.question_tree.heading("region", text="Region")
        self.question_tree.heading("client", text="Client")
        self.question_tree.heading("surveys", text="Surveys")
        self.question_tree.heading("occurrences", text="Occurrences")
        self.question_tree.heading("question", text="Question")
        self.question_tree.heading("item", text="Item")
        self.question_tree.column("mode", width=70, anchor="w")
        self.question_tree.column("survey", width=160, anchor="w")
        self.question_tree.column("qid", width=80, anchor="w")
        self.question_tree.column("qtype", width=160, anchor="w")
        self.question_tree.column("qfamily", width=180, anchor="w")
        self.question_tree.column("qgroup", width=160, anchor="w")
        self.question_tree.column("qtarget", width=200, anchor="w")
        self.question_tree.column("state", width=70, anchor="w")
        self.question_tree.column("region", width=140, anchor="w")
        self.question_tree.column("client", width=180, anchor="w")
        self.question_tree.column("surveys", width=90, anchor="e")
        self.question_tree.column("occurrences", width=100, anchor="e")
        self.question_tree.column("question", width=480, anchor="w")
        self.question_tree.column("item", width=280, anchor="w")
        self.question_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(6, 0), pady=6)
        self.question_tree.bind("<Double-1>", self.preview_selected_question_row)

        y_scroll = ttk.Scrollbar(result_frame, orient=tk.VERTICAL, command=self.question_tree.yview)
        y_scroll.pack(side=tk.RIGHT, fill=tk.Y, padx=(0, 6), pady=6)
        self.question_tree.configure(yscrollcommand=y_scroll.set)

        x_scroll = ttk.Scrollbar(tab, orient=tk.HORIZONTAL, command=self.question_tree.xview)
        x_scroll.pack(fill=tk.X, padx=(6, 6), pady=(0, 2))
        self.question_tree.configure(xscrollcommand=x_scroll.set)

        ttk.Label(tab, textvariable=self.question_status_var).pack(anchor="w", pady=(2, 0))
        self._update_question_pagination_state(mode="", filters=None, total_count=0, page=1)

    def _browse_repo_root(self) -> None:
        initial = self.repo_root_var.get().strip() or str(DEFAULT_REPO_ROOT)
        chosen = filedialog.askdirectory(initialdir=initial, title="Choose repository root")
        if chosen:
            self.repo_root_var.set(chosen)

    def _browse_source_root(self) -> None:
        initial = self.source_root_var.get().strip() or str(BASE_DIR)
        chosen = filedialog.askdirectory(initialdir=initial, title="Choose source/archive root")
        if chosen:
            self.source_root_var.set(chosen)

    def _cli_path(self) -> Path:
        return BASE_DIR / "survey_tool_cli.py"

    def _selected_survey_id(self) -> str | None:
        survey_id = self.selected_survey_id_var.get().strip()
        if not survey_id:
            survey_id = self.main_survey_id_var.get().strip() or self.survey_selector_var.get().strip()
        if survey_id:
            return survey_id
        messagebox.showwarning("Survey Platform", "Select a survey first.")
        return None

    def _tree_values_for_row(self, row: dict[str, str]) -> tuple[str, ...]:
        return (
            row.get("survey_id", ""),
            row.get("stage", ""),
            row.get("wincross", ""),
            row.get("catalog", ""),
            row.get("merge", ""),
            row.get("oe", ""),
            row.get("tomas", ""),
            row.get("results", ""),
            row.get("region", ""),
            row.get("client", ""),
            row.get("last_run", "") or "(unknown)",
        )

    def _format_platform_statuses(self, row: dict[str, str]) -> str:
        parts: list[str] = []
        for stage_id in REPO_PLATFORM_STAGE_IDS:
            status = str(row.get(stage_id, "")).strip() or "unknown"
            parts.append(f"{stage_id}={status}")
        return " | ".join(parts)

    def _parse_run_args(self) -> list[str]:
        raw = self.run_args_var.get().strip()
        if not raw:
            return []
        try:
            return shlex.split(raw, posix=False)
        except ValueError as exc:
            raise ValueError(f"Failed to parse run args: {exc}") from exc

    def _show_command_output(self, title: str, body_text: str) -> None:
        top = tk.Toplevel(self.root)
        top.title(title)
        top.geometry("1080x760")

        container = ttk.Frame(top, padding=10)
        container.pack(fill=tk.BOTH, expand=True)

        body = tk.Text(container, wrap=tk.NONE)
        body.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        y_scroll = ttk.Scrollbar(container, orient=tk.VERTICAL, command=body.yview)
        y_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        x_scroll = ttk.Scrollbar(top, orient=tk.HORIZONTAL, command=body.xview)
        x_scroll.pack(fill=tk.X, padx=10, pady=(0, 10))

        body.configure(yscrollcommand=y_scroll.set, xscrollcommand=x_scroll.set)
        body.insert("1.0", body_text or "(no output)")
        body.configure(state=tk.DISABLED)

    def _run_cli_command(
        self,
        cmd: list[str],
        *,
        title: str,
        running_text: str,
        success_callback=None,
        status_target: tk.StringVar | None = None,
        show_output_on_success: bool = True,
    ) -> bool:
        cli_path = self._cli_path()
        if not cli_path.exists():
            messagebox.showerror("Survey Platform", f"CLI not found:\n{cli_path}")
            return False

        target = status_target or self.status_var
        target.set(running_text)
        self.root.update_idletasks()

        result = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            cwd=str(BASE_DIR),
        )
        stdout = (result.stdout or "").strip()
        stderr = (result.stderr or "").strip()
        detail = "\n\n".join([part for part in (stdout, stderr) if part])

        if result.returncode != 0:
            target.set(f"{title} failed.")
            self._show_command_output(f"{title} (failed)", detail or "(no output)")
            messagebox.showerror("Survey Platform", f"{title} failed.")
            return False

        if show_output_on_success:
            self._show_command_output(title, detail or "(no output)")
        target.set(stdout.splitlines()[-1] if stdout else f"{title} complete.")
        if callable(success_callback):
            success_callback()
        return True

    def _run_cli_command_async(
        self,
        cmd: list[str],
        *,
        title: str,
        running_text: str,
        on_complete=None,
        success_callback=None,
        status_target: tk.StringVar | None = None,
        show_output_on_success: bool = True,
    ) -> None:
        def _finish_callback(ok: bool, detail: str = "") -> None:
            if not callable(on_complete):
                return
            try:
                signature = inspect.signature(on_complete)
                params = list(signature.parameters.values())
                positional = [
                    param
                    for param in params
                    if param.kind
                    in (inspect.Parameter.POSITIONAL_ONLY, inspect.Parameter.POSITIONAL_OR_KEYWORD)
                ]
                accepts_varargs = any(param.kind == inspect.Parameter.VAR_POSITIONAL for param in params)
                if accepts_varargs or len(positional) >= 2:
                    on_complete(ok, detail)
                else:
                    on_complete(ok)
            except (TypeError, ValueError):
                on_complete(ok)

        cli_path = self._cli_path()
        if not cli_path.exists():
            messagebox.showerror("Survey Platform", f"CLI not found:\n{cli_path}")
            _finish_callback(False)
            return

        target = status_target or self.status_var
        target.set(running_text)
        self.root.update_idletasks()

        def worker() -> None:
            try:
                result = subprocess.run(
                    cmd,
                    capture_output=True,
                    text=True,
                    cwd=str(BASE_DIR),
                )
                stdout = (result.stdout or "").strip()
                stderr = (result.stderr or "").strip()
                detail = "\n\n".join([part for part in (stdout, stderr) if part])
            except Exception as exc:
                self.root.after(0, lambda exc=exc: finish_exception(exc))
                return
            self.root.after(0, lambda: finish_result(result.returncode, stdout, stderr, detail))

        def finish_exception(exc: Exception) -> None:
            target.set(f"{title} failed.")
            self._show_command_output(f"{title} (failed)", str(exc))
            messagebox.showerror("Survey Platform", f"{title} failed.")
            _finish_callback(False, str(exc))

        def finish_result(returncode: int, stdout: str, stderr: str, detail: str) -> None:
            if returncode != 0:
                target.set(f"{title} failed.")
                self._show_command_output(f"{title} (failed)", detail or "(no output)")
                messagebox.showerror("Survey Platform", f"{title} failed.")
                _finish_callback(False, detail)
                return

            if show_output_on_success:
                self._show_command_output(title, detail or "(no output)")
            target.set(stdout.splitlines()[-1] if stdout else f"{title} complete.")
            if callable(success_callback):
                success_callback()
            _finish_callback(True, detail)

        threading.Thread(target=worker, daemon=True).start()

    def _repo_root(self) -> Path:
        raw = self.repo_root_var.get().strip()
        if not raw:
            return DEFAULT_REPO_ROOT
        return Path(os.path.expanduser(os.path.expandvars(raw))).resolve()

    def _index_path(self) -> Path:
        return self._repo_root() / REPO_INDEX_FILENAME

    def _meta_path(self, survey_id: str) -> Path:
        return self._repo_root() / REPO_SURVEYS_DIRNAME / survey_id / REPO_META_FILENAME

    def _script_path(self, filename: str) -> Path:
        return BASE_DIR / filename

    def _config_path(self) -> Path:
        return _package_config_path()

    def _read_config(self) -> dict:
        path = self._config_path()
        if not path.exists():
            return {}
        try:
            return _read_json_dict(path)
        except Exception:
            return {}

    def _write_config(self, payload: dict) -> None:
        _write_json(self._config_path(), payload)

    def _read_dashboard_preferences(self) -> dict[str, object]:
        config = self._read_config()
        prefs = config.get(DASHBOARD_CONFIG_KEY, {})
        return prefs if isinstance(prefs, dict) else {}

    def _save_dashboard_survey_preference(self, survey_id: str, *, confirmed: bool = False) -> None:
        survey_id = survey_id.strip()
        if not survey_id:
            return
        config = self._read_config()
        prefs = config.get(DASHBOARD_CONFIG_KEY, {})
        if not isinstance(prefs, dict):
            prefs = {}
        prefs["last_selected_survey_id"] = survey_id
        if confirmed:
            prefs["last_main_survey_id"] = survey_id
        prefs["updated_at"] = _utc_now_iso()
        config[DASHBOARD_CONFIG_KEY] = prefs
        try:
            self._write_config(config)
        except Exception as exc:
            if hasattr(self, "status_var"):
                self.status_var.set(f"Could not save dashboard preference: {exc}")

    def _preferred_dashboard_survey_id(
        self,
        available_survey_ids: set[str],
        current_main_survey_id: str = "",
    ) -> str:
        current_main_survey_id = current_main_survey_id.strip()
        if current_main_survey_id in available_survey_ids:
            return current_main_survey_id
        prefs = self._read_dashboard_preferences()
        for key in ("last_main_survey_id", "last_selected_survey_id"):
            survey_id = str(prefs.get(key, "")).strip()
            if survey_id in available_survey_ids:
                return survey_id
        return ""

    def _workspace_root(self) -> Path:
        return BASE_DIR / WORKSPACE_DIRNAME

    def _survey_workspace_dir(self, survey_id: str) -> Path:
        return self._workspace_root() / survey_id

    def _workflow_status_path(self, survey_id: str) -> Path:
        return self._survey_workspace_dir(survey_id) / WORKFLOW_STATUS_FILENAME

    def _workspace_manifest_path(self, survey_id: str) -> Path:
        return self._survey_workspace_dir(survey_id) / WORKSPACE_MANIFEST_FILENAME

    def _survey_state_path(self, survey_id: str) -> Path:
        return self._survey_workspace_dir(survey_id) / SURVEY_STATE_FILENAME

    def _default_survey_state(self, survey_id: str) -> dict[str, object]:
        return {
            "schema_version": 1,
            "survey_id": survey_id,
            "updated_at": "",
            "data_intake": {},
        }

    def _read_survey_state(self, survey_id: str) -> dict[str, object]:
        if not survey_id:
            return self._default_survey_state("")
        path = self._survey_state_path(survey_id)
        if not path.exists():
            return self._default_survey_state(survey_id)
        try:
            payload = _read_json_dict(path)
        except Exception:
            return self._default_survey_state(survey_id)
        if str(payload.get("survey_id", "")).strip() != survey_id:
            payload["survey_id"] = survey_id
        if not isinstance(payload.get("data_intake"), dict):
            payload["data_intake"] = {}
        return payload

    def _write_survey_state(self, survey_id: str, payload: dict[str, object]) -> None:
        if not survey_id:
            return
        payload["schema_version"] = 1
        payload["survey_id"] = survey_id
        payload["updated_at"] = _utc_now_iso()
        _write_json(self._survey_state_path(survey_id), payload)

    def _default_workspace_manifest(self, survey_id: str) -> dict[str, object]:
        return {
            "schema_version": 1,
            "survey_id": survey_id,
            "updated_at": "",
            "files": {},
        }

    def _read_workspace_manifest(self, survey_id: str) -> dict[str, object]:
        if not survey_id:
            return self._default_workspace_manifest("")
        path = self._workspace_manifest_path(survey_id)
        if not path.exists():
            return self._default_workspace_manifest(survey_id)
        try:
            payload = _read_json_dict(path)
        except Exception:
            return self._default_workspace_manifest(survey_id)
        if str(payload.get("survey_id", "")).strip() != survey_id:
            payload["survey_id"] = survey_id
        if not isinstance(payload.get("files"), dict):
            payload["files"] = {}
        return payload

    def _write_workspace_manifest(self, survey_id: str, payload: dict[str, object]) -> None:
        if not survey_id:
            return
        payload["schema_version"] = 1
        payload["survey_id"] = survey_id
        payload["updated_at"] = _utc_now_iso()
        _write_json(self._workspace_manifest_path(survey_id), payload)

    def _workspace_role_def(self, role_id: str) -> tuple[str, str, tuple[str, ...]]:
        try:
            return WORKSPACE_FILE_ROLE_BY_ID[role_id]
        except KeyError as exc:
            raise ValueError(f"Unsupported workspace file role: {role_id}") from exc

    def _workspace_role_dir(self, survey_id: str, role_id: str) -> Path:
        _label, rel_dir, _extensions = self._workspace_role_def(role_id)
        return self._survey_workspace_dir(survey_id).joinpath(*Path(rel_dir).parts)

    def _workspace_role_file_candidates(self, survey_id: str, role_id: str) -> list[Path]:
        _label, _rel_dir, extensions = self._workspace_role_def(role_id)
        role_dir = self._workspace_role_dir(survey_id, role_id)
        if not role_dir.exists() or not role_dir.is_dir():
            return []
        paths = [path for path in role_dir.iterdir() if path.is_file()]
        if extensions:
            allowed = {ext.casefold() for ext in extensions}
            paths = [path for path in paths if path.suffix.casefold() in allowed]
        paths.sort(key=lambda path: (path.stat().st_mtime if path.exists() else 0, path.name), reverse=True)
        return paths

    def _record_workspace_file(
        self,
        survey_id: str,
        role_id: str,
        dest_path: Path,
        *,
        source_path: Path | None = None,
        copied: bool = True,
        workspace_dir: str | None = None,
    ) -> None:
        manifest = self._read_workspace_manifest(survey_id)
        files = manifest.get("files", {})
        if not isinstance(files, dict):
            files = {}
            manifest["files"] = files
        label, rel_dir, _extensions = self._workspace_role_def(role_id)
        entries = files.get(role_id, [])
        if not isinstance(entries, list):
            entries = []
        dest_resolved = str(dest_path)
        entries = [
            entry
            for entry in entries
            if not (isinstance(entry, dict) and str(entry.get("path", "")).strip() == dest_resolved)
        ]
        entries.insert(
            0,
            {
                "role": role_id,
                "label": label,
                "path": dest_resolved,
                "source_path": str(source_path) if source_path is not None else "",
                "workspace_dir": workspace_dir or rel_dir,
                "copied": bool(copied),
                "added_at": _utc_now_iso(),
            },
        )
        files[role_id] = entries[:20]
        self._write_workspace_manifest(survey_id, manifest)

    def _preferred_workspace_role_file(self, survey_id: str, *role_ids: str) -> str:
        if not survey_id:
            return ""
        manifest = self._read_workspace_manifest(survey_id)
        files = manifest.get("files", {})
        if not isinstance(files, dict):
            files = {}
        for role_id in role_ids:
            entries = files.get(role_id, [])
            if isinstance(entries, list):
                for entry in entries:
                    if not isinstance(entry, dict):
                        continue
                    raw_path = str(entry.get("path", "")).strip()
                    if not raw_path:
                        continue
                    path = self._resolve_gui_path(raw_path)
                    if path.exists() and path.is_file():
                        return str(path)
            for path in self._workspace_role_file_candidates(survey_id, role_id):
                return str(path)
        return ""

    def _default_workflow_status(self, survey_id: str) -> dict[str, object]:
        return {
            "schema_version": 1,
            "survey_id": survey_id,
            "current_phase": PHASE_VALUES[0],
            "data_day": "",
            "updated_at": "",
            "steps": {},
        }

    def _read_workflow_status(self, survey_id: str) -> dict[str, object]:
        if not survey_id:
            return self._default_workflow_status("")
        path = self._workflow_status_path(survey_id)
        if not path.exists():
            return self._default_workflow_status(survey_id)
        try:
            payload = _read_json_dict(path)
        except Exception:
            return self._default_workflow_status(survey_id)
        if str(payload.get("survey_id", "")).strip() != survey_id:
            payload["survey_id"] = survey_id
        if str(payload.get("current_phase", "")).strip() not in PHASE_VALUES:
            payload["current_phase"] = PHASE_VALUES[0]
        if not isinstance(payload.get("steps"), dict):
            payload["steps"] = {}
        return payload

    def _write_workflow_status(self, survey_id: str, payload: dict[str, object]) -> None:
        if not survey_id:
            return
        payload["schema_version"] = 1
        payload["survey_id"] = survey_id
        payload["updated_at"] = _utc_now_iso()
        _write_json(self._workflow_status_path(survey_id), payload)

    def _workflow_step_payload(self, payload: dict[str, object], step_id: str) -> dict[str, object]:
        steps = payload.get("steps", {})
        if not isinstance(steps, dict):
            steps = {}
            payload["steps"] = steps
        step_payload = steps.get(step_id, {})
        if not isinstance(step_payload, dict):
            step_payload = {}
        return step_payload

    def _set_workflow_phase(self, survey_id: str, phase: str) -> None:
        if not survey_id:
            return
        clean_phase = phase if phase in PHASE_VALUES else PHASE_VALUES[0]
        payload = self._read_workflow_status(survey_id)
        payload["current_phase"] = clean_phase
        self._write_workflow_status(survey_id, payload)
        self.phase_var.set(clean_phase)
        self._refresh_workflow_home()

    def _mark_workflow_step(
        self,
        step_id: str,
        status: str,
        *,
        phase: str | None = None,
        note: str = "",
        survey_id: str = "",
    ) -> None:
        current_survey_id = self.main_survey_id_var.get().strip() or self.selected_survey_id_var.get().strip()
        survey_id = survey_id.strip() or current_survey_id
        if not survey_id:
            return
        clean_status = str(status or "pending").strip().casefold()
        if clean_status not in WORKFLOW_STEP_STATUS_VALUES:
            clean_status = "pending"
        payload = self._read_workflow_status(survey_id)
        if phase in PHASE_VALUES:
            payload["current_phase"] = phase
            if survey_id == current_survey_id:
                self.phase_var.set(phase)
        steps = payload.get("steps", {})
        if not isinstance(steps, dict):
            steps = {}
            payload["steps"] = steps
        steps[step_id] = {
            "status": clean_status,
            "updated_at": _utc_now_iso(),
            "note": note,
        }
        self._write_workflow_status(survey_id, payload)
        self._refresh_workflow_home()

    def _record_merge_output(
        self,
        survey_id: str,
        output_path: Path | None,
        *,
        phase: str | None = "In-field Daily",
    ) -> None:
        # Guided and advanced-editor merges both record here so the workspace
        # manifest and merge_generated step stay in sync with outputs/merge.
        if not survey_id:
            return
        if output_path is not None and output_path.exists() and output_path.is_file():
            self._record_workspace_file(survey_id, "merge_output", output_path, copied=False)
        self._mark_workflow_step("merge_generated", "done", phase=phase, survey_id=survey_id)

    def _set_workflow_data_day(self, survey_id: str, day: int) -> None:
        if not survey_id:
            return
        payload = self._read_workflow_status(survey_id)
        payload["data_day"] = max(1, int(day))
        payload["current_phase"] = "In-field Daily"
        self._write_workflow_status(survey_id, payload)
        self.phase_var.set("In-field Daily")
        self._refresh_workflow_home()

    def _refresh_workflow_home(self) -> None:
        survey_id = self.main_survey_id_var.get().strip()
        phase = self.phase_var.get().strip() or PHASE_VALUES[0]
        self.phase_primary_action_var.set(self._phase_primary_action_label(phase))
        if not survey_id:
            self.workflow_status_summary_var.set("Select a survey from the repo index, then Confirm Survey.")
            self.workflow_workspace_var.set("")
            self.workflow_detail_var.set("")
            if self.workflow_status_tree is not None:
                for item in self.workflow_status_tree.get_children():
                    self.workflow_status_tree.delete(item)
            return

        workspace = self._survey_workspace_dir(survey_id)
        self.workflow_workspace_var.set(str(workspace))
        payload = self._read_workflow_status(survey_id)
        saved_phase = str(payload.get("current_phase", "")).strip()
        if saved_phase in PHASE_VALUES and saved_phase != phase:
            self.phase_var.set(saved_phase)
            phase = saved_phase
            self.phase_primary_action_var.set(self._phase_primary_action_label(phase))

        steps_payload = payload.get("steps", {})
        if not isinstance(steps_payload, dict):
            steps_payload = {}
        done_count = 0
        total_count = 0
        for step_defs in self._workflow_step_defs().values():
            total_count += len(step_defs)
            for step_id, _label in step_defs:
                step_payload = steps_payload.get(step_id, {})
                status = ""
                if isinstance(step_payload, dict):
                    status = str(step_payload.get("status", "")).strip()
                if status in {"done", "reviewed", "skipped"}:
                    done_count += 1

        updated_at = str(payload.get("updated_at", "")).strip() or "(not saved yet)"
        data_day = str(payload.get("data_day", "") or "").strip()
        day_text = f" | data_day={data_day}" if data_day else ""
        self.workflow_status_summary_var.set(
            f"{survey_id} | phase={phase}{day_text} | checklist={done_count}/{total_count} set | updated={updated_at}"
        )
        self.workflow_detail_var.set("Pipeline buttons stay manual; review steps are marked here.")

        if self.workflow_status_tree is None:
            return
        for item in self.workflow_status_tree.get_children():
            self.workflow_status_tree.delete(item)
        for phase_name, step_defs in self._workflow_step_defs().items():
            for step_id, label in step_defs:
                step_payload = steps_payload.get(step_id, {})
                if not isinstance(step_payload, dict):
                    step_payload = {}
                status = str(step_payload.get("status", "pending") or "pending")
                updated = str(step_payload.get("updated_at", "") or "")
                self.workflow_status_tree.insert(
                    "",
                    tk.END,
                    iid=step_id,
                    values=(phase_name, label, status, updated),
                )

    def _workflow_step_defs(self) -> dict[str, tuple[tuple[str, str], ...]]:
        """Checklist steps, without TOMAS steps while TOMAS is turned off."""
        if self.tomas_enabled:
            return WORKFLOW_STEP_DEFS
        return {
            phase: tuple(step for step in steps if not step[0].startswith("tomas_"))
            for phase, steps in WORKFLOW_STEP_DEFS.items()
        }

    def _phase_primary_action_label(self, phase: str) -> str:
        if phase == "Pre-fielding":
            return "Generate WinCross Tables"
        if phase == "In-field Daily":
            return "Track Today's Data"
        if phase == "Final Results":
            return "Generate Empty Results"
        if phase == "Archive/Search":
            return "Open Question Browser"
        return "Run Phase Action"

    def save_current_phase(self) -> None:
        survey_id = self.main_survey_id_var.get().strip() or self.selected_survey_id_var.get().strip()
        if not survey_id:
            messagebox.showinfo("Survey Platform", "Confirm a survey first.")
            return
        self._set_workflow_phase(survey_id, self.phase_var.get().strip() or PHASE_VALUES[0])
        self.status_var.set(f"Saved phase for {survey_id}: {self.phase_var.get().strip()}")

    def mark_selected_workflow_step(self, status: str) -> None:
        if self.workflow_status_tree is None:
            return
        survey_id = self.main_survey_id_var.get().strip()
        if not survey_id:
            messagebox.showinfo("Survey Platform", "Confirm a survey first.")
            return
        selected = self.workflow_status_tree.selection()
        if not selected:
            messagebox.showinfo("Survey Platform", "Select a checklist item first.")
            return
        step_id = selected[0]
        values = self.workflow_status_tree.item(step_id, "values")
        phase = str(values[0]) if values else self.phase_var.get().strip()
        self._mark_workflow_step(step_id, status, phase=phase if phase in PHASE_VALUES else None)
        self.status_var.set(f"Marked {step_id} as {status}.")

    def run_phase_primary_action(self) -> None:
        phase = self.phase_var.get().strip() or PHASE_VALUES[0]
        if phase == "Pre-fielding":
            self.run_guided_wincross()
            return
        if phase == "In-field Daily":
            self.track_data_delivery()
            return
        if phase == "Final Results":
            self.run_empty_results()
            return
        if phase == "Archive/Search":
            self._select_top_tab("Question Browser")
            return
        self.go_to_selected_phase()

    def _survey_metadata_dir(self, survey_id: str) -> Path:
        return self._meta_path(survey_id).parent

    def _candidate_materialized_roots(self) -> list[Path]:
        repo_root = self._repo_root()
        explicit_names = (
            "materialized_surveys_csv_priority",
            "materialized_surveys",
        )
        roots: list[Path] = []
        seen: set[Path] = set()
        for name in explicit_names:
            candidate = repo_root / name
            if candidate in seen:
                continue
            seen.add(candidate)
            roots.append(candidate)
        for candidate in sorted(repo_root.glob("materialized*")):
            if not candidate.is_dir() or candidate in seen:
                continue
            seen.add(candidate)
            roots.append(candidate)
        return roots

    def _survey_materialized_dir(self, survey_id: str) -> Path | None:
        for root in self._candidate_materialized_roots():
            candidate = root / survey_id
            if candidate.exists() and candidate.is_dir():
                return candidate
        return None

    def _survey_working_dir(self, survey_id: str) -> Path | None:
        workspace = self._survey_workspace_dir(survey_id)
        if workspace.exists() and workspace.is_dir():
            return workspace
        return self._survey_materialized_dir(survey_id)

    def _default_materialized_root(self) -> Path | None:
        for root in self._candidate_materialized_roots():
            if root.exists() and root.is_dir():
                return root
        return None

    def _ensure_materialized_root(self, *, title: str) -> Path | None:
        root = self._default_materialized_root()
        if root is None:
            messagebox.showwarning(
                "Survey Platform",
                "No materialized survey root was found.\n\nMaterialize surveys first.",
            )
            return None
        survey_dirs = [p for p in root.iterdir() if p.is_dir() and re.match(r"^NPA\d+$", p.name)]
        if not survey_dirs:
            messagebox.showwarning(
                "Survey Platform",
                f"No materialized survey folders were found under:\n{root}",
            )
            return None
        if not messagebox.askyesno(
            "Survey Platform",
            f"{title} for all materialized surveys under:\n{root}\n\nThis may take a while. Continue?",
        ):
            return None
        return root

    def _resolve_gui_path(self, raw: str) -> Path:
        path = Path(str(raw).strip())
        if path.is_absolute():
            return path
        return (BASE_DIR / path).resolve()

    def _meta_payload_for_survey(self, survey_id: str) -> dict:
        meta_path = self._meta_path(survey_id)
        if not meta_path.exists():
            return {}
        try:
            return _read_json_dict(meta_path)
        except Exception:
            return {}

    def _survey_file_paths(self, survey_id: str, kind: str) -> list[str]:
        meta_payload = self._meta_payload_for_survey(survey_id)
        files = meta_payload.get("files", {})
        if not isinstance(files, dict):
            return []
        values = files.get(kind, [])
        if not isinstance(values, list):
            return []
        return [str(item).strip() for item in values if str(item).strip()]

    def _survey_stage_paths(self, survey_id: str, stage_id: str, block: str, kind: str) -> list[str]:
        meta_payload = self._meta_payload_for_survey(survey_id)
        stages = meta_payload.get("stages", {})
        if not isinstance(stages, dict):
            return []
        stage_payload = stages.get(stage_id, {})
        if not isinstance(stage_payload, dict):
            return []
        block_payload = stage_payload.get(block, {})
        if not isinstance(block_payload, dict):
            return []
        values = block_payload.get(kind, [])
        if not isinstance(values, list):
            return []
        return [str(item).strip() for item in values if str(item).strip()]

    def _first_stage_path(self, survey_id: str, stage_id: str, block: str, kind: str) -> str:
        paths = self._survey_stage_paths(survey_id, stage_id, block, kind)
        return paths[0] if paths else ""

    def _preferred_survey_file(self, survey_id: str, kind: str) -> str:
        paths = self._survey_file_paths(survey_id, kind)
        if not paths:
            return ""
        if kind == "csv":
            preferred = [p for p in paths if "ts" not in Path(p).name.casefold()]
            return preferred[0] if preferred else paths[0]
        if kind == "tables_txt":
            checked = [p for p in paths if "autogen" not in Path(p).name.casefold()]
            return checked[0] if checked else paths[0]
        return paths[0]

    def _preferred_variant_survey_file(self, survey_id: str, kind: str, tokens: tuple[str, ...]) -> str:
        paths = self._survey_file_paths(survey_id, kind)
        if not paths:
            return ""
        scored: list[tuple[int, str]] = []
        for path in paths:
            name = Path(path).name.casefold()
            score = sum(1 for token in tokens if token in name)
            if "final" in name:
                score += 1
            if "spanish" in name or "_es" in name or " es " in name:
                score -= 2
            if score > 0:
                scored.append((score, path))
        if not scored:
            return ""
        scored.sort(key=lambda item: (-item[0], item[1]))
        return scored[0][1]

    def _detect_primary_csv_for_survey(self, survey_id: str) -> str:
        return (
            self._preferred_workspace_role_file(survey_id, "current_data")
            or self._preferred_workspace_role_file(survey_id, "timeseries_current")
            or self._preferred_survey_file(survey_id, "csv")
        )

    def _default_workspace_file(self, survey_id: str, *parts: str) -> str:
        if not survey_id:
            return ""
        return str(self._survey_workspace_dir(survey_id).joinpath(*parts))

    def _default_wincross_raw_path(self, survey_id: str) -> str:
        return self._default_workspace_file(survey_id, "outputs", "wincross", "raw", f"{survey_id}RAW.txt")

    def _default_wincross_json_workspace_path(self, survey_id: str) -> str:
        return self._default_workspace_file(survey_id, "outputs", "wincross", "json", f"{survey_id}JSON.json")

    def _default_wincross_tables_workspace_path(self, survey_id: str) -> str:
        return self._default_workspace_file(survey_id, "outputs", "wincross", "tables", f"{survey_id}TABLES.autogen.txt")

    def _current_data_day(self, survey_id: str) -> int | None:
        status_day = str(self._read_workflow_status(survey_id).get("data_day", "")).strip() if survey_id else ""
        for raw in (status_day,):
            try:
                day = int(raw)
            except Exception:
                continue
            if day >= 1:
                return day
        latest_day = self._latest_data_delivery_day(survey_id)
        if latest_day is not None:
            return latest_day
        if self.phase_var.get().strip() == "In-field Daily":
            try:
                day = int(self.data_delivery_day_var.get().strip())
            except Exception:
                return None
            return day if day >= 1 else None
        return None

    def _default_merge_output_path(self, survey_id: str, secondary_raw: str = "", day: int | None = None) -> str:
        if not survey_id:
            return ""
        day = day if day is not None else self._current_data_day(survey_id)
        day_part = f"_day{day}" if day is not None and day >= 1 else ""
        suffix = "ts" if "ts" in Path(secondary_raw).stem.casefold() else "merged"
        name = f"{survey_id}{day_part}_{suffix}.csv"
        return self._default_workspace_file(survey_id, "outputs", "merge", name)

    def _merge_map_secondary_key(self, secondary_raw: str) -> str:
        secondary_id = self._survey_id_from_text(secondary_raw)
        if secondary_id:
            return secondary_id
        stem = self._resolve_gui_path(secondary_raw).stem if secondary_raw.strip() else "secondary"
        safe = re.sub(r"[^A-Za-z0-9]+", "_", stem).strip("_")
        return safe or "secondary"

    def _default_merge_map_path(self, survey_id: str, secondary_raw: str) -> Path:
        secondary_key = self._merge_map_secondary_key(secondary_raw)
        filename = f"{survey_id}__{secondary_key}.map.txt"
        return self._workspace_role_dir(survey_id, "merge_map") / filename

    def _merge_primary_json_for_survey(self, survey_id: str) -> str:
        if not survey_id:
            return ""
        return (
            self._preferred_workspace_role_file(survey_id, "wincross_json")
            or self._preferred_survey_file(survey_id, "survey_json")
        )

    def _merge_secondary_json_for_path(self, secondary_raw: str) -> str:
        secondary_survey_id = self._survey_id_from_text(secondary_raw)
        if not secondary_survey_id:
            return ""
        return (
            self._preferred_workspace_role_file(secondary_survey_id, "wincross_json")
            or self._preferred_survey_file(secondary_survey_id, "survey_json")
        )

    def _merge_map_records(self, survey_id: str) -> dict[str, dict]:
        state = self._read_survey_state(survey_id)
        merge_maps = state.get("merge_maps", {})
        if not isinstance(merge_maps, dict):
            return {}
        return {
            str(key): value
            for key, value in merge_maps.items()
            if isinstance(value, dict)
        }

    def _saved_guided_merge_record(
        self,
        survey_id: str,
        secondary_raw: str = "",
    ) -> tuple[str, dict[str, object] | None, str]:
        records = self._merge_map_records(survey_id)
        if not records:
            return "", None, ""
        if secondary_raw.strip():
            secondary_key = self._merge_map_secondary_key(secondary_raw)
            record = records.get(secondary_key)
            if record is not None:
                return secondary_key, record, ""
        if len(records) == 1:
            key, record = next(iter(records.items()))
            return key, record, f"Using saved merge setup for {key}."
        return "", None, "Choose a secondary file with a saved merge map, or open the Advanced Merge Editor."

    def _latest_merge_run_path_for_output(self, output_raw: str, primary_raw: str = "") -> Path | None:
        candidate_dirs: list[Path] = []
        for raw in (output_raw, primary_raw):
            text = str(raw or "").strip()
            if not text:
                continue
            path = self._resolve_gui_path(text)
            parent = path.parent if path.suffix else path
            candidate_dirs.append(parent / "merge_runs")
        for run_dir in candidate_dirs:
            if not run_dir.exists() or not run_dir.is_dir():
                continue
            candidates = [path for path in run_dir.glob("*.json") if path.is_file()]
            if candidates:
                candidates.sort(key=lambda path: (path.stat().st_mtime, path.name), reverse=True)
                return candidates[0]
        return None

    def _merge_run_path_from_output_text(
        self,
        detail: str,
        *,
        output_raw: str,
        primary_raw: str,
    ) -> Path | None:
        for line in str(detail or "").splitlines():
            if not line.casefold().startswith("catalog run json:"):
                continue
            raw_path = line.split(":", 1)[1].strip()
            if not raw_path:
                continue
            path = self._resolve_gui_path(raw_path)
            if path.exists() and path.is_file():
                return path
        return self._latest_merge_run_path_for_output(output_raw, primary_raw)

    def _merge_map_header_snapshot(self, raw_path: str) -> list[str]:
        return self._table_headers(raw_path)

    def _saved_guided_merge_setup(
        self,
        survey_id: str,
        *,
        primary_raw: str,
        secondary_raw: str,
    ) -> tuple[Path | None, str, dict[str, object]]:
        if not survey_id:
            return None, "", {}
        _secondary_key, record, note = self._saved_guided_merge_record(survey_id, secondary_raw)
        if record is None:
            return None, note, {}
        raw_map_path = str(record.get("map_path", "")).strip()
        if not raw_map_path:
            return None, "", {}
        map_path = self._resolve_gui_path(raw_map_path)
        if not map_path.exists() or not map_path.is_file():
            return None, f"Saved merge map is missing: {map_path}", {}

        locked_primary = str(record.get("primary_path", "") or primary_raw).strip()
        locked_secondary = str(record.get("secondary_path", "") or secondary_raw).strip()
        if not locked_primary or not locked_secondary:
            return None, "Saved merge setup is missing primary or secondary paths.", {}
        for label, raw_path in (("primary", locked_primary), ("secondary", locked_secondary)):
            path = self._resolve_gui_path(raw_path)
            if not path.exists() or not path.is_file():
                return None, f"Saved {label} merge file is missing: {path}", {}

        primary_headers = self._merge_map_header_snapshot(locked_primary)
        secondary_headers = self._merge_map_header_snapshot(locked_secondary)
        saved_primary_headers = record.get("primary_headers", [])
        saved_secondary_headers = record.get("secondary_headers", [])
        if isinstance(saved_primary_headers, list) and saved_primary_headers and saved_primary_headers != primary_headers:
            return None, "Saved merge map not reused because primary headers changed.", {}
        if isinstance(saved_secondary_headers, list) and saved_secondary_headers and saved_secondary_headers != secondary_headers:
            return None, "Saved merge map not reused because secondary headers changed.", {}
        locked_record = dict(record)
        locked_record["primary_path"] = locked_primary
        locked_record["secondary_path"] = locked_secondary
        reuse_note = note or f"Reusing saved merge map: {map_path}"
        return map_path, reuse_note, locked_record

    def _saved_guided_merge_map(
        self,
        survey_id: str,
        *,
        primary_raw: str,
        secondary_raw: str,
    ) -> tuple[Path | None, str]:
        map_path, note, _record = self._saved_guided_merge_setup(
            survey_id,
            primary_raw=primary_raw,
            secondary_raw=secondary_raw,
        )
        return map_path, note

    def _save_guided_merge_map_from_run(
        self,
        survey_id: str,
        *,
        primary_raw: str,
        secondary_raw: str,
        output_raw: str,
        command_output: str,
        primary_json_raw: str = "",
        secondary_json_raw: str = "",
    ) -> Path | None:
        run_path = self._merge_run_path_from_output_text(
            command_output,
            output_raw=output_raw,
            primary_raw=primary_raw,
        )
        if run_path is None:
            return None
        try:
            payload = _read_json_dict(run_path)
        except Exception:
            return None

        mapping_rows = payload.get("column_mappings", [])
        if not isinstance(mapping_rows, list):
            return None
        pairs: list[tuple[str, str]] = []
        seen_primary: set[str] = set()
        for row in mapping_rows:
            if not isinstance(row, dict) or not row.get("linked", True):
                continue
            primary_col = str(row.get("primary_col", "")).strip()
            secondary_col = str(row.get("secondary_col", "")).strip()
            if not primary_col or not secondary_col or primary_col in seen_primary:
                continue
            pairs.append((primary_col, secondary_col))
            seen_primary.add(primary_col)

        blocked_primary: list[str] = []
        raw_blocked = payload.get("blocked_primary_columns", [])
        if isinstance(raw_blocked, list):
            seen_blocked: set[str] = set()
            for value in raw_blocked:
                primary_col = str(value).strip()
                if (
                    not primary_col
                    or primary_col in seen_primary
                    or primary_col in seen_blocked
                ):
                    continue
                blocked_primary.append(primary_col)
                seen_blocked.add(primary_col)
        if not pairs and not blocked_primary:
            return None

        secondary_key = self._merge_map_secondary_key(secondary_raw)
        map_path = self._default_merge_map_path(survey_id, secondary_raw)
        map_path.parent.mkdir(parents=True, exist_ok=True)
        lines = [
            "# Survey Platform guided merge map",
            f"# primary={survey_id}",
            f"# secondary={secondary_key}",
            f"# saved_at={_utc_now_iso()}",
            f"# source_run={run_path}",
            "",
        ]
        lines.extend(f"{primary_col}={secondary_col}" for primary_col, secondary_col in pairs)
        lines.extend(f"{primary_col}=" for primary_col in blocked_primary)
        map_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

        primary_headers = self._merge_map_header_snapshot(primary_raw)
        secondary_headers = self._merge_map_header_snapshot(secondary_raw)
        state = self._read_survey_state(survey_id)
        merge_maps = state.get("merge_maps")
        if not isinstance(merge_maps, dict):
            merge_maps = {}
        state["merge_maps"] = merge_maps
        merge_maps[secondary_key] = {
            "updated_at": _utc_now_iso(),
            "primary_survey_id": survey_id,
            "secondary_key": secondary_key,
            "primary_path": primary_raw,
            "secondary_path": secondary_raw,
            "primary_survey_json": primary_json_raw,
            "secondary_survey_json": secondary_json_raw,
            "map_path": str(map_path),
            "source_run_path": str(run_path),
            "output_path": output_raw,
            "mapping_count": len(pairs),
            "blocked_primary_count": len(blocked_primary),
            "primary_headers": primary_headers,
            "secondary_headers": secondary_headers,
        }
        self._write_survey_state(survey_id, state)
        self._record_workspace_file(survey_id, "merge_map", map_path, copied=False)
        return map_path

    def _is_default_merge_output_path(self, survey_id: str, raw_path: str) -> bool:
        if not survey_id or not raw_path.strip():
            return False
        path = self._resolve_gui_path(raw_path)
        expected_dir = self._survey_workspace_dir(survey_id) / "outputs" / "merge"
        try:
            path.parent.resolve().relative_to(expected_dir.resolve())
        except Exception:
            return False
        pattern = rf"^{re.escape(survey_id)}(?:_day\d+)?_(?:merged|ts)\.(?:csv|xlsx|xlsm)$"
        return re.match(pattern, path.name, re.IGNORECASE) is not None

    def _default_oe_output_path(self, survey_id: str, question_col: str) -> str:
        safe_q = re.sub(r"[^A-Za-z0-9_-]+", "_", question_col.strip()) if question_col.strip() else "oe_codebook"
        return self._default_workspace_file(survey_id, "outputs", "oe", "codebooks", f"{safe_q}_codebook.csv")

    def _default_oe_render_output_path(self, survey_id: str, question_col: str) -> str:
        safe_q = re.sub(r"[^A-Za-z0-9_-]+", "_", question_col.strip()) if question_col.strip() else "oe"
        return self._default_workspace_file(survey_id, "outputs", "oe", "wincross_blocks", f"{safe_q}_wincross_oe.txt")

    def _default_results_output_path(self, survey_id: str, docx_raw: str) -> str:
        if not survey_id or not docx_raw.strip():
            return ""
        docx_path = self._resolve_gui_path(docx_raw)
        stem = docx_path.stem
        if stem.upper().endswith(" FILLED"):
            stem = stem[:-7].rstrip()
        if not stem.upper().endswith(" EMPTY"):
            stem = f"{stem} EMPTY"
        return self._default_workspace_file(survey_id, "outputs", "results", f"{stem}{docx_path.suffix}")

    def _default_tomas_out_dir(self, survey_id: str) -> str:
        if not survey_id:
            return ""
        return self._default_workspace_file(survey_id, "outputs", "tomas")

    def _latest_data_delivery_day(self, survey_id: str) -> int | None:
        if not survey_id:
            return None
        manifest_path = self._survey_workspace_dir(survey_id) / "data_delivery_manifest.json"
        if not manifest_path.exists():
            return None
        try:
            payload = _read_json_dict(manifest_path)
        except Exception:
            return None
        entries = payload.get("entries", [])
        if not isinstance(entries, list):
            return None
        days: list[int] = []
        for entry in entries:
            if not isinstance(entry, dict):
                continue
            try:
                days.append(int(entry.get("day", 0) or 0))
            except Exception:
                continue
        return max(days) if days else None

    def _default_wincross_json_path(self, docx_raw: str, raw_input_raw: str) -> str:
        raw_source = raw_input_raw.strip()
        if raw_source:
            raw_path = self._resolve_gui_path(raw_source)
            return str(raw_path.parent / f"{raw_path.stem}JSON.json")
        docx_source = docx_raw.strip()
        if docx_source:
            docx_path = self._resolve_gui_path(docx_source)
            return str(docx_path.parent / f"{docx_path.stem}JSON.json")
        return ""

    def _default_wincross_tables_path(self, docx_raw: str, raw_input_raw: str) -> str:
        raw_source = raw_input_raw.strip()
        if raw_source:
            raw_path = self._resolve_gui_path(raw_source)
            return str(raw_path.parent / f"{raw_path.stem}TABLES.autogen.txt")
        docx_source = docx_raw.strip()
        if docx_source:
            docx_path = self._resolve_gui_path(docx_source)
            return str(docx_path.parent / f"{docx_path.stem}TABLES.autogen.txt")
        return ""

    def _default_empty_results_output(self, docx_raw: str) -> str:
        if not docx_raw.strip():
            return ""
        docx_path = self._resolve_gui_path(docx_raw)
        stem = docx_path.stem
        if stem.upper().endswith(" FILLED"):
            stem = stem[:-7].rstrip()
        if not stem.upper().endswith(" EMPTY"):
            stem = f"{stem} EMPTY"
        return str(docx_path.with_name(f"{stem}{docx_path.suffix}"))

    def _initial_dir_from_candidates(self, *raw_candidates: str) -> str:
        for raw in raw_candidates:
            text = str(raw or "").strip()
            if not text:
                continue
            path = self._resolve_gui_path(text)
            target = path if path.is_dir() else path.parent
            if target.exists():
                return str(target)

        main_source = self.main_survey_source_var.get().strip()
        if main_source:
            path = self._resolve_gui_path(main_source)
            target = path if path.is_dir() else path.parent
            if target.exists():
                return str(target)
        return str(BASE_DIR)

    def _survey_id_from_text(self, raw: str) -> str:
        code = _normalize_survey_code(raw)
        return f"NPA{code}" if code else ""

    def _preferred_existing_oe_codebook(self) -> str:
        for candidate in (
            self.oe_existing_path_var.get().strip(),
            self.oe_render_codebook_var.get().strip(),
            self.oe_extract_out_var.get().strip(),
        ):
            if candidate:
                return candidate
        return ""

    def _specialist_gui_path(self, *parts: str) -> Path:
        return BASE_DIR.joinpath(*parts)

    def _load_specialist_gui_module(self, module_name: str, module_path: Path):
        if not module_path.exists():
            raise FileNotFoundError(f"Specialist GUI not found: {module_path}")
        module_dir = str(module_path.parent)
        if module_dir not in sys.path:
            sys.path.insert(0, module_dir)
        spec = importlib.util.spec_from_file_location(module_name, module_path)
        if spec is None or spec.loader is None:
            raise ImportError(f"Could not load specialist GUI module: {module_path}")
        module = importlib.util.module_from_spec(spec)
        sys.modules[module_name] = module
        spec.loader.exec_module(module)
        return module

    def _bind_refresh_on_close(self, window: tk.Toplevel) -> None:
        def _on_destroy(event) -> None:
            if event.widget is window:
                try:
                    self.load_index()
                except Exception:
                    pass

        window.bind("<Destroy>", _on_destroy, add="+")

    def _open_specialist_window(self, *, title: str, module_name: str, module_path: Path, class_name: str):
        module = self._load_specialist_gui_module(module_name, module_path)
        app_cls = getattr(module, class_name, None)
        if app_cls is None:
            raise AttributeError(f"{module_path.name} does not expose {class_name}")
        window = tk.Toplevel(self.root)
        window.title(title)
        app = app_cls(window)
        setattr(window, "_survey_platform_app", app)
        self._bind_refresh_on_close(window)
        return window, app

    def _table_headers(self, raw_path: str) -> list[str]:
        if not raw_path.strip():
            return []
        path = self._resolve_gui_path(raw_path)
        if not path.exists() or not path.is_file():
            return []
        suffix = path.suffix.casefold()
        headers: list[str] = []
        try:
            if suffix == ".csv":
                for enc in ("utf-8-sig", "utf-8", "latin1"):
                    try:
                        with path.open("r", encoding=enc, newline="") as handle:
                            headers = next(csv.reader(handle), [])
                        break
                    except Exception:
                        continue
            elif suffix in {".xlsx", ".xlsm"}:
                try:
                    from openpyxl import load_workbook  # type: ignore

                    wb = load_workbook(path, read_only=True, data_only=True)
                    try:
                        ws = wb[wb.sheetnames[0]]
                        headers = [str(cell.value or "").strip() for cell in next(ws.iter_rows(max_row=1))]
                    finally:
                        wb.close()
                except Exception:
                    headers = []
        except Exception:
            return []
        return [str(item).strip() for item in headers if str(item).strip()]

    def _data_files_in_dir(self, directory: Path) -> list[Path]:
        if not directory.exists() or not directory.is_dir():
            return []
        return [
            path
            for path in sorted(directory.iterdir())
            if path.is_file() and path.suffix.casefold() in {".csv", ".xlsx", ".xlsm"}
        ]

    def _header_cell_text(self, value: object) -> str:
        if value is None:
            return ""
        text = str(value).strip()
        return text

    def _identity_cell_text(self, value: object) -> str:
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    def _canonical_data_header(self, header: str) -> str:
        key = header.strip().casefold()
        if key == "survnum":
            return "survnum"
        if key == "timeseries":
            return "timeseries"
        return header.strip()

    def _data_intake_sort_key(self, value: str) -> tuple[int, int | str]:
        text = str(value).strip()
        if re.fullmatch(r"\d+", text):
            return (0, int(text))
        return (1, text.casefold())

    def _read_data_intake_rows(self, path: Path) -> list[list[object]]:
        suffix = path.suffix.casefold()
        if suffix == ".csv":
            last_error: Exception | None = None
            for enc in ("utf-8-sig", "utf-8", "latin1"):
                try:
                    with path.open("r", encoding=enc, newline="") as handle:
                        return [list(row) for row in csv.reader(handle)]
                except Exception as exc:
                    last_error = exc
            raise ValueError(f"Could not read CSV data delivery:\n{path}\n\n{last_error}")
        if suffix in {".xlsx", ".xlsm"}:
            try:
                from openpyxl import load_workbook  # type: ignore
            except Exception as exc:
                raise ValueError("Reading Excel data deliveries requires openpyxl.") from exc
            try:
                wb = load_workbook(path, read_only=True, data_only=True)
                try:
                    ws = wb[wb.sheetnames[0]]
                    return [list(row) for row in ws.iter_rows(values_only=True)]
                finally:
                    wb.close()
            except Exception as exc:
                raise ValueError(f"Could not read Excel data delivery:\n{path}\n\n{exc}") from exc
        raise ValueError("Daily data delivery must be CSV or XLSX/XLSM.")

    def _prepare_data_intake_table(self, source_path: Path) -> DataIntakeTable:
        raw_rows = self._read_data_intake_rows(source_path)
        if not raw_rows:
            raise ValueError(f"Data delivery is empty:\n{source_path}")

        raw_headers = [self._header_cell_text(value) for value in raw_rows[0]]
        while raw_headers and not raw_headers[-1]:
            raw_headers.pop()
        if not raw_headers:
            raise ValueError(f"Data delivery has no column headers:\n{source_path}")
        blank_header_positions = [str(index + 1) for index, header in enumerate(raw_headers) if not header]
        if blank_header_positions:
            raise ValueError(
                "Data delivery has blank column headers.\n\n"
                f"Blank column number(s): {', '.join(blank_header_positions[:8])}"
            )

        survnum_indexes = [
            index for index, header in enumerate(raw_headers) if header.strip().casefold() == "survnum"
        ]
        timeseries_indexes = [
            index for index, header in enumerate(raw_headers) if header.strip().casefold() == "timeseries"
        ]
        if not survnum_indexes:
            raise ValueError("Data delivery must include a survnum column.")
        if len(survnum_indexes) > 1:
            raise ValueError("Data delivery has more than one survnum column.")
        if len(timeseries_indexes) > 1:
            raise ValueError("Data delivery has more than one timeseries column.")

        survnum_index = survnum_indexes[0]
        timeseries_index = timeseries_indexes[0] if timeseries_indexes else None
        non_timeseries_indexes = [
            index for index in range(len(raw_headers)) if index != timeseries_index
        ]
        normalized_headers = [
            self._canonical_data_header(raw_headers[index]) for index in non_timeseries_indexes
        ] + ["timeseries"]

        header_counts = Counter(header.casefold() for header in normalized_headers)
        duplicate_headers = [
            header for header, count in header_counts.items() if count > 1
        ]
        if duplicate_headers:
            raise ValueError(
                "Data delivery has duplicate column headers after normalization.\n\n"
                f"Duplicate header(s): {', '.join(sorted(duplicate_headers)[:8])}"
            )

        normalized_rows: list[list[object]] = []
        pair_counts: Counter[tuple[str, str]] = Counter()
        blank_survnum_rows: list[int] = []
        blank_timeseries_filled = 0
        header_width = len(raw_headers)
        for row_number, raw_row in enumerate(raw_rows[1:], start=2):
            row = list(raw_row[:header_width])
            if len(row) < header_width:
                row.extend([""] * (header_width - len(row)))
            if not any(self._identity_cell_text(value) for value in row):
                continue

            survnum = self._identity_cell_text(row[survnum_index])
            if not survnum:
                blank_survnum_rows.append(row_number)
                continue

            if timeseries_index is None:
                timeseries = "1"
            else:
                timeseries = self._identity_cell_text(row[timeseries_index])
                if not timeseries:
                    timeseries = "1"
                    blank_timeseries_filled += 1

            normalized_row = [
                row[index] if index < len(row) else "" for index in non_timeseries_indexes
            ]
            normalized_row.append(timeseries)
            normalized_rows.append(normalized_row)
            pair_counts[(survnum, timeseries)] += 1

        if blank_survnum_rows:
            examples = ", ".join(str(row) for row in blank_survnum_rows[:8])
            raise ValueError(
                "Data delivery has rows with blank survnum values.\n\n"
                f"Example row number(s): {examples}"
            )
        if not normalized_rows:
            raise ValueError("Data delivery has no response rows.")

        duplicate_pairs = [
            (pair, count) for pair, count in pair_counts.items() if count > 1
        ]
        duplicate_pairs.sort(key=lambda item: (-item[1], item[0][0], item[0][1]))
        duplicate_examples = [
            f"survnum={survnum}, timeseries={timeseries} ({count} rows)"
            for (survnum, timeseries), count in duplicate_pairs[:10]
        ]
        timeseries_values = sorted(
            {timeseries for _survnum, timeseries in pair_counts},
            key=self._data_intake_sort_key,
        )
        return DataIntakeTable(
            headers=normalized_headers,
            rows=normalized_rows,
            pairs=set(pair_counts),
            duplicate_count=sum(count - 1 for _pair, count in duplicate_pairs),
            duplicate_examples=duplicate_examples,
            timeseries_added=timeseries_index is None,
            blank_timeseries_filled=blank_timeseries_filled,
            timeseries_values=timeseries_values,
        )

    def _write_data_intake_table(self, dest_path: Path, table: DataIntakeTable) -> None:
        dest_path.parent.mkdir(parents=True, exist_ok=True)
        suffix = dest_path.suffix.casefold()
        if suffix == ".csv":
            with dest_path.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(table.headers)
                writer.writerows(table.rows)
            return
        if suffix in {".xlsx", ".xlsm"}:
            try:
                from openpyxl import Workbook  # type: ignore
            except Exception as exc:
                raise ValueError("Writing Excel data deliveries requires openpyxl.") from exc
            wb = Workbook()
            ws = wb.active
            ws.title = "Data"
            ws.append(table.headers)
            for row in table.rows:
                ws.append(row)
            wb.save(dest_path)
            wb.close()
            return
        raise ValueError("Daily data delivery must be CSV or XLSX/XLSM.")

    def _data_intake_identity_summary(self, table: DataIntakeTable) -> dict[str, object]:
        return {
            "key_columns": ["survnum", "timeseries"],
            "row_count": table.row_count,
            "unique_pair_count": table.pair_count,
            "duplicate_pair_count": table.duplicate_count,
            "timeseries_added": table.timeseries_added,
            "timeseries_default": "1" if table.timeseries_added else "",
            "blank_timeseries_filled": table.blank_timeseries_filled,
            "timeseries_values": table.timeseries_values,
        }

    def _data_intake_pair_changes(
        self,
        previous: DataIntakeTable | None,
        current: DataIntakeTable,
    ) -> dict[str, int]:
        if previous is None:
            return {
                "retained_identity_pairs": 0,
                "new_identity_pairs": current.pair_count,
                "removed_identity_pairs": 0,
            }
        return {
            "retained_identity_pairs": len(previous.pairs & current.pairs),
            "new_identity_pairs": len(current.pairs - previous.pairs),
            "removed_identity_pairs": len(previous.pairs - current.pairs),
        }

    def _update_data_intake_state(
        self,
        survey_id: str,
        *,
        role_id: str,
        day: int,
        source_path: Path,
        current_path: Path,
        table: DataIntakeTable,
        pair_changes: dict[str, int],
    ) -> None:
        state = self._read_survey_state(survey_id)
        data_intake = state.get("data_intake", {})
        if not isinstance(data_intake, dict):
            data_intake = {}
            state["data_intake"] = data_intake
        data_intake["latest_role"] = role_id
        data_intake[role_id] = {
            "updated_at": _utc_now_iso(),
            "day": day,
            "source_path": str(source_path),
            "current_path": str(current_path),
            "identity": self._data_intake_identity_summary(table),
            "pair_changes": pair_changes,
            "headers": table.headers,
        }
        self._write_survey_state(survey_id, state)

    def _is_timeseries_data_file(self, raw_path: str) -> bool:
        if not raw_path.strip():
            return False
        path = self._resolve_gui_path(raw_path)
        if "ts" not in path.stem.casefold():
            return False
        return any(header.casefold() == "timeseries" for header in self._table_headers(raw_path))

    def _oe_question_candidates(self, input_path_raw: str) -> list[str]:
        input_raw = input_path_raw.strip()
        if not input_raw:
            return []
        path = self._resolve_gui_path(input_raw)
        if not path.exists():
            return []

        headers: list[str] = []
        suffix = path.suffix.casefold()
        try:
            if suffix == ".csv":
                for enc in ("utf-8-sig", "utf-8", "latin1"):
                    try:
                        with path.open("r", encoding=enc, newline="") as handle:
                            reader = csv.reader(handle)
                            headers = next(reader, [])
                        break
                    except Exception:
                        continue
            elif suffix in {".xlsx", ".xlsm"}:
                try:
                    from openpyxl import load_workbook  # type: ignore
                except Exception:
                    headers = []
                else:
                    wb = load_workbook(path, read_only=True, data_only=True)
                    try:
                        ws = wb[wb.sheetnames[0]]
                        row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
                        headers = [str(cell).strip() for cell in row if str(cell or "").strip()]
                    finally:
                        wb.close()
        except Exception:
            headers = []

        candidates: list[str] = []
        for header in headers:
            text = str(header).strip()
            if not text:
                continue
            if OE_QUESTION_COL_RE.match(text):
                candidates.append(text)
        return candidates

    def _suggest_oe_selector(self, question_col: str) -> str:
        match = OE_QUESTION_COL_RE.match(str(question_col).strip())
        if not match:
            return ""
        qnum = str(int(match.group("num")))
        return f"c{qnum}_1 S3"

    def _suggest_oe_extract_out(self, input_path_raw: str, question_col: str) -> str:
        input_raw = input_path_raw.strip()
        if not input_raw or not question_col.strip():
            return ""
        input_path = self._resolve_gui_path(input_raw)
        safe_q = re.sub(r"[^A-Za-z0-9_-]+", "_", question_col.strip())
        return str(input_path.with_name(f"{input_path.stem}_{safe_q}_codebook.csv"))

    def _suggest_oe_render_out(self, codebook_raw: str, question_col: str) -> str:
        codebook = codebook_raw.strip()
        if not codebook:
            return ""
        codebook_path = self._resolve_gui_path(codebook)
        safe_q = re.sub(r"[^A-Za-z0-9_-]+", "_", question_col.strip()) if question_col.strip() else codebook_path.stem
        return str(codebook_path.with_name(f"{safe_q}_wincross_oe.txt"))

    def _apply_main_survey_defaults(self, survey_id: str) -> None:
        survey_id = survey_id.strip()
        if not survey_id:
            return
        meta_payload = self._meta_payload_for_survey(survey_id)
        self.main_survey_id_var.set(survey_id)
        self.survey_selector_var.set(survey_id)
        self.main_survey_source_var.set(str(meta_payload.get("primary_source_dir", "")).strip())

        docx = self._preferred_workspace_role_file(
            survey_id,
            "final_questionnaire_docx",
            "phone_docx",
        ) or self._first_stage_path(survey_id, "wincross", "inputs", "docx") or self._preferred_survey_file(
            survey_id,
            "docx",
        )
        pdf = self._preferred_workspace_role_file(
            survey_id,
            "final_questionnaire_pdf",
            "phone_pdf",
        ) or self._first_stage_path(survey_id, "wincross", "inputs", "pdf") or self._preferred_survey_file(
            survey_id,
            "pdf",
        )
        raw_txt = self._preferred_workspace_role_file(survey_id, "wincross_raw") or self._preferred_survey_file(survey_id, "raw_txt")
        survey_json = self._preferred_workspace_role_file(survey_id, "wincross_json") or self._preferred_survey_file(survey_id, "survey_json")
        tables_txt = self._preferred_workspace_role_file(
            survey_id,
            "reviewed_tables",
            "wincross_tables",
        ) or self._preferred_survey_file(survey_id, "tables_txt")
        primary_csv = self._detect_primary_csv_for_survey(survey_id)
        results_docx = self._preferred_workspace_role_file(
            survey_id,
            "results_template_docx",
        ) or self._first_stage_path(survey_id, "results", "inputs", "report_docx")
        tomas_phone_docx = self._preferred_workspace_role_file(
            survey_id,
            "phone_docx",
            "final_questionnaire_docx",
        ) or self._preferred_variant_survey_file(survey_id, "docx", ("phone", "cati"))
        tomas_phone_pdf = self._preferred_workspace_role_file(
            survey_id,
            "phone_pdf",
            "final_questionnaire_pdf",
        ) or self._preferred_variant_survey_file(survey_id, "pdf", ("phone", "cati"))
        tomas_sms_docx = self._preferred_workspace_role_file(survey_id, "sms_docx") or self._preferred_variant_survey_file(
            survey_id,
            "docx",
            ("sms", "text"),
        )
        tomas_sms_pdf = self._preferred_workspace_role_file(survey_id, "sms_pdf") or self._preferred_variant_survey_file(
            survey_id,
            "pdf",
            ("sms", "text"),
        )

        self.wincross_docx_var.set(docx)
        self.wincross_pdf_var.set(pdf)
        self.wincross_raw_input_var.set("")
        self.wincross_json_out_var.set(survey_json or self._default_wincross_json_workspace_path(survey_id))
        self.wincross_tables_out_var.set(tables_txt or self._default_wincross_tables_workspace_path(survey_id))
        self.wincross_translate_var.set(False)
        self.wincross_skip_tables_var.set(False)
        self.wincross_status_var.set(f"Loaded WinCross defaults for {survey_id}.")

        self.tomas_phone_docx_var.set(tomas_phone_docx)
        self.tomas_phone_pdf_var.set(tomas_phone_pdf)
        self.tomas_sms_docx_var.set(tomas_sms_docx)
        self.tomas_sms_pdf_var.set(tomas_sms_pdf)
        self.tomas_out_dir_var.set(self._default_tomas_out_dir(survey_id))
        self.tomas_status_var.set(f"Loaded optional TOMAS defaults for {survey_id}.")

        self.merge_primary_survey_var.set(survey_id)
        self.merge_primary_csv_var.set(primary_csv)
        merge_output_raw = self.merge_output_path_var.get().strip()
        if not merge_output_raw or self._is_default_merge_output_path(survey_id, merge_output_raw):
            self.merge_output_path_var.set(self._default_merge_output_path(survey_id))
        self.merge_status_var.set(f"Loaded merge defaults for {survey_id}.")

        self.data_delivery_path_var.set(primary_csv)
        self.data_delivery_codelist_path_var.set(self._preferred_workspace_role_file(survey_id, "daily_codelist"))
        self.data_delivery_timeseries_var.set(self._is_timeseries_data_file(primary_csv))
        survey_state = self._read_survey_state(survey_id)
        daily_codelist_state = survey_state.get("daily_codelist", {})
        if not isinstance(daily_codelist_state, dict):
            daily_codelist_state = {}
        delivered_codes_path = str(daily_codelist_state.get("delivered_codes_current_path", "")).strip()
        user_codes_path = str(daily_codelist_state.get("user_codes_current_path", "")).strip()
        if delivered_codes_path:
            self.oe_defaults_path_var.set(delivered_codes_path)
        if user_codes_path and not self.oe_existing_path_var.get().strip():
            self.oe_status_var.set(f"Loaded OE user-code reserve file: {user_codes_path}")
        latest_day = self._latest_data_delivery_day(survey_id)
        if latest_day is not None:
            self.data_delivery_day_var.set(str(latest_day))
        self.data_delivery_status_var.set(f"Loaded data tracking defaults for {survey_id}.")

        self.oe_input_path_var.set(primary_csv)
        question_candidates = self._oe_question_candidates(primary_csv)
        if self.oe_question_combo is not None:
            self.oe_question_combo["values"] = question_candidates
        question_choice = question_candidates[0] if question_candidates else ""
        self.oe_question_col_var.set(question_choice)
        self.oe_extract_out_var.set(self._default_oe_output_path(survey_id, question_choice))
        self.oe_render_codebook_var.set(self.oe_extract_out_var.get().strip())
        self.oe_render_selector_var.set(self._suggest_oe_selector(question_choice))
        self.oe_render_out_var.set(self._default_oe_render_output_path(survey_id, question_choice))
        self.oe_status_var.set(f"Loaded OE defaults for {survey_id}.")

        self.results_docx_var.set(results_docx)
        self.results_tables_var.set(tables_txt)
        self.results_json_var.set(survey_json)
        self.results_output_var.set(self._default_results_output_path(survey_id, results_docx) if results_docx else "")
        self.results_status_var.set(f"Loaded results defaults for {survey_id}.")

    def set_main_survey_from_selection(self) -> None:
        survey_id = self.selected_survey_id_var.get().strip()
        if not survey_id:
            messagebox.showinfo("Survey Platform", "Select a survey in the Survey Workspaces tab first.")
            return
        self._apply_main_survey_defaults(survey_id)
        status_payload = self._read_workflow_status(survey_id)
        saved_phase = str(status_payload.get("current_phase", "")).strip()
        if saved_phase in PHASE_VALUES:
            self.phase_var.set(saved_phase)
        self._refresh_workflow_home()
        self._save_dashboard_survey_preference(survey_id, confirmed=True)
        self.status_var.set(f"Main survey set to {survey_id}.")

    def confirm_survey_selection(self) -> None:
        survey_id = self.survey_selector_var.get().strip() or self.selected_survey_id_var.get().strip()
        if not survey_id:
            messagebox.showinfo("Survey Platform", "Select a survey from the repo index first.")
            return
        if survey_id not in self.rows_by_id:
            messagebox.showinfo("Survey Platform", "Survey must come from the loaded repo index.")
            return
        self.selected_survey_id_var.set(survey_id)
        self._apply_main_survey_defaults(survey_id)
        status_payload = self._read_workflow_status(survey_id)
        saved_phase = str(status_payload.get("current_phase", "")).strip()
        if saved_phase in PHASE_VALUES:
            self.phase_var.set(saved_phase)
        else:
            self._set_workflow_phase(survey_id, self.phase_var.get().strip() or PHASE_VALUES[0])
        self._refresh_workflow_home()
        self._save_dashboard_survey_preference(survey_id, confirmed=True)
        self.status_var.set(f"Confirmed main survey: {survey_id}.")

    def select_survey_from_selector(self, _event=None) -> None:
        survey_id = self.survey_selector_var.get().strip()
        if not survey_id:
            return
        if survey_id in self.rows_by_id and self.tree is not None:
            self.tree.selection_set(survey_id)
            self.tree.see(survey_id)
            self._on_select_row()
        else:
            self._save_dashboard_survey_preference(survey_id)
        self.status_var.set(f"Selected {survey_id}. Click Confirm Survey to preload workspace defaults.")

    def _select_top_tab(self, label: str) -> None:
        if self.notebook is None:
            return
        for tab_id in self.notebook.tabs():
            if self.notebook.tab(tab_id, "text") == label:
                self.notebook.select(tab_id)
                return

    def go_to_selected_phase(self) -> None:
        phase = self.phase_var.get().strip() or PHASE_VALUES[0]
        survey_id = self.main_survey_id_var.get().strip()
        if survey_id:
            self._set_workflow_phase(survey_id, phase)
        else:
            self.phase_primary_action_var.set(self._phase_primary_action_label(phase))
        self._select_top_tab("Phase Workflows")
        if self.workflow_book is None:
            return
        phase_to_tab = {
            "Pre-fielding": "Pre-fielding",
            "In-field Daily": "In-field Daily",
            "Final Results": "Final Results",
            "Archive/Search": "Archive/Search",
        }
        wanted = phase_to_tab.get(phase, "Pre-fielding")
        for tab_id in self.workflow_book.tabs():
            if self.workflow_book.tab(tab_id, "text") == wanted:
                self.workflow_book.select(tab_id)
                break

    def _questions_db_path(self) -> Path:
        return self._repo_root() / REPO_QUESTIONS_DB_FILENAME

    def _questions_db_open(self) -> sqlite3.Connection:
        conn = sqlite3.connect(self._questions_db_path())
        conn.row_factory = sqlite3.Row
        conn.create_function(
            "question_ballot_options_key",
            1,
            _question_frequency_ballot_options_key,
        )
        conn.create_function(
            "question_browser_is_english",
            3,
            _question_browser_is_english,
        )
        return conn

    def sync_merge_from_main(self) -> None:
        survey_id = self.main_survey_id_var.get().strip()
        if not survey_id:
            messagebox.showinfo("Survey Platform", "Set a Main Survey first.")
            return
        self.merge_primary_survey_var.set(survey_id)
        self.merge_primary_csv_var.set(self._detect_primary_csv_for_survey(survey_id))
        if not self.merge_output_path_var.get().strip():
            self.merge_output_path_var.set(self._default_merge_output_path(survey_id))
        self.merge_status_var.set(f"Merge defaults refreshed from {survey_id}.")

    def add_file_to_workspace_dialog(self) -> None:
        survey_id = self.main_survey_id_var.get().strip() or self.selected_survey_id_var.get().strip()
        if not survey_id:
            messagebox.showinfo("Survey Platform", "Confirm or select a survey first.")
            return
        if survey_id not in self.rows_by_id:
            messagebox.showinfo("Survey Platform", "Workspace files must be attached to a survey from the repo index.")
            return

        role_labels = [label for _role_id, label, _rel_dir, _extensions in WORKSPACE_FILE_ROLES]
        top = tk.Toplevel(self.root)
        top.title(f"Add File to Workspace ({survey_id})")
        top.geometry("900x260")
        top.transient(self.root)

        frame = ttk.Frame(top, padding=12)
        frame.pack(fill=tk.BOTH, expand=True)

        role_var = tk.StringVar(value=role_labels[0])
        survey_var = tk.StringVar(value=survey_id)
        source_var = tk.StringVar(value="")
        dest_var = tk.StringVar(value="")

        def role_id_from_label() -> str:
            return WORKSPACE_FILE_ROLE_LABEL_TO_ID.get(role_var.get(), "other")

        def dialog_day() -> int:
            try:
                return max(1, int(self.data_delivery_day_var.get().strip() or "1"))
            except ValueError:
                return 1

        def destination_for(role_id: str, source_text: str) -> Path:
            dest_dir = self._workspace_role_dir(survey_id, role_id)
            if role_id == "daily_codelist":
                dest_dir = dest_dir / f"day_{dialog_day()}"
            if source_text:
                return dest_dir / Path(source_text).name
            return dest_dir

        def update_dest() -> None:
            role_id = role_id_from_label()
            source_text = source_var.get().strip()
            dest_var.set(str(destination_for(role_id, source_text)))

        def browse_source() -> None:
            role_id = role_id_from_label()
            label, _rel_dir, extensions = self._workspace_role_def(role_id)
            filetypes: list[tuple[str, str]] = []
            if extensions:
                pattern = " ".join(f"*{ext}" for ext in extensions)
                filetypes.append((label, pattern))
            filetypes.append(("All files", "*.*"))
            chosen = filedialog.askopenfilename(
                title=f"Choose {label}",
                initialdir=self._initial_dir_from_candidates(self.main_survey_source_var.get()),
                filetypes=filetypes,
            )
            if chosen:
                source_var.set(chosen)
                update_dest()

        def copy_file() -> None:
            role_id = role_id_from_label()
            label, _rel_dir, extensions = self._workspace_role_def(role_id)
            source_raw = source_var.get().strip()
            if not source_raw:
                messagebox.showinfo("Survey Platform", "Choose a source file first.")
                return
            source_path = self._resolve_gui_path(source_raw)
            if not source_path.exists() or not source_path.is_file():
                messagebox.showerror("Survey Platform", f"Source file not found:\n{source_path}")
                return
            if extensions and source_path.suffix.casefold() not in {ext.casefold() for ext in extensions}:
                allowed = ", ".join(extensions)
                messagebox.showerror("Survey Platform", f"{label} expects one of: {allowed}")
                return
            if role_id in {"current_data", "timeseries_current"}:
                try:
                    day = int(self.data_delivery_day_var.get().strip() or "1")
                except ValueError:
                    day = 1
                day = max(1, day)
                is_timeseries = bool(role_id == "timeseries_current" or self._is_timeseries_data_file(str(source_path)))
                try:
                    current_path, new_headers = self._track_data_delivery_file(
                        survey_id,
                        source_path,
                        day=day,
                        is_timeseries=is_timeseries,
                    )
                except ValueError as exc:
                    messagebox.showerror("Survey Platform", str(exc))
                    self.data_delivery_status_var.set("Data delivery blocked.")
                    return
                except Exception as exc:
                    messagebox.showerror("Survey Platform", f"Failed tracking data delivery:\n{exc}")
                    self.data_delivery_status_var.set("Data delivery tracking failed.")
                    return

                self.data_delivery_path_var.set(str(current_path))
                self.data_delivery_day_var.set(str(day))
                self.data_delivery_timeseries_var.set(is_timeseries)
                self.merge_primary_survey_var.set(survey_id)
                self.merge_primary_csv_var.set(str(current_path))
                self.oe_input_path_var.set(str(current_path))
                question_candidates = self._oe_question_candidates(str(current_path))
                if self.oe_question_combo is not None:
                    self.oe_question_combo["values"] = question_candidates
                self._mark_workflow_step("daily_data_current", "done", phase="In-field Daily")
                if new_headers:
                    self._mark_workflow_step("daily_data_headers_checked", "done", phase="In-field Daily")
                self._set_workflow_data_day(survey_id, day)
                self._apply_main_survey_defaults(survey_id)
                self._refresh_workflow_home()
                self.status_var.set(f"Tracked day {day} data delivery for {survey_id}.")
                self.data_delivery_status_var.set(f"Tracked day {day} delivery to {current_path}.")
                top.destroy()
                return
            dest_path = destination_for(role_id, str(source_path))
            if dest_path.exists() and not messagebox.askyesno(
                "Survey Platform",
                f"Replace existing workspace file?\n\n{dest_path}",
            ):
                return
            try:
                dest_path.parent.mkdir(parents=True, exist_ok=True)
                if source_path.resolve() != dest_path.resolve():
                    shutil.copy2(source_path, dest_path)
                workspace_dir = None
                if role_id == "daily_codelist":
                    workspace_dir = f"inputs/data/codelists/day_{dialog_day()}"
                self._record_workspace_file(
                    survey_id,
                    role_id,
                    dest_path,
                    source_path=source_path,
                    workspace_dir=workspace_dir,
                )
            except Exception as exc:
                messagebox.showerror("Survey Platform", f"Failed adding workspace file:\n{exc}")
                return

            if role_id in {"final_questionnaire_docx", "final_questionnaire_pdf", "phone_docx", "phone_pdf"}:
                self._mark_workflow_step("final_questionnaire_selected", "done", phase="Pre-fielding")
            if role_id in {"current_data", "timeseries_current"}:
                self._mark_workflow_step("daily_data_current", "done", phase="In-field Daily")
            if role_id == "daily_codelist":
                self._mark_workflow_step(
                    "daily_codelist_tracked",
                    "done",
                    phase="In-field Daily",
                    note=f"Day {dialog_day()}",
                )
            if role_id == "reviewed_tables":
                self._mark_workflow_step("wincross_reviewed", "reviewed", phase="Pre-fielding")
            if role_id == "results_template_docx":
                self._mark_workflow_step("final_data_confirmed", "done", phase="Final Results")

            self._apply_main_survey_defaults(survey_id)
            self._refresh_workflow_home()
            self.status_var.set(f"Added {label} to {survey_id} workspace.")
            top.destroy()

        ttk.Label(frame, text="Survey").grid(row=0, column=0, sticky="w", padx=6, pady=6)
        ttk.Entry(frame, textvariable=survey_var, width=18, state="readonly").grid(
            row=0,
            column=1,
            sticky="w",
            padx=6,
            pady=6,
        )
        ttk.Label(frame, text="File role").grid(row=1, column=0, sticky="w", padx=6, pady=6)
        role_combo = ttk.Combobox(frame, textvariable=role_var, values=role_labels, state="readonly", width=34)
        role_combo.grid(row=1, column=1, sticky="w", padx=6, pady=6)
        role_combo.bind("<<ComboboxSelected>>", lambda _event: update_dest())

        ttk.Label(frame, text="Source file").grid(row=2, column=0, sticky="w", padx=6, pady=6)
        ttk.Entry(frame, textvariable=source_var, width=90).grid(row=2, column=1, sticky="we", padx=6, pady=6)
        ttk.Button(frame, text="Browse", command=browse_source).grid(row=2, column=2, sticky="w", padx=6, pady=6)

        ttk.Label(frame, text="Destination").grid(row=3, column=0, sticky="w", padx=6, pady=6)
        ttk.Entry(frame, textvariable=dest_var, width=90, state="readonly").grid(
            row=3,
            column=1,
            columnspan=2,
            sticky="we",
            padx=6,
            pady=6,
        )

        buttons = ttk.Frame(frame)
        buttons.grid(row=4, column=0, columnspan=3, sticky="w", padx=6, pady=(12, 0))
        ttk.Button(buttons, text="Add File", command=copy_file).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(buttons, text="Cancel", command=top.destroy).pack(side=tk.LEFT, padx=6)

        frame.grid_columnconfigure(1, weight=1)
        update_dest()

    def browse_tomas_input(self, target_var: tk.StringVar, label: str) -> None:
        ext = "*.pdf" if label.casefold() == "pdf" else "*.docx"
        chosen = filedialog.askopenfilename(
            title=f"Choose TOMAS {label}",
            initialdir=self._initial_dir_from_candidates(target_var.get(), self.main_survey_source_var.get()),
            filetypes=[
                (label, ext),
                ("All files", "*.*"),
            ],
        )
        if chosen:
            target_var.set(chosen)

    def browse_tomas_out_dir(self) -> None:
        initial = self._initial_dir_from_candidates(self.tomas_out_dir_var.get(), self.main_survey_source_var.get())
        chosen = filedialog.askdirectory(initialdir=initial, title="Choose TOMAS output folder")
        if chosen:
            self.tomas_out_dir_var.set(chosen)

    def open_advanced_tomas_editor(self) -> None:
        script_path = BASE_DIR / "pipelines" / "thomas_in_a_box" / "tomas_pipeline_gui.py"
        if not script_path.exists():
            messagebox.showerror("Survey Platform", f"TOMAS GUI not found:\n{script_path}")
            return
        try:
            subprocess.Popen([sys.executable, str(script_path)], cwd=str(BASE_DIR))
        except Exception as exc:
            messagebox.showerror("Survey Platform", f"Could not open TOMAS GUI:\n{exc}")
            self.tomas_status_var.set("Advanced TOMAS failed to open.")
            return
        self.tomas_status_var.set("Opened advanced TOMAS GUI.")

    def run_guided_tomas(self) -> None:
        survey_id = self.main_survey_id_var.get().strip() or self.selected_survey_id_var.get().strip()
        if not survey_id:
            messagebox.showwarning("Survey Platform", "Confirm a survey first.")
            return

        phone_docx = self.tomas_phone_docx_var.get().strip()
        phone_pdf = self.tomas_phone_pdf_var.get().strip()
        sms_docx = self.tomas_sms_docx_var.get().strip()
        sms_pdf = self.tomas_sms_pdf_var.get().strip()
        missing = [
            label
            for label, value in (
                ("Phone DOCX", phone_docx),
                ("Phone PDF", phone_pdf),
                ("SMS DOCX", sms_docx),
                ("SMS PDF", sms_pdf),
            )
            if not value
        ]
        if missing:
            self.show_advanced_details_var.set(True)
            self.toggle_advanced_details()
            messagebox.showwarning(
                "Survey Platform",
                "TOMAS needs four inputs before it can run:\n" + "\n".join(missing),
            )
            return

        out_dir = self.tomas_out_dir_var.get().strip() or self._default_tomas_out_dir(survey_id)
        if not out_dir:
            messagebox.showwarning("Survey Platform", "Choose a TOMAS output folder first.")
            return
        self.tomas_out_dir_var.set(out_dir)
        out_path = self._resolve_gui_path(out_dir)
        manifest_out = out_path / "tomas_pipeline_manifest.json"
        script_path = BASE_DIR / "pipelines" / "thomas_in_a_box" / "tomas_pipeline.py"

        cmd = [
            sys.executable,
            str(script_path),
            "--phone-docx",
            phone_docx,
            "--phone-pdf",
            phone_pdf,
            "--sms-docx",
            sms_docx,
            "--sms-pdf",
            sms_pdf,
            "--out-dir",
            str(out_path),
            "--manifest-out",
            str(manifest_out),
        ]
        glossary = BASE_DIR / "pipelines" / "thomas_in_a_box" / "glossary.json"
        scale_map = BASE_DIR / "pipelines" / "thomas_in_a_box" / "scale_map.json"
        if glossary.exists():
            cmd.extend(["--glossary", str(glossary)])
        if scale_map.exists():
            cmd.extend(["--scale-map", str(scale_map)])
        if self.tomas_build_formatted_var.get():
            cmd.append("--build-formatted-docx")

        ok = self._run_cli_command(
            cmd,
            title=f"{survey_id} TOMAS",
            running_text=f"Running optional TOMAS for {survey_id}...",
            status_target=self.tomas_status_var,
        )
        if ok:
            self._mark_workflow_step("tomas_generated", "done", phase="Pre-fielding")
            self.tomas_status_var.set(f"TOMAS complete for {survey_id}.")
        else:
            self.tomas_status_var.set(f"TOMAS failed for {survey_id}.")

    def mark_final_data_confirmed(self) -> None:
        self._mark_workflow_step("final_data_confirmed", "done", phase="Final Results")
        self.results_status_var.set("Final data confirmed.")

    def browse_data_delivery(self) -> None:
        initial = self._initial_dir_from_candidates(
            self.data_delivery_path_var.get(),
            self.main_survey_source_var.get(),
        )
        chosen = filedialog.askopenfilename(
            initialdir=initial,
            title="Choose daily data delivery",
            filetypes=[
                ("Data files", "*.csv *.xlsx *.xlsm"),
                ("CSV", "*.csv"),
                ("Excel", "*.xlsx *.xlsm"),
                ("All files", "*.*"),
            ],
        )
        if chosen:
            self.data_delivery_path_var.set(chosen)
            self.data_delivery_timeseries_var.set(self._is_timeseries_data_file(chosen))

    def browse_data_delivery_codelist(self) -> None:
        initial = self._initial_dir_from_candidates(
            self.data_delivery_codelist_path_var.get(),
            self.data_delivery_path_var.get(),
            self.main_survey_source_var.get(),
        )
        chosen = filedialog.askopenfilename(
            initialdir=initial,
            title="Choose optional daily codelist",
            filetypes=[
                ("Codelist files", "*.csv *.xlsx *.xlsm"),
                ("CSV", "*.csv"),
                ("Excel", "*.xlsx *.xlsm"),
                ("All files", "*.*"),
            ],
        )
        if chosen:
            self.data_delivery_codelist_path_var.set(chosen)

    def _normalize_codelist_code(self, value: object) -> str:
        text = self._identity_cell_text(value)
        if not text:
            return ""
        if re.fullmatch(r"-?\d+(\.0+)?", text):
            return str(int(float(text)))
        return text.casefold()

    def _valid_codelist_code(self, value: str) -> bool:
        return re.fullmatch(r"\d+[a-z]?", value, re.IGNORECASE) is not None

    def _normalize_codelist_label(self, value: object) -> str:
        return re.sub(r"\s+", " ", str(value or "")).strip()

    def _codelist_label_key(self, value: str) -> str:
        return re.sub(r"[^a-z0-9]+", " ", value.casefold()).strip()

    def _read_codelist_sheets(self, path: Path) -> list[tuple[str, list[list[object]]]]:
        suffix = path.suffix.casefold()
        if suffix == ".csv":
            return [("", self._read_data_intake_rows(path))]
        if suffix in {".xlsx", ".xlsm"}:
            try:
                from openpyxl import load_workbook  # type: ignore
            except Exception as exc:
                raise ValueError("Reading Excel codelists requires openpyxl.") from exc
            try:
                wb = load_workbook(path, read_only=True, data_only=True)
                try:
                    return [
                        (name, [list(row) for row in wb[name].iter_rows(values_only=True)])
                        for name in wb.sheetnames
                    ]
                finally:
                    wb.close()
            except Exception as exc:
                raise ValueError(f"Could not read Excel codelist:\n{path}\n\n{exc}") from exc
        raise ValueError("Daily codelist must be CSV or XLSX/XLSM.")

    def _codelist_header_roles(self, headers: list[str]) -> dict[str, int]:
        normalized = [re.sub(r"[^a-z0-9]+", "_", header.casefold()).strip("_") for header in headers]

        def first_match(candidates: tuple[str, ...]) -> int | None:
            for candidate in candidates:
                if candidate in normalized:
                    return normalized.index(candidate)
            return None

        roles: dict[str, int] = {}
        question_idx = first_match(("question_col", "question", "qid", "q", "variable"))
        code_idx = first_match(("code", "default_code", "prelim_code", "final_code", "value"))
        label_idx = first_match(
            (
                "label",
                "description",
                "text",
                "name",
                "final_label",
                "original_label",
                "category",
            )
        )
        if question_idx is not None:
            roles["question"] = question_idx
        if code_idx is not None:
            roles["code"] = code_idx
        if label_idx is not None:
            roles["label"] = label_idx
        return roles

    def _question_from_codelist_sheet(self, sheet_name: str) -> str:
        sheet_name = self._normalize_codelist_label(sheet_name)
        m = re.fullmatch(r"(?i)q\s*0*(\d{1,3})(?:[_ -]?[a-z])?", sheet_name)
        if m:
            return f"q{int(m.group(1))}"
        return ""

    def _parse_codelist_records(self, source_path: Path) -> tuple[list[CodelistRecord], list[str]]:
        records_by_key: dict[tuple[str, str], CodelistRecord] = {}
        warnings: list[str] = []
        for sheet_name, rows in self._read_codelist_sheets(source_path):
            if not rows:
                continue
            sheet_question = self._question_from_codelist_sheet(sheet_name)
            header_index: int | None = None
            header_roles: dict[str, int] = {}
            for idx, raw_row in enumerate(rows[:10]):
                headers = [self._normalize_codelist_label(value) for value in raw_row]
                roles = self._codelist_header_roles(headers)
                if "code" in roles and "label" in roles:
                    header_index = idx
                    header_roles = roles
                    break

            if header_index is not None:
                parse_rows = rows[header_index + 1 :]
                code_idx = header_roles["code"]
                label_idx = header_roles["label"]
                question_idx = header_roles.get("question")
            else:
                parse_rows = rows
                code_idx = 0
                label_idx = 1
                question_idx = None

            for row in parse_rows:
                if not row or not any(self._normalize_codelist_label(value) for value in row):
                    continue
                code = self._normalize_codelist_code(row[code_idx] if code_idx < len(row) else "")
                label = self._normalize_codelist_label(row[label_idx] if label_idx < len(row) else "")
                if not code or not label or not self._valid_codelist_code(code):
                    continue
                question = sheet_question
                if question_idx is not None and question_idx < len(row):
                    question = self._normalize_codelist_label(row[question_idx]) or question
                key = (question.casefold(), code.casefold())
                prior = records_by_key.get(key)
                if prior is not None:
                    if prior.label != label:
                        warnings.append(
                            f"Duplicate uploaded code {code} has conflicting definitions; using latest: {label}"
                        )
                    else:
                        warnings.append(f"Duplicate uploaded code {code} repeated; keeping one copy.")
                records_by_key[key] = CodelistRecord(
                    question_col=question,
                    code=code,
                    label=label,
                )

        records = sorted(
            records_by_key.values(),
            key=lambda item: (item.question_col.casefold(), self._data_intake_sort_key(item.code)),
        )
        if not records:
            raise ValueError(
                "No usable codelist rows found.\n\n"
                "Expected code/label columns, or an XLSX Q-sheet with code in column A and label in column B."
            )
        return records, warnings

    def _read_codelist_records_csv(self, path: Path, *, source: str) -> list[CodelistRecord]:
        if not path.exists() or not path.is_file():
            return []
        rows = self._read_data_intake_rows(path)
        if not rows:
            return []
        headers = [self._normalize_codelist_label(value) for value in rows[0]]
        roles = self._codelist_header_roles(headers)
        code_idx = roles.get("code")
        label_idx = roles.get("label")
        if code_idx is None or label_idx is None:
            return []
        question_idx = roles.get("question")
        records: list[CodelistRecord] = []
        for row in rows[1:]:
            code = self._normalize_codelist_code(row[code_idx] if code_idx < len(row) else "")
            label = self._normalize_codelist_label(row[label_idx] if label_idx < len(row) else "")
            if not code or not label:
                continue
            question = ""
            if question_idx is not None and question_idx < len(row):
                question = self._normalize_codelist_label(row[question_idx])
            records.append(CodelistRecord(question_col=question, code=code, label=label, source=source))
        return records

    def _write_codelist_records_csv(self, path: Path, records: list[CodelistRecord], *, include_created_at: bool = False) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        fieldnames = ["question_col", "code", "label", "source"]
        if include_created_at:
            fieldnames.extend(["created_at", "notes"])
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            writer.writeheader()
            for record in sorted(
                records,
                key=lambda item: (item.question_col.casefold(), self._data_intake_sort_key(item.code)),
            ):
                row = {
                    "question_col": record.question_col,
                    "code": record.code,
                    "label": record.label,
                    "source": record.source,
                }
                if include_created_at:
                    row["created_at"] = ""
                    row["notes"] = ""
                writer.writerow(row)

    def _codelist_records_by_key(self, records: list[CodelistRecord]) -> dict[tuple[str, str], CodelistRecord]:
        return {(record.question_col.casefold(), record.code.casefold()): record for record in records}

    def _next_user_code(self, user_records: list[CodelistRecord]) -> int:
        numeric = [
            int(record.code)
            for record in user_records
            if record.code.isdigit() and int(record.code) >= USER_CODE_START
        ]
        return (max(numeric) + 1) if numeric else USER_CODE_START

    def _write_codelist_reconciliation_report(
        self,
        report_path: Path,
        rows: list[dict[str, str]],
    ) -> None:
        report_path.parent.mkdir(parents=True, exist_ok=True)
        fieldnames = ["status", "question_col", "code", "old_label", "new_label", "note"]
        with report_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)

    def _track_daily_codelist_file(
        self,
        survey_id: str,
        source_path: Path,
        *,
        day: int,
    ) -> CodelistReconciliationResult:
        if source_path.suffix.casefold() not in {".csv", ".xlsx", ".xlsm"}:
            raise ValueError("Daily codelist must be CSV or XLSX/XLSM.")
        workspace = self._survey_workspace_dir(survey_id)
        current_dir = workspace / "inputs" / "data" / "codelists" / "current"
        current_path = current_dir / f"{survey_id}_current_codelist{source_path.suffix.casefold()}"
        reconciled_dir = workspace / "inputs" / "data" / "codelists" / "reconciled"
        delivered_path = reconciled_dir / f"{survey_id}_delivered_codes_current.csv"
        user_codes_path = reconciled_dir / f"{survey_id}_user_codes_current.csv"
        report_stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        report_path = reconciled_dir / "reports" / f"{survey_id}_day{day}_codelist_reconciliation_{report_stamp}.csv"

        current_dir.mkdir(parents=True, exist_ok=True)
        previous_current_paths = [str(path) for path in self._data_files_in_dir(current_dir)]
        for path in self._data_files_in_dir(current_dir):
            if path.resolve() == source_path.resolve():
                continue
            path.unlink()
        if source_path.resolve() != current_path.resolve():
            shutil.copy2(source_path, current_path)

        delivered_records, parse_warnings = self._parse_codelist_records(source_path)
        previous_records = self._read_codelist_records_csv(delivered_path, source="delivered")
        user_records = self._read_codelist_records_csv(user_codes_path, source="user")
        if not user_codes_path.exists():
            self._write_codelist_records_csv(user_codes_path, [], include_created_at=True)

        previous_by_key = self._codelist_records_by_key(previous_records)
        delivered_by_key = self._codelist_records_by_key(delivered_records)
        user_by_label: dict[tuple[str, str], list[CodelistRecord]] = {}
        user_code_warnings: list[str] = []
        for record in user_records:
            if record.code.isdigit() and int(record.code) < USER_CODE_START:
                user_code_warnings.append(
                    f"User code {record.code} is below reserved range {USER_CODE_START}+."
                )
            label_key = self._codelist_label_key(record.label)
            if label_key:
                user_by_label.setdefault((record.question_col.casefold(), label_key), []).append(record)

        added_keys = sorted(set(delivered_by_key) - set(previous_by_key))
        removed_keys = sorted(set(previous_by_key) - set(delivered_by_key))
        shared_keys = sorted(set(delivered_by_key) & set(previous_by_key))
        changed_keys = [
            key
            for key in shared_keys
            if previous_by_key[key].label != delivered_by_key[key].label
        ]

        report_rows: list[dict[str, str]] = []
        for key in added_keys:
            record = delivered_by_key[key]
            report_rows.append(
                {
                    "status": "ADDED",
                    "question_col": record.question_col,
                    "code": record.code,
                    "old_label": "",
                    "new_label": record.label,
                    "note": "New delivered code.",
                }
            )
        for key in removed_keys:
            record = previous_by_key[key]
            report_rows.append(
                {
                    "status": "REMOVED",
                    "question_col": record.question_col,
                    "code": record.code,
                    "old_label": record.label,
                    "new_label": "",
                    "note": "Removed because latest delivered codelist is authoritative.",
                }
            )
        for key in changed_keys:
            old_record = previous_by_key[key]
            new_record = delivered_by_key[key]
            report_rows.append(
                {
                    "status": "CHANGED",
                    "question_col": new_record.question_col,
                    "code": new_record.code,
                    "old_label": old_record.label,
                    "new_label": new_record.label,
                    "note": "Definition changed; latest delivered definition is used.",
                }
            )

        manual_review_count = 0
        for record in delivered_records:
            label_key = self._codelist_label_key(record.label)
            if not label_key:
                continue
            matches = user_by_label.get((record.question_col.casefold(), label_key), [])
            if not matches and record.question_col:
                matches = user_by_label.get(("", label_key), [])
            for user_record in matches:
                manual_review_count += 1
                report_rows.append(
                    {
                        "status": "MANUAL_REVIEW",
                        "question_col": record.question_col,
                        "code": record.code,
                        "old_label": user_record.label,
                        "new_label": record.label,
                        "note": f"Delivered code resembles user code {user_record.code}; no automatic merge.",
                    }
                )

        for warning in parse_warnings + user_code_warnings:
            report_rows.append(
                {
                    "status": "WARNING",
                    "question_col": "",
                    "code": "",
                    "old_label": "",
                    "new_label": "",
                    "note": warning,
                }
            )

        self._write_codelist_records_csv(delivered_path, delivered_records)
        self._write_codelist_reconciliation_report(report_path, report_rows)
        next_user_code = self._next_user_code(user_records)
        warning_count = len(removed_keys) + len(changed_keys) + manual_review_count + len(parse_warnings) + len(user_code_warnings)

        self._record_workspace_file(
            survey_id,
            "daily_codelist",
            current_path,
            source_path=source_path,
            workspace_dir="inputs/data/codelists/current",
        )

        state = self._read_survey_state(survey_id)
        state["daily_codelist"] = {
            "day": day,
            "source_path": str(source_path),
            "current_path": str(current_path),
            "delivered_codes_current_path": str(delivered_path),
            "user_codes_current_path": str(user_codes_path),
            "reconciliation_report_path": str(report_path),
            "previous_current_paths": previous_current_paths,
            "delivered_code_count": len(delivered_records),
            "added_codes": len(added_keys),
            "removed_codes": len(removed_keys),
            "changed_definitions": len(changed_keys),
            "manual_review_items": manual_review_count,
            "warning_count": warning_count,
            "user_code_start": USER_CODE_START,
            "next_user_code": next_user_code,
            "updated_at": _utc_now_iso(),
        }
        self._write_survey_state(survey_id, state)

        manifest_path = workspace / "codelist_delivery_manifest.json"
        manifest = _read_json_dict(manifest_path) if manifest_path.exists() else {}
        entries = manifest.get("entries", []) if isinstance(manifest, dict) else []
        if not isinstance(entries, list):
            entries = []
        entries.insert(
            0,
            {
                "tracked_at": _utc_now_iso(),
                "survey_id": survey_id,
                "day": day,
                "source_path": str(source_path),
                "current_path": str(current_path),
                "delivered_codes_current_path": str(delivered_path),
                "user_codes_current_path": str(user_codes_path),
                "reconciliation_report_path": str(report_path),
                "previous_current_paths": previous_current_paths,
                "delivered_code_count": len(delivered_records),
                "added_codes": len(added_keys),
                "removed_codes": len(removed_keys),
                "changed_definitions": len(changed_keys),
                "manual_review_items": manual_review_count,
                "warning_count": warning_count,
                "user_code_start": USER_CODE_START,
                "next_user_code": next_user_code,
            },
        )
        _write_json(
            manifest_path,
            {
                "schema_version": 1,
                "survey_id": survey_id,
                "updated_at": _utc_now_iso(),
                "entries": entries[:100],
            },
        )
        return CodelistReconciliationResult(
            current_path=current_path,
            delivered_path=delivered_path,
            user_codes_path=user_codes_path,
            report_path=report_path,
            delivered_count=len(delivered_records),
            added_count=len(added_keys),
            removed_count=len(removed_keys),
            changed_count=len(changed_keys),
            manual_review_count=manual_review_count,
            warning_count=warning_count,
            next_user_code=next_user_code,
        )

    def _track_data_delivery_file(
        self,
        survey_id: str,
        source_path: Path,
        *,
        day: int,
        is_timeseries: bool,
    ) -> tuple[Path, list[str]]:
        if source_path.suffix.casefold() not in {".csv", ".xlsx", ".xlsm"}:
            raise ValueError("Daily data delivery must be CSV or XLSX/XLSM.")
        workspace = self._survey_workspace_dir(survey_id)
        kind_dir = "timeseries_current" if is_timeseries else "current"
        current_dir = workspace / "inputs" / "data" / kind_dir
        stable_name = f"{survey_id}_{'timeseries_current' if is_timeseries else 'current_data'}.csv"
        current_path = current_dir / stable_name

        new_table = self._prepare_data_intake_table(source_path)
        if new_table.duplicate_count:
            examples = "\n".join(f"  - {item}" for item in new_table.duplicate_examples)
            raise ValueError(
                "Duplicate response identity pairs found in the new delivery.\n\n"
                "Each survnum + timeseries pair must appear only once.\n\n"
                f"Examples:\n{examples}"
            )

        existing_current_files = self._data_files_in_dir(current_dir)
        existing_header_path = existing_current_files[0] if existing_current_files else current_path
        existing_table: DataIntakeTable | None = None
        if existing_header_path.exists():
            try:
                existing_table = self._prepare_data_intake_table(existing_header_path)
            except ValueError as exc:
                raise ValueError(
                    "Could not validate the current tracked data file.\n\n"
                    f"{existing_header_path}\n\n{exc}"
                ) from exc
        existing_headers = existing_table.headers if existing_table is not None else []
        if existing_headers and existing_headers != new_table.headers:
            raise ValueError(
                "Column headers differ from the current tracked file.\n\n"
                "Expected the same columns after applying the standard survnum + timeseries normalization.\n\n"
                "The current local copy was not overwritten."
            )

        current_dir.mkdir(parents=True, exist_ok=True)
        previous_current_paths = [str(path) for path in existing_current_files]
        for path in existing_current_files:
            if path.resolve() == source_path.resolve():
                continue
            path.unlink()
        self._write_data_intake_table(current_path, new_table)
        role_id = "timeseries_current" if is_timeseries else "current_data"
        self._record_workspace_file(
            survey_id,
            role_id,
            current_path,
            source_path=source_path,
        )
        pair_changes = self._data_intake_pair_changes(existing_table, new_table)
        self._update_data_intake_state(
            survey_id,
            role_id=role_id,
            day=day,
            source_path=source_path,
            current_path=current_path,
            table=new_table,
            pair_changes=pair_changes,
        )
        manifest_path = workspace / "data_delivery_manifest.json"
        manifest = _read_json_dict(manifest_path) if manifest_path.exists() else {}
        entries = manifest.get("entries", []) if isinstance(manifest, dict) else []
        if not isinstance(entries, list):
            entries = []
        entries.insert(
            0,
            {
                "tracked_at": _utc_now_iso(),
                "survey_id": survey_id,
                "day": day,
                "source_path": str(source_path),
                "current_path": str(current_path),
                "previous_current_paths": previous_current_paths,
                "timeseries": is_timeseries,
                "headers": new_table.headers,
                "identity": self._data_intake_identity_summary(new_table),
                "pair_changes": pair_changes,
            },
        )
        _write_json(
            manifest_path,
            {
                "schema_version": 2,
                "survey_id": survey_id,
                "updated_at": _utc_now_iso(),
                "entries": entries[:100],
            },
        )
        return current_path, new_table.headers

    def track_data_delivery(self) -> None:
        survey_id = self.main_survey_id_var.get().strip() or self.selected_survey_id_var.get().strip()
        if not survey_id:
            messagebox.showwarning("Survey Platform", "Choose a survey first.")
            return
        source_raw = self.data_delivery_path_var.get().strip()
        if not source_raw:
            messagebox.showwarning("Survey Platform", "Choose a CSV/XLSX delivery first.")
            return
        source_path = self._resolve_gui_path(source_raw)
        if not source_path.exists() or not source_path.is_file():
            messagebox.showerror("Survey Platform", f"Data delivery file not found:\n{source_path}")
            return
        codelist_source_path: Path | None = None
        codelist_raw = self.data_delivery_codelist_path_var.get().strip()
        if codelist_raw:
            codelist_source_path = self._resolve_gui_path(codelist_raw)
            if not codelist_source_path.exists() or not codelist_source_path.is_file():
                messagebox.showerror("Survey Platform", f"Codelist file not found:\n{codelist_source_path}")
                return
            if codelist_source_path.suffix.casefold() not in {".csv", ".xlsx", ".xlsm"}:
                messagebox.showerror("Survey Platform", "Daily codelist must be CSV or XLSX/XLSM.")
                return
        try:
            day = int(self.data_delivery_day_var.get().strip() or "1")
        except ValueError:
            messagebox.showerror("Survey Platform", f"Invalid day: {self.data_delivery_day_var.get()}")
            return
        day = max(1, day)
        self.data_delivery_day_var.set(str(day))

        auto_timeseries = self._is_timeseries_data_file(str(source_path))
        is_timeseries = bool(self.data_delivery_timeseries_var.get() or auto_timeseries)
        try:
            current_path, new_headers = self._track_data_delivery_file(
                survey_id,
                source_path,
                day=day,
                is_timeseries=is_timeseries,
            )
        except ValueError as exc:
            messagebox.showerror("Survey Platform", str(exc))
            self.data_delivery_status_var.set("Data delivery blocked.")
            return
        except Exception as exc:
            messagebox.showerror("Survey Platform", f"Failed tracking data delivery:\n{exc}")
            self.data_delivery_status_var.set("Data delivery tracking failed.")
            return
        codelist_result: CodelistReconciliationResult | None = None
        if codelist_source_path is not None:
            try:
                codelist_result = self._track_daily_codelist_file(
                    survey_id,
                    codelist_source_path,
                    day=day,
                )
            except ValueError as exc:
                messagebox.showerror("Survey Platform", str(exc))
                self.data_delivery_status_var.set("Codelist tracking blocked after data delivery.")
                return
            except Exception as exc:
                messagebox.showerror("Survey Platform", f"Failed tracking codelist:\n{exc}")
                self.data_delivery_status_var.set("Codelist tracking failed after data delivery.")
                return

        self.merge_primary_survey_var.set(survey_id)
        self.merge_primary_csv_var.set(str(current_path))
        self.oe_input_path_var.set(str(current_path))
        question_candidates = self._oe_question_candidates(str(current_path))
        if self.oe_question_combo is not None:
            self.oe_question_combo["values"] = question_candidates
        self.data_delivery_timeseries_var.set(is_timeseries)
        self._mark_workflow_step("daily_data_current", "done", phase="In-field Daily")
        if new_headers:
            self._mark_workflow_step("daily_data_headers_checked", "done", phase="In-field Daily")
        if codelist_result is not None:
            self.data_delivery_codelist_path_var.set(str(codelist_result.current_path))
            self.oe_defaults_path_var.set(str(codelist_result.delivered_path))
            self._mark_workflow_step(
                "daily_codelist_tracked",
                "done",
                phase="In-field Daily",
                note=f"Day {day}",
            )
            if codelist_result.warning_count:
                messagebox.showwarning(
                    "Survey Platform",
                    (
                        "Codelist tracked with items to review.\n\n"
                        f"Removed delivered codes: {codelist_result.removed_count}\n"
                        f"Changed definitions: {codelist_result.changed_count}\n"
                        f"Manual review items: {codelist_result.manual_review_count}\n\n"
                        f"Report:\n{codelist_result.report_path}"
                    ),
                )
        self._set_workflow_data_day(survey_id, day)
        merge_output_raw = self.merge_output_path_var.get().strip()
        if not merge_output_raw or self._is_default_merge_output_path(survey_id, merge_output_raw):
            self.merge_output_path_var.set(
                self._default_merge_output_path(
                    survey_id,
                    self.merge_secondary_path_var.get().strip(),
                    day=day,
                )
            )
        codelist_note = ""
        if codelist_result is not None:
            codelist_note = (
                f" Codelist: {codelist_result.current_path}. "
                f"Warnings: {codelist_result.warning_count}. "
                f"Next user code: {codelist_result.next_user_code}."
            )
        self.data_delivery_status_var.set(f"Tracked day {day} delivery to {current_path}.{codelist_note}")

    def browse_merge_secondary(self) -> None:
        chosen = filedialog.askopenfilename(
            title="Choose secondary merge file",
            initialdir=self._initial_dir_from_candidates(self.merge_primary_csv_var.get()),
            filetypes=[
                ("CSV and Excel", "*.csv *.xlsx"),
                ("CSV", "*.csv"),
                ("Excel", "*.xlsx"),
                ("All files", "*.*"),
            ],
        )
        if chosen:
            self.merge_secondary_path_var.set(chosen)
            survey_id = self.merge_primary_survey_var.get().strip() or self.main_survey_id_var.get().strip()
            merge_output_raw = self.merge_output_path_var.get().strip()
            if survey_id and (not merge_output_raw or self._is_default_merge_output_path(survey_id, merge_output_raw)):
                self.merge_output_path_var.set(self._default_merge_output_path(survey_id, chosen))
            self.merge_status_var.set("Secondary merge file selected.")

    def browse_merge_output(self) -> None:
        chosen = filedialog.asksaveasfilename(
            title="Choose merged output file",
            initialdir=self._initial_dir_from_candidates(
                self.merge_output_path_var.get(),
                self.merge_secondary_path_var.get(),
                self.merge_primary_csv_var.get(),
            ),
            defaultextension=".csv",
            filetypes=[
                ("CSV", "*.csv"),
                ("Excel", "*.xlsx"),
                ("All files", "*.*"),
            ],
        )
        if chosen:
            self.merge_output_path_var.set(chosen)
            self.merge_status_var.set("Merge output path selected.")

    def clear_merge_form(self) -> None:
        self.merge_primary_survey_var.set(self.main_survey_id_var.get().strip())
        self.merge_primary_csv_var.set(self._detect_primary_csv_for_survey(self.merge_primary_survey_var.get().strip()))
        self.merge_secondary_path_var.set("")
        self.merge_output_mode_var.set("new_file")
        self.merge_output_path_var.set(self._default_merge_output_path(self.merge_primary_survey_var.get().strip()))
        self.merge_mapping_mode_var.set("auto")
        self.merge_threshold_var.set("85")
        self.merge_similarity_policy_var.set("drop")
        self.merge_status_var.set("Cleared guided merge form.")

    def browse_wincross_docx(self) -> None:
        chosen = filedialog.askopenfilename(
            title="Choose survey DOCX",
            initialdir=self._initial_dir_from_candidates(
                self.wincross_docx_var.get(),
                self.main_survey_source_var.get(),
            ),
            filetypes=[("DOCX", "*.docx"), ("All files", "*.*")],
        )
        if chosen:
            self.wincross_docx_var.set(chosen)
            if not self.wincross_pdf_var.get().strip():
                candidate_pdf = str(Path(chosen).with_suffix(".pdf"))
                if Path(candidate_pdf).exists():
                    self.wincross_pdf_var.set(candidate_pdf)
            survey_id = self.main_survey_id_var.get().strip() or self._survey_id_from_text(chosen)
            self.wincross_json_out_var.set(
                self._default_wincross_json_workspace_path(survey_id)
                or self._default_wincross_json_path(chosen, self.wincross_raw_input_var.get())
            )
            self.wincross_tables_out_var.set(
                self._default_wincross_tables_workspace_path(survey_id)
                or self._default_wincross_tables_path(chosen, self.wincross_raw_input_var.get())
            )

    def browse_wincross_pdf(self) -> None:
        chosen = filedialog.askopenfilename(
            title="Choose survey PDF",
            initialdir=self._initial_dir_from_candidates(
                self.wincross_pdf_var.get(),
                self.wincross_docx_var.get(),
            ),
            filetypes=[("PDF", "*.pdf"), ("All files", "*.*")],
        )
        if chosen:
            self.wincross_pdf_var.set(chosen)

    def browse_wincross_raw_input(self) -> None:
        chosen = filedialog.askopenfilename(
            title="Choose pre-numbered RAW input",
            initialdir=self._initial_dir_from_candidates(
                self.wincross_raw_input_var.get(),
                self.wincross_docx_var.get(),
            ),
            filetypes=[("Text", "*.txt"), ("All files", "*.*")],
        )
        if chosen:
            self.wincross_raw_input_var.set(chosen)
            survey_id = self.main_survey_id_var.get().strip() or self._survey_id_from_text(chosen)
            self.wincross_json_out_var.set(
                self._default_wincross_json_workspace_path(survey_id)
                or self._default_wincross_json_path(self.wincross_docx_var.get(), chosen)
            )
            self.wincross_tables_out_var.set(
                self._default_wincross_tables_workspace_path(survey_id)
                or self._default_wincross_tables_path(self.wincross_docx_var.get(), chosen)
            )

    def browse_wincross_json_out(self) -> None:
        chosen = filedialog.asksaveasfilename(
            title="Choose JSON output path",
            initialdir=self._initial_dir_from_candidates(
                self.wincross_json_out_var.get(),
                self.wincross_docx_var.get(),
                self.wincross_raw_input_var.get(),
            ),
            defaultextension=".json",
            filetypes=[("JSON", "*.json"), ("All files", "*.*")],
        )
        if chosen:
            self.wincross_json_out_var.set(chosen)

    def browse_wincross_tables_out(self) -> None:
        chosen = filedialog.asksaveasfilename(
            title="Choose TABLES output path",
            initialdir=self._initial_dir_from_candidates(
                self.wincross_tables_out_var.get(),
                self.wincross_docx_var.get(),
                self.wincross_raw_input_var.get(),
            ),
            defaultextension=".txt",
            filetypes=[("Text", "*.txt"), ("All files", "*.*")],
        )
        if chosen:
            self.wincross_tables_out_var.set(chosen)

    def reset_wincross_form(self) -> None:
        survey_id = self.main_survey_id_var.get().strip()
        if survey_id:
            self._apply_main_survey_defaults(survey_id)
        else:
            self.wincross_docx_var.set("")
            self.wincross_pdf_var.set("")
            self.wincross_raw_input_var.set("")
            self.wincross_json_out_var.set("")
            self.wincross_tables_out_var.set("")
            self.wincross_translate_var.set(False)
            self.wincross_skip_tables_var.set(False)
        self.wincross_status_var.set("Reset WinCross form.")

    def browse_oe_input(self) -> None:
        chosen = filedialog.askopenfilename(
            title="Choose survey CSV/XLSX for OE extract",
            initialdir=self._initial_dir_from_candidates(
                self.oe_input_path_var.get(),
                self.main_survey_source_var.get(),
            ),
            filetypes=[("CSV/XLSX", "*.csv *.xlsx"), ("CSV", "*.csv"), ("Excel", "*.xlsx"), ("All files", "*.*")],
        )
        if chosen:
            self.oe_input_path_var.set(chosen)
            self.sync_oe_from_current_input()

    def browse_oe_defaults(self) -> None:
        chosen = filedialog.askopenfilename(
            title="Choose defaults map",
            initialdir=self._initial_dir_from_candidates(
                self.oe_defaults_path_var.get(),
                self.oe_input_path_var.get(),
            ),
            filetypes=[("CSV/XLSX", "*.csv *.xlsx"), ("All files", "*.*")],
        )
        if chosen:
            self.oe_defaults_path_var.set(chosen)

    def browse_oe_existing(self) -> None:
        chosen = filedialog.askopenfilename(
            title="Choose existing codebook",
            initialdir=self._initial_dir_from_candidates(
                self.oe_existing_path_var.get(),
                self.oe_input_path_var.get(),
            ),
            filetypes=[("CSV/XLSX", "*.csv *.xlsx"), ("All files", "*.*")],
        )
        if chosen:
            self.oe_existing_path_var.set(chosen)

    def browse_oe_extract_out(self) -> None:
        chosen = filedialog.asksaveasfilename(
            title="Choose extracted codebook output",
            initialdir=self._initial_dir_from_candidates(
                self.oe_extract_out_var.get(),
                self.oe_input_path_var.get(),
            ),
            defaultextension=".csv",
            filetypes=[("CSV", "*.csv"), ("All files", "*.*")],
        )
        if chosen:
            self.oe_extract_out_var.set(chosen)
            self.oe_render_codebook_var.set(chosen)
            if not self.oe_render_out_var.get().strip():
                self.oe_render_out_var.set(self._suggest_oe_render_out(chosen, self.oe_question_col_var.get()))

    def browse_oe_render_codebook(self) -> None:
        chosen = filedialog.askopenfilename(
            title="Choose codebook to render",
            initialdir=self._initial_dir_from_candidates(
                self.oe_render_codebook_var.get(),
                self.oe_extract_out_var.get(),
                self.oe_input_path_var.get(),
            ),
            filetypes=[("CSV/XLSX", "*.csv *.xlsx"), ("All files", "*.*")],
        )
        if chosen:
            self.oe_render_codebook_var.set(chosen)
            if not self.oe_render_out_var.get().strip():
                self.oe_render_out_var.set(self._suggest_oe_render_out(chosen, self.oe_question_col_var.get()))

    def browse_oe_render_out(self) -> None:
        chosen = filedialog.asksaveasfilename(
            title="Choose rendered OE TXT output",
            initialdir=self._initial_dir_from_candidates(
                self.oe_render_out_var.get(),
                self.oe_render_codebook_var.get(),
            ),
            defaultextension=".txt",
            filetypes=[("Text", "*.txt"), ("All files", "*.*")],
        )
        if chosen:
            self.oe_render_out_var.set(chosen)

    def sync_oe_from_current_input(self) -> None:
        input_path = self.oe_input_path_var.get().strip()
        question_candidates = self._oe_question_candidates(input_path)
        if self.oe_question_combo is not None:
            self.oe_question_combo["values"] = question_candidates
        current = self.oe_question_col_var.get().strip()
        if current not in question_candidates:
            current = question_candidates[0] if question_candidates else ""
            self.oe_question_col_var.set(current)
        self._on_oe_question_selected()
        self.oe_status_var.set("Updated OE question choices from current input.")

    def sync_oe_from_main(self) -> None:
        survey_id = self.main_survey_id_var.get().strip()
        if not survey_id:
            messagebox.showinfo("Survey Platform", "Set a Main Survey first.")
            return
        self.oe_input_path_var.set(self._detect_primary_csv_for_survey(survey_id))
        self.sync_oe_from_current_input()
        self.oe_status_var.set(f"Loaded OE defaults from {survey_id}.")

    def _on_oe_question_selected(self, _event=None) -> None:
        question_col = self.oe_question_col_var.get().strip()
        if question_col:
            survey_id = self.main_survey_id_var.get().strip() or self.selected_survey_id_var.get().strip()
            if not self.oe_extract_out_var.get().strip():
                self.oe_extract_out_var.set(
                    self._default_oe_output_path(survey_id, question_col)
                    or self._suggest_oe_extract_out(self.oe_input_path_var.get(), question_col)
                )
            self.oe_render_codebook_var.set(self.oe_extract_out_var.get().strip())
            self.oe_render_selector_var.set(self._suggest_oe_selector(question_col))
            if not self.oe_render_out_var.get().strip():
                self.oe_render_out_var.set(
                    self._default_oe_render_output_path(survey_id, question_col)
                    or self._suggest_oe_render_out(self.oe_render_codebook_var.get(), question_col)
                )

    def reset_oe_form(self) -> None:
        survey_id = self.main_survey_id_var.get().strip()
        self.oe_defaults_path_var.set("")
        self.oe_existing_path_var.set("")
        self.oe_include_blank_var.set(False)
        self.oe_render_title_var.set("ABLE TO MENTION")
        self.oe_render_item_label_var.set("short")
        self.oe_render_out_var.set("")
        if survey_id:
            self.sync_oe_from_main()
        else:
            self.oe_input_path_var.set("")
            self.oe_question_col_var.set("")
            self.oe_extract_out_var.set("")
            self.oe_render_codebook_var.set("")
            self.oe_render_selector_var.set("")
        self.oe_status_var.set("Reset OE form.")

    def browse_results_docx(self) -> None:
        chosen = filedialog.askopenfilename(
            title="Choose results DOCX template",
            initialdir=self._initial_dir_from_candidates(
                self.results_docx_var.get(),
                self.main_survey_source_var.get(),
            ),
            filetypes=[("DOCX", "*.docx"), ("All files", "*.*")],
        )
        if chosen:
            self.results_docx_var.set(chosen)
            if not self.results_output_var.get().strip():
                survey_id = self.main_survey_id_var.get().strip() or self._survey_id_from_text(chosen)
                self.results_output_var.set(
                    self._default_results_output_path(survey_id, chosen)
                    or self._default_empty_results_output(chosen)
                )

    def browse_results_tables(self) -> None:
        chosen = filedialog.askopenfilename(
            title="Choose TABLES file",
            initialdir=self._initial_dir_from_candidates(
                self.results_tables_var.get(),
                self.results_docx_var.get(),
            ),
            filetypes=[("Text", "*.txt"), ("All files", "*.*")],
        )
        if chosen:
            self.results_tables_var.set(chosen)

    def browse_results_json(self) -> None:
        chosen = filedialog.askopenfilename(
            title="Choose survey JSON",
            initialdir=self._initial_dir_from_candidates(
                self.results_json_var.get(),
                self.results_docx_var.get(),
            ),
            filetypes=[("JSON", "*.json"), ("All files", "*.*")],
        )
        if chosen:
            self.results_json_var.set(chosen)

    def browse_results_output(self) -> None:
        chosen = filedialog.asksaveasfilename(
            title="Choose empty-results output DOCX",
            initialdir=self._initial_dir_from_candidates(
                self.results_output_var.get(),
                self.results_docx_var.get(),
            ),
            defaultextension=".docx",
            filetypes=[("DOCX", "*.docx"), ("All files", "*.*")],
        )
        if chosen:
            self.results_output_var.set(chosen)

    def sync_results_from_main(self) -> None:
        survey_id = self.main_survey_id_var.get().strip()
        if not survey_id:
            messagebox.showinfo("Survey Platform", "Set a Main Survey first.")
            return
        docx = self._preferred_workspace_role_file(
            survey_id,
            "results_template_docx",
        ) or self._first_stage_path(survey_id, "results", "inputs", "report_docx")
        self.results_docx_var.set(docx)
        self.results_tables_var.set(
            self._preferred_workspace_role_file(survey_id, "reviewed_tables", "wincross_tables")
            or self._preferred_survey_file(survey_id, "tables_txt")
        )
        self.results_json_var.set(
            self._preferred_workspace_role_file(survey_id, "wincross_json")
            or self._preferred_survey_file(survey_id, "survey_json")
        )
        self.results_output_var.set(self._default_results_output_path(survey_id, docx) if docx else "")
        self.results_status_var.set(f"Loaded results defaults from {survey_id}.")

    def _capture_advanced_merge_result(
        self,
        survey_id: str,
        *,
        opened_at_ts: float,
        primary_raw: str,
        secondary_raw: str,
        output_raw: str,
        primary_json_raw: str = "",
        secondary_json_raw: str = "",
    ) -> Path | None:
        """Record a merge run made in the advanced editor and save its map."""
        if not survey_id or not primary_raw or not secondary_raw:
            return None
        run_path = self._latest_merge_run_path_for_output(output_raw, primary_raw)
        if run_path is None:
            return None
        try:
            if run_path.stat().st_mtime < opened_at_ts - 2:
                return None
        except OSError:
            return None
        self._record_merge_output(survey_id, self._resolve_gui_path(output_raw) if output_raw else None)
        return self._save_guided_merge_map_from_run(
            survey_id,
            primary_raw=primary_raw,
            secondary_raw=secondary_raw,
            output_raw=output_raw,
            command_output=f"Catalog run JSON: {run_path}",
            primary_json_raw=primary_json_raw,
            secondary_json_raw=secondary_json_raw,
        )

    def open_advanced_merge_editor(self) -> None:
        survey_id = (
            self.merge_primary_survey_var.get().strip()
            or self.main_survey_id_var.get().strip()
            or self.selected_survey_id_var.get().strip()
        )
        primary_csv = self.merge_primary_csv_var.get().strip()
        if not primary_csv and survey_id:
            primary_csv = self._detect_primary_csv_for_survey(survey_id)
        secondary_path = self.merge_secondary_path_var.get().strip()
        output_path = self.merge_output_path_var.get().strip()
        saved_map_path, _saved_map_note, saved_record = self._saved_guided_merge_setup(
            survey_id,
            primary_raw=primary_csv,
            secondary_raw=secondary_path,
        )
        if saved_record:
            primary_csv = str(saved_record.get("primary_path", "") or primary_csv).strip()
            secondary_path = str(saved_record.get("secondary_path", "") or secondary_path).strip()
            if not output_path:
                output_path = self._default_merge_output_path(survey_id, secondary_path)
        primary_json = str(saved_record.get("primary_survey_json", "") if saved_record else "").strip()
        if not primary_json:
            primary_json = self._merge_primary_json_for_survey(survey_id)
        secondary_json = str(saved_record.get("secondary_survey_json", "") if saved_record else "").strip()
        if not secondary_json:
            secondary_json = self._merge_secondary_json_for_path(secondary_path)
        opened_at_ts = datetime.now().timestamp()

        try:
            window, app = self._open_specialist_window(
                title="Advanced Table Merge Editor",
                module_name="surveyplatform_advanced_merge_gui",
                module_path=self._specialist_gui_path("pipelines", "table_merger", "sheetMerging_gui.py"),
                class_name="SheetMergeGUI",
            )
            if primary_csv:
                app.primary_path_var.set(primary_csv)
            if secondary_path:
                app.secondary_path_var.set(secondary_path)
            if output_path:
                app.output_path_var.set(output_path)
            app.output_mode_var.set(str(self.merge_output_mode_var.get() or "new_file"))
            app.mapping_mode_var.set("manual_only" if saved_map_path is not None else str(self.merge_mapping_mode_var.get() or "auto"))
            app.threshold_var.set(str(self.merge_threshold_var.get() or "85"))
            app.question_similarity_policy_var.set(str(self.merge_similarity_policy_var.get() or "drop"))
            if primary_json:
                app.primary_survey_json_var.set(primary_json)
            if secondary_json:
                app.secondary_survey_json_var.set(secondary_json)
            if hasattr(app, "_sync_output_state"):
                app._sync_output_state()
            if hasattr(app, "refresh_inputs"):
                app.refresh_inputs()
            if saved_map_path is not None and hasattr(app, "load_mapping_path"):
                try:
                    loaded = app.load_mapping_path(saved_map_path)
                    if hasattr(app, "log"):
                        app.log(f"Loaded saved merge map: {saved_map_path} ({loaded} mapped row(s))")
                except Exception as exc:
                    if hasattr(app, "log"):
                        app.log(f"Could not load saved merge map {saved_map_path}: {exc}")
            if hasattr(app, "backend_status_var"):
                app.backend_status_var.set("Backend: seeded from Survey Platform")
            if hasattr(app, "log"):
                app.log(
                    f"Seeded from Survey Platform for {survey_id or '(manual merge)'}"
                )
            def _capture_advanced_merge_map(event) -> None:
                if event.widget is not window:
                    return
                try:
                    primary_after = str(app.primary_path_var.get()).strip()
                    secondary_after = str(app.secondary_path_var.get()).strip()
                    output_after = str(app.output_path_var.get()).strip()
                    primary_json_after = str(app.primary_survey_json_var.get()).strip()
                    secondary_json_after = str(app.secondary_survey_json_var.get()).strip()
                    if primary_after:
                        self.merge_primary_csv_var.set(primary_after)
                    if secondary_after:
                        self.merge_secondary_path_var.set(secondary_after)
                    if output_after:
                        self.merge_output_path_var.set(output_after)
                    map_path = self._capture_advanced_merge_result(
                        survey_id,
                        opened_at_ts=opened_at_ts,
                        primary_raw=primary_after,
                        secondary_raw=secondary_after,
                        output_raw=output_after,
                        primary_json_raw=primary_json_after,
                        secondary_json_raw=secondary_json_after,
                    )
                    if map_path is not None:
                        self.merge_status_var.set(f"Saved merge map from advanced editor: {map_path}")
                except Exception as exc:
                    self.merge_status_var.set(f"Advanced merge closed; map capture warning: {exc}")

            window.bind("<Destroy>", _capture_advanced_merge_map, add="+")
            self.merge_status_var.set("Opened advanced merge editor.")
        except Exception as exc:
            messagebox.showerror("Survey Platform", f"Could not open advanced merge editor:\n{exc}")
            self.merge_status_var.set("Advanced merge editor failed to open.")

    def open_advanced_oe_editor(self) -> None:
        survey_id = self.main_survey_id_var.get().strip() or self.selected_survey_id_var.get().strip()
        survey_input = self.oe_input_path_var.get().strip() or self._detect_primary_csv_for_survey(
            survey_id
        )
        defaults_path = self.oe_defaults_path_var.get().strip()
        user_codes_path = ""
        if survey_id:
            codelist_state = self._read_survey_state(survey_id).get("daily_codelist", {})
            if isinstance(codelist_state, dict):
                defaults_path = defaults_path or str(codelist_state.get("delivered_codes_current_path", "")).strip()
                user_codes_path = str(codelist_state.get("user_codes_current_path", "")).strip()
        existing_codebook = self._preferred_existing_oe_codebook()
        question_col = self.oe_question_col_var.get().strip()
        title = self.oe_render_title_var.get().strip() or "ABLE TO MENTION"
        selector = self.oe_render_selector_var.get().strip()
        item_label = self.oe_render_item_label_var.get().strip() or "short"

        try:
            _window, app = self._open_specialist_window(
                title="Advanced OE Codebook Editor",
                module_name="surveyplatform_advanced_oe_gui",
                module_path=self._specialist_gui_path("pipelines", "open_end_coder", "oe_codebook_gui.py"),
                class_name="App",
            )
            if survey_input:
                app.survey_var.set(survey_input)
            if defaults_path:
                app.defaults_var.set(defaults_path)
            if user_codes_path and hasattr(app, "user_codes_var"):
                app.user_codes_var.set(user_codes_path)
            if existing_codebook:
                app.existing_var.set(existing_codebook)
            if survey_input and hasattr(app, "load_survey"):
                app.load_survey()
            if question_col and getattr(app, "q_combo", None) is not None:
                values = [str(value) for value in app.q_combo["values"]]
                if any(value.casefold() == question_col.casefold() for value in values):
                    app.question_var.set(question_col)
                    app.question_changed(show_errors=False)
            app.title_var.set(title)
            if selector:
                app.selector_var.set(selector)
            app.label_mode_var.set(item_label)
            if hasattr(app, "refresh_preview"):
                app.refresh_preview()
            if hasattr(app, "set_status"):
                app.set_status("Seeded from Survey Platform.")
            self.oe_status_var.set("Opened advanced OE editor.")
        except Exception as exc:
            messagebox.showerror("Survey Platform", f"Could not open advanced OE editor:\n{exc}")
            self.oe_status_var.set("Advanced OE editor failed to open.")

    def _questions_table_exists(self, conn: sqlite3.Connection, name: str) -> bool:
        row = conn.execute(
            """
            SELECT name
            FROM sqlite_master
            WHERE type='table' AND name=?
            """,
            (name,),
        ).fetchone()
        return row is not None

    def _questions_column_names(self, conn: sqlite3.Connection, table: str) -> set[str]:
        rows = conn.execute(f"PRAGMA table_info({table})").fetchall()
        out: set[str] = set()
        for row in rows:
            if isinstance(row, sqlite3.Row):
                out.add(str(row["name"]))
            else:
                out.add(str(row[1]))
        return out

    def _questions_attribute_columns_available(self, conn: sqlite3.Connection) -> bool:
        cols = self._questions_column_names(conn, "repo_questions")
        required = {
            "question_type",
            "question_family",
            "question_group",
            "question_target",
            "target_source",
            "question_attributes_json",
            "classification_groups_json",
        }
        return required.issubset(cols)

    def _question_page_size(self) -> int:
        limit_raw = str(self.question_limit_var.get() or "").strip()
        if not limit_raw:
            return 30
        try:
            limit = int(limit_raw)
        except Exception as exc:
            raise ValueError("Per page must be an integer.") from exc
        if limit <= 0:
            raise ValueError("Per page must be > 0.")
        return limit

    def _question_filters_snapshot(self) -> dict[str, object]:
        (
            query_text,
            match_mode,
            state,
            region,
            client,
            q_type_filter,
            q_family_filter,
            q_group_filter,
            q_target_filter,
            survey_codes,
            page_size,
        ) = self._questions_filters()
        return {
            "query_text": query_text,
            "match_mode": match_mode,
            "state": state,
            "region": region,
            "client": client,
            "q_type_filter": q_type_filter,
            "q_family_filter": q_family_filter,
            "q_group_filter": q_group_filter,
            "q_target_filter": q_target_filter,
            "survey_codes": set(survey_codes) if survey_codes else None,
            "page_size": page_size,
        }

    def _on_question_group_preset_selected(self, _event=None) -> None:
        self.question_group_filter_var.set(str(self.question_group_preset_var.get() or "").strip())

    def _question_total_pages(self, total_count: int, page_size: int) -> int:
        if total_count <= 0 or page_size <= 0:
            return 0
        return (total_count + page_size - 1) // page_size

    def _update_question_pagination_state(
        self,
        *,
        mode: str,
        filters: dict[str, object] | None,
        total_count: int,
        page: int,
    ) -> None:
        self.question_last_mode = mode
        self.question_last_filters = dict(filters) if filters is not None else None
        self.question_last_total_count = max(0, int(total_count))
        page_size = 0
        if filters is not None:
            try:
                page_size = int(filters.get("page_size", 0) or 0)
            except Exception:
                page_size = 0
        total_pages = self._question_total_pages(self.question_last_total_count, page_size)
        current_page = max(1, min(page, total_pages)) if total_pages else 1
        self.question_page_var.set(str(current_page))
        if total_pages:
            start = ((current_page - 1) * page_size) + 1
            end = min(current_page * page_size, self.question_last_total_count)
            self.question_page_summary_var.set(
                f"Page {current_page} of {total_pages} | showing {start}-{end} of {self.question_last_total_count}"
            )
        else:
            self.question_page_summary_var.set("Page 0 of 0 | showing 0 of 0")

        if self.question_prev_button is not None:
            self.question_prev_button.configure(state=tk.NORMAL if total_pages > 1 and current_page > 1 else tk.DISABLED)
        if self.question_next_button is not None:
            self.question_next_button.configure(
                state=tk.NORMAL if total_pages > 1 and current_page < total_pages else tk.DISABLED
            )

    def _rerun_question_page(self, target_page: int) -> None:
        if not self.question_last_mode or not self.question_last_filters:
            messagebox.showinfo("Question Browser", "Run Search or Frequency first.")
            return
        page = max(1, int(target_page))
        if self.question_last_mode == "search":
            self._run_question_search_page(self.question_last_filters, page)
            return
        if self.question_last_mode == "frequency":
            self._run_question_frequency_page(self.question_last_filters, page)
            return
        messagebox.showinfo("Question Browser", "Unsupported question result mode.")

    def show_previous_question_page(self) -> None:
        current_page = max(1, int(str(self.question_page_var.get() or "1").strip() or "1"))
        if current_page <= 1:
            return
        self._rerun_question_page(current_page - 1)

    def show_next_question_page(self) -> None:
        if not self.question_last_filters:
            messagebox.showinfo("Question Browser", "Run Search or Frequency first.")
            return
        page_size = int(self.question_last_filters.get("page_size", 0) or 0)
        total_pages = self._question_total_pages(self.question_last_total_count, page_size)
        current_page = max(1, int(str(self.question_page_var.get() or "1").strip() or "1"))
        if total_pages and current_page >= total_pages:
            return
        self._rerun_question_page(current_page + 1)

    def _questions_filters(self) -> tuple[str, str, str, str, str, str, str, str, str, set[str] | None, int]:
        query_text = str(self.question_query_var.get() or "").strip()
        match_mode = _normalize_question_match_mode(self.question_match_mode_var.get())
        state = str(self.question_state_var.get() or "").strip().upper()
        region = str(self.question_region_var.get() or "").strip()
        client = str(self.question_client_var.get() or "").strip()
        q_type_filter = str(self.question_type_filter_var.get() or "").strip()
        q_family_filter = str(self.question_family_filter_var.get() or "").strip()
        q_group_filter = str(self.question_group_filter_var.get() or "").strip()
        q_target_filter = str(self.question_target_filter_var.get() or "").strip()
        limit = self._question_page_size()
        if state and not re.fullmatch(r"[A-Z]{2}", state):
            raise ValueError("State must be a 2-letter code like IL/WI/CA.")

        survey_codes = self._parse_survey_codes(self.question_survey_var.get())
        return (
            query_text,
            match_mode,
            state,
            region,
            client,
            q_type_filter,
            q_family_filter,
            q_group_filter,
            q_target_filter,
            survey_codes,
            limit,
        )

    def _parse_survey_codes(self, raw: str) -> set[str] | None:
        text = str(raw or "").strip()
        if not text:
            return None
        out: set[str] = set()
        for token in re.split(r"[\s,]+", text):
            token = token.strip()
            if not token:
                continue
            code = _normalize_survey_code(token)
            if not code:
                raise ValueError(f"Invalid survey code filter: {token}")
            out.add(code)
        return out if out else None

    def _question_filter_sql(
        self,
        *,
        survey_codes: set[str] | None,
        region: str,
        state: str,
        client: str,
        q_type_filter: str,
        q_family_filter: str,
        q_group_filter: str,
        q_target_filter: str,
        attrs_enabled: bool,
    ) -> tuple[str, list[object]]:
        clauses: list[str] = [
            "question_browser_is_english(q.question_text, q.item_text, q.source_json_path) = 1"
        ]
        params: list[object] = []

        if survey_codes:
            codes = sorted(survey_codes, key=int)
            placeholders = ",".join(["?"] * len(codes))
            clauses.append(f"q.survey_code IN ({placeholders})")
            params.extend(codes)
        if region:
            clauses.append("LOWER(q.region) = LOWER(?)")
            params.append(region)
        if state:
            clauses.append("q.state_code = ?")
            params.append(state)
        if client:
            clauses.append("LOWER(q.client) = LOWER(?)")
            params.append(client)
        if q_type_filter:
            if not attrs_enabled:
                raise ValueError(
                    "Question Type/Family/Group/Target filters require attribute columns. Rebuild index first."
                )
            clauses.append("LOWER(q.question_type) LIKE ?")
            params.append(f"%{q_type_filter.casefold()}%")
        if q_family_filter:
            if not attrs_enabled:
                raise ValueError(
                    "Question Type/Family/Group/Target filters require attribute columns. Rebuild index first."
                )
            clauses.append("LOWER(q.question_family) LIKE ?")
            params.append(f"%{q_family_filter.casefold()}%")
        if q_group_filter:
            if not attrs_enabled:
                raise ValueError(
                    "Question Type/Family/Group/Target filters require attribute columns. Rebuild index first."
                )
            clauses.append(
                "("
                "LOWER(q.question_family) LIKE ? "
                "OR LOWER(q.question_group) LIKE ? "
                "OR LOWER(q.classification_groups_json) LIKE ? "
                "OR LOWER(q.question_topics_json) LIKE ? "
                "OR LOWER(q.structure_tags_json) LIKE ? "
                "OR LOWER(q.scale_tags_json) LIKE ?"
                ")"
            )
            group_pattern = f"%{q_group_filter.casefold()}%"
            params.extend(
                [
                    group_pattern,
                    group_pattern,
                    group_pattern,
                    group_pattern,
                    group_pattern,
                    group_pattern,
                ]
            )
        if q_target_filter:
            if not attrs_enabled:
                raise ValueError(
                    "Question Type/Family/Group/Target filters require attribute columns. Rebuild index first."
                )
            clauses.append("LOWER(q.question_target) LIKE ?")
            params.append(f"%{q_target_filter.casefold()}%")

        return " AND " + " AND ".join(clauses), params

    def refresh_question_summary(self) -> None:
        db_path = self._questions_db_path()
        if not db_path.exists():
            self.question_summary_var.set(
                f"Question index: missing ({db_path.name}). Click 'Rebuild Index' to create it."
            )
            return
        try:
            db_mtime = db_path.stat().st_mtime
        except OSError:
            db_mtime = 0.0

        stale_sources: list[str] = []
        try:
            index_path = self._index_path()
            if index_path.exists() and index_path.stat().st_mtime > db_mtime + 1:
                stale_sources.append("repo index")
        except OSError:
            pass
        workspace_root = BASE_DIR / WORKSPACE_DIRNAME
        latest_workspace_json = 0.0
        if workspace_root.exists():
            for json_path in workspace_root.glob("NPA*/outputs/wincross/json/*.json"):
                try:
                    latest_workspace_json = max(latest_workspace_json, json_path.stat().st_mtime)
                except OSError:
                    continue
        if latest_workspace_json > db_mtime + 1:
            stale_sources.append("workspace JSON")

        try:
            conn = self._questions_db_open()
        except Exception as exc:
            self.question_summary_var.set(f"Question index: failed to open ({exc})")
            return

        try:
            if not self._questions_table_exists(conn, "repo_questions"):
                self.question_summary_var.set(
                    f"Question index: schema missing in {db_path.name}. Run rebuild."
                )
                return
            question_count = int(conn.execute("SELECT COUNT(*) FROM repo_questions").fetchone()[0])
            survey_count = int(conn.execute("SELECT COUNT(*) FROM repo_questions_surveys").fetchone()[0])
            fts_enabled = self._questions_table_exists(conn, "repo_questions_fts")
            attrs_enabled = self._questions_attribute_columns_available(conn)
            meta_rows = conn.execute(
                """
                SELECT key, value
                FROM repo_questions_meta
                WHERE key IN ('rebuilt_at', 'workspace_surveys_indexed', 'workspace_questions_indexed')
                """
            ).fetchall()
            meta = {str(row[0]): str(row[1]) for row in meta_rows}
            rebuilt_at = meta.get("rebuilt_at", "(unknown)")
            workspace_surveys = meta.get("workspace_surveys_indexed", "")
        finally:
            conn.close()

        workspace_note = f" | workspace={workspace_surveys}" if workspace_surveys else ""
        stale_note = f" | stale: {', '.join(stale_sources)} newer than index" if stale_sources else ""
        self.question_summary_var.set(
            f"Question index: {db_path} | surveys={survey_count} | questions={question_count} | "
            f"fts={'on' if fts_enabled else 'off'} | attrs={'on' if attrs_enabled else 'off'}"
            f"{workspace_note}{stale_note} | rebuilt_at={rebuilt_at}"
        )

    def _clear_question_tree(self) -> None:
        self.question_result_rows.clear()
        if self.question_tree is None:
            return
        for item in self.question_tree.get_children():
            self.question_tree.delete(item)

    def _populate_question_search_rows(self, rows: list[sqlite3.Row]) -> None:
        self._clear_question_tree()
        if self.question_tree is None:
            return
        for idx, row in enumerate(rows, start=1):
            row_id = f"search:{idx}"
            question_text = str(row["question_text"] or "")
            item_text = str(row["item_text"] or "")
            option_texts = self._parse_option_texts_json(str(row["option_texts_json"] or ""))
            self.question_result_rows[row_id] = {
                "mode": "search",
                "survey_id": str(row["survey_id"] or ""),
                "qid": str(row["qid"] or ""),
                "question_type": str(row["question_type"] or ""),
                "question_family": str(row["question_family"] or ""),
                "question_group": str(row["question_group"] or ""),
                "question_target": str(row["question_target"] or ""),
                "target_source": str(row["target_source"] or ""),
                "state_code": str(row["state_code"] or ""),
                "region": str(row["region"] or ""),
                "client": str(row["client"] or ""),
                "question_text": question_text,
                "item_text": item_text,
                "option_texts": option_texts,
                "canonical_key": str(row["canonical_key"] or ""),
            }
            self.question_tree.insert(
                "",
                tk.END,
                iid=row_id,
                values=(
                    "search",
                    str(row["survey_id"] or ""),
                    str(row["qid"] or ""),
                    str(row["question_type"] or ""),
                    str(row["question_family"] or ""),
                    str(row["question_group"] or ""),
                    str(row["question_target"] or ""),
                    str(row["state_code"] or ""),
                    str(row["region"] or ""),
                    str(row["client"] or ""),
                    "",
                    "",
                    _shorten(question_text, 280),
                    _shorten(item_text, 180),
                ),
            )

    def _populate_question_frequency_rows(self, rows: list[sqlite3.Row]) -> None:
        self._clear_question_tree()
        if self.question_tree is None:
            return
        for idx, row in enumerate(rows, start=1):
            survey_ids_raw = str(row["survey_ids"] or "")
            survey_ids = [token.strip() for token in survey_ids_raw.split(",") if token.strip()]
            survey_preview = ", ".join(survey_ids[:8])
            if len(survey_ids) > 8:
                survey_preview = f"{survey_preview}, ..."
            row_id = f"freq:{idx}"
            question_text = str(row["sample_question_text"] or "")
            item_text = str(row["sample_item_text"] or "")
            option_texts = self._parse_option_texts_json(str(row["sample_option_texts_json"] or ""))
            canonical_keys_raw = str(row["canonical_keys"] or "")
            canonical_keys = [token.strip() for token in canonical_keys_raw.split(",") if token.strip()]
            self.question_result_rows[row_id] = {
                "mode": "freq",
                "survey_count": int(row["survey_count"]),
                "occurrence_count": int(row["occurrence_count"]),
                "frequency_key": str(row["frequency_key"] or ""),
                "canonical_key": str(row["sample_canonical_key"] or ""),
                "canonical_key_count": int(row["canonical_key_count"] or 0),
                "canonical_keys": canonical_keys,
                "question_type": str(row["sample_question_type"] or ""),
                "question_family": str(row["sample_question_family"] or ""),
                "question_group": str(row["sample_question_group"] or ""),
                "question_target": str(row["sample_question_target"] or ""),
                "question_text": question_text,
                "item_text": item_text,
                "option_texts": option_texts,
                "survey_ids": survey_ids,
            }
            self.question_tree.insert(
                "",
                tk.END,
                iid=row_id,
                values=(
                    "freq",
                    survey_preview,
                    "",
                    str(row["sample_question_type"] or ""),
                    str(row["sample_question_family"] or ""),
                    str(row["sample_question_group"] or ""),
                    str(row["sample_question_target"] or ""),
                    "",
                    "",
                    "",
                    int(row["survey_count"]),
                    int(row["occurrence_count"]),
                    _shorten(question_text, 280),
                    _shorten(item_text, 180),
                ),
            )

    def clear_question_results(self) -> None:
        self._clear_question_tree()
        self._update_question_pagination_state(mode="", filters=None, total_count=0, page=1)
        self.question_status_var.set("Question results cleared.")

    def show_question_classification_help(self) -> None:
        top = tk.Toplevel(self.root)
        top.title("Question Classification Help")
        top.geometry("920x700")

        body = tk.Text(top, wrap=tk.WORD)
        body.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

        lines = [
            "Question Browser Classification Help",
            "",
            "Use the filters this way:",
            "Type: exact subtype like favorability_4pt or head_to_head_ballot",
            "Family: older mid-level bucket like candidate_evaluation or ballot_preference",
            "Group/Tag: broader search layer for group names, family aliases, cross-cutting tags, topics, and scale/structure tags",
            "",
            "Common Groups",
            "favorability: combines favorability_4pt and favorability_0_10",
            "approval: approve/disapprove style questions and job-performance approval-style ratings",
            "ballot: head-to-head ballot, generic ballot, undecided follow-up, past-vote recall",
            "demographics: age, race, gender, education, zip, state, occupation, density, union, etc.",
            "issue_position: support/oppose, favor/oppose, agree/disagree, statement-choice",
            "job_performance: explicit excellent/good/poor performance ratings",
            "likelihood: vote likelihood, generic likelihood, vote certainty",
            "party_identity: party ID, leaner, strength, registration",
            "directional_mood: right track / wrong track and similar mood questions",
            "",
            "Common Tags",
            "evaluation: broad attitude/evaluation layer across favorability, approval, traits, tone, consideration",
            "voting: vote-choice, vote-likelihood, vote-certainty, persuasion, past-vote",
            "party: party-ID and party-comparison style questions",
            "comparison: ballot matchups, trait forced-choice, party issue comparisons",
            "position: issue-position questions like support/oppose or favor/oppose",
            "opinion: similar to position, useful for policy-opinion questions",
            "salience: issue-priority / importance questions",
            "probability: generic non-vote likelihood questions",
            "performance: job performance ratings",
            "election: ballot and turnout style questions",
            "",
            "Good Search Examples",
            "Group/Tag = favorability",
            "Group/Tag = evaluation",
            "Group/Tag = voting",
            "Group/Tag = party",
            "Group/Tag = comparison",
            "Group/Tag = salience",
            "Group/Tag = probability",
            "",
            "Useful combinations",
            "Type = favorability_4pt",
            "Family = candidate_evaluation",
            "Group/Tag = evaluation",
            "Target = Trump",
            "",
            "The Group/Tag filter checks:",
            "question_family",
            "question_group",
            "classification_groups (cross-cutting tags like evaluation/voting/comparison)",
            "question_topics",
            "structure_tags",
            "scale_tags",
        ]
        body.insert("1.0", "\n".join(lines))
        body.configure(state=tk.DISABLED)

    def _parse_option_texts_json(self, raw: str) -> list[str]:
        text = str(raw or "").strip()
        if not text:
            return []
        try:
            payload = json.loads(text)
        except Exception:
            return []
        if not isinstance(payload, list):
            return []
        out: list[str] = []
        for item in payload:
            val = str(item).strip()
            if val:
                out.append(val)
        return out

    def _question_payload_match_fields(self, payload: dict[str, object]) -> dict[str, str]:
        options = payload.get("option_texts", [])
        options_text = ""
        if isinstance(options, list):
            options_text = " ".join(str(item) for item in options if str(item).strip())
        return {
            "question": str(payload.get("question_text", "") or ""),
            "item": str(payload.get("item_text", "") or ""),
            "answers": options_text,
            "type": str(payload.get("question_type", "") or ""),
            "family": str(payload.get("question_family", "") or ""),
            "group": str(payload.get("question_group", "") or ""),
            "target": str(payload.get("question_target", "") or ""),
        }

    def _question_preview_match_lines(self, payload: dict[str, object]) -> list[str]:
        filters = self.question_last_filters or {}
        query_text = str(filters.get("query_text", "") or "").strip()
        if not query_text:
            return []
        try:
            match_mode = _normalize_question_match_mode(filters.get("match_mode", "all_words"))
        except Exception:
            match_mode = "all_words"
        fields = self._question_payload_match_fields(payload)
        field_norms = {name: value.casefold() for name, value in fields.items() if value.strip()}

        lines = ["", "Match Details:", f"Query: {query_text}", f"Mode: {match_mode}"]
        if match_mode == "exact":
            phrase = query_text.casefold()
            matched_fields = [name for name, text in field_norms.items() if phrase in text]
            if matched_fields:
                lines.append("Exact phrase found in: " + ", ".join(matched_fields))
            else:
                lines.append("Exact phrase was not found in the preview fields.")
            return lines

        tokens = _question_query_tokens(query_text)
        if not tokens:
            lines.append("No searchable words were found in the query.")
            return lines
        for token in tokens:
            matched_fields = [name for name, text in field_norms.items() if token in text]
            label = ", ".join(matched_fields) if matched_fields else "(not in preview fields)"
            lines.append(f"{token}: {label}")
        return lines

    def _append_question_option_lines(self, lines: list[str], payload: dict[str, object]) -> None:
        options = payload.get("option_texts", [])
        if not isinstance(options, list) or not options:
            return
        lines.append("")
        lines.append("Answer Options:")
        for idx, opt in enumerate(options, start=1):
            text = str(opt).strip()
            if text:
                lines.append(f"{idx}. {text}")

    def preview_selected_question_row(self, _event=None) -> None:
        if self.question_tree is None:
            return
        selected = self.question_tree.selection()
        if not selected:
            messagebox.showinfo("Question Browser", "Select a result row first.")
            return
        row_id = selected[0]
        payload = self.question_result_rows.get(row_id)
        if payload is None:
            messagebox.showinfo("Question Browser", "No detailed payload available for this row.")
            return

        mode = str(payload.get("mode", "")).strip() or "search"
        top = tk.Toplevel(self.root)
        top.title(f"Question Preview ({mode})")
        top.geometry("980x720")

        body = tk.Text(top, wrap=tk.WORD)
        body.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

        lines: list[str] = []
        lines.append(f"Mode: {mode}")
        if mode == "search":
            lines.append(f"Survey: {payload.get('survey_id', '')}")
            lines.append(f"QID: {payload.get('qid', '')}")
            lines.append(f"Type: {payload.get('question_type', '')}")
            lines.append(f"Family: {payload.get('question_family', '')}")
            lines.append(f"Group: {payload.get('question_group', '')}")
            lines.append(f"Target: {payload.get('question_target', '')}")
            lines.append(f"Target source: {payload.get('target_source', '')}")
            lines.append(f"State: {payload.get('state_code', '')}")
            lines.append(f"Region: {payload.get('region', '')}")
            lines.append(f"Client: {payload.get('client', '')}")
            lines.extend(self._question_preview_match_lines(payload))
            lines.append("")
            lines.append("Question:")
            lines.append(str(payload.get("question_text", "")).strip())
            item = str(payload.get("item_text", "")).strip()
            if item:
                lines.append("")
                lines.append("Item:")
                lines.append(item)
            self._append_question_option_lines(lines, payload)
        elif mode == "audit":
            lines.append(f"Audit check: {payload.get('audit_check', '')}")
            lines.append(f"Count: {payload.get('occurrence_count', 0)}")
            lines.append("")
            details = payload.get("audit_details", [])
            if isinstance(details, list) and details:
                lines.append("Details:")
                for detail in details:
                    lines.append(str(detail))
            else:
                lines.append("No row-level details for this check.")
        else:
            frequency_key = str(payload.get("frequency_key", "") or "").strip()
            if frequency_key:
                lines.append(f"Frequency key: {frequency_key}")
            lines.append(f"Sample canonical key: {payload.get('canonical_key', '')}")
            canonical_key_count = int(payload.get("canonical_key_count", 0) or 0)
            if canonical_key_count > 1:
                lines.append(f"Canonical variants merged: {canonical_key_count}")
            lines.append(f"Survey count: {payload.get('survey_count', 0)}")
            lines.append(f"Occurrence count: {payload.get('occurrence_count', 0)}")
            lines.append(f"Type: {payload.get('question_type', '')}")
            lines.append(f"Family: {payload.get('question_family', '')}")
            lines.append(f"Group: {payload.get('question_group', '')}")
            lines.append(f"Target: {payload.get('question_target', '')}")
            lines.extend(self._question_preview_match_lines(payload))
            lines.append("")
            lines.append("Sample Question:")
            lines.append(str(payload.get("question_text", "")).strip())
            item = str(payload.get("item_text", "")).strip()
            if item:
                lines.append("")
                lines.append("Sample Item:")
                lines.append(item)
            self._append_question_option_lines(lines, payload)
            survey_ids = payload.get("survey_ids", [])
            if isinstance(survey_ids, list) and survey_ids:
                lines.append("")
                lines.append("Surveys:")
                lines.append(", ".join(survey_ids))

        body.insert("1.0", "\n".join(lines))
        body.configure(state=tk.DISABLED)

    def export_question_results_csv(self) -> None:
        if self.question_tree is None:
            return
        row_ids = list(self.question_tree.get_children())
        if not row_ids:
            messagebox.showinfo("Question Browser", "No results to export.")
            return

        default_name = f"question_browser_results_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        out_path = filedialog.asksaveasfilename(
            title="Export Question Results CSV",
            defaultextension=".csv",
            initialdir=str(self._repo_root()),
            initialfile=default_name,
            filetypes=[("CSV files", "*.csv"), ("All files", "*.*")],
        )
        if not out_path:
            return

        fieldnames = [
            "mode",
            "survey_id",
            "qid",
            "question_type",
            "question_family",
            "question_group",
            "question_target",
            "target_source",
            "state_code",
            "region",
            "client",
            "survey_count",
            "occurrence_count",
            "frequency_key",
            "canonical_key",
            "canonical_key_count",
            "audit_check",
            "question_text",
            "item_text",
            "option_texts",
            "survey_ids",
            "audit_details",
        ]
        rows_out: list[dict[str, str]] = []
        for row_id in row_ids:
            payload = self.question_result_rows.get(row_id, {})
            option_texts = payload.get("option_texts", [])
            survey_ids = payload.get("survey_ids", [])
            rows_out.append(
                {
                    "mode": str(payload.get("mode", "")),
                    "survey_id": str(payload.get("survey_id", "")),
                    "qid": str(payload.get("qid", "")),
                    "question_type": str(payload.get("question_type", "")),
                    "question_family": str(payload.get("question_family", "")),
                    "question_group": str(payload.get("question_group", "")),
                    "question_target": str(payload.get("question_target", "")),
                    "target_source": str(payload.get("target_source", "")),
                    "state_code": str(payload.get("state_code", "")),
                    "region": str(payload.get("region", "")),
                    "client": str(payload.get("client", "")),
                    "survey_count": str(payload.get("survey_count", "")),
                    "occurrence_count": str(payload.get("occurrence_count", "")),
                    "frequency_key": str(payload.get("frequency_key", "")),
                    "canonical_key": str(payload.get("canonical_key", "")),
                    "canonical_key_count": str(payload.get("canonical_key_count", "")),
                    "audit_check": str(payload.get("audit_check", "")),
                    "question_text": str(payload.get("question_text", "")),
                    "item_text": str(payload.get("item_text", "")),
                    "option_texts": " | ".join([str(x) for x in option_texts]) if isinstance(option_texts, list) else "",
                    "survey_ids": ",".join([str(x) for x in survey_ids]) if isinstance(survey_ids, list) else "",
                    "audit_details": " | ".join([str(x) for x in payload.get("audit_details", [])])
                    if isinstance(payload.get("audit_details", []), list)
                    else "",
                }
            )

        try:
            with open(out_path, "w", newline="", encoding="utf-8-sig") as fh:
                writer = csv.DictWriter(fh, fieldnames=fieldnames)
                writer.writeheader()
                writer.writerows(rows_out)
        except Exception as exc:
            messagebox.showerror("Question Browser", f"CSV export failed:\n{exc}")
            return

        self.question_status_var.set(f"Exported {len(rows_out)} row(s) to {out_path}")

    def _set_selected_survey_details(self, survey_id: str, row: dict, *, status_text: str | None = None) -> None:
        self.selected_survey_id_var.set(row.get("survey_id", survey_id))
        self.selected_stage_var.set(row.get("stage", ""))
        self.selected_last_run_var.set(row.get("last_run", "") or "(unknown)")
        self.selected_platform_status_var.set(self._format_platform_statuses(row))
        available = row.get("available_actions", "") or "(not indexed)"
        self.selected_available_actions_var.set(str(available))
        self.region_var.set(row.get("region", ""))
        self.client_var.set(row.get("client", ""))
        if status_text is not None:
            self.status_var.set(status_text)

    def _select_survey_in_ui(
        self,
        survey_id: str,
        *,
        status_text: str | None = None,
        persist: bool = False,
    ) -> bool:
        row = self.rows_by_id.get(survey_id)
        if not row:
            return False
        if self.tree is not None:
            self.tree.selection_set(survey_id)
            self.tree.focus(survey_id)
            self.tree.see(survey_id)
        self.survey_selector_var.set(survey_id)
        self._set_selected_survey_details(survey_id, row, status_text=status_text)
        if persist:
            self._save_dashboard_survey_preference(survey_id)
        return True

    def load_index(self) -> None:
        if self.tree is None:
            return

        repo_root = self._repo_root()
        index_path = self._index_path()
        current_main = self.main_survey_id_var.get().strip()
        if not index_path.exists():
            messagebox.showerror(
                "Repository Dashboard",
                f"Index not found:\n{index_path}\n\nRun:\npy -3 survey_tool_cli.py repo scan --root {repo_root}",
            )
            return

        try:
            index_payload = _read_json_dict(index_path)
        except Exception as exc:
            messagebox.showerror("Repository Dashboard", f"Failed to read index:\n{exc}")
            return

        surveys = index_payload.get("surveys", [])
        if not isinstance(surveys, list):
            surveys = []

        self.rows_by_id.clear()
        for item in self.tree.get_children():
            self.tree.delete(item)

        sorted_rows = sorted(
            [row for row in surveys if isinstance(row, dict)],
            key=_survey_index_row_desc_sort_key,
        )

        source_root_display = str(index_payload.get("source_root", "")).strip()
        discovery_mode = str(index_payload.get("discovery_mode", "")).strip() or "standard"
        if source_root_display:
            self.source_root_var.set(source_root_display)
        if discovery_mode in REPO_DISCOVERY_MODES:
            self.discovery_mode_var.set(discovery_mode)
        rows_with_platform = 0

        for row in sorted_rows:
            survey_id = str(row.get("survey_id", "")).strip()
            if not survey_id:
                code = str(row.get("survey_code", "")).strip()
                survey_id = f"NPA{code}" if code else ""
            if not survey_id:
                continue

            meta_payload: dict = {}
            meta_path = self._meta_path(survey_id)
            if meta_path.exists():
                try:
                    meta_payload = _read_json_dict(meta_path)
                except Exception:
                    meta_payload = {}

            categorization = _normalize_categorization(
                {"region": row.get("region", ""), "client": row.get("client", "")}
            )
            if not categorization["region"] and not categorization["client"]:
                categorization = _normalize_categorization(meta_payload.get("categorization"))

            meta_stages = meta_payload.get("stages", {})
            if not isinstance(meta_stages, dict):
                meta_stages = {}
            stage_statuses: dict[str, str] = {}
            for stage_id in REPO_PLATFORM_STAGE_IDS:
                status = str(row.get(f"{stage_id}_status", "")).strip()
                if not status:
                    stage_payload = meta_stages.get(stage_id, {})
                    if isinstance(stage_payload, dict):
                        status = str(stage_payload.get("status", "")).strip()
                stage_statuses[stage_id] = status or "unknown"

            if any(status != "unknown" for status in stage_statuses.values()):
                rows_with_platform += 1

            available_actions = row.get("available_actions")
            if isinstance(available_actions, list):
                available_actions_text = ", ".join(
                    [str(item).strip() for item in available_actions if str(item).strip()]
                )
            else:
                available_actions_text = ""

            display_row = {
                "survey_id": survey_id,
                "survey_code": str(row.get("survey_code", "")).strip(),
                "stage": str(row.get("stage", "")).strip(),
                "last_run": str(row.get("last_run", "")).strip(),
                "region": categorization["region"],
                "client": categorization["client"],
                "available_actions": available_actions_text,
            }
            display_row.update(stage_statuses)
            self.rows_by_id[survey_id] = display_row
            self.tree.insert(
                "",
                tk.END,
                iid=survey_id,
                values=self._tree_values_for_row(display_row),
            )

        if self.survey_selector_combo is not None:
            self.survey_selector_combo["values"] = sorted(self.rows_by_id.keys(), key=_survey_id_desc_sort_key)

        status_counts = index_payload.get("status_counts", {})
        counts_summary = ""
        if isinstance(status_counts, dict):
            parts = [f"{key}: {status_counts[key]}" for key in sorted(status_counts.keys())]
            counts_summary = " | ".join(parts)
        updated_at = str(index_payload.get("updated_at", "")).strip()
        self.summary_var.set(
            f"Repository: {repo_root} | Source root: {source_root_display or '(unknown)'} | "
            f"Discovery: {discovery_mode} | "
            f"Surveys: {len(self.rows_by_id)} | Updated: {updated_at or '(unknown)'}"
        )
        self.status_var.set(
            f"Legacy counts: {counts_summary or '(none)'} | Platform stage metadata rows: {rows_with_platform}/{len(self.rows_by_id)}"
        )

        self.selected_survey_id_var.set("")
        self.selected_stage_var.set("")
        self.selected_last_run_var.set("")
        self.selected_platform_status_var.set("")
        self.selected_available_actions_var.set("")
        self.region_var.set("")
        self.client_var.set("")
        restore_survey_id = self._preferred_dashboard_survey_id(set(self.rows_by_id.keys()), current_main)
        if restore_survey_id:
            self._select_survey_in_ui(
                restore_survey_id,
                status_text=f"Restored last survey: {restore_survey_id}.",
            )
            self._apply_main_survey_defaults(restore_survey_id)
            status_payload = self._read_workflow_status(restore_survey_id)
            saved_phase = str(status_payload.get("current_phase", "")).strip()
            if saved_phase in PHASE_VALUES:
                self.phase_var.set(saved_phase)
        else:
            self.main_survey_id_var.set("")
            self.survey_selector_var.set("")
            self.main_survey_source_var.set("")
            self.merge_primary_survey_var.set("")
            self.merge_primary_csv_var.set("")
            self.merge_status_var.set("")
        self._refresh_workflow_home()
        self.refresh_question_summary()

    def _on_select_row(self, _event=None) -> None:
        if self.tree is None:
            return
        selected = self.tree.selection()
        if not selected:
            return
        survey_id = selected[0]
        row = self.rows_by_id.get(survey_id)
        if not row:
            return
        self._set_selected_survey_details(
            survey_id,
            row,
            status_text=f"Selected: {survey_id}. Confirm Survey to make it the active workflow survey.",
        )
        self._save_dashboard_survey_preference(survey_id)

    def save_selected(self) -> None:
        survey_id = self.selected_survey_id_var.get().strip()
        if not survey_id:
            messagebox.showwarning("Repository Dashboard", "Select a survey row first.")
            return

        region = self.region_var.get().strip()
        client = self.client_var.get().strip()
        now_iso = _utc_now_iso()
        repo_root = self._repo_root()

        meta_path = self._meta_path(survey_id)
        meta_payload: dict
        if meta_path.exists():
            try:
                meta_payload = _read_json_dict(meta_path)
            except Exception as exc:
                messagebox.showerror("Repository Dashboard", f"Failed reading meta file:\n{exc}")
                return
        else:
            code = survey_id.removeprefix("NPA")
            meta_payload = {
                "schema_version": 1,
                "survey_code": code,
                "survey_id": survey_id,
            }

        meta_payload["categorization"] = {
            "region": region,
            "client": client,
        }
        meta_payload["updated_at"] = now_iso
        try:
            _write_json(meta_path, meta_payload)
        except Exception as exc:
            messagebox.showerror("Repository Dashboard", f"Failed writing meta file:\n{exc}")
            return

        index_path = self._index_path()
        if not index_path.exists():
            messagebox.showerror(
                "Repository Dashboard",
                f"Index not found:\n{index_path}\n\nRun scan before saving categorization.",
            )
            return
        try:
            index_payload = _read_json_dict(index_path)
        except Exception as exc:
            messagebox.showerror("Repository Dashboard", f"Failed reading index:\n{exc}")
            return

        surveys = index_payload.get("surveys", [])
        if not isinstance(surveys, list):
            surveys = []
            index_payload["surveys"] = surveys

        found = False
        for row in surveys:
            if not isinstance(row, dict):
                continue
            row_survey_id = str(row.get("survey_id", "")).strip()
            if row_survey_id != survey_id:
                continue
            row["region"] = region
            row["client"] = client
            found = True
            break

        if not found:
            code = survey_id.removeprefix("NPA")
            surveys.append(
                {
                    "survey_code": code,
                    "survey_id": survey_id,
                    "stage": meta_payload.get("status", {}).get("stage", ""),
                    "last_run": meta_payload.get("last_run", ""),
                    "primary_source_dir": meta_payload.get("primary_source_dir", ""),
                    "pipelines_completed": meta_payload.get("pipelines_completed", []),
                    "has_json": bool(meta_payload.get("status", {}).get("has_json", False)),
                    "has_tables_txt": bool(meta_payload.get("status", {}).get("has_tables_txt", False)),
                    "region": region,
                    "client": client,
                }
            )

        index_payload["survey_count"] = len([row for row in surveys if isinstance(row, dict)])
        index_payload["updated_at"] = now_iso
        try:
            _write_json(index_path, index_payload)
        except Exception as exc:
            messagebox.showerror("Repository Dashboard", f"Failed writing index:\n{exc}")
            return

        row = self.rows_by_id.get(survey_id)
        if row:
            row["region"] = region
            row["client"] = client

        if self.tree is not None and self.tree.exists(survey_id):
            if row:
                self.tree.item(survey_id, values=self._tree_values_for_row(row))

        self.status_var.set(f"Saved categorization for {survey_id} at {now_iso}")
        self.summary_var.set(f"Repository: {repo_root} | Surveys: {len(self.rows_by_id)} | Updated: {now_iso}")

    def scan_repository(self) -> None:
        cmd = [
            sys.executable,
            str(self._cli_path()),
            "repo",
            "scan",
            "--root",
            str(self._repo_root()),
        ]
        source_root = self.source_root_var.get().strip()
        if source_root:
            cmd.extend(["--source-root", source_root])
        discovery_mode = str(self.discovery_mode_var.get() or "").strip().casefold() or "standard"
        if discovery_mode in REPO_DISCOVERY_MODES:
            cmd.extend(["--discovery-mode", discovery_mode])
        ok = self._run_cli_command(
            cmd,
            title="Repository Scan",
            running_text="Scanning repository...",
            success_callback=self.load_index,
        )
        if ok:
            self._mark_workflow_step("repo_scan_reviewed", "done", phase="Archive/Search")

    def show_selected_survey(self, _event=None) -> None:
        survey_id = self._selected_survey_id()
        if survey_id is None:
            return
        cmd = [
            sys.executable,
            str(self._cli_path()),
            "survey",
            "show",
            "--root",
            str(self._repo_root()),
            survey_id,
        ]
        self._run_cli_command(
            cmd,
            title=f"{survey_id} Survey Show",
            running_text=f"Loading {survey_id} details...",
        )

    def detect_selected_survey(self) -> None:
        survey_id = self._selected_survey_id()
        if survey_id is None:
            return
        cmd = [
            sys.executable,
            str(self._cli_path()),
            "survey",
            "detect",
            "--root",
            str(self._repo_root()),
            survey_id,
        ]
        self._run_cli_command(
            cmd,
            title=f"{survey_id} Detect",
            running_text=f"Detecting live stage status for {survey_id}...",
        )

    def show_selected_history(self) -> None:
        survey_id = self._selected_survey_id()
        if survey_id is None:
            return
        cmd = [
            sys.executable,
            str(self._cli_path()),
            "survey",
            "history",
            "--root",
            str(self._repo_root()),
            survey_id,
            "--limit",
            "50",
        ]
        self._run_cli_command(
            cmd,
            title=f"{survey_id} History",
            running_text=f"Loading {survey_id} history...",
        )

    def open_selected_survey_files(self) -> None:
        survey_id = self._selected_survey_id()
        if survey_id is None:
            return

        survey_dir = self._survey_working_dir(survey_id)
        if survey_dir is None:
            survey_dir = self._survey_metadata_dir(survey_id)
        if not survey_dir.exists() or not survey_dir.is_dir():
            messagebox.showerror(
                "Survey Platform",
                f"Survey workspace folder not found:\n{survey_dir}\n\nRun a local workflow or materialize the survey first if needed.",
            )
            return

        try:
            os.startfile(str(survey_dir))
        except AttributeError:
            subprocess.Popen(["explorer", str(survey_dir)])
        except Exception as exc:
            messagebox.showerror("Survey Platform", f"Could not open survey workspace:\n{exc}")

    def open_selected_survey_metadata(self) -> None:
        survey_id = self._selected_survey_id()
        if survey_id is None:
            return

        survey_dir = self._survey_metadata_dir(survey_id)
        if not survey_dir.exists() or not survey_dir.is_dir():
            messagebox.showerror(
                "Survey Platform",
                f"Survey metadata folder not found:\n{survey_dir}\n\nRun Repo Scan first if needed.",
            )
            return

        try:
            os.startfile(str(survey_dir))
        except AttributeError:
            subprocess.Popen(["explorer", str(survey_dir)])
        except Exception as exc:
            messagebox.showerror("Survey Platform", f"Could not open survey metadata:\n{exc}")

    def materialize_selected_survey(self) -> None:
        survey_id = self._selected_survey_id()
        if survey_id is None:
            return

        default_root = self._default_materialized_root()
        initial = default_root if default_root is not None and default_root.exists() else self._repo_root()
        chosen = filedialog.askdirectory(
            initialdir=str(initial),
            title=f"Choose local materialized output root for {survey_id}",
        )
        if not chosen:
            return

        cmd = [
            sys.executable,
            str(self._cli_path()),
            "repo",
            "materialize",
            "--root",
            str(self._repo_root()),
            "--out",
            chosen,
            "--survey",
            survey_id,
        ]
        self._run_cli_command(
            cmd,
            title=f"{survey_id} Materialize",
            running_text=f"Copying curated files for {survey_id}...",
        )

    def materialize_all_surveys(self) -> None:
        default_root = self._default_materialized_root()
        initial = default_root if default_root is not None and default_root.exists() else self._repo_root()
        chosen = filedialog.askdirectory(
            initialdir=str(initial),
            title="Choose local materialized output root for all surveys",
        )
        if not chosen:
            return

        if not messagebox.askyesno(
            "Survey Platform",
            (
                "Materialize all indexed surveys?\n\n"
                f"Output root:\n{chosen}\n\n"
                "This copies the curated local working files for the whole repository."
            ),
        ):
            return

        cmd = [
            sys.executable,
            str(self._cli_path()),
            "repo",
            "materialize",
            "--root",
            str(self._repo_root()),
            "--out",
            chosen,
        ]
        self._run_cli_command(
            cmd,
            title="Materialize All Surveys",
            running_text="Copying curated files for all indexed surveys...",
        )

    def export_selected_survey_pdf(self) -> None:
        survey_id = self._selected_survey_id()
        if survey_id is None:
            return

        survey_dir = self._survey_materialized_dir(survey_id)
        if survey_dir is None:
            messagebox.showwarning(
                "Survey Platform",
                f"No materialized survey folder found for {survey_id}.\n\nMaterialize it first.",
            )
            return

        script_path = self._script_path("export_materialized_docx_pdfs.ps1")
        if not script_path.exists():
            messagebox.showerror("Survey Platform", f"Script not found:\n{script_path}")
            return

        cmd = [
            "powershell",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script_path),
            "-Root",
            str(survey_dir.parent),
            "-Survey",
            survey_id,
        ]
        self._run_cli_command(
            cmd,
            title=f"{survey_id} Export PDF",
            running_text=f"Exporting PDF for {survey_id}...",
            show_output_on_success=True,
        )

    def jsonify_selected_survey_files(self) -> None:
        survey_id = self._selected_survey_id()
        if survey_id is None:
            return

        survey_dir = self._survey_materialized_dir(survey_id)
        if survey_dir is None:
            messagebox.showwarning(
                "Survey Platform",
                f"No materialized survey folder found for {survey_id}.\n\nMaterialize it first.",
            )
            return

        script_path = self._script_path("jsonify_materialized_surveys.ps1")
        if not script_path.exists():
            messagebox.showerror("Survey Platform", f"Script not found:\n{script_path}")
            return

        cmd = [
            "powershell",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script_path),
            "-Root",
            str(survey_dir.parent),
            "-Survey",
            survey_id,
        ]
        self._run_cli_command(
            cmd,
            title=f"{survey_id} JSONify Files",
            running_text=f"Generating JSON for {survey_id}...",
            show_output_on_success=True,
        )

    def export_all_materialized_pdfs(self) -> None:
        root = self._ensure_materialized_root(title="Export PDFs")
        if root is None:
            return

        script_path = self._script_path("export_materialized_docx_pdfs.ps1")
        if not script_path.exists():
            messagebox.showerror("Survey Platform", f"Script not found:\n{script_path}")
            return

        cmd = [
            "powershell",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script_path),
            "-Root",
            str(root),
        ]
        self._run_cli_command(
            cmd,
            title="Export All PDFs",
            running_text=f"Exporting PDFs under {root.name}...",
            show_output_on_success=True,
        )

    def jsonify_all_materialized_files(self) -> None:
        root = self._ensure_materialized_root(title="Generate JSON")
        if root is None:
            return

        script_path = self._script_path("jsonify_materialized_surveys.ps1")
        if not script_path.exists():
            messagebox.showerror("Survey Platform", f"Script not found:\n{script_path}")
            return

        cmd = [
            "powershell",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script_path),
            "-Root",
            str(root),
        ]
        self._run_cli_command(
            cmd,
            title="JSONify All Materialized Surveys",
            running_text=f"Generating JSON under {root.name}...",
            show_output_on_success=True,
        )

    def run_selected_stage(self) -> None:
        survey_id = self._selected_survey_id()
        if survey_id is None:
            return

        stage = str(self.run_stage_var.get() or "").strip() or SURVEY_RUN_STAGE_IDS[0]
        try:
            stage_args = self._parse_run_args()
        except Exception as exc:
            messagebox.showerror("Survey Platform", str(exc))
            return

        if stage == "merge" and "--secondary" not in stage_args:
            messagebox.showwarning(
                "Survey Platform",
                "Merge runs need forwarded args including --secondary <path>.",
            )
            return

        cmd = [
            sys.executable,
            str(self._cli_path()),
            "survey",
            "run",
            "--root",
            str(self._repo_root()),
            survey_id,
            stage,
        ]
        if stage_args:
            cmd.append("--")
            cmd.extend(stage_args)

        self._run_cli_command(
            cmd,
            title=f"{survey_id} Run {stage}",
            running_text=f"Running {survey_id} stage={stage}...",
            success_callback=self.load_index,
        )

    def run_guided_wincross(self) -> None:
        docx_path = self.wincross_docx_var.get().strip()
        raw_input = self.wincross_raw_input_var.get().strip()
        pdf_path = self.wincross_pdf_var.get().strip()
        if not docx_path and not raw_input:
            messagebox.showwarning(
                "Survey Platform",
                "Choose a survey DOCX or a pre-numbered RAW input first.",
            )
            return

        cmd = [
            sys.executable,
            str(self._cli_path()),
            "pipeline",
            "--root",
            str(BASE_DIR),
        ]
        if docx_path:
            cmd.extend(["--docx", docx_path])
        if pdf_path:
            cmd.extend(["--pdf", pdf_path])
        if raw_input:
            cmd.extend(["--raw-input", raw_input])
        survey_id = (
            self.main_survey_id_var.get().strip()
            or self.selected_survey_id_var.get().strip()
            or self._survey_id_from_text(docx_path or raw_input)
        )
        raw_out = ""
        if survey_id:
            raw_out = self._default_wincross_raw_path(survey_id)
            if raw_out:
                cmd.extend(["--raw-out", raw_out])
        json_out = self.wincross_json_out_var.get().strip()
        if not json_out and survey_id:
            json_out = self._default_wincross_json_workspace_path(survey_id)
        if json_out:
            cmd.extend(["--json-out", json_out])
        tables_out = self.wincross_tables_out_var.get().strip()
        if not tables_out and survey_id:
            tables_out = self._default_wincross_tables_workspace_path(survey_id)
        if tables_out:
            cmd.extend(["--tables-out", tables_out])
        if self.wincross_translate_var.get():
            cmd.append("--translate-spanish")
        if self.wincross_skip_tables_var.get():
            cmd.append("--skip-tables")
        cmd.append("--no-repo-sync")

        ok = self._run_cli_command(
            cmd,
            title="WinCross Pipeline",
            running_text="Running WinCross pipeline...",
            success_callback=self.load_index,
            status_target=self.wincross_status_var,
        )
        if ok:
            if survey_id:
                for role_id, raw_path in (
                    ("wincross_raw", raw_out),
                    ("wincross_json", json_out),
                    ("wincross_tables", tables_out),
                ):
                    if raw_path:
                        path = self._resolve_gui_path(raw_path)
                        if path.exists() and path.is_file():
                            self._record_workspace_file(survey_id, role_id, path, copied=False)
            self._mark_workflow_step("wincross_generated", "done", phase="Pre-fielding")
            self.wincross_status_var.set("WinCross pipeline complete.")
        else:
            self.wincross_status_var.set("WinCross pipeline failed.")

    def run_guided_merge_with_saved_map(self) -> None:
        self.run_guided_merge(require_saved_map=True)

    def run_guided_merge(self, *, require_saved_map: bool = False) -> None:
        survey_id = (
            self.merge_primary_survey_var.get().strip()
            or self.main_survey_id_var.get().strip()
            or self.selected_survey_id_var.get().strip()
        )
        if not survey_id:
            messagebox.showwarning("Survey Platform", "Choose a primary survey first.")
            return

        primary = self.merge_primary_csv_var.get().strip()
        if not primary:
            primary = self._detect_primary_csv_for_survey(survey_id)
            if primary:
                self.merge_primary_csv_var.set(primary)
        secondary = self.merge_secondary_path_var.get().strip()

        saved_map_path, saved_map_note, saved_record = self._saved_guided_merge_setup(
            survey_id,
            primary_raw=primary,
            secondary_raw=secondary,
        )
        if saved_record:
            primary = str(saved_record.get("primary_path", "") or primary).strip()
            secondary = str(saved_record.get("secondary_path", "") or secondary).strip()
            if primary:
                self.merge_primary_csv_var.set(primary)
            if secondary:
                self.merge_secondary_path_var.set(secondary)
        if require_saved_map and saved_map_path is None:
            messagebox.showinfo(
                "Survey Platform",
                (
                    "No reusable merge map is ready for this primary/secondary pair.\n\n"
                    "Open the Advanced Merge Editor, review or run the first merge, then close it. "
                    "Survey Platform will save the map for quick daily runs when a new merge run is found."
                ),
            )
            self.merge_status_var.set(saved_map_note or "Open Advanced Merge Editor to create or select a merge map.")
            return
        if not secondary:
            messagebox.showwarning("Survey Platform", "Choose a secondary file first.")
            return
        if primary:
            primary_path = self._resolve_gui_path(primary)
            if not primary_path.exists() or not primary_path.is_file():
                messagebox.showerror("Survey Platform", f"Primary merge file not found:\n{primary_path}")
                return
        secondary_path = self._resolve_gui_path(secondary)
        if not secondary_path.exists() or not secondary_path.is_file():
            messagebox.showerror("Survey Platform", f"Secondary merge file not found:\n{secondary_path}")
            return

        threshold_raw = self.merge_threshold_var.get().strip() or "85"
        try:
            int(threshold_raw)
        except ValueError:
            messagebox.showerror("Survey Platform", f"Invalid threshold: {threshold_raw}")
            return

        primary_json = str(saved_record.get("primary_survey_json", "") if saved_record else "").strip()
        if not primary_json:
            primary_json = self._merge_primary_json_for_survey(survey_id)
        secondary_json = str(saved_record.get("secondary_survey_json", "") if saved_record else "").strip()
        if not secondary_json:
            secondary_json = self._merge_secondary_json_for_path(secondary)
        mapping_mode = str(self.merge_mapping_mode_var.get() or "auto")
        if saved_map_path is not None:
            mapping_mode = "manual_only"

        cmd = [
            sys.executable,
            str(self._cli_path()),
            "survey",
            "run",
            "--root",
            str(self._repo_root()),
            survey_id,
            "merge",
            "--",
            "--secondary",
            secondary,
        ]
        if primary:
            cmd.extend(["--primary", primary])
        if primary_json:
            cmd.extend(["--primary-survey-json", primary_json])
        if secondary_json:
            cmd.extend(["--secondary-survey-json", secondary_json])
        cmd.extend(
            [
                "--mapping-mode",
                mapping_mode,
                "--threshold",
                threshold_raw,
                "--row-order",
                "secondary_first",
                "--question-similarity-policy",
                str(self.merge_similarity_policy_var.get() or "drop"),
            ]
        )
        if saved_map_path is not None:
            cmd.extend(["--map-file", str(saved_map_path)])

        output_mode = str(self.merge_output_mode_var.get() or "new_file").strip() or "new_file"
        output_path = ""
        if output_mode == "overwrite":
            cmd.extend(["--output-mode", "overwrite"])
        else:
            cmd.extend(["--output-mode", "new_file"])
            output_path = self.merge_output_path_var.get().strip() or self._default_merge_output_path(survey_id, secondary)
            if output_path:
                cmd.extend(["--output", output_path])

        def finish_guided_merge(ok: bool, command_output: str = "") -> None:
            if ok:
                saved_map_after_run: Path | None = None
                map_warning = ""
                try:
                    saved_map_after_run = self._save_guided_merge_map_from_run(
                        survey_id,
                        primary_raw=primary,
                        secondary_raw=secondary,
                        output_raw=output_path,
                        command_output=command_output,
                        primary_json_raw=primary_json,
                        secondary_json_raw=secondary_json,
                    )
                except Exception as exc:
                    map_warning = f" Merge map save warning: {exc}"
                self._record_merge_output(survey_id, self._resolve_gui_path(output_path) if output_path else None)
                map_note = ""
                if saved_map_path is not None:
                    map_note = " Reused saved map."
                if saved_map_after_run is not None:
                    map_note += f" Saved map: {saved_map_after_run}."
                elif map_warning:
                    map_note += map_warning
                self.merge_status_var.set(f"Guided merge finished for {survey_id}.{map_note}")
            else:
                self.merge_status_var.set(f"Guided merge failed for {survey_id}.")

        running_text = f"Running guided merge for {survey_id}..."
        if saved_map_path is not None:
            running_text = f"Running guided merge for {survey_id} with saved map..."
        elif saved_map_note:
            self.merge_status_var.set(saved_map_note)
        self.merge_status_var.set(running_text)
        self._run_cli_command_async(
            cmd,
            title=f"{survey_id} Guided Merge",
            running_text=running_text,
            success_callback=self.load_index,
            status_target=self.merge_status_var,
            on_complete=finish_guided_merge,
        )

    def run_oe_extract(self) -> None:
        input_path = self.oe_input_path_var.get().strip()
        question_col = self.oe_question_col_var.get().strip()
        if not input_path:
            messagebox.showwarning("Survey Platform", "Choose a survey CSV/XLSX input first.")
            return
        if not question_col:
            messagebox.showwarning("Survey Platform", "Choose an OE question column first.")
            return

        cmd = [
            sys.executable,
            str(BASE_DIR / "pipelines" / "open_end_coder" / "oe_codebook_tool.py"),
            "extract",
            "--input",
            input_path,
            "--question-col",
            question_col,
        ]
        defaults_path = self.oe_defaults_path_var.get().strip()
        if defaults_path:
            cmd.extend(["--defaults", defaults_path])
        existing_path = self.oe_existing_path_var.get().strip()
        if existing_path:
            cmd.extend(["--existing", existing_path])
        survey_id = self.main_survey_id_var.get().strip() or self.selected_survey_id_var.get().strip()
        extract_out = self.oe_extract_out_var.get().strip() or self._default_oe_output_path(survey_id, question_col)
        if extract_out:
            cmd.extend(["--out", extract_out])
        if self.oe_include_blank_var.get():
            cmd.append("--include-blank")

        ok = self._run_cli_command(
            cmd,
            title="OE Codebook Extract",
            running_text="Extracting OE codebook...",
            status_target=self.oe_status_var,
        )
        if ok:
            if extract_out:
                self.oe_render_codebook_var.set(extract_out)
                if not self.oe_render_out_var.get().strip():
                    self.oe_render_out_var.set(self._suggest_oe_render_out(extract_out, question_col))
                path = self._resolve_gui_path(extract_out)
                if path.exists() and path.is_file():
                    self._record_workspace_file(survey_id, "oe_codebook", path, copied=False)
            self._mark_workflow_step("oe_codebook_generated", "done", phase="In-field Daily")
            self.oe_status_var.set("OE codebook extraction complete.")
        else:
            self.oe_status_var.set("OE codebook extraction failed.")

    def run_oe_render(self) -> None:
        codebook_path = self.oe_render_codebook_var.get().strip()
        selector = self.oe_render_selector_var.get().strip()
        if not codebook_path:
            messagebox.showwarning("Survey Platform", "Choose a codebook file to render first.")
            return
        if not selector:
            messagebox.showwarning("Survey Platform", "Enter an OE selector first.")
            return

        cmd = [
            sys.executable,
            str(BASE_DIR / "pipelines" / "open_end_coder" / "oe_codebook_tool.py"),
            "render",
            "--codebook",
            codebook_path,
            "--selector",
            selector,
            "--title",
            self.oe_render_title_var.get().strip() or "ABLE TO MENTION",
            "--item-label",
            self.oe_render_item_label_var.get().strip() or "short",
        ]
        survey_id = self.main_survey_id_var.get().strip() or self.selected_survey_id_var.get().strip()
        render_out = self.oe_render_out_var.get().strip() or self._default_oe_render_output_path(survey_id, selector)
        if render_out:
            cmd.extend(["--out", render_out])

        ok = self._run_cli_command(
            cmd,
            title="OE WinCross Render",
            running_text="Rendering OE WinCross block...",
            status_target=self.oe_status_var,
        )
        if ok:
            if render_out:
                path = self._resolve_gui_path(render_out)
                if path.exists() and path.is_file():
                    self._record_workspace_file(survey_id, "oe_wincross_block", path, copied=False)
            self._mark_workflow_step("oe_wincross_rendered", "done", phase="In-field Daily")
            self.oe_status_var.set("OE WinCross block render complete.")
        else:
            self.oe_status_var.set("OE WinCross block render failed.")

    def run_empty_results(self) -> None:
        docx_path = self.results_docx_var.get().strip()
        if not docx_path:
            messagebox.showwarning("Survey Platform", "Choose a results DOCX template first.")
            return

        cmd = [
            sys.executable,
            str(BASE_DIR / "pipelines" / "auto_results" / "empty_results_generator.py"),
            "--docx",
            docx_path,
        ]
        tables_path = self.results_tables_var.get().strip()
        if tables_path:
            cmd.extend(["--tables", tables_path])
        json_path = self.results_json_var.get().strip()
        if json_path:
            cmd.extend(["--json", json_path])
        survey_id = self.main_survey_id_var.get().strip() or self.selected_survey_id_var.get().strip()
        output_path = self.results_output_var.get().strip() or self._default_results_output_path(survey_id, docx_path)
        if output_path:
            cmd.extend(["--output", output_path])
        if not self.results_insert_missing_var.get():
            cmd.append("--no-insert-missing")

        ok = self._run_cli_command(
            cmd,
            title="Empty Results Generation",
            running_text="Generating empty results DOCX...",
            status_target=self.results_status_var,
        )
        if ok:
            if output_path:
                path = self._resolve_gui_path(output_path)
                if path.exists() and path.is_file():
                    self._record_workspace_file(survey_id, "results_output_docx", path, copied=False)
            self._mark_workflow_step("empty_results_generated", "done", phase="Final Results")
            self.results_status_var.set("Empty results generation complete.")
        else:
            self.results_status_var.set("Empty results generation failed.")

    def rebuild_question_index(self) -> None:
        cli_path = BASE_DIR / "survey_tool_cli.py"
        if not cli_path.exists():
            messagebox.showerror("Question Browser", f"CLI not found:\n{cli_path}")
            return

        try:
            (
                _query_text,
                _match_mode,
                _state,
                _region,
                _client,
                _q_type,
                _q_family,
                _q_group,
                _q_target,
                survey_codes,
                _limit,
            ) = self._questions_filters()
        except Exception:
            survey_codes = None

        cmd = [
            sys.executable,
            str(cli_path),
            "repo",
            "questions-rebuild",
            "--root",
            str(self._repo_root()),
            "--db",
            str(self._questions_db_path()),
        ]
        if survey_codes:
            for code in sorted(survey_codes, key=int):
                cmd.extend(["--survey", code])

        self.question_status_var.set("Rebuilding question index...")
        self.root.update_idletasks()
        result = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            cwd=str(BASE_DIR),
        )
        if result.returncode != 0:
            detail = (result.stderr or result.stdout or "").strip()
            messagebox.showerror("Question Browser", f"Question index rebuild failed:\n{detail}")
            self.question_status_var.set("Question index rebuild failed.")
            return

        stdout = (result.stdout or "").strip()
        last_line = stdout.splitlines()[-1] if stdout else "Question index rebuild complete."
        self.question_status_var.set(last_line)
        self._mark_workflow_step("question_index_rebuilt", "done", phase="Archive/Search")
        self.refresh_question_summary()

    def run_question_audit(self) -> None:
        db_path = self._questions_db_path()
        if not db_path.exists():
            messagebox.showerror(
                "Question Browser",
                f"Question index not found:\n{db_path}\n\nClick Rebuild Index first.",
            )
            return

        try:
            (
                query_text,
                match_mode,
                state,
                region,
                client,
                q_type_filter,
                q_family_filter,
                q_group_filter,
                q_target_filter,
                survey_codes,
                limit,
            ) = self._questions_filters()
        except Exception as exc:
            messagebox.showerror("Question Browser", str(exc))
            return

        sample_limit = max(1, min(int(limit), 500))

        try:
            conn = self._questions_db_open()
        except Exception as exc:
            messagebox.showerror("Question Browser", f"Failed to open question index:\n{exc}")
            return

        audit_rows: list[dict[str, object]] = []
        scoped_question_count = 0
        attrs_enabled = False

        try:
            if not self._questions_table_exists(conn, "repo_questions"):
                raise ValueError("Question DB schema missing `repo_questions`. Rebuild index first.")

            attrs_enabled = self._questions_attribute_columns_available(conn)
            fts_enabled = self._questions_table_exists(conn, "repo_questions_fts")
            filter_sql, filter_params = self._question_filter_sql(
                survey_codes=survey_codes,
                region=region,
                state=state,
                client=client,
                q_type_filter=q_type_filter,
                q_family_filter=q_family_filter,
                q_group_filter=q_group_filter,
                q_target_filter=q_target_filter,
                attrs_enabled=attrs_enabled,
            )

            search_sql, search_params = _question_text_match_sql(
                query_text,
                match_mode,
                attrs_enabled=attrs_enabled,
            )

            scope_where_sql = "WHERE 1=1" + filter_sql + search_sql
            scope_params = [*filter_params, *search_params]

            def count_in_scope(condition_sql: str = "", condition_params: list[object] | None = None) -> int:
                sql = "SELECT COUNT(*) AS c FROM repo_questions q " + scope_where_sql
                params = list(scope_params)
                if condition_sql:
                    sql += " AND (" + condition_sql + ")"
                    if condition_params:
                        params.extend(condition_params)
                row = conn.execute(sql, params).fetchone()
                if row is None:
                    return 0
                return int(row["c"] if isinstance(row, sqlite3.Row) else row[0])

            def add_audit_row(
                check: str,
                count: int,
                title: str,
                summary: str = "",
                details: list[str] | None = None,
            ) -> None:
                clean_details = [str(line) for line in (details or []) if str(line).strip()]
                audit_rows.append(
                    {
                        "mode": "audit",
                        "survey_id": "",
                        "qid": "",
                        "question_type": "",
                        "question_family": "",
                        "question_group": "",
                        "question_target": "",
                        "target_source": "",
                        "state_code": "",
                        "region": "",
                        "client": "",
                        "survey_count": "",
                        "occurrence_count": max(int(count), 0),
                        "canonical_key": "",
                        "audit_check": check,
                        "question_text": title,
                        "item_text": summary,
                        "option_texts": [],
                        "survey_ids": [],
                        "audit_details": clean_details,
                    }
                )

            scoped_question_count = count_in_scope()

            table_counts: dict[str, str] = {}
            for table_name in ("repo_questions_surveys", "repo_questions", "repo_questions_fts"):
                if self._questions_table_exists(conn, table_name):
                    val = int(conn.execute(f"SELECT COUNT(*) AS c FROM {table_name}").fetchone()["c"])
                    table_counts[table_name] = str(val)
                else:
                    table_counts[table_name] = "(missing)"
            row_count_details = [
                f"repo_questions_surveys: {table_counts['repo_questions_surveys']}",
                f"repo_questions: {table_counts['repo_questions']}",
                f"repo_questions_fts: {table_counts['repo_questions_fts']}",
                f"scope_rows: {scoped_question_count}",
            ]
            if query_text:
                row_count_details.append(f"query_filter: {query_text}")
            if survey_codes:
                row_count_details.append(f"survey_filter_count: {len(survey_codes)}")
            if state:
                row_count_details.append(f"state_filter: {state}")
            if region:
                row_count_details.append(f"region_filter: {region}")
            if client:
                row_count_details.append(f"client_filter: {client}")
            if q_type_filter:
                row_count_details.append(f"type_filter: {q_type_filter}")
            if q_family_filter:
                row_count_details.append(f"family_filter: {q_family_filter}")
            if q_group_filter:
                row_count_details.append(f"group_filter: {q_group_filter}")
            if q_target_filter:
                row_count_details.append(f"target_filter: {q_target_filter}")
            add_audit_row(
                "row_counts",
                scoped_question_count,
                "Row counts and scope",
                (
                    f"questions={table_counts['repo_questions']}, surveys={table_counts['repo_questions_surveys']}, "
                    f"fts={table_counts['repo_questions_fts']}"
                ),
                row_count_details,
            )

            type_bucket_count = int(
                conn.execute(
                    "SELECT COUNT(DISTINCT q.question_type) AS c FROM repo_questions q " + scope_where_sql,
                    scope_params,
                ).fetchone()["c"]
            )
            type_rows = list(
                conn.execute(
                    """
                    SELECT q.question_type, COUNT(*) AS c
                    FROM repo_questions q
                    """
                    + scope_where_sql
                    + """
                    GROUP BY q.question_type
                    ORDER BY c DESC, q.question_type ASC
                    LIMIT ?
                    """,
                    [*scope_params, sample_limit],
                ).fetchall()
            )
            type_details: list[str] = [f"type_buckets: {type_bucket_count}"]
            for row in type_rows:
                q_type = str(row["question_type"] or "").strip() or "(blank)"
                type_details.append(f"{q_type}: {int(row['c'])}")
            type_summary = ", ".join(type_details[1:4]) if len(type_details) > 1 else "No question_type values."
            add_audit_row("question_type_counts", type_bucket_count, "Question type buckets", type_summary, type_details)

            family_bucket_count = int(
                conn.execute(
                    "SELECT COUNT(DISTINCT q.question_family) AS c FROM repo_questions q " + scope_where_sql,
                    scope_params,
                ).fetchone()["c"]
            )
            family_rows = list(
                conn.execute(
                    """
                    SELECT q.question_family, COUNT(*) AS c
                    FROM repo_questions q
                    """
                    + scope_where_sql
                    + """
                    GROUP BY q.question_family
                    ORDER BY c DESC, q.question_family ASC
                    LIMIT ?
                    """,
                    [*scope_params, sample_limit],
                ).fetchall()
            )
            family_details: list[str] = [f"family_buckets: {family_bucket_count}"]
            for row in family_rows:
                q_family = str(row["question_family"] or "").strip() or "(blank)"
                family_details.append(f"{q_family}: {int(row['c'])}")
            family_summary = ", ".join(family_details[1:4]) if len(family_details) > 1 else "No question_family values."
            add_audit_row(
                "question_family_counts",
                family_bucket_count,
                "Question family buckets",
                family_summary,
                family_details,
            )

            unknown_total = count_in_scope("q.question_type = 'unknown'")
            unknown_rows = list(
                conn.execute(
                    """
                    SELECT q.survey_id, COUNT(*) AS c
                    FROM repo_questions q
                    """
                    + scope_where_sql
                    + """
                    AND q.question_type = 'unknown'
                    GROUP BY q.survey_id
                    ORDER BY c DESC, q.survey_code DESC
                    LIMIT ?
                    """,
                    [*scope_params, sample_limit],
                ).fetchall()
            )
            unknown_details = [f"unknown_total_rows: {unknown_total}"]
            for row in unknown_rows:
                unknown_details.append(f"{row['survey_id']}: {int(row['c'])}")
            if not unknown_rows:
                unknown_details.append("No unknown question_type rows in scope.")
            unknown_summary = ", ".join(unknown_details[1:4]) if len(unknown_details) > 1 else ""
            add_audit_row(
                "unknown_by_survey",
                unknown_total,
                "Unknown question_type rows by survey",
                unknown_summary,
                unknown_details,
            )

            blank_checks = {
                "blank_question_type": "q.question_type IS NULL OR TRIM(q.question_type) = ''",
                "blank_question_text": "q.question_text IS NULL OR TRIM(q.question_text) = ''",
                "blank_canonical_key": "q.canonical_key IS NULL OR TRIM(q.canonical_key) = ''",
                "blank_qid": "q.qid IS NULL OR TRIM(q.qid) = ''",
            }
            blank_details: list[str] = []
            blank_total = 0
            for name, sql_condition in blank_checks.items():
                n = count_in_scope(sql_condition)
                blank_total += n
                blank_details.append(f"{name}: {n}")
            blank_sample_rows = list(
                conn.execute(
                    """
                    SELECT q.survey_id, q.qid, q.question_type, q.canonical_key
                    FROM repo_questions q
                    """
                    + scope_where_sql
                    + """
                    AND (
                        q.question_type IS NULL OR TRIM(q.question_type) = ''
                        OR q.question_text IS NULL OR TRIM(q.question_text) = ''
                        OR q.canonical_key IS NULL OR TRIM(q.canonical_key) = ''
                        OR q.qid IS NULL OR TRIM(q.qid) = ''
                    )
                    ORDER BY q.survey_code DESC, q.question_ord ASC
                    LIMIT ?
                    """,
                    [*scope_params, sample_limit],
                ).fetchall()
            )
            for row in blank_sample_rows:
                blank_details.append(
                    f"{row['survey_id']} {row['qid'] or '(blank qid)'} "
                    f"type={str(row['question_type'] or '(blank)').strip()} "
                    f"key={str(row['canonical_key'] or '(blank)').strip()}"
                )
            blank_summary = ", ".join(blank_details[:3]) if blank_details else ""
            add_audit_row("blank_required_fields", blank_total, "Blank required fields", blank_summary, blank_details)

            if attrs_enabled:
                suspicious_condition = (
                    "TRIM(q.question_target) <> '' "
                    "AND (INSTR(q.question_target, '[') > 0 "
                    "OR INSTR(q.question_target, ']') > 0 "
                    "OR LENGTH(q.question_target) > 80)"
                )
                suspicious_total = count_in_scope(suspicious_condition)
                suspicious_rows = list(
                    conn.execute(
                        """
                        SELECT q.survey_id, q.qid, q.question_type, q.question_target
                        FROM repo_questions q
                        """
                        + scope_where_sql
                        + """
                        AND (
                        """
                        + suspicious_condition
                        + """
                        )
                        ORDER BY q.survey_code DESC, q.question_ord ASC
                        LIMIT ?
                        """,
                        [*scope_params, sample_limit],
                    ).fetchall()
                )
                suspicious_details = [f"suspicious_rows: {suspicious_total}"]
                for row in suspicious_rows:
                    suspicious_details.append(
                        f"{row['survey_id']} {row['qid']} "
                        f"{str(row['question_type'] or '').strip() or 'unknown'} -> "
                        f"{str(row['question_target'] or '').strip()}"
                    )
                if not suspicious_rows:
                    suspicious_details.append("No suspicious target strings in scope.")
                suspicious_summary = ", ".join(suspicious_details[1:3]) if len(suspicious_details) > 1 else ""
                add_audit_row(
                    "suspicious_targets",
                    suspicious_total,
                    "Suspicious question_target strings",
                    suspicious_summary,
                    suspicious_details,
                )
            else:
                add_audit_row(
                    "suspicious_targets",
                    0,
                    "Suspicious question_target strings",
                    "Skipped (attribute columns missing)",
                    ["Question attribute columns are missing. Run Rebuild Index."],
                )

            duplicate_group_count = int(
                conn.execute(
                    """
                    SELECT COUNT(*) AS c
                    FROM (
                        SELECT 1
                        FROM repo_questions q
                    """
                    + scope_where_sql
                    + """
                        GROUP BY q.survey_id, q.qid
                        HAVING COUNT(*) > 1
                    ) d
                    """,
                    scope_params,
                ).fetchone()["c"]
            )
            duplicate_row_total = int(
                conn.execute(
                    """
                    SELECT COALESCE(SUM(group_count), 0) AS c
                    FROM (
                        SELECT COUNT(*) AS group_count
                        FROM repo_questions q
                    """
                    + scope_where_sql
                    + """
                        GROUP BY q.survey_id, q.qid
                        HAVING COUNT(*) > 1
                    ) d
                    """,
                    scope_params,
                ).fetchone()["c"]
            )
            duplicate_rows = list(
                conn.execute(
                    """
                    SELECT q.survey_id, q.qid, COUNT(*) AS c
                    FROM repo_questions q
                    """
                    + scope_where_sql
                    + """
                    GROUP BY q.survey_id, q.qid
                    HAVING COUNT(*) > 1
                    ORDER BY c DESC, q.survey_id ASC, q.qid ASC
                    LIMIT ?
                    """,
                    [*scope_params, sample_limit],
                ).fetchall()
            )
            duplicate_details = [
                f"duplicate_groups: {duplicate_group_count}",
                f"rows_in_duplicate_groups: {duplicate_row_total}",
            ]
            for row in duplicate_rows:
                duplicate_details.append(f"{row['survey_id']} {row['qid']} x{int(row['c'])}")
            if not duplicate_rows:
                duplicate_details.append("No duplicate (survey_id, qid) groups in scope.")
            duplicate_summary = ", ".join(duplicate_details[2:5]) if len(duplicate_details) > 2 else ""
            add_audit_row(
                "duplicate_survey_qid",
                duplicate_group_count,
                "Duplicate (survey_id, qid) groups",
                duplicate_summary,
                duplicate_details,
            )

            if attrs_enabled:
                type_conflict_rows = list(
                    conn.execute(
                        """
                        SELECT
                            q.canonical_key,
                            COUNT(DISTINCT q.question_type) AS type_count,
                            COUNT(*) AS row_count,
                            GROUP_CONCAT(DISTINCT q.question_type) AS types
                        FROM repo_questions q
                        """
                        + scope_where_sql
                        + """
                        GROUP BY q.canonical_key
                        HAVING COUNT(DISTINCT q.question_type) > 1
                        ORDER BY type_count DESC, row_count DESC, q.canonical_key ASC
                        LIMIT ?
                        """,
                        [*scope_params, sample_limit],
                    ).fetchall()
                )
                type_conflict_count = int(
                    conn.execute(
                        """
                        SELECT COUNT(*) AS c
                        FROM (
                            SELECT 1
                            FROM repo_questions q
                        """
                        + scope_where_sql
                        + """
                            GROUP BY q.canonical_key
                            HAVING COUNT(DISTINCT q.question_type) > 1
                        ) t
                        """,
                        scope_params,
                    ).fetchone()["c"]
                )
                type_conflict_details = [f"canonical_keys_with_type_conflicts: {type_conflict_count}"]
                for row in type_conflict_rows:
                    type_conflict_details.append(
                        f"{str(row['canonical_key'])[:28]}... "
                        f"types={str(row['types'] or '')} rows={int(row['row_count'])}"
                    )
                if not type_conflict_rows:
                    type_conflict_details.append("No canonical key type conflicts in scope.")
                add_audit_row(
                    "canonical_type_conflicts",
                    type_conflict_count,
                    "Canonical key type conflicts",
                    ", ".join(type_conflict_details[1:3]) if len(type_conflict_details) > 1 else "",
                    type_conflict_details,
                )

                target_conflict_rows = list(
                    conn.execute(
                        """
                        SELECT
                            q.canonical_key,
                            COUNT(DISTINCT q.question_target) AS target_count,
                            COUNT(*) AS row_count,
                            GROUP_CONCAT(DISTINCT q.question_target) AS targets
                        FROM repo_questions q
                        """
                        + scope_where_sql
                        + """
                        AND TRIM(q.question_target) <> ''
                        GROUP BY q.canonical_key
                        HAVING COUNT(DISTINCT q.question_target) > 1
                        ORDER BY target_count DESC, row_count DESC, q.canonical_key ASC
                        LIMIT ?
                        """,
                        [*scope_params, sample_limit],
                    ).fetchall()
                )
                target_conflict_count = int(
                    conn.execute(
                        """
                        SELECT COUNT(*) AS c
                        FROM (
                            SELECT 1
                            FROM repo_questions q
                        """
                        + scope_where_sql
                        + """
                            AND TRIM(q.question_target) <> ''
                            GROUP BY q.canonical_key
                            HAVING COUNT(DISTINCT q.question_target) > 1
                        ) t
                        """,
                        scope_params,
                    ).fetchone()["c"]
                )
                target_conflict_details = [f"canonical_keys_with_target_conflicts: {target_conflict_count}"]
                for row in target_conflict_rows:
                    targets = str(row["targets"] or "")
                    if len(targets) > 180:
                        targets = targets[:177] + "..."
                    target_conflict_details.append(
                        f"{str(row['canonical_key'])[:28]}... targets={targets} rows={int(row['row_count'])}"
                    )
                if not target_conflict_rows:
                    target_conflict_details.append("No canonical key target conflicts in scope.")
                add_audit_row(
                    "canonical_target_conflicts",
                    target_conflict_count,
                    "Canonical key target conflicts",
                    ", ".join(target_conflict_details[1:3]) if len(target_conflict_details) > 1 else "",
                    target_conflict_details,
                )

                source_rows = list(
                    conn.execute(
                        """
                        SELECT q.target_source, COUNT(*) AS c
                        FROM repo_questions q
                        """
                        + scope_where_sql
                        + """
                        GROUP BY q.target_source
                        ORDER BY c DESC, q.target_source ASC
                        LIMIT ?
                        """,
                        [*scope_params, sample_limit],
                    ).fetchall()
                )
                source_bucket_count = int(
                    conn.execute(
                        "SELECT COUNT(DISTINCT q.target_source) AS c FROM repo_questions q " + scope_where_sql,
                        scope_params,
                    ).fetchone()["c"]
                )
                source_details = [f"target_source_buckets: {source_bucket_count}"]
                for row in source_rows:
                    source_key = str(row["target_source"] or "").strip() or "(blank)"
                    source_details.append(f"{source_key}: {int(row['c'])}")
                add_audit_row(
                    "target_source_counts",
                    source_bucket_count,
                    "Target source buckets",
                    ", ".join(source_details[1:4]) if len(source_details) > 1 else "",
                    source_details,
                )
            else:
                add_audit_row(
                    "canonical_type_conflicts",
                    0,
                    "Canonical key type conflicts",
                    "Skipped (attribute columns missing)",
                    ["Question attribute columns are missing. Run Rebuild Index."],
                )
                add_audit_row(
                    "canonical_target_conflicts",
                    0,
                    "Canonical key target conflicts",
                    "Skipped (attribute columns missing)",
                    ["Question attribute columns are missing. Run Rebuild Index."],
                )
                add_audit_row(
                    "target_source_counts",
                    0,
                    "Target source buckets",
                    "Skipped (attribute columns missing)",
                    ["Question attribute columns are missing. Run Rebuild Index."],
                )
        except Exception as exc:
            messagebox.showerror("Question Browser", str(exc))
            return
        finally:
            conn.close()

        self._clear_question_tree()
        if self.question_tree is not None:
            for idx, row in enumerate(audit_rows, start=1):
                row_id = f"audit:{idx}"
                self.question_result_rows[row_id] = row
                self.question_tree.insert(
                    "",
                    tk.END,
                    iid=row_id,
                    values=(
                        "audit",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        int(row.get("occurrence_count", 0) or 0),
                        _shorten(str(row.get("question_text", "")), 280),
                        _shorten(str(row.get("item_text", "")), 180),
                    ),
                )

        finding_checks = {
            "unknown_by_survey",
            "blank_required_fields",
            "suspicious_targets",
            "duplicate_survey_qid",
            "canonical_type_conflicts",
            "canonical_target_conflicts",
        }
        finding_count = 0
        for row in audit_rows:
            check = str(row.get("audit_check", ""))
            count = int(row.get("occurrence_count", 0) or 0)
            if check in finding_checks and count > 0:
                finding_count += 1

        self.question_status_var.set(
            f"Audit complete: checks={len(audit_rows)}, findings={finding_count}, "
            f"scope_rows={scoped_question_count}, attrs={'on' if attrs_enabled else 'off'}, "
            f"fts={'on' if fts_enabled else 'off'}."
        )

    def run_question_search(self) -> None:
        try:
            filters = self._question_filters_snapshot()
        except Exception as exc:
            messagebox.showerror("Question Browser", str(exc))
            return
        self._run_question_search_page(filters, 1)

    def _run_question_search_page(self, filters: dict[str, object], page: int) -> None:
        db_path = self._questions_db_path()
        if not db_path.exists():
            messagebox.showerror(
                "Question Browser",
                f"Question index not found:\n{db_path}\n\nClick Rebuild Index first.",
            )
            return

        query_text = str(filters.get("query_text", "") or "")
        match_mode = _normalize_question_match_mode(filters.get("match_mode", "all_words"))
        state = str(filters.get("state", "") or "")
        region = str(filters.get("region", "") or "")
        client = str(filters.get("client", "") or "")
        q_type_filter = str(filters.get("q_type_filter", "") or "")
        q_family_filter = str(filters.get("q_family_filter", "") or "")
        q_group_filter = str(filters.get("q_group_filter", "") or "")
        q_target_filter = str(filters.get("q_target_filter", "") or "")
        survey_codes = filters.get("survey_codes")
        if not isinstance(survey_codes, set):
            survey_codes = None
        page_size = int(filters.get("page_size", 30) or 30)

        try:
            conn = self._questions_db_open()
        except Exception as exc:
            messagebox.showerror("Question Browser", f"Failed to open question index:\n{exc}")
            return

        try:
            if not self._questions_table_exists(conn, "repo_questions"):
                raise ValueError("Question DB schema missing `repo_questions`. Rebuild index first.")
            attrs_enabled = self._questions_attribute_columns_available(conn)
            type_select = "q.question_type AS question_type" if attrs_enabled else "'' AS question_type"
            family_select = "q.question_family AS question_family" if attrs_enabled else "'' AS question_family"
            group_select = "q.question_group AS question_group" if attrs_enabled else "'' AS question_group"
            target_select = "q.question_target AS question_target" if attrs_enabled else "'' AS question_target"
            source_select = "q.target_source AS target_source" if attrs_enabled else "'' AS target_source"
            filter_sql, filter_params = self._question_filter_sql(
                survey_codes=survey_codes,
                region=region,
                state=state,
                client=client,
                q_type_filter=q_type_filter,
                q_family_filter=q_family_filter,
                q_group_filter=q_group_filter,
                q_target_filter=q_target_filter,
                attrs_enabled=attrs_enabled,
            )

            rows: list[sqlite3.Row] = []
            search_mode = match_mode if query_text else "filters_only"
            total_count = 0
            page = max(1, int(page))
            match_clause, match_params = _question_text_match_sql(
                query_text,
                match_mode,
                attrs_enabled=attrs_enabled,
            )
            count_sql = (
                """
                SELECT COUNT(*)
                FROM repo_questions q
                WHERE 1=1
                """
                + filter_sql
                + match_clause
            )
            total_count = int(conn.execute(count_sql, [*filter_params, *match_params]).fetchone()[0])
            total_pages = self._question_total_pages(total_count, page_size)
            page = min(page, total_pages) if total_pages else 1
            offset = (page - 1) * page_size

            select_sql = (
                """
                SELECT
                    q.survey_id,
                    q.qid,
                    """
                + type_select
                + """,
                    """
                + family_select
                + """,
                    """
                + group_select
                + """,
                    """
                + target_select
                + """,
                    """
                + source_select
                + """,
                    q.state_code,
                    q.region,
                    q.client,
                    q.question_text,
                    q.item_text,
                    q.option_texts_json,
                    q.canonical_key
                FROM repo_questions q
                WHERE 1=1
                """
                + filter_sql
                + match_clause
                + """
                ORDER BY q.survey_code DESC, q.question_ord ASC
                LIMIT ? OFFSET ?
                """
            )
            rows = list(
                conn.execute(select_sql, [*filter_params, *match_params, page_size, offset]).fetchall()
            )

        except Exception as exc:
            messagebox.showerror("Question Browser", str(exc))
            return
        finally:
            conn.close()

        self._populate_question_search_rows(rows)
        self._update_question_pagination_state(mode="search", filters=filters, total_count=total_count, page=page)
        self.question_status_var.set(
            f"Search complete: showing {len(rows)} result(s), total={total_count}, mode={search_mode}, page={self.question_page_var.get()}, per_page={page_size}."
        )

    def run_question_frequency(self) -> None:
        try:
            filters = self._question_filters_snapshot()
        except Exception as exc:
            messagebox.showerror("Question Browser", str(exc))
            return
        self._run_question_frequency_page(filters, 1)

    def _question_frequency_key_sql(self) -> str:
        return """
            CASE
                WHEN TRIM(COALESCE(q.normalized_item_text, '')) <> ''
                    AND LOWER(COALESCE(q.question_type, '')) = 'approve_disapprove_4pt'
                    AND TRIM(COALESCE(q.question_target, '')) <> ''
                THEN 'approval_target|'
                    || LOWER(COALESCE(q.question_type, '')) || '|'
                    || LOWER(COALESCE(q.response_style, '')) || '|'
                    || CASE
                        WHEN LOWER(COALESCE(q.normalized_item_text, '')) LIKE '% did when %'
                            OR LOWER(COALESCE(q.normalized_item_text, '')) LIKE '% did as %'
                            OR LOWER(COALESCE(q.normalized_item_text, '')) LIKE '% when he was %'
                            OR LOWER(COALESCE(q.normalized_item_text, '')) LIKE '% when she was %'
                            OR LOWER(COALESCE(q.normalized_item_text, '')) LIKE '% when they were %'
                        THEN 'past|'
                        ELSE 'current|'
                    END
                    || TRIM(
                        REPLACE(
                            REPLACE(
                                REPLACE(
                                    REPLACE(
                                        LOWER(COALESCE(q.question_target, '')),
                                        ' is doing now',
                                        ''
                                    ),
                                    ' is doing as president',
                                    ''
                                ),
                                ' when he was president',
                                ''
                            ),
                            ' when she was president',
                            ''
                        )
                    ) || '|'
                    || LOWER(COALESCE(q.option_texts_json, ''))
                WHEN TRIM(COALESCE(q.normalized_item_text, '')) = ''
                    AND LOWER(COALESCE(q.question_type, '')) = 'head_to_head_ballot'
                THEN 'ballot_question|'
                    || LOWER(COALESCE(q.question_type, '')) || '|'
                    || LOWER(COALESCE(q.response_style, '')) || '|'
                    || LOWER(COALESCE(q.normalized_question_text, '')) || '|'
                    || question_ballot_options_key(q.option_texts_json)
                WHEN TRIM(COALESCE(q.normalized_item_text, '')) <> ''
                    AND LOWER(COALESCE(q.question_type, '')) NOT IN ('', 'unknown')
                THEN 'item|'
                    || LOWER(COALESCE(q.question_type, '')) || '|'
                    || LOWER(COALESCE(q.response_style, '')) || '|'
                    || REPLACE(
                        REPLACE(
                            LOWER(COALESCE(q.normalized_item_text, '')),
                            ' is doing as ',
                            ' as '
                        ),
                        ' is doing ',
                        ' '
                    ) || '|'
                    || LOWER(COALESCE(q.option_texts_json, ''))
                ELSE 'exact|' || COALESCE(q.canonical_key, '')
            END
        """

    def _run_question_frequency_page(self, filters: dict[str, object], page: int) -> None:
        db_path = self._questions_db_path()
        if not db_path.exists():
            messagebox.showerror(
                "Question Browser",
                f"Question index not found:\n{db_path}\n\nClick Rebuild Index first.",
            )
            return

        query_text = str(filters.get("query_text", "") or "")
        match_mode = _normalize_question_match_mode(filters.get("match_mode", "all_words"))
        state = str(filters.get("state", "") or "")
        region = str(filters.get("region", "") or "")
        client = str(filters.get("client", "") or "")
        q_type_filter = str(filters.get("q_type_filter", "") or "")
        q_family_filter = str(filters.get("q_family_filter", "") or "")
        q_group_filter = str(filters.get("q_group_filter", "") or "")
        q_target_filter = str(filters.get("q_target_filter", "") or "")
        survey_codes = filters.get("survey_codes")
        if not isinstance(survey_codes, set):
            survey_codes = None
        page_size = int(filters.get("page_size", 30) or 30)

        try:
            conn = self._questions_db_open()
        except Exception as exc:
            messagebox.showerror("Question Browser", f"Failed to open question index:\n{exc}")
            return

        try:
            if not self._questions_table_exists(conn, "repo_questions"):
                raise ValueError("Question DB schema missing `repo_questions`. Rebuild index first.")
            attrs_enabled = self._questions_attribute_columns_available(conn)
            type_select = "MIN(q.question_type) AS sample_question_type" if attrs_enabled else "'' AS sample_question_type"
            family_select = (
                "MIN(q.question_family) AS sample_question_family"
                if attrs_enabled
                else "'' AS sample_question_family"
            )
            group_select = "MIN(q.question_group) AS sample_question_group" if attrs_enabled else "'' AS sample_question_group"
            target_select = "MIN(q.question_target) AS sample_question_target" if attrs_enabled else "'' AS sample_question_target"
            filter_sql, filter_params = self._question_filter_sql(
                survey_codes=survey_codes,
                region=region,
                state=state,
                client=client,
                q_type_filter=q_type_filter,
                q_family_filter=q_family_filter,
                q_group_filter=q_group_filter,
                q_target_filter=q_target_filter,
                attrs_enabled=attrs_enabled,
            )

            rows: list[sqlite3.Row] = []
            grouping_mode = match_mode if query_text else "filters_only"
            total_count = 0
            page = max(1, int(page))
            match_clause, match_params = _question_text_match_sql(
                query_text,
                match_mode,
                attrs_enabled=attrs_enabled,
            )
            frequency_key_sql = self._question_frequency_key_sql()
            count_sql = (
                """
                SELECT COUNT(*) FROM (
                    SELECT
                """
                + frequency_key_sql
                + """
                    AS frequency_key
                    FROM repo_questions q
                    WHERE 1=1
                """
                + filter_sql
                + match_clause
                + """
                    GROUP BY frequency_key
                ) grouped
                """
            )
            total_count = int(conn.execute(count_sql, [*filter_params, *match_params]).fetchone()[0])
            total_pages = self._question_total_pages(total_count, page_size)
            page = min(page, total_pages) if total_pages else 1
            offset = (page - 1) * page_size

            select_sql = (
                """
                SELECT
                """
                + frequency_key_sql
                + """
                    AS frequency_key,
                    MIN(q.canonical_key) AS sample_canonical_key,
                    COUNT(DISTINCT q.canonical_key) AS canonical_key_count,
                    GROUP_CONCAT(DISTINCT q.canonical_key) AS canonical_keys,
                    MIN(q.question_text) AS sample_question_text,
                    MIN(q.item_text) AS sample_item_text,
                    MIN(q.option_texts_json) AS sample_option_texts_json,
                    """
                + type_select
                + """,
                    """
                + family_select
                + """,
                    """
                + group_select
                + """,
                    """
                + target_select
                + """,
                    COUNT(*) AS occurrence_count,
                    COUNT(DISTINCT q.survey_id) AS survey_count,
                    GROUP_CONCAT(DISTINCT q.survey_id) AS survey_ids
                FROM repo_questions q
                WHERE 1=1
                """
                + filter_sql
                + match_clause
                + """
                GROUP BY frequency_key
                ORDER BY survey_count DESC, occurrence_count DESC
                LIMIT ? OFFSET ?
                """
            )
            rows = list(
                conn.execute(select_sql, [*filter_params, *match_params, page_size, offset]).fetchall()
            )

        except Exception as exc:
            messagebox.showerror("Question Browser", str(exc))
            return
        finally:
            conn.close()

        self._populate_question_frequency_rows(rows)
        self._update_question_pagination_state(mode="frequency", filters=filters, total_count=total_count, page=page)
        self.question_status_var.set(
            f"Frequency complete: showing {len(rows)} concept group(s), total={total_count}, mode={grouping_mode}, page={self.question_page_var.get()}, per_page={page_size}."
        )


def main() -> None:
    _configure_tcl_env()
    root = tk.Tk()
    app = RepositoryDashboardGUI(root)
    app.load_index()
    root.mainloop()


if __name__ == "__main__":
    main()
