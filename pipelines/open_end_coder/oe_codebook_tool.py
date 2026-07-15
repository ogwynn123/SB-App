#!/usr/bin/env python3
"""
Standalone open-ended codebook helper.

Subcommands:
1) extract: build/update a codebook from qN + cN_1..3 rows.
2) render: render a draft WinCross-style table block from a codebook.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
import xml.etree.ElementTree as ET


QCOL_RE = re.compile(
    r"^(?:(?P<prefix>\d{4,})[_-])?q(?P<num>\d{1,3})(?:[_-]?oe)?$",
    re.IGNORECASE,
)
COMPANION_RE = re.compile(
    r"^(?:(?P<prefix>\d{4,})[_-])?c(?P<num>\d{1,3})_(?P<slot>[123])$",
    re.IGNORECASE,
)
NON_ALNUM_RE = re.compile(r"[^a-z0-9 ]+")
WS_RE = re.compile(r"\s+")
SINGLE_CODE_RE = re.compile(r"^[0-9]+[a-z]?$", re.IGNORECASE)
DIRECTION_TO_WORD = {
    "positive": "positive",
    "negative": "negative",
    "neutral": "neutral",
    "+": "positive",
    "-": "negative",
    "o": "neutral",
    "0": "neutral",
}


@dataclass(frozen=True)
class QuestionSpec:
    prefix: str | None
    num: str
    question_col: str
    companion_cols: tuple[str, ...]


@dataclass(frozen=True)
class ExtractRow:
    question_col: str
    code_filter: str
    full_string: str
    short_string: str
    short_collision: str
    count: int
    default_code: str
    recode: str
    category: str
    supercategory: str
    direction: str


@dataclass(frozen=True)
class PrelimCodeRow:
    question_col: str
    question_text: str
    prelim_code: str
    original_label: str
    final_label: str
    supercategory: str
    final_code: str
    direction: str
    include: str
    sort_order: str
    notes: str


@dataclass(frozen=True)
class PrelimQuestion:
    question_col: str
    question_text: str
    rows: tuple[PrelimCodeRow, ...]


@dataclass(frozen=True)
class FinalCodelistRow:
    code: str
    label: str


@dataclass(frozen=True)
class FinalCodelistQuestion:
    question_col: str
    question_text: str
    rows: tuple[FinalCodelistRow, ...]


def safe_console(value: str) -> str:
    enc = sys.stdout.encoding or "utf-8"
    try:
        return value.encode(enc, errors="replace").decode(enc, errors="replace")
    except Exception:  # pragma: no cover
        return value


def clean_text(value: str | None) -> str:
    if value is None:
        return ""
    return WS_RE.sub(" ", str(value)).strip()


def clean_open_end_string(value: str | None) -> str:
    text = clean_text(value)
    if not text:
        return ""
    return clean_text(text.replace(",", " "))


def normalize_for_match(value: str | None) -> str:
    text = clean_text(value).casefold()
    text = NON_ALNUM_RE.sub(" ", text)
    text = WS_RE.sub(" ", text).strip()
    return text


def short12(value: str) -> str:
    return clean_text(value)[:12]


def normalize_code(value: str | None) -> str:
    text = clean_text(value)
    if not text:
        return ""
    if re.fullmatch(r"-?\d+(\.0+)?", text):
        return str(int(float(text)))
    return text.casefold()


def normalize_direction_word(value: str | None, default: str = "") -> str:
    tok = clean_text(value).casefold()
    if not tok:
        return default
    mapped = DIRECTION_TO_WORD.get(tok)
    if mapped:
        return mapped
    return default


def read_csv_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    encodings = ("utf-8-sig", "utf-8", "latin1")
    last_exc: Exception | None = None
    for enc in encodings:
        try:
            with path.open("r", encoding=enc, newline="") as f:
                reader = csv.DictReader(f)
                if reader.fieldnames is None:
                    raise ValueError(f"No header row in CSV: {path}")
                rows = [dict(row) for row in reader]
                return list(reader.fieldnames), rows
        except Exception as exc:
            last_exc = exc
    raise ValueError(f"Could not read CSV {path}: {last_exc}")


def _xlsx_col_idx(cell_ref: str) -> int:
    m = re.match(r"([A-Z]+)", cell_ref)
    if not m:
        return -1
    idx = 0
    for ch in m.group(1):
        idx = idx * 26 + (ord(ch) - 64)
    return idx - 1


def _read_xlsx_sheets_stdlib(path: Path) -> list[tuple[str, list[list[str]]]]:
    ns = {
        "a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    with zipfile.ZipFile(path, "r") as zf:
        shared: list[str] = []
        if "xl/sharedStrings.xml" in zf.namelist():
            root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
            for si in root.findall(".//a:si", ns):
                txt = "".join((t.text or "") for t in si.findall(".//a:t", ns))
                shared.append(clean_text(txt))

        wb = ET.fromstring(zf.read("xl/workbook.xml"))
        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        rid_to_target = {r.attrib["Id"]: r.attrib["Target"] for r in rels}

        out: list[tuple[str, list[list[str]]]] = []
        for sheet in wb.findall(".//a:sheets/a:sheet", ns):
            title = sheet.attrib.get("name", "")
            rid = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", "")
            target = rid_to_target.get(rid, "")
            if not target:
                continue
            if not target.startswith("xl/"):
                target = f"xl/{target}"
            if target not in zf.namelist():
                continue
            root = ET.fromstring(zf.read(target))
            rows_out: list[list[str]] = []
            for row in root.findall(".//a:sheetData/a:row", ns):
                vals: dict[int, str] = {}
                for cell in row.findall("a:c", ns):
                    ref = cell.attrib.get("r", "")
                    col_idx = _xlsx_col_idx(ref)
                    if col_idx < 0:
                        continue
                    ctype = cell.attrib.get("t", "")
                    v = cell.find("a:v", ns)
                    inline = cell.find("a:is", ns)
                    text = ""
                    if ctype == "s" and v is not None and v.text is not None:
                        try:
                            sidx = int(v.text)
                            if 0 <= sidx < len(shared):
                                text = shared[sidx]
                        except Exception:
                            text = ""
                    elif ctype == "inlineStr" and inline is not None:
                        text = "".join((t.text or "") for t in inline.findall(".//a:t", ns))
                    elif v is not None and v.text is not None:
                        text = v.text
                    text = clean_text(text)
                    if text:
                        vals[col_idx] = text
                if vals:
                    max_idx = max(vals)
                    arr = [""] * (max_idx + 1)
                    for idx, text in vals.items():
                        arr[idx] = text
                    rows_out.append(arr)
            out.append((title, rows_out))
    return out


def read_xlsx_sheets(path: Path) -> list[tuple[str, list[list[str]]]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception:
        return _read_xlsx_sheets_stdlib(path)

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        out: list[tuple[str, list[list[str]]]] = []
        for ws in wb.worksheets:
            rows_out: list[list[str]] = []
            for raw in ws.iter_rows(values_only=True):
                arr = [clean_text(v) for v in raw]
                while arr and not arr[-1]:
                    arr.pop()
                if any(arr):
                    rows_out.append(arr)
            out.append((ws.title, rows_out))
        return out
    finally:
        close = getattr(wb, "close", None)
        if callable(close):
            close()


def normalize_include_flag(value: str | None, default: str = "Y") -> str:
    tok = clean_text(value).casefold()
    if not tok:
        return default
    if tok in {"n", "no", "false", "0"}:
        return "N"
    return "Y"


def prelim_sort_key(value: str | None) -> tuple[int, str]:
    tok = clean_text(value)
    if tok.isdigit():
        return (int(tok), "")
    return (10**9, tok.casefold())


def load_prelim_codebook(path: Path) -> list[PrelimQuestion]:
    suffix = path.suffix.casefold()
    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError("Preliminary codelists must be Excel workbooks (.xlsx or .xlsm).")

    out: list[PrelimQuestion] = []
    for sheet_name, raw_rows in read_xlsx_sheets(path):
        question_col = clean_text(sheet_name)
        if not question_col or not QCOL_RE.match(question_col):
            continue
        question_text = clean_text(raw_rows[0][0] if raw_rows and raw_rows[0] else "") or question_col
        parsed_rows: list[PrelimCodeRow] = []
        next_sort = 1
        for raw in raw_rows[1:]:
            code_raw = clean_text(raw[0] if len(raw) > 0 else "")
            label_raw = clean_open_end_string(raw[1] if len(raw) > 1 else "")
            if not code_raw or not label_raw:
                continue
            try:
                prelim_code = validate_single_code(code_raw, "prelim_code")
            except Exception:
                continue
            final_label = clean_text(raw[2] if len(raw) > 2 else "") or label_raw
            supercategory = clean_text(raw[3] if len(raw) > 3 else "")
            final_code_raw = clean_text(raw[4] if len(raw) > 4 else "")
            final_code = prelim_code
            if final_code_raw:
                final_code = validate_single_code(final_code_raw, "final_code")
            direction = clean_text(raw[5] if len(raw) > 5 else "")
            include = normalize_include_flag(raw[6] if len(raw) > 6 else "")
            sort_order = clean_text(raw[7] if len(raw) > 7 else "") or str(next_sort)
            notes = clean_text(raw[8] if len(raw) > 8 else "")
            parsed_rows.append(
                PrelimCodeRow(
                    question_col=question_col,
                    question_text=question_text,
                    prelim_code=prelim_code,
                    original_label=label_raw,
                    final_label=final_label,
                    supercategory=supercategory,
                    final_code=final_code,
                    direction=direction,
                    include=include,
                    sort_order=sort_order,
                    notes=notes,
                )
            )
            next_sort += 1
        if parsed_rows:
            out.append(
                PrelimQuestion(
                    question_col=question_col,
                    question_text=question_text,
                    rows=tuple(parsed_rows),
                )
            )

    if out:
        return out
    raise ValueError("No Q-sheet preliminary codelists found in workbook.")


def _sorted_prelim_rows(rows: Iterable[PrelimCodeRow]) -> list[PrelimCodeRow]:
    return sorted(
        rows,
        key=lambda row: (
            prelim_sort_key(row.sort_order),
            code_sort_key(row.prelim_code),
            row.original_label.casefold(),
        ),
    )


def write_prelim_codebook_csv(path: Path, questions: Iterable[PrelimQuestion]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "question_col",
                "question_text",
                "prelim_code",
                "original_label",
                "final_label",
                "supercategory",
                "final_code",
                "direction",
                "include",
                "sort_order",
                "notes",
            ],
        )
        writer.writeheader()
        for question in questions:
            for row in _sorted_prelim_rows(question.rows):
                writer.writerow(
                    {
                        "question_col": question.question_col,
                        "question_text": question.question_text,
                        "prelim_code": row.prelim_code,
                        "original_label": row.original_label,
                        "final_label": row.final_label,
                        "supercategory": row.supercategory,
                        "final_code": row.final_code,
                        "direction": row.direction,
                        "include": row.include,
                        "sort_order": row.sort_order,
                        "notes": row.notes,
                    }
                )
    return path


def write_prelim_codebook_workbook(path: Path, questions: Iterable[PrelimQuestion]) -> Path:
    try:
        from openpyxl import Workbook  # type: ignore
    except Exception as exc:
        raise ValueError("Writing preliminary codelist workbooks requires openpyxl.") from exc

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    try:
        header = [
            "Prelim Code",
            "Original Label",
            "Final Label",
            "Supercategory",
            "Final Code",
            "Direction",
            "Include",
            "Sort Order",
            "Notes",
        ]
        first_sheet = True
        for question in questions:
            ws = wb.active if first_sheet else wb.create_sheet(question.question_col)
            first_sheet = False
            ws.title = question.question_col
            ws.cell(row=1, column=1, value=question.question_text or question.question_col)
            for col_idx, label in enumerate(header, start=1):
                ws.cell(row=2, column=col_idx, value=label)
            for row_idx, row in enumerate(_sorted_prelim_rows(question.rows), start=3):
                values = [
                    row.prelim_code,
                    row.original_label,
                    row.final_label,
                    row.supercategory,
                    row.final_code,
                    row.direction,
                    row.include,
                    row.sort_order,
                    row.notes,
                ]
                for col_idx, value in enumerate(values, start=1):
                    ws.cell(row=row_idx, column=col_idx, value=value)
            ws.freeze_panes = "A3"
        wb.save(path)
        return path
    finally:
        close = getattr(wb, "close", None)
        if callable(close):
            close()


def write_final_codelist_workbook(path: Path, questions: Iterable[FinalCodelistQuestion]) -> Path:
    try:
        from openpyxl import Workbook  # type: ignore
    except Exception as exc:
        raise ValueError("Writing final codelist workbooks requires openpyxl.") from exc

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    try:
        first_sheet = True
        for question in questions:
            ws = wb.active if first_sheet else wb.create_sheet(question.question_col)
            first_sheet = False
            ws.title = question.question_col
            ws.cell(row=1, column=1, value=question.question_text or question.question_col)
            for row_idx, row in enumerate(question.rows, start=3):
                code_value: object = row.code
                if clean_text(row.code).isdigit():
                    code_value = int(clean_text(row.code))
                ws.cell(row=row_idx, column=1, value=code_value)
                ws.cell(row=row_idx, column=2, value=row.label)
            ws.freeze_panes = "A3"
        wb.save(path)
        return path
    finally:
        close = getattr(wb, "close", None)
        if callable(close):
            close()


def read_tabular_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        return read_csv_rows(path)
    if suffix in {".xlsx", ".xlsm"}:
        sheets = read_xlsx_sheets(path)
        if not sheets:
            return [], []
        _, raw_rows = sheets[0]
        if not raw_rows:
            return [], []
        headers = [clean_text(v) for v in raw_rows[0]]
        rows: list[dict[str, str]] = []
        for raw in raw_rows[1:]:
            rec: dict[str, str] = {}
            for idx, h in enumerate(headers):
                if not h:
                    continue
                val = raw[idx] if idx < len(raw) else ""
                rec[h] = clean_text(val)
            rows.append(rec)
        return headers, rows
    raise ValueError(f"Unsupported input format: {path.suffix}")


def write_csv_rows(path: Path, headers: list[str], rows: Iterable[dict[str, str]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({header: clean_text(row.get(header)) for header in headers})
    return path


def write_xlsx_rows(path: Path, headers: list[str], rows: Iterable[dict[str, str]]) -> Path:
    try:
        from openpyxl import Workbook  # type: ignore
    except Exception as exc:
        raise ValueError("Writing Excel datafiles requires openpyxl.") from exc

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    try:
        ws = wb.active
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)
        for row_idx, row in enumerate(rows, start=2):
            for col_idx, header in enumerate(headers, start=1):
                ws.cell(row=row_idx, column=col_idx, value=clean_text(row.get(header)))
        ws.freeze_panes = "A2"
        wb.save(path)
        return path
    finally:
        close = getattr(wb, "close", None)
        if callable(close):
            close()


def write_tabular_rows(path: Path, headers: list[str], rows: Iterable[dict[str, str]]) -> Path:
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        return write_csv_rows(path, headers, rows)
    if suffix in {".xlsx", ".xlsm"}:
        return write_xlsx_rows(path, headers, rows)
    raise ValueError(f"Unsupported output format: {path.suffix}")


def resolve_col(headers: Iterable[str], wanted: str) -> str | None:
    want = wanted.casefold()
    for h in headers:
        if h.casefold() == want:
            return h
    return None


def parse_question_spec(headers: list[str], question_col: str) -> QuestionSpec:
    q_actual = resolve_col(headers, question_col)
    if not q_actual:
        raise ValueError(f"Question column not found: {question_col}")
    m = QCOL_RE.match(q_actual.strip())
    if not m:
        raise ValueError(
            f"Question column must look like qN / qNoe / XXXX_qN: {q_actual}"
        )
    prefix = m.group("prefix")
    num = str(int(m.group("num")))

    def slot_candidate(slot: str) -> list[str]:
        if prefix:
            return [
                f"{prefix}_c{num}_{slot}",
                f"{prefix}-c{num}_{slot}",
            ]
        return [f"c{num}_{slot}"]

    companions: list[str] = []
    for slot in ("1", "2", "3"):
        found = None
        for candidate in slot_candidate(slot):
            found = resolve_col(headers, candidate)
            if found:
                break
        if found:
            companions.append(found)
    if not companions:
        raise ValueError(f"No companion columns found for {q_actual} (expected c{num}_1..3)")

    return QuestionSpec(
        prefix=prefix,
        num=num,
        question_col=q_actual,
        companion_cols=tuple(companions),
    )


def load_defaults_map(
    path: Path | None,
    match_mode: str,
    question_col: str | None = None,
) -> dict[str, str]:
    if path is None:
        return {}
    headers, rows = read_tabular_rows(path)
    full_col = resolve_col(headers, "full_string")
    code_col = resolve_col(headers, "default_code") or resolve_col(headers, "code")
    label_col = (
        resolve_col(headers, "label")
        or resolve_col(headers, "description")
        or resolve_col(headers, "text")
        or resolve_col(headers, "name")
    )
    question_col_header = resolve_col(headers, "question_col")
    requested_question = clean_text(question_col).casefold() if question_col else ""

    def row_question_matches(row: dict[str, str]) -> bool:
        if not requested_question or not question_col_header:
            return True
        row_question = clean_text(row.get(question_col_header)).casefold()
        return not row_question or row_question == requested_question

    if full_col and code_col:
        out: dict[str, str] = {}
        for row in rows:
            if not row_question_matches(row):
                continue
            full = clean_open_end_string(row.get(full_col))
            code = clean_text(row.get(code_col))
            if not full or not code:
                continue
            key = full if match_mode == "exact" else normalize_for_match(full)
            out[key] = code
        return out

    if code_col and label_col:
        out = {}
        for row in rows:
            if not row_question_matches(row):
                continue
            code = normalize_code(row.get(code_col))
            label = clean_open_end_string(row.get(label_col))
            if not code or not label:
                continue
            key = label if match_mode == "exact" else normalize_for_match(label)
            out[key] = code
        return out

    suffix = path.suffix.casefold()
    if suffix in {".xlsx", ".xlsm"}:
        pairs = load_qsheet_code_label_pairs(path, question_col=question_col)
        out: dict[str, str] = {}
        for code, label in pairs:
            cleaned_label = clean_open_end_string(label)
            key = cleaned_label if match_mode == "exact" else normalize_for_match(cleaned_label)
            out[key] = code
        return out

    raise ValueError(
        "Defaults file must be either: "
        "(1) CSV/XLSX with full_string,default_code columns, or "
        "(2) CSV/XLSX with code,label columns, or "
        "(3) XLSX Q-sheet codelist with code in col A and label in col B."
    )


def load_qsheet_code_label_pairs(
    path: Path,
    question_col: str | None = None,
) -> list[tuple[str, str]]:
    qnum = None
    if question_col:
        m = QCOL_RE.match(clean_text(question_col))
        if m:
            qnum = str(int(m.group("num")))

    sheets = read_xlsx_sheets(path)
    by_name = {name.casefold(): (name, rows_) for name, rows_ in sheets}
    sheet_names = [name for name, _ in sheets]
    chosen_name = ""
    if qnum:
        preferred = f"q{qnum}".casefold()
        chosen_name = by_name.get(preferred, ("", []))[0]
    if not chosen_name:
        if not qnum:
            return []
        if len(sheet_names) == 1:
            chosen_name = sheet_names[0]
        else:
            joined = ", ".join(sheet_names)
            raise ValueError(
                f"Defaults workbook has no sheet for Q{qnum}. Available: {joined}"
            )

    sheet_rows = by_name.get(chosen_name.casefold(), ("", []))[1]
    out: list[tuple[str, str]] = []
    for row in sheet_rows:
        code_raw = clean_text(row[0] if len(row) > 0 else "")
        label_raw = clean_text(row[1] if len(row) > 1 else "")
        if not code_raw or not label_raw:
            continue
        try:
            code = validate_single_code(code_raw, "default_code")
        except Exception:
            # Skip prompt/header/free-text rows.
            continue
        out.append((code, label_raw))
    return out


def load_code_labels(
    path: Path | None,
    question_col: str | None = None,
) -> dict[str, str]:
    if path is None:
        return {}
    headers, rows = read_tabular_rows(path)
    code_col = (
        resolve_col(headers, "code")
        or resolve_col(headers, "default_code")
    )
    label_col = (
        resolve_col(headers, "label")
        or resolve_col(headers, "description")
        or resolve_col(headers, "text")
        or resolve_col(headers, "name")
    )
    question_col_header = resolve_col(headers, "question_col")
    requested_question = clean_text(question_col).casefold() if question_col else ""

    def row_question_matches(row: dict[str, str]) -> bool:
        if not requested_question or not question_col_header:
            return True
        row_question = clean_text(row.get(question_col_header)).casefold()
        return not row_question or row_question == requested_question

    out: dict[str, str] = {}
    if code_col and label_col:
        for row in rows:
            if not row_question_matches(row):
                continue
            code = normalize_code(row.get(code_col))
            label = clean_text(row.get(label_col))
            if not code or not label:
                continue
            out[code] = label
        return out

    suffix = path.suffix.casefold()
    if suffix in {".xlsx", ".xlsm"}:
        for code, label in load_qsheet_code_label_pairs(path, question_col=question_col):
            out[normalize_code(code)] = label
        return out

    return {}


def load_existing_codebook(path: Path | None) -> dict[tuple[str, str, str, str], dict[str, str]]:
    if path is None or not path.exists():
        return {}
    headers, rows = read_tabular_rows(path)
    full_col = resolve_col(headers, "full_string")
    if not full_col:
        return {}
    code_filter_col = resolve_col(headers, "code_filter")
    question_col = resolve_col(headers, "question_col")
    timeseries_col = resolve_col(headers, "timeseries")
    out: dict[tuple[str, str, str, str], dict[str, str]] = {}
    for row in rows:
        full = clean_open_end_string(row.get(full_col))
        if not full:
            continue
        question = clean_text(row.get(question_col)) if question_col else ""
        code_filter = clean_text(row.get(code_filter_col)) if code_filter_col else ""
        timeseries = clean_text(row.get(timeseries_col)) if timeseries_col else ""
        rec = {
            "question_col": question,
            "timeseries": timeseries,
            "default_code": clean_text(row.get(resolve_col(headers, "default_code") or "")),
            "recode": clean_text(row.get(resolve_col(headers, "recode") or "")),
            "category": clean_text(row.get(resolve_col(headers, "category") or "")),
            "supercategory": clean_text(row.get(resolve_col(headers, "supercategory") or "")),
            "direction": clean_text(row.get(resolve_col(headers, "direction") or "")),
        }
        qkey = question.casefold()
        tskey = timeseries.casefold()
        if code_filter:
            out[(qkey, code_filter.casefold(), tskey, full)] = rec
            out.setdefault((qkey, "", tskey, full), rec)
        else:
            out[(qkey, "", tskey, full)] = rec
    return out


def pick_existing_entry(
    existing: dict[tuple[str, str, str, str], dict[str, str]],
    question_col: str | None,
    full_string: str,
    code_filter: str | None,
    timeseries: str | None = None,
) -> dict[str, str]:
    qf = clean_text(question_col).casefold() if question_col else ""
    full = clean_open_end_string(full_string)
    cf = clean_text(code_filter).casefold() if code_filter else ""
    ts = clean_text(timeseries).casefold() if timeseries else ""
    if cf:
        hit = existing.get((qf, cf, ts, full))
        if hit:
            return hit
    hit = existing.get((qf, "", ts, full))
    if hit:
        return hit
    if ts:
        if cf:
            hit = existing.get((qf, cf, "", full))
            if hit:
                return hit
        hit = existing.get((qf, "", "", full))
        if hit:
            return hit
    # Legacy fallback for older codebooks that were not question-scoped.
    if cf:
        hit = existing.get(("", cf, ts, full))
        if hit:
            return hit
    if ts:
        if cf:
            hit = existing.get(("", cf, "", full))
            if hit:
                return hit
        hit = existing.get(("", "", "", full))
        if hit:
            return hit
    if cf:
        hit = existing.get(("", cf, "", full))
        if hit:
            return hit
    return existing.get(("", "", "", full), {})


def code_sort_key(code: str) -> tuple[int, str, str]:
    m = re.match(r"^(\d+)([a-z]?)$", code.casefold())
    if m:
        return (int(m.group(1)), m.group(2), code.casefold())
    return (10**9, "", code.casefold())


def parse_codes_filter(raw: str | None) -> set[str]:
    if not raw:
        return set()
    out: set[str] = set()
    for part in raw.split(","):
        tok = normalize_code(part)
        if tok:
            out.add(tok)
    return out


def extract_codebook(args: argparse.Namespace) -> int:
    input_path = Path(args.input)
    if not input_path.exists():
        raise ValueError(f"Input file not found: {input_path}")
    headers, rows = read_tabular_rows(input_path)
    qspec = parse_question_spec(headers, args.question_col)
    target_codes = parse_codes_filter(args.codes)
    defaults = load_defaults_map(
        Path(args.defaults) if args.defaults else None,
        match_mode=args.defaults_match,
        question_col=args.question_col,
    )
    existing = load_existing_codebook(Path(args.existing) if args.existing else None)
    existing_code_context = ""
    if len(target_codes) == 1:
        existing_code_context = next(iter(target_codes))

    counts: Counter[str] = Counter()
    origin_code_counts: dict[str, Counter[str]] = defaultdict(Counter)
    for row in rows:
        companions = [normalize_code(row.get(c)) for c in qspec.companion_cols]
        matched_codes: set[str] = set()
        for tok in companions:
            if not tok:
                continue
            if target_codes and tok not in target_codes:
                continue
            matched_codes.add(tok)
        if not matched_codes:
            continue
        full = clean_open_end_string(row.get(qspec.question_col))
        if not full and not args.include_blank:
            continue
        counts[full] += 1
        for tok in matched_codes:
            origin_code_counts[full][tok] += 1

    short_counts: Counter[str] = Counter(short12(full) for full in counts.keys())
    out_rows: list[ExtractRow] = []

    for full, cnt in sorted(counts.items(), key=lambda item: (-item[1], item[0].casefold())):
        short = short12(full)
        code_filter = existing_code_context
        if not code_filter:
            code_counter = origin_code_counts.get(full, Counter())
            if code_counter:
                code_filter = sorted(
                    code_counter.items(),
                    key=lambda item: (-item[1], code_sort_key(item[0])),
                )[0][0]

        default_code = code_filter
        key = full if args.defaults_match == "exact" else normalize_for_match(full)
        if key in defaults:
            default_code = defaults[key]

        prior = pick_existing_entry(existing, qspec.question_col, full, code_filter)
        if prior.get("default_code"):
            default_code = prior["default_code"]

        direction = normalize_direction_word(prior.get("direction"), default="")

        out_rows.append(
            ExtractRow(
                question_col=qspec.question_col,
                code_filter=code_filter,
                full_string=full,
                short_string=short,
                short_collision="Y" if short and short_counts[short] > 1 else "",
                count=cnt,
                default_code=default_code,
                recode=prior.get("recode", ""),
                category=prior.get("category", ""),
                supercategory=prior.get("supercategory", ""),
                direction=direction,
            )
        )

    out_path = (
        Path(args.out)
        if args.out
        else input_path.with_name(f"{input_path.stem}_{qspec.question_col}_codebook.csv")
    )
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "question_col",
                "code_filter",
                "full_string",
                "short_string",
                "short_collision",
                "count",
                "default_code",
                "recode",
                "category",
                "supercategory",
                "direction",
            ],
        )
        writer.writeheader()
        for r in out_rows:
            writer.writerow(
                {
                    "question_col": r.question_col,
                    "code_filter": r.code_filter,
                    "full_string": r.full_string,
                    "short_string": r.short_string,
                    "short_collision": r.short_collision,
                    "count": r.count,
                    "default_code": r.default_code,
                    "recode": r.recode,
                    "category": r.category,
                    "supercategory": r.supercategory,
                    "direction": r.direction,
                }
            )

    collision_count = sum(1 for r in out_rows if r.short_collision == "Y")
    print(f"Codebook rows: {len(out_rows)}")
    print(f"Question column: {qspec.question_col}")
    print(f"Output: {out_path}")
    if target_codes:
        print(f"Filtered codes: {', '.join(sorted(target_codes, key=code_sort_key))}")
    print(f"Short-string collisions flagged: {collision_count}")
    return 0


def validate_single_code(code: str, label: str) -> str:
    cleaned = clean_text(code)
    if not cleaned:
        return ""
    if "," in cleaned:
        raise ValueError(f"{label} must be exactly one code (no comma lists): {cleaned}")
    if not SINGLE_CODE_RE.match(cleaned):
        raise ValueError(f"{label} is not a valid single code token: {cleaned}")
    return cleaned


def render_wincross(args: argparse.Namespace) -> int:
    codebook_path = Path(args.codebook)
    if not codebook_path.exists():
        raise ValueError(f"Codebook file not found: {codebook_path}")

    headers, rows = read_tabular_rows(codebook_path)
    needed = ["full_string", "short_string", "default_code", "recode", "category", "direction"]
    cols: dict[str, str] = {}
    for need in needed:
        hit = resolve_col(headers, need)
        if not hit:
            raise ValueError(f"Codebook missing required column: {need}")
        cols[need] = hit

    records: list[dict[str, str]] = []
    for row in rows:
        full = clean_open_end_string(row.get(cols["full_string"]))
        short = clean_text(row.get(cols["short_string"]))
        default_code = validate_single_code(row.get(cols["default_code"]) or "", "default_code")
        recode = validate_single_code(row.get(cols["recode"]) or "", "recode")
        code = recode or default_code
        if not code:
            continue
        category = clean_text(row.get(cols["category"])) or "UNCATEGORIZED"
        direction = normalize_direction_word(row.get(cols["direction"]), default="neutral")
        label = short if args.item_label == "short" else full
        label = clean_text(label).replace("^", " ")
        records.append(
            {
                "full": full,
                "short": short,
                "code": code,
                "category": category,
                "direction": direction,
                "label": label,
            }
        )

    if not records:
        raise ValueError("No coded rows found (need default_code or recode values).")

    # preserve input order for category and item ordering
    category_order: list[str] = []
    for rec in records:
        if rec["category"] not in category_order:
            category_order.append(rec["category"])

    all_codes = sorted({r["code"] for r in records}, key=code_sort_key)
    pos_codes = sorted({r["code"] for r in records if r["direction"] == "positive"}, key=code_sort_key)
    neg_codes = sorted({r["code"] for r in records if r["direction"] == "negative"}, key=code_sort_key)
    neu_codes = sorted({r["code"] for r in records if r["direction"] == "neutral"}, key=code_sort_key)

    lines: list[str] = []
    lines.append(f"{args.title}^{args.selector} ({','.join(all_codes)})^{args.total_tail}")
    if pos_codes:
        lines.append(f"POSITIVE MENTIONS^({','.join(pos_codes)})^{args.total_tail}")
    if neg_codes:
        lines.append(f"NEGATIVE MENTIONS^({','.join(neg_codes)})^{args.total_tail}")
    if neu_codes and args.include_neutral_summary:
        lines.append(f"NEUTRAL MENTIONS^({','.join(neu_codes)})^{args.total_tail}")
    lines.append(f" ^ ^{args.total_tail}")

    by_category: dict[str, list[dict[str, str]]] = defaultdict(list)
    for rec in records:
        by_category[rec["category"]].append(rec)

    for category in category_order:
        items = by_category[category]
        unique_codes = sorted({it["code"] for it in items}, key=code_sort_key)
        lines.append(f"{category}^NET {len(unique_codes)}^{args.category_tail}")
        for it in items:
            lines.append(f"{it['label']}^({it['code']})^{args.category_tail}")
        lines.append(f" ^ ^{args.break_tail}")

    text = "\n".join(lines).rstrip() + "\n"
    if args.out:
        out_path = Path(args.out)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(text, encoding="utf-8")
        print(f"WinCross draft written: {out_path}")
    else:
        print(safe_console(text))
    print(f"Rows rendered: {len(records)}")
    print(f"Categories: {len(category_order)}")
    return 0


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description=(
            "Standalone OE codebook helper: extract q/c strings and render draft WinCross blocks."
        )
    )
    sub = p.add_subparsers(dest="cmd", required=True)

    ex = sub.add_parser("extract", help="Extract codebook rows from qN + cN_1..3 columns.")
    ex.add_argument("--input", required=True, help="Survey CSV/XLSX path")
    ex.add_argument("--question-col", required=True, help="Question column (e.g., q16, 3731_q15oe)")
    ex.add_argument(
        "--codes",
        default="98",
        help="Filter to rows where companion code matches these tokens (comma-separated). Empty => any nonblank companion.",
    )
    ex.add_argument(
        "--defaults",
        help=(
            "Optional defaults map. Supported formats: "
            "CSV/XLSX with full_string,default_code columns; or "
            "XLSX Q-sheet codelist (Q14/Q15...) with code in col A and label in col B."
        ),
    )
    ex.add_argument(
        "--defaults-match",
        choices=["exact", "normalized"],
        default="exact",
        help="How to match defaults full_string values (default: exact).",
    )
    ex.add_argument(
        "--existing",
        help="Optional existing codebook CSV/XLSX to carry forward recode/category/direction/default_code",
    )
    ex.add_argument("--out", help="Output codebook CSV path")
    ex.add_argument(
        "--include-blank",
        action=argparse.BooleanOptionalAction,
        default=False,
        help="Include blank q-text rows in extraction.",
    )

    rn = sub.add_parser("render", help="Render draft WinCross block from a codebook.")
    rn.add_argument("--codebook", required=True, help="Codebook CSV/XLSX path")
    rn.add_argument("--selector", required=True, help="WinCross selector for 'ABLE TO MENTION' line (e.g., c16_1 S3)")
    rn.add_argument("--title", default="ABLE TO MENTION", help="Top line title")
    rn.add_argument("--item-label", choices=["short", "full"], default="short")
    rn.add_argument("--total-tail", default="NS,RH", help="Tail for total/summary lines")
    rn.add_argument("--category-tail", default="NS,SK", help="Tail for category/item lines")
    rn.add_argument("--break-tail", default="NS", help="Tail for separator blank lines")
    rn.add_argument(
        "--include-neutral-summary",
        action=argparse.BooleanOptionalAction,
        default=False,
        help="Include NEUTRAL MENTIONS summary line.",
    )
    rn.add_argument("--out", help="Output TXT path (prints to stdout if omitted)")

    return p


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    if args.cmd == "extract":
        return extract_codebook(args)
    if args.cmd == "render":
        return render_wincross(args)
    raise ValueError(f"Unknown command: {args.cmd}")


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("\nCancelled.", file=sys.stderr)
        raise SystemExit(130)
