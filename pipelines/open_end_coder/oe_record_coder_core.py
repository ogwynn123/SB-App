#!/usr/bin/env python3
"""Record-level open-end coding session, import, and export logic."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import sqlite3
from typing import Iterable, Sequence


QCOL_RE = re.compile(
    r"^(?:(?P<prefix>\d{4,})[_-])?q(?P<num>\d{1,3})(?:[_-]?oe)?$",
    re.IGNORECASE,
)
COMPANION_RE = re.compile(
    r"^(?:(?P<prefix>\d{4,})[_-])?c(?P<num>\d{1,3})_(?P<slot>[123])$",
    re.IGNORECASE,
)
CODE_RE = re.compile(r"^[0-9]+[a-z]?$", re.IGNORECASE)
KEY_CANDIDATES = (
    "survnum",
    "record",
    "respondent_id",
    "respondentid",
    "respid",
    "caseid",
    "id",
)
VALID_DIRECTIONS = {"+", "-", "o"}
VALID_IMPORT_SCOPES = {"all", "new", "changed", "new_or_changed"}


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def clean_text(value: object | None) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def clean_verbatim(value: object | None) -> str:
    return clean_text(value).replace(",", " ")


def normalize_code(value: object | None) -> str:
    text = clean_text(value)
    if not text:
        return ""
    if re.fullmatch(r"-?\d+(?:\.0+)?", text):
        return str(int(float(text)))
    return text.casefold()


def normalize_record_key(value: object | None) -> str:
    text = clean_text(value)
    if re.fullmatch(r"-?\d+\.0+", text):
        return str(int(float(text)))
    return text


def validate_code(value: object | None) -> str:
    code = normalize_code(value)
    if not code or not CODE_RE.fullmatch(code):
        raise ValueError(f"Invalid code: {clean_text(value) or '(blank)'}")
    return code


def code_sort_key(code: str) -> tuple[int, str, str]:
    match = re.fullmatch(r"(\d+)([a-z]?)", normalize_code(code))
    if match:
        return int(match.group(1)), match.group(2), normalize_code(code)
    return 10**9, "", normalize_code(code)


def file_fingerprint(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


@dataclass(frozen=True)
class QuestionSpec:
    name: str
    text_column: str
    companion_columns: tuple[str, ...]
    ordinal: int


@dataclass(frozen=True)
class ImportSummary:
    dataset_id: int
    source_path: Path
    key_column: str
    row_count: int
    question_count: int
    response_count: int
    new_response_count: int
    source_changed_response_count: int
    preserved_override_count: int
    changed_override_count: int


@dataclass(frozen=True)
class QuestionInfo:
    id: int
    name: str
    text_column: str
    companion_columns: tuple[str, ...]
    ordinal: int


@dataclass(frozen=True)
class ResponseInfo:
    id: int
    question_id: int
    record_key: str
    row_number: int
    timeseries: int | None
    verbatim: str
    original_codes: tuple[str, ...]
    final_codes: tuple[str, ...]
    override_active: bool
    review_status: str
    flagged: bool
    note: str
    import_status: str


@dataclass(frozen=True)
class CodeInfo:
    question_id: int
    code: str
    label: str
    supercategory: str
    direction: str
    source: str
    active: bool


@dataclass(frozen=True)
class QuestionReview:
    question_id: int
    question_name: str
    total: int
    coded: int
    pending: int
    flagged: int
    blank_final: int


@dataclass(frozen=True)
class ExportSummary:
    output_path: Path
    row_count: int
    modified_records: int
    modified_cells: int
    overridden_responses: int


@dataclass(frozen=True)
class WinCrossExportSummary:
    output_path: Path
    question_count: int
    code_count: int
    unlabeled_codes: tuple[str, ...]


@dataclass(frozen=True)
class WinCrossCodebookEntry:
    code: str
    label: str
    supercategory: str


@dataclass(frozen=True)
class WinCrossMetadataConflict:
    code: str
    field: str
    existing: str
    imported: str


@dataclass(frozen=True)
class WinCrossCodebookImportSummary:
    source_path: Path
    question_name: str
    overwrite_existing: bool
    parsed_code_count: int
    added_code_count: int
    updated_code_count: int
    filled_label_count: int
    filled_supercategory_count: int
    unchanged_code_count: int
    conflicts: tuple[WinCrossMetadataConflict, ...]
    overwrites: tuple[WinCrossMetadataConflict, ...]


@dataclass(frozen=True)
class CodelistMetadataConflict:
    question_name: str
    code: str
    field: str
    existing: str
    imported: str


@dataclass(frozen=True)
class CodelistImportSummary:
    source_path: Path
    overwrite_existing: bool
    parsed_code_count: int
    added_code_count: int
    updated_code_count: int
    filled_label_count: int
    filled_supercategory_count: int
    filled_direction_count: int
    unchanged_code_count: int
    conflicts: tuple[CodelistMetadataConflict, ...]
    overwrites: tuple[CodelistMetadataConflict, ...]


@dataclass(frozen=True)
class RedundantCodeCleanupSummary:
    question_id: int
    question_name: str
    code: str
    timeseries: int | None
    matching_response_count: int
    changed_response_count: int
    retained_only_code_count: int


def _repair_mojibake(value: str) -> str:
    markers = ("\u00c2", "\u00c3", "\u00e2\u20ac")
    if not any(marker in value for marker in markers):
        return value
    try:
        return value.encode("cp1252").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return value


def _read_text_file(path: Path) -> str:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin1"):
        try:
            return _repair_mojibake(path.read_text(encoding=encoding))
        except UnicodeError as exc:
            last_error = exc
    raise ValueError(f"Could not read text file {path}: {last_error}")


def _wincross_records(text: str) -> list[tuple[str, str, str]]:
    records: list[tuple[str, str, str]] = []
    for raw_line in text.splitlines():
        parts = raw_line.split("^")
        if len(parts) == 3:
            records.append((clean_text(parts[0]), clean_text(parts[1]), clean_text(parts[2])))
    if any(re.fullmatch(r"NET(?:\s+\d+)?", expression, re.IGNORECASE) for _, expression, _ in records):
        return records

    first_net = re.search(r"(?:^|\s)[^^]+\^NET(?:\s+\d+)?\^[A-Za-z0-9%,]+", text, re.IGNORECASE)
    if first_net is None:
        return records
    flattened = text[first_net.start() :]
    pattern = re.compile(
        r"(?:^|\s)(?P<label>[^^]*?)\^(?P<expression>[^\^]*)\^"
        r"(?P<flags>[A-Za-z0-9%]+(?:,[A-Za-z0-9%]+)*)(?=\s|$)"
    )
    return [
        (
            clean_text(match.group("label")),
            clean_text(match.group("expression")),
            clean_text(match.group("flags")),
        )
        for match in pattern.finditer(flattened)
    ]


def parse_wincross_codebook(text: str) -> list[WinCrossCodebookEntry]:
    """Extract code labels and NET membership from a WinCross table or snippet."""
    current_supercategory = ""
    seen_net = False
    entries: list[WinCrossCodebookEntry] = []
    for label, expression, flags in _wincross_records(_repair_mojibake(text)):
        if re.fullmatch(r"NET(?:\s+\d+)?", expression, re.IGNORECASE):
            current_supercategory = label
            seen_net = True
            continue
        if not label or not expression:
            current_supercategory = ""
            continue
        flag_values = {item.strip().upper() for item in flags.split(",")}
        if "HR" in flag_values or label.casefold() in {"hide", "hidden"}:
            continue
        normalized_label = label.casefold().replace("\u2019", "'")
        ungrouped_review_label = (
            seen_net
            and "RL" in flag_values
            and normalized_label
            in {"none", "(none)", "don't know", "(don't know)", "refused", "(refused)"}
        )
        if not current_supercategory and not ungrouped_review_label:
            continue
        parentheticals = re.findall(r"\(([^()]*)\)", expression)
        if not parentheticals:
            continue
        raw_codes = [part.strip() for part in parentheticals[-1].split(",")]
        codes = [normalize_code(part) for part in raw_codes if CODE_RE.fullmatch(part)]
        for code in codes:
            entries.append(
                WinCrossCodebookEntry(
                    code=code,
                    label=label,
                    supercategory=current_supercategory,
                )
            )
    return entries


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8", "latin1"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                reader = csv.DictReader(handle)
                if not reader.fieldnames:
                    raise ValueError(f"No header row found in {path}")
                headers = [str(header) for header in reader.fieldnames]
                rows = [
                    {header: "" if row.get(header) is None else str(row.get(header)) for header in headers}
                    for row in reader
                ]
                return headers, rows
        except UnicodeError as exc:
            last_error = exc
    raise ValueError(f"Could not read CSV {path}: {last_error}")


def _read_xlsx(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception as exc:  # pragma: no cover - deployment dependency
        raise ValueError("Reading Excel files requires openpyxl.") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        first = next(iterator, None)
        if first is None:
            raise ValueError(f"No header row found in {path}")
        headers = [clean_text(value) for value in first]
        while headers and not headers[-1]:
            headers.pop()
        if not headers or any(not header for header in headers):
            raise ValueError(f"Excel header row contains blank column names: {path}")
        rows: list[dict[str, str]] = []
        for raw in iterator:
            record = {
                header: clean_text(raw[index] if index < len(raw) else "")
                for index, header in enumerate(headers)
            }
            rows.append(record)
        return headers, rows
    finally:
        workbook.close()


def read_tabular(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        return _read_csv(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(path)
    raise ValueError(f"Unsupported data format: {path.suffix}")


def resolve_column(headers: Iterable[str], wanted: str) -> str | None:
    target = wanted.casefold()
    for header in headers:
        if header.casefold() == target:
            return header
    return None


def suggest_key_column(headers: Sequence[str]) -> str | None:
    for candidate in KEY_CANDIDATES:
        hit = resolve_column(headers, candidate)
        if hit:
            return hit
    return None


def validate_record_keys(rows: Sequence[dict[str, str]], key_column: str) -> None:
    seen: dict[str, int] = {}
    blanks: list[int] = []
    duplicates: list[tuple[str, int, int]] = []
    for row_number, row in enumerate(rows, start=2):
        key = normalize_record_key(row.get(key_column))
        if not key:
            blanks.append(row_number)
            continue
        if key in seen:
            duplicates.append((key, seen[key], row_number))
        else:
            seen[key] = row_number
    if blanks:
        shown = ", ".join(str(value) for value in blanks[:8])
        raise ValueError(f"Key column {key_column} is blank on row(s): {shown}")
    if duplicates:
        shown = ", ".join(
            f"{key} (rows {first}/{second})" for key, first, second in duplicates[:8]
        )
        raise ValueError(f"Key column {key_column} is not unique: {shown}")


def detect_open_end_questions(headers: Sequence[str]) -> list[QuestionSpec]:
    lower = {header.casefold(): header for header in headers}
    companions: dict[tuple[str | None, str], dict[int, str]] = {}
    for header in headers:
        match = COMPANION_RE.fullmatch(header.strip())
        if not match:
            continue
        key = (match.group("prefix"), str(int(match.group("num"))))
        companions.setdefault(key, {})[int(match.group("slot"))] = header

    specs: list[QuestionSpec] = []
    for ordinal, header in enumerate(headers):
        match = QCOL_RE.fullmatch(header.strip())
        if not match:
            continue
        key = (match.group("prefix"), str(int(match.group("num"))))
        slot_map = companions.get(key, {})
        if not slot_map:
            continue
        actual = lower.get(header.casefold(), header)
        specs.append(
            QuestionSpec(
                name=actual,
                text_column=actual,
                companion_columns=tuple(slot_map[slot] for slot in sorted(slot_map)),
                ordinal=ordinal,
            )
        )
    return specs


def _json_codes(raw: str | None) -> tuple[str, ...]:
    try:
        values = json.loads(raw or "[]")
    except json.JSONDecodeError:
        return ()
    if not isinstance(values, list):
        return ()
    return tuple(code for code in (normalize_code(value) for value in values) if code)


class CodingSession:
    """Persistent record-level coding session backed by SQLite."""

    def __init__(
        self,
        path: Path,
        *,
        survey_id: str = "",
        project_root: Path | None = None,
    ) -> None:
        self.path = path.resolve()
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.connection = sqlite3.connect(self.path)
        self.connection.row_factory = sqlite3.Row
        self.connection.execute("PRAGMA foreign_keys = ON")
        self.connection.execute("PRAGMA journal_mode = WAL")
        self._create_schema()
        if survey_id:
            self.set_meta("survey_id", survey_id)
        if project_root is not None:
            self.set_meta("project_root", str(project_root.resolve()))

    def close(self) -> None:
        self.connection.close()

    def __enter__(self) -> "CodingSession":
        return self

    def __exit__(self, _exc_type, _exc, _tb) -> None:
        self.close()

    def _create_schema(self) -> None:
        self.connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS meta (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS datasets (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                source_path TEXT NOT NULL,
                fingerprint TEXT NOT NULL,
                imported_at TEXT NOT NULL,
                key_column TEXT NOT NULL,
                row_count INTEGER NOT NULL,
                headers_json TEXT NOT NULL,
                active INTEGER NOT NULL DEFAULT 1
            );

            CREATE TABLE IF NOT EXISTS questions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE COLLATE NOCASE,
                text_column TEXT NOT NULL,
                companion_columns_json TEXT NOT NULL,
                ordinal INTEGER NOT NULL
            );

            CREATE TABLE IF NOT EXISTS responses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                question_id INTEGER NOT NULL REFERENCES questions(id) ON DELETE CASCADE,
                record_key TEXT NOT NULL,
                row_number INTEGER NOT NULL,
                timeseries INTEGER,
                verbatim TEXT NOT NULL,
                original_codes_json TEXT NOT NULL,
                dataset_id INTEGER NOT NULL REFERENCES datasets(id),
                first_seen_dataset_id INTEGER REFERENCES datasets(id),
                last_source_change_dataset_id INTEGER REFERENCES datasets(id),
                present INTEGER NOT NULL DEFAULT 1,
                UNIQUE(question_id, record_key)
            );

            CREATE TABLE IF NOT EXISTS response_state (
                response_id INTEGER PRIMARY KEY REFERENCES responses(id) ON DELETE CASCADE,
                override_active INTEGER NOT NULL DEFAULT 0,
                review_status TEXT NOT NULL DEFAULT 'pending',
                flagged INTEGER NOT NULL DEFAULT 0,
                note TEXT NOT NULL DEFAULT '',
                updated_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS assignments (
                response_id INTEGER NOT NULL REFERENCES responses(id) ON DELETE CASCADE,
                slot INTEGER NOT NULL,
                code TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                PRIMARY KEY(response_id, slot)
            );

            CREATE TABLE IF NOT EXISTS codes (
                question_id INTEGER NOT NULL REFERENCES questions(id) ON DELETE CASCADE,
                code TEXT NOT NULL,
                label TEXT NOT NULL DEFAULT '',
                supercategory TEXT NOT NULL DEFAULT '',
                direction TEXT NOT NULL DEFAULT 'o',
                source TEXT NOT NULL DEFAULT 'imported',
                active INTEGER NOT NULL DEFAULT 1,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                PRIMARY KEY(question_id, code)
            );

            CREATE TABLE IF NOT EXISTS audit_log (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                changed_at TEXT NOT NULL,
                action TEXT NOT NULL,
                question_name TEXT NOT NULL DEFAULT '',
                record_key TEXT NOT NULL DEFAULT '',
                old_value TEXT NOT NULL DEFAULT '',
                new_value TEXT NOT NULL DEFAULT '',
                detail TEXT NOT NULL DEFAULT ''
            );

            CREATE INDEX IF NOT EXISTS idx_responses_question_present
                ON responses(question_id, present, row_number);
            CREATE INDEX IF NOT EXISTS idx_responses_record_key
                ON responses(record_key);
            CREATE INDEX IF NOT EXISTS idx_codes_question
                ON codes(question_id, active, code);
            """
        )
        response_columns = {
            str(row["name"]).casefold()
            for row in self.connection.execute("PRAGMA table_info(responses)").fetchall()
        }
        if "timeseries" not in response_columns:
            self.connection.execute("ALTER TABLE responses ADD COLUMN timeseries INTEGER")
        if "first_seen_dataset_id" not in response_columns:
            self.connection.execute(
                "ALTER TABLE responses ADD COLUMN first_seen_dataset_id INTEGER"
            )
        if "last_source_change_dataset_id" not in response_columns:
            self.connection.execute(
                "ALTER TABLE responses ADD COLUMN last_source_change_dataset_id INTEGER"
            )
        self.connection.execute(
            "UPDATE responses SET first_seen_dataset_id=dataset_id "
            "WHERE first_seen_dataset_id IS NULL"
        )
        self.connection.commit()

    def set_meta(self, key: str, value: str) -> None:
        self.connection.execute(
            "INSERT INTO meta(key, value) VALUES(?, ?) "
            "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (key, value),
        )
        self.connection.commit()

    def get_meta(self, key: str, default: str = "") -> str:
        row = self.connection.execute("SELECT value FROM meta WHERE key=?", (key,)).fetchone()
        return str(row["value"]) if row else default

    def active_dataset(self) -> sqlite3.Row | None:
        return self.connection.execute(
            "SELECT * FROM datasets WHERE active=1 ORDER BY id DESC LIMIT 1"
        ).fetchone()

    def import_data(self, source_path: Path, *, key_column: str | None = None) -> ImportSummary:
        source_path = source_path.resolve()
        if not source_path.exists():
            raise ValueError(f"Datafile not found: {source_path}")
        headers, rows = read_tabular(source_path)
        actual_key = resolve_column(headers, key_column or "") if key_column else suggest_key_column(headers)
        if not actual_key:
            raise ValueError("No record key was selected. Choose Survnum or another unique column.")
        validate_record_keys(rows, actual_key)
        timeseries_column = resolve_column(headers, "Timeseries")
        timeseries_values: list[int | None] = []
        for row_number, row in enumerate(rows, start=2):
            if not timeseries_column:
                timeseries_values.append(None)
                continue
            raw_timeseries = clean_text(row.get(timeseries_column))
            if not re.fullmatch(r"-?\d+(?:\.0+)?", raw_timeseries):
                raise ValueError(
                    f"Timeseries must contain integers. Row {row_number} has: "
                    f"{raw_timeseries or '(blank)'}"
                )
            timeseries_values.append(int(float(raw_timeseries)))
        maximum_timeseries = (
            max(value for value in timeseries_values if value is not None)
            if timeseries_column and timeseries_values
            else None
        )
        specs = detect_open_end_questions(headers)
        if not specs:
            raise ValueError("No qN/qNoe columns with matching cN_1..3 columns were found.")

        imported_at = now_iso()
        fingerprint = file_fingerprint(source_path)
        new_responses = 0
        source_changed_responses = 0
        preserved_overrides = 0
        changed_overrides = 0
        response_count = 0
        with self.connection:
            self.connection.execute("UPDATE datasets SET active=0 WHERE active=1")
            cursor = self.connection.execute(
                "INSERT INTO datasets(source_path, fingerprint, imported_at, key_column, row_count, headers_json, active) "
                "VALUES(?, ?, ?, ?, ?, ?, 1)",
                (str(source_path), fingerprint, imported_at, actual_key, len(rows), json.dumps(headers)),
            )
            dataset_id = int(cursor.lastrowid)
            self.connection.execute("UPDATE responses SET present=0")
            self.connection.execute(
                "INSERT INTO meta(key, value) VALUES('timeseries_column', ?) "
                "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                (timeseries_column or "",),
            )
            self.connection.execute(
                "INSERT INTO meta(key, value) VALUES('timeseries_max', ?) "
                "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                ("" if maximum_timeseries is None else str(maximum_timeseries),),
            )

            for spec in specs:
                self.connection.execute(
                    "INSERT INTO questions(name, text_column, companion_columns_json, ordinal) VALUES(?, ?, ?, ?) "
                    "ON CONFLICT(name) DO UPDATE SET text_column=excluded.text_column, "
                    "companion_columns_json=excluded.companion_columns_json, ordinal=excluded.ordinal",
                    (spec.name, spec.text_column, json.dumps(spec.companion_columns), spec.ordinal),
                )
                question_row = self.connection.execute(
                    "SELECT id FROM questions WHERE name=? COLLATE NOCASE", (spec.name,)
                ).fetchone()
                assert question_row is not None
                question_id = int(question_row["id"])

                for row_index, source_row in enumerate(rows):
                    row_number = row_index + 2
                    timeseries = timeseries_values[row_index]
                    record_key = normalize_record_key(source_row.get(actual_key))
                    verbatim = clean_verbatim(source_row.get(spec.text_column))
                    original_codes = tuple(
                        code
                        for code in (normalize_code(source_row.get(column)) for column in spec.companion_columns)
                        if code
                    )
                    if not verbatim and not original_codes:
                        continue
                    response_count += 1
                    existing = self.connection.execute(
                        "SELECT r.id, r.verbatim, r.original_codes_json, "
                        "COALESCE(s.override_active, 0) AS override_active "
                        "FROM responses r LEFT JOIN response_state s ON s.response_id=r.id "
                        "WHERE r.question_id=? AND r.record_key=?",
                        (question_id, record_key),
                    ).fetchone()
                    if existing:
                        response_id = int(existing["id"])
                        source_changed = (
                            str(existing["verbatim"]) != verbatim
                            or _json_codes(existing["original_codes_json"]) != original_codes
                        )
                        if source_changed:
                            source_changed_responses += 1
                            self.connection.execute(
                                "INSERT INTO audit_log(changed_at, action, question_name, record_key, "
                                "old_value, new_value) VALUES(?, 'source_changed', ?, ?, ?, ?)",
                                (
                                    imported_at,
                                    spec.name,
                                    record_key,
                                    json.dumps(
                                        {
                                            "verbatim": str(existing["verbatim"]),
                                            "codes": _json_codes(existing["original_codes_json"]),
                                        }
                                    ),
                                    json.dumps({"verbatim": verbatim, "codes": original_codes}),
                                ),
                            )
                        if bool(existing["override_active"]):
                            preserved_overrides += 1
                            if source_changed:
                                changed_overrides += 1
                                self.connection.execute(
                                    "UPDATE response_state SET review_status='flagged', flagged=1, "
                                    "note=CASE WHEN note='' THEN ? ELSE note || ' | ' || ? END, updated_at=? "
                                    "WHERE response_id=?",
                                    (
                                        f"Source response changed during import {imported_at}",
                                        f"Source response changed during import {imported_at}",
                                        imported_at,
                                        response_id,
                                    ),
                                )
                        self.connection.execute(
                            "UPDATE responses SET row_number=?, timeseries=?, verbatim=?, original_codes_json=?, "
                            "dataset_id=?, last_source_change_dataset_id="
                            "CASE WHEN ? THEN ? ELSE last_source_change_dataset_id END, "
                            "present=1 WHERE id=?",
                            (
                                row_number,
                                timeseries,
                                verbatim,
                                json.dumps(original_codes),
                                dataset_id,
                                int(source_changed),
                                dataset_id,
                                response_id,
                            ),
                        )
                    else:
                        response_cursor = self.connection.execute(
                            "INSERT INTO responses(question_id, record_key, row_number, timeseries, verbatim, "
                            "original_codes_json, dataset_id, first_seen_dataset_id, present) "
                            "VALUES(?, ?, ?, ?, ?, ?, ?, ?, 1)",
                            (
                                question_id,
                                record_key,
                                row_number,
                                timeseries,
                                verbatim,
                                json.dumps(original_codes),
                                dataset_id,
                                dataset_id,
                            ),
                        )
                        response_id = int(response_cursor.lastrowid)
                        self.connection.execute(
                            "INSERT INTO response_state(response_id, updated_at) VALUES(?, ?)",
                            (response_id, imported_at),
                        )
                        new_responses += 1

                    for code in original_codes:
                        self._upsert_code_sql(question_id, code, source="imported", preserve_metadata=True)

            self.connection.execute(
                "INSERT INTO audit_log(changed_at, action, detail) VALUES(?, 'import_data', ?)",
                (
                    imported_at,
                    json.dumps(
                        {
                            "source_path": str(source_path),
                            "key_column": actual_key,
                            "rows": len(rows),
                            "questions": len(specs),
                            "responses": response_count,
                        }
                    ),
                ),
            )
        return ImportSummary(
            dataset_id=dataset_id,
            source_path=source_path,
            key_column=actual_key,
            row_count=len(rows),
            question_count=len(specs),
            response_count=response_count,
            new_response_count=new_responses,
            source_changed_response_count=source_changed_responses,
            preserved_override_count=preserved_overrides,
            changed_override_count=changed_overrides,
        )

    def questions(self) -> list[QuestionInfo]:
        rows = self.connection.execute(
            "SELECT q.* FROM questions q WHERE EXISTS "
            "(SELECT 1 FROM responses r WHERE r.question_id=q.id AND r.present=1) "
            "ORDER BY q.ordinal, q.name"
        ).fetchall()
        return [
            QuestionInfo(
                id=int(row["id"]),
                name=str(row["name"]),
                text_column=str(row["text_column"]),
                companion_columns=tuple(json.loads(row["companion_columns_json"])),
                ordinal=int(row["ordinal"]),
            )
            for row in rows
        ]

    def question(self, question_id: int) -> QuestionInfo:
        row = self.connection.execute("SELECT * FROM questions WHERE id=?", (question_id,)).fetchone()
        if row is None:
            raise ValueError(f"Question not found: {question_id}")
        return QuestionInfo(
            id=int(row["id"]),
            name=str(row["name"]),
            text_column=str(row["text_column"]),
            companion_columns=tuple(json.loads(row["companion_columns_json"])),
            ordinal=int(row["ordinal"]),
        )

    def _assignment_codes(self, response_id: int) -> tuple[str, ...]:
        rows = self.connection.execute(
            "SELECT code FROM assignments WHERE response_id=? ORDER BY slot", (response_id,)
        ).fetchall()
        return tuple(str(row["code"]) for row in rows)

    def response(self, response_id: int) -> ResponseInfo:
        row = self.connection.execute(
            "SELECT r.*, s.override_active, s.review_status, s.flagged, s.note, "
            "(SELECT id FROM datasets WHERE active=1 ORDER BY id DESC LIMIT 1) AS active_dataset_id "
            "FROM responses r JOIN response_state s ON s.response_id=r.id WHERE r.id=?",
            (response_id,),
        ).fetchone()
        if row is None:
            raise ValueError(f"Response not found: {response_id}")
        original = _json_codes(row["original_codes_json"])
        override = bool(row["override_active"])
        assigned = self._assignment_codes(response_id)
        active_dataset_id = (
            int(row["active_dataset_id"]) if row["active_dataset_id"] is not None else None
        )
        if active_dataset_id is not None and row["first_seen_dataset_id"] == active_dataset_id:
            import_status = "New"
        elif (
            active_dataset_id is not None
            and row["last_source_change_dataset_id"] == active_dataset_id
        ):
            import_status = "Source changed"
        else:
            import_status = "Existing"
        return ResponseInfo(
            id=int(row["id"]),
            question_id=int(row["question_id"]),
            record_key=str(row["record_key"]),
            row_number=int(row["row_number"]),
            timeseries=int(row["timeseries"]) if row["timeseries"] is not None else None,
            verbatim=str(row["verbatim"]),
            original_codes=original,
            final_codes=assigned if override else original,
            override_active=override,
            review_status=str(row["review_status"]),
            flagged=bool(row["flagged"]),
            note=str(row["note"]),
            import_status=import_status,
        )

    def list_responses(
        self,
        question_id: int,
        *,
        status: str = "all",
        search: str = "",
        code_filter: str = "",
        code_scope: str = "either",
        timeseries: int | None = None,
        import_scope: str = "all",
    ) -> list[ResponseInfo]:
        clauses = ["r.question_id=?", "r.present=1"]
        params: list[object] = [question_id]
        normalized_status = clean_text(status).casefold()
        if normalized_status == "pending":
            clauses.append("s.override_active=0")
        elif normalized_status == "coded":
            clauses.append("s.override_active=1")
        elif normalized_status == "flagged":
            clauses.append("s.flagged=1")
        if clean_text(search):
            clauses.append("(r.verbatim LIKE ? OR r.record_key LIKE ?)")
            pattern = f"%{clean_text(search)}%"
            params.extend([pattern, pattern])
        if timeseries is not None:
            clauses.append("r.timeseries=?")
            params.append(int(timeseries))
        normalized_import_scope = clean_text(import_scope).casefold() or "all"
        if normalized_import_scope not in VALID_IMPORT_SCOPES:
            raise ValueError("Import filter must be All, New, Changed, or New or Changed.")
        if normalized_import_scope != "all":
            dataset = self.active_dataset()
            if dataset is None:
                return []
            active_dataset_id = int(dataset["id"])
            if normalized_import_scope == "new":
                clauses.append("r.first_seen_dataset_id=?")
                params.append(active_dataset_id)
            elif normalized_import_scope == "changed":
                clauses.append(
                    "r.last_source_change_dataset_id=? AND "
                    "COALESCE(r.first_seen_dataset_id, -1)<>?"
                )
                params.extend([active_dataset_id, active_dataset_id])
            else:
                clauses.append(
                    "(r.first_seen_dataset_id=? OR r.last_source_change_dataset_id=?)"
                )
                params.extend([active_dataset_id, active_dataset_id])
        rows = self.connection.execute(
            "SELECT r.id FROM responses r JOIN response_state s ON s.response_id=r.id WHERE "
            + " AND ".join(clauses)
            + " ORDER BY r.row_number, r.record_key",
            params,
        ).fetchall()
        responses = [self.response(int(row["id"])) for row in rows]
        raw_filter = clean_text(code_filter)
        if not raw_filter:
            return responses
        match_blank = raw_filter.casefold() == "(blank)"
        wanted = "" if match_blank else validate_code(raw_filter)
        scope = clean_text(code_scope).casefold() or "either"
        if scope not in {"original", "final", "either"}:
            raise ValueError("Code filter scope must be Original, Final, or Either.")

        def matches(codes: tuple[str, ...]) -> bool:
            return not codes if match_blank else wanted in codes

        output: list[ResponseInfo] = []
        for response in responses:
            original_match = matches(response.original_codes)
            final_match = matches(response.final_codes)
            if (
                (scope == "original" and original_match)
                or (scope == "final" and final_match)
                or (scope == "either" and (original_match or final_match))
            ):
                output.append(response)
        return output

    def available_response_codes(self, question_id: int, *, timeseries: int | None = None) -> list[str]:
        codes: set[str] = set()
        for response in self.list_responses(question_id, timeseries=timeseries):
            codes.update(response.original_codes)
            codes.update(response.final_codes)
        codes.update(item.code for item in self.list_codes(question_id, include_inactive=True))
        return sorted(codes, key=code_sort_key)

    def max_timeseries(self) -> int | None:
        if not self.get_meta("timeseries_column"):
            return None
        saved_maximum = clean_text(self.get_meta("timeseries_max"))
        if saved_maximum:
            return int(saved_maximum)
        row = self.connection.execute(
            "SELECT MAX(timeseries) AS maximum FROM responses WHERE present=1"
        ).fetchone()
        if row is None or row["maximum"] is None:
            return None
        return int(row["maximum"])

    def set_assignment(
        self,
        response_id: int,
        codes: Sequence[str],
        *,
        flagged: bool = False,
        note: str = "",
        action: str = "set_assignment",
    ) -> ResponseInfo:
        response = self.response(response_id)
        question = self.question(response.question_id)
        normalized: list[str] = []
        for raw in codes:
            code = validate_code(raw)
            if code not in normalized:
                normalized.append(code)
        if len(normalized) > len(question.companion_columns):
            raise ValueError(
                f"{question.name} has {len(question.companion_columns)} code slot(s); "
                f"{len(normalized)} codes were assigned."
            )
        old_value = json.dumps(response.final_codes)
        timestamp = now_iso()
        with self.connection:
            self.connection.execute("DELETE FROM assignments WHERE response_id=?", (response_id,))
            for slot, code in enumerate(normalized, start=1):
                self.connection.execute(
                    "INSERT INTO assignments(response_id, slot, code, updated_at) VALUES(?, ?, ?, ?)",
                    (response_id, slot, code, timestamp),
                )
                self._upsert_code_sql(response.question_id, code, source="user", preserve_metadata=True)
            self.connection.execute(
                "UPDATE response_state SET override_active=1, review_status=?, flagged=?, note=?, updated_at=? "
                "WHERE response_id=?",
                ("flagged" if flagged else "coded", int(flagged), clean_text(note), timestamp, response_id),
            )
            self.connection.execute(
                "INSERT INTO audit_log(changed_at, action, question_name, record_key, old_value, new_value) "
                "VALUES(?, ?, ?, ?, ?, ?)",
                (timestamp, action, question.name, response.record_key, old_value, json.dumps(normalized)),
            )
        return self.response(response_id)

    def reset_assignment(self, response_id: int) -> ResponseInfo:
        response = self.response(response_id)
        question = self.question(response.question_id)
        timestamp = now_iso()
        with self.connection:
            self.connection.execute("DELETE FROM assignments WHERE response_id=?", (response_id,))
            self.connection.execute(
                "UPDATE response_state SET override_active=0, review_status='pending', flagged=0, "
                "note='', updated_at=? WHERE response_id=?",
                (timestamp, response_id),
            )
            self.connection.execute(
                "INSERT INTO audit_log(changed_at, action, question_name, record_key, old_value, new_value) "
                "VALUES(?, 'reset_assignment', ?, ?, ?, ?)",
                (timestamp, question.name, response.record_key, json.dumps(response.final_codes), json.dumps(response.original_codes)),
            )
        return self.response(response_id)

    def apply_to_matching_text(
        self,
        response_id: int,
        codes: Sequence[str],
        *,
        timeseries: int | None = None,
    ) -> int:
        response = self.response(response_id)
        clauses = ["question_id=?", "present=1", "verbatim=?"]
        params: list[object] = [response.question_id, response.verbatim]
        if timeseries is not None:
            clauses.append("timeseries=?")
            params.append(int(timeseries))
        matches = self.connection.execute(
            "SELECT id FROM responses WHERE " + " AND ".join(clauses) + " ORDER BY row_number",
            params,
        ).fetchall()
        for match in matches:
            self.set_assignment(int(match["id"]), codes, action="bulk_matching_text")
        return len(matches)

    def preview_redundant_code_cleanup(
        self,
        question_id: int,
        code: str,
        *,
        timeseries: int | None = None,
    ) -> RedundantCodeCleanupSummary:
        question = self.question(question_id)
        normalized_code = validate_code(code)
        matching = 0
        changed = 0
        retained = 0
        for response in self.list_responses(question.id, timeseries=timeseries):
            if normalized_code not in response.final_codes:
                continue
            matching += 1
            if any(item != normalized_code for item in response.final_codes):
                changed += 1
            else:
                retained += 1
        return RedundantCodeCleanupSummary(
            question_id=question.id,
            question_name=question.name,
            code=normalized_code,
            timeseries=timeseries,
            matching_response_count=matching,
            changed_response_count=changed,
            retained_only_code_count=retained,
        )

    def cleanup_redundant_code(
        self,
        question_id: int,
        code: str,
        *,
        timeseries: int | None = None,
    ) -> RedundantCodeCleanupSummary:
        preview = self.preview_redundant_code_cleanup(
            question_id,
            code,
            timeseries=timeseries,
        )
        changed = 0
        for response in self.list_responses(question_id, timeseries=timeseries):
            if preview.code not in response.final_codes:
                continue
            remaining = [item for item in response.final_codes if item != preview.code]
            if not remaining:
                continue
            self.set_assignment(
                response.id,
                remaining,
                flagged=response.flagged,
                note=response.note,
                action="cleanup_redundant_code",
            )
            changed += 1
        summary = RedundantCodeCleanupSummary(
            question_id=preview.question_id,
            question_name=preview.question_name,
            code=preview.code,
            timeseries=preview.timeseries,
            matching_response_count=preview.matching_response_count,
            changed_response_count=changed,
            retained_only_code_count=preview.retained_only_code_count,
        )
        with self.connection:
            self.connection.execute(
                "INSERT INTO audit_log(changed_at, action, question_name, detail) "
                "VALUES(?, 'cleanup_redundant_code_summary', ?, ?)",
                (
                    now_iso(),
                    preview.question_name,
                    json.dumps(
                        {
                            "code": preview.code,
                            "timeseries": timeseries,
                            "matching_responses": preview.matching_response_count,
                            "changed_responses": changed,
                            "retained_only_code": preview.retained_only_code_count,
                        }
                    ),
                ),
            )
        return summary

    def _upsert_code_sql(
        self,
        question_id: int,
        code: str,
        *,
        label: str = "",
        supercategory: str = "",
        direction: str = "o",
        source: str = "user",
        active: bool = True,
        preserve_metadata: bool = False,
    ) -> None:
        code = validate_code(code)
        timestamp = now_iso()
        if preserve_metadata:
            self.connection.execute(
                "INSERT INTO codes(question_id, code, label, supercategory, direction, source, active, created_at, updated_at) "
                "VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?) "
                "ON CONFLICT(question_id, code) DO NOTHING",
                (question_id, code, clean_text(label), clean_text(supercategory), direction, source, int(active), timestamp, timestamp),
            )
            return
        self.connection.execute(
            "INSERT INTO codes(question_id, code, label, supercategory, direction, source, active, created_at, updated_at) "
            "VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?) "
            "ON CONFLICT(question_id, code) DO UPDATE SET label=excluded.label, "
            "supercategory=excluded.supercategory, direction=excluded.direction, source=excluded.source, "
            "active=excluded.active, updated_at=excluded.updated_at",
            (question_id, code, clean_text(label), clean_text(supercategory), direction, source, int(active), timestamp, timestamp),
        )

    def upsert_code(
        self,
        question_id: int,
        code: str,
        *,
        label: str,
        supercategory: str = "",
        direction: str = "o",
        source: str = "user",
        active: bool = True,
    ) -> CodeInfo:
        normalized_direction = clean_text(direction).casefold() or "o"
        if normalized_direction not in VALID_DIRECTIONS:
            raise ValueError("Direction must be +, -, or o.")
        normalized_code = validate_code(code)
        with self.connection:
            self._upsert_code_sql(
                question_id,
                normalized_code,
                label=label,
                supercategory=supercategory,
                direction=normalized_direction,
                source=source,
                active=active,
            )
            question = self.question(question_id)
            self.connection.execute(
                "INSERT INTO audit_log(changed_at, action, question_name, new_value, detail) "
                "VALUES(?, 'upsert_code', ?, ?, ?)",
                (
                    now_iso(),
                    question.name,
                    normalized_code,
                    json.dumps({"label": clean_text(label), "supercategory": clean_text(supercategory)}),
                ),
            )
        return self.get_code(question_id, normalized_code)

    def get_code(self, question_id: int, code: str) -> CodeInfo:
        row = self.connection.execute(
            "SELECT * FROM codes WHERE question_id=? AND code=?", (question_id, normalize_code(code))
        ).fetchone()
        if row is None:
            raise ValueError(f"Code not found: {code}")
        return CodeInfo(
            question_id=int(row["question_id"]),
            code=str(row["code"]),
            label=str(row["label"]),
            supercategory=str(row["supercategory"]),
            direction=str(row["direction"]),
            source=str(row["source"]),
            active=bool(row["active"]),
        )

    def list_codes(self, question_id: int, *, include_inactive: bool = False) -> list[CodeInfo]:
        sql = "SELECT * FROM codes WHERE question_id=?"
        if not include_inactive:
            sql += " AND active=1"
        rows = self.connection.execute(sql, (question_id,)).fetchall()
        result = [
            CodeInfo(
                question_id=int(row["question_id"]),
                code=str(row["code"]),
                label=str(row["label"]),
                supercategory=str(row["supercategory"]),
                direction=str(row["direction"]),
                source=str(row["source"]),
                active=bool(row["active"]),
            )
            for row in rows
        ]
        return sorted(result, key=lambda item: code_sort_key(item.code))

    def codes_sharing_label(self, question_id: int, label: str) -> list[CodeInfo]:
        label_key = self._wincross_label_key(label)
        if not label_key:
            return []
        return [
            item
            for item in self.list_codes(question_id)
            if self._wincross_label_key(item.label) == label_key
        ]

    def next_user_code(self, question_id: int, *, start: int = 500) -> str:
        used = {
            int(item.code)
            for item in self.list_codes(question_id, include_inactive=True)
            if item.code.isdigit()
        }
        candidate = start
        while candidate in used:
            candidate += 1
        return str(candidate)

    def load_codelist(
        self,
        path: Path,
        *,
        overwrite_existing: bool = False,
    ) -> CodelistImportSummary:
        path = path.resolve()
        if not path.exists():
            raise ValueError(f"Codelist not found: {path}")
        suffix = path.suffix.casefold()
        entries: list[tuple[str, str, str, str, str]] = []
        question_by_name = {question.name.casefold(): question for question in self.questions()}
        if suffix == ".csv":
            headers, rows = _read_csv(path)
            code_col = resolve_column(headers, "code") or resolve_column(headers, "final_code")
            label_col = (
                resolve_column(headers, "label")
                or resolve_column(headers, "category")
                or resolve_column(headers, "final_label")
            )
            question_col = resolve_column(headers, "question_col") or resolve_column(headers, "question")
            super_col = resolve_column(headers, "supercategory")
            direction_col = resolve_column(headers, "direction")
            if not code_col or not label_col:
                raise ValueError("Codelist CSV must contain code and label/category columns.")
            if not question_col and len(question_by_name) != 1:
                raise ValueError("Codelist CSV needs question_col when the session has multiple questions.")
            only_question = next(iter(question_by_name.values()), None)
            for row in rows:
                question_name = clean_text(row.get(question_col)) if question_col else (only_question.name if only_question else "")
                code = normalize_code(row.get(code_col))
                label = clean_text(row.get(label_col))
                if question_name and code and label and CODE_RE.fullmatch(code):
                    entries.append(
                        (
                            question_name,
                            code,
                            label,
                            clean_text(row.get(super_col)) if super_col else "",
                            clean_text(row.get(direction_col)) if direction_col else "",
                        )
                    )
        elif suffix in {".xlsx", ".xlsm"}:
            try:
                from openpyxl import load_workbook  # type: ignore
            except Exception as exc:  # pragma: no cover
                raise ValueError("Reading codelist workbooks requires openpyxl.") from exc
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                for worksheet in workbook.worksheets:
                    question = question_by_name.get(worksheet.title.casefold())
                    if question is None:
                        continue
                    for raw in worksheet.iter_rows(values_only=True):
                        code = normalize_code(raw[0] if len(raw) > 0 else "")
                        label = clean_text(raw[1] if len(raw) > 1 else "")
                        if code and label and CODE_RE.fullmatch(code):
                            supercategory = clean_text(raw[2] if len(raw) > 2 else "")
                            direction = clean_text(raw[3] if len(raw) > 3 else "")
                            entries.append((question.name, code, label, supercategory, direction))
            finally:
                workbook.close()
        else:
            raise ValueError(f"Unsupported codelist format: {path.suffix}")

        added = 0
        updated = 0
        filled_labels = 0
        filled_supercategories = 0
        filled_directions = 0
        unchanged = 0
        conflicts: list[CodelistMetadataConflict] = []
        overwrites: list[CodelistMetadataConflict] = []
        timestamp = now_iso()
        with self.connection:
            for question_name, code, label, supercategory, direction in entries:
                question = question_by_name.get(question_name.casefold())
                if question is None:
                    continue
                imported_direction = direction.casefold()
                if imported_direction not in VALID_DIRECTIONS:
                    imported_direction = ""
                row = self.connection.execute(
                    "SELECT label, supercategory, direction, source FROM codes "
                    "WHERE question_id=? AND code=?",
                    (question.id, code),
                ).fetchone()
                if row is None:
                    self._upsert_code_sql(
                        question.id,
                        code,
                        label=label,
                        supercategory=supercategory,
                        direction=imported_direction or "o",
                        source="delivered",
                    )
                    added += 1
                    continue

                existing_source = clean_text(row["source"])
                updates: dict[str, str] = {}
                incoming = {
                    "label": clean_text(label),
                    "supercategory": clean_text(supercategory),
                    "direction": imported_direction,
                }
                existing = {
                    "label": clean_text(row["label"]),
                    "supercategory": clean_text(row["supercategory"]),
                    "direction": clean_text(row["direction"]).casefold() or "o",
                }
                for field in ("label", "supercategory", "direction"):
                    imported_value = incoming[field]
                    existing_value = existing[field]
                    if not imported_value or existing_value.casefold() == imported_value.casefold():
                        continue
                    fillable = not existing_value
                    if (
                        field == "direction"
                        and existing_source == "imported"
                        and existing_value == "o"
                    ):
                        fillable = True
                    if fillable:
                        updates[field] = imported_value
                        if field == "label":
                            filled_labels += 1
                        elif field == "supercategory":
                            filled_supercategories += 1
                        else:
                            filled_directions += 1
                        continue

                    difference = CodelistMetadataConflict(
                        question_name=question.name,
                        code=code,
                        field=field,
                        existing=existing_value,
                        imported=imported_value,
                    )
                    if overwrite_existing:
                        updates[field] = imported_value
                        overwrites.append(difference)
                    else:
                        conflicts.append(difference)

                if updates:
                    assignments = [f"{field}=?" for field in updates]
                    values: list[object] = list(updates.values())
                    if existing_source == "imported":
                        assignments.append("source=?")
                        values.append("delivered")
                    assignments.append("updated_at=?")
                    values.extend([timestamp, question.id, code])
                    self.connection.execute(
                        f"UPDATE codes SET {', '.join(assignments)} "
                        "WHERE question_id=? AND code=?",
                        values,
                    )
                    updated += 1
                else:
                    unchanged += 1

            detail = {
                "path": str(path),
                "overwrite_existing": overwrite_existing,
                "parsed_codes": len(entries),
                "added_codes": added,
                "updated_codes": updated,
                "conflicts": [
                    {
                        "question": item.question_name,
                        "code": item.code,
                        "field": item.field,
                        "existing": item.existing,
                        "imported": item.imported,
                    }
                    for item in conflicts
                ],
                "overwrites": [
                    {
                        "question": item.question_name,
                        "code": item.code,
                        "field": item.field,
                        "existing": item.existing,
                        "imported": item.imported,
                    }
                    for item in overwrites
                ],
            }
            self.connection.execute(
                "INSERT INTO audit_log(changed_at, action, detail) VALUES(?, 'load_codelist', ?)",
                (timestamp, json.dumps(detail)),
            )
        self.set_meta("codelist_path", str(path))
        return CodelistImportSummary(
            source_path=path,
            overwrite_existing=overwrite_existing,
            parsed_code_count=len(entries),
            added_code_count=added,
            updated_code_count=updated,
            filled_label_count=filled_labels,
            filled_supercategory_count=filled_supercategories,
            filled_direction_count=filled_directions,
            unchanged_code_count=unchanged,
            conflicts=tuple(conflicts),
            overwrites=tuple(overwrites),
        )

    def import_wincross_codebook(
        self,
        path: Path,
        question_id: int,
        *,
        overwrite_existing: bool = False,
    ) -> WinCrossCodebookImportSummary:
        """Merge WinCross metadata, preserving existing nonblank values by default."""
        path = path.resolve()
        if not path.exists():
            raise ValueError(f"WinCross table not found: {path}")
        question = self.question(question_id)
        parsed_entries = parse_wincross_codebook(_read_text_file(path))
        if not parsed_entries:
            raise ValueError("No WinCross NET groups with coded category rows were found.")

        entries_by_code: dict[str, WinCrossCodebookEntry] = {}
        conflicts: list[WinCrossMetadataConflict] = []
        overwrites: list[WinCrossMetadataConflict] = []
        for entry in parsed_entries:
            prior = entries_by_code.get(entry.code)
            if prior is None:
                entries_by_code[entry.code] = entry
                continue
            for field in ("label", "supercategory"):
                existing_value = getattr(prior, field)
                imported_value = getattr(entry, field)
                if imported_value and existing_value.casefold() != imported_value.casefold():
                    conflicts.append(
                        WinCrossMetadataConflict(
                            code=entry.code,
                            field=field,
                            existing=existing_value,
                            imported=imported_value,
                        )
                    )

        added = 0
        updated = 0
        filled_labels = 0
        filled_supercategories = 0
        unchanged = 0
        timestamp = now_iso()
        with self.connection:
            for entry in entries_by_code.values():
                row = self.connection.execute(
                    "SELECT label, supercategory FROM codes WHERE question_id=? AND code=?",
                    (question.id, entry.code),
                ).fetchone()
                if row is None:
                    self._upsert_code_sql(
                        question.id,
                        entry.code,
                        label=entry.label,
                        supercategory=entry.supercategory,
                        direction="o",
                        source="wincross",
                    )
                    added += 1
                    continue

                existing_label = clean_text(row["label"])
                existing_supercategory = clean_text(row["supercategory"])
                updates: dict[str, str] = {}
                if not existing_label and entry.label:
                    updates["label"] = entry.label
                    filled_labels += 1
                elif existing_label and existing_label.casefold() != entry.label.casefold():
                    difference = WinCrossMetadataConflict(
                        code=entry.code,
                        field="label",
                        existing=existing_label,
                        imported=entry.label,
                    )
                    if overwrite_existing:
                        updates["label"] = entry.label
                        overwrites.append(difference)
                    else:
                        conflicts.append(difference)
                if not existing_supercategory and entry.supercategory:
                    updates["supercategory"] = entry.supercategory
                    filled_supercategories += 1
                elif (
                    entry.supercategory
                    and existing_supercategory
                    and existing_supercategory.casefold() != entry.supercategory.casefold()
                ):
                    difference = WinCrossMetadataConflict(
                        code=entry.code,
                        field="supercategory",
                        existing=existing_supercategory,
                        imported=entry.supercategory,
                    )
                    if overwrite_existing:
                        updates["supercategory"] = entry.supercategory
                        overwrites.append(difference)
                    else:
                        conflicts.append(difference)

                if updates:
                    assignments = [f"{field}=?" for field in updates]
                    values: list[object] = list(updates.values())
                    assignments.append("updated_at=?")
                    values.extend([timestamp, question.id, entry.code])
                    self.connection.execute(
                        f"UPDATE codes SET {', '.join(assignments)} WHERE question_id=? AND code=?",
                        values,
                    )
                    updated += 1
                else:
                    unchanged += 1

            detail = {
                "path": str(path),
                "question": question.name,
                "overwrite_existing": overwrite_existing,
                "parsed_codes": len(entries_by_code),
                "added_codes": added,
                "updated_codes": updated,
                "conflicts": [
                    {
                        "code": item.code,
                        "field": item.field,
                        "existing": item.existing,
                        "imported": item.imported,
                    }
                    for item in conflicts
                ],
                "overwrites": [
                    {
                        "code": item.code,
                        "field": item.field,
                        "existing": item.existing,
                        "imported": item.imported,
                    }
                    for item in overwrites
                ],
            }
            self.connection.execute(
                "INSERT INTO audit_log(changed_at, action, question_name, detail) "
                "VALUES(?, 'import_wincross_codebook', ?, ?)",
                (timestamp, question.name, json.dumps(detail)),
            )
            self.connection.execute(
                "INSERT INTO meta(key, value) VALUES('wincross_codebook_path', ?) "
                "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                (str(path),),
            )

        return WinCrossCodebookImportSummary(
            source_path=path,
            question_name=question.name,
            overwrite_existing=overwrite_existing,
            parsed_code_count=len(entries_by_code),
            added_code_count=added,
            updated_code_count=updated,
            filled_label_count=filled_labels,
            filled_supercategory_count=filled_supercategories,
            unchanged_code_count=unchanged,
            conflicts=tuple(conflicts),
            overwrites=tuple(overwrites),
        )

    def review_summary(self) -> list[QuestionReview]:
        rows = self.connection.execute(
            "SELECT q.id, q.name, COUNT(r.id) AS total, "
            "SUM(CASE WHEN s.override_active=1 THEN 1 ELSE 0 END) AS coded, "
            "SUM(CASE WHEN s.override_active=0 THEN 1 ELSE 0 END) AS pending, "
            "SUM(CASE WHEN s.flagged=1 THEN 1 ELSE 0 END) AS flagged "
            "FROM questions q JOIN responses r ON r.question_id=q.id AND r.present=1 "
            "JOIN response_state s ON s.response_id=r.id GROUP BY q.id, q.name ORDER BY q.ordinal, q.name"
        ).fetchall()
        output: list[QuestionReview] = []
        for row in rows:
            blank_final = 0
            for response in self.list_responses(int(row["id"])):
                if not response.final_codes:
                    blank_final += 1
            output.append(
                QuestionReview(
                    question_id=int(row["id"]),
                    question_name=str(row["name"]),
                    total=int(row["total"] or 0),
                    coded=int(row["coded"] or 0),
                    pending=int(row["pending"] or 0),
                    flagged=int(row["flagged"] or 0),
                    blank_final=blank_final,
                )
            )
        return output

    def export_coded_csv(self, output_path: Path) -> ExportSummary:
        dataset = self.active_dataset()
        if dataset is None:
            raise ValueError("Import survey data before exporting.")
        source_path = Path(str(dataset["source_path"]))
        if not source_path.exists():
            raise ValueError(f"Imported source datafile no longer exists: {source_path}")
        if file_fingerprint(source_path) != str(dataset["fingerprint"]):
            raise ValueError("The source datafile changed after import. Re-import it before exporting.")
        output_path = output_path.resolve()
        if output_path == source_path.resolve():
            raise ValueError("The coded export cannot overwrite the original source datafile.")
        headers, rows = read_tabular(source_path)
        key_column = str(dataset["key_column"])
        actual_key = resolve_column(headers, key_column)
        if not actual_key:
            raise ValueError(f"Key column {key_column} is missing from the source datafile.")
        validate_record_keys(rows, actual_key)

        override_rows = self.connection.execute(
            "SELECT r.id, r.record_key, q.name, q.companion_columns_json "
            "FROM responses r JOIN response_state s ON s.response_id=r.id "
            "JOIN questions q ON q.id=r.question_id "
            "WHERE r.present=1 AND s.override_active=1"
        ).fetchall()
        overrides: dict[tuple[str, str], tuple[tuple[str, ...], tuple[str, ...]]] = {}
        for row in override_rows:
            columns = tuple(json.loads(row["companion_columns_json"]))
            codes = self._assignment_codes(int(row["id"]))
            if len(codes) > len(columns):
                raise ValueError(
                    f"{row['name']} / {row['record_key']} has more codes than available companion columns."
                )
            overrides[(str(row["name"]).casefold(), str(row["record_key"]))] = (columns, codes)

        modified_records: set[str] = set()
        modified_cells = 0
        questions = self.questions()
        for source_row in rows:
            record_key = normalize_record_key(source_row.get(actual_key))
            for question in questions:
                override = overrides.get((question.name.casefold(), record_key))
                if override is None:
                    continue
                columns, codes = override
                for index, column in enumerate(columns):
                    desired = codes[index] if index < len(codes) else ""
                    current = clean_text(source_row.get(column))
                    if current != desired:
                        source_row[column] = desired
                        modified_cells += 1
                        modified_records.add(record_key)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow({header: row.get(header, "") for header in headers})
        with self.connection:
            self.connection.execute(
                "INSERT INTO audit_log(changed_at, action, detail) VALUES(?, 'export_coded_csv', ?)",
                (
                    now_iso(),
                    json.dumps(
                        {
                            "path": str(output_path),
                            "rows": len(rows),
                            "modified_records": len(modified_records),
                            "modified_cells": modified_cells,
                        }
                    ),
                ),
            )
        return ExportSummary(
            output_path=output_path,
            row_count=len(rows),
            modified_records=len(modified_records),
            modified_cells=modified_cells,
            overridden_responses=len(override_rows),
        )

    def export_codebook_xlsx(self, output_path: Path) -> Path:
        try:
            from openpyxl import Workbook  # type: ignore
        except Exception as exc:  # pragma: no cover
            raise ValueError("Writing the codebook workbook requires openpyxl.") from exc
        output_path = output_path.resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        try:
            first = True
            for question in self.questions():
                worksheet = workbook.active if first else workbook.create_sheet()
                first = False
                worksheet.title = question.name[:31]
                worksheet.append([question.name])
                worksheet.append(["Code", "Category", "Supercategory", "Direction", "Source", "Active"])
                for item in self.list_codes(question.id, include_inactive=True):
                    worksheet.append(
                        [item.code, item.label, item.supercategory, item.direction, item.source, "Y" if item.active else "N"]
                    )
                worksheet.freeze_panes = "A3"
            workbook.save(output_path)
        finally:
            workbook.close()
        return output_path

    def _effective_codes_for_question(self, question_id: int) -> set[str]:
        used: set[str] = set()
        for response in self.list_responses(question_id):
            used.update(response.final_codes)
        return used

    @staticmethod
    def _wincross_label(value: str) -> str:
        return clean_text(value).replace("^", " ")

    @staticmethod
    def _wincross_label_key(value: str) -> str:
        normalized = clean_text(value).casefold().replace("'", "").replace("\u2019", "")
        return re.sub(r"[^a-z0-9]+", " ", normalized).strip()

    def render_wincross_tables(self) -> tuple[str, tuple[str, ...], int]:
        blocks: list[str] = []
        unlabeled: list[str] = []
        rendered_code_count = 0
        for question in self.questions():
            used_codes = self._effective_codes_for_question(question.id)
            metadata = {
                item.code: item
                for item in self.list_codes(question.id, include_inactive=True)
            }
            included_codes = {
                code
                for code, item in metadata.items()
                if code in used_codes or (item.active and bool(clean_text(item.label)))
            }
            included_codes.update(used_codes)
            ordered_codes = sorted(included_codes, key=code_sort_key)
            if not ordered_codes:
                continue

            rows: list[dict[str, str]] = []
            for code in ordered_codes:
                item = metadata.get(code)
                label = clean_text(item.label) if item is not None else ""
                if not label:
                    unlabeled.append(f"{question.name}:{code}")
                    label = f"(Label not set - code {code})"
                rows.append(
                    {
                        "code": code,
                        "label": self._wincross_label(label),
                        "supercategory": self._wincross_label(item.supercategory if item else ""),
                        "direction": item.direction if item else "o",
                    }
                )

            summary_codes = [
                row["code"]
                for row in rows
                if self._wincross_label_key(row["label"]) not in {"none", "dont know"}
            ]
            positive_codes = [row["code"] for row in rows if row["direction"] == "+"]
            negative_codes = [row["code"] for row in rows if row["direction"] == "-"]
            selector = f"{question.companion_columns[0]} S3"
            lines = [f"ABLE TO MENTION^{selector} ({','.join(summary_codes)})^NS,RH"]
            if positive_codes:
                lines.append(f"POSITIVE MENTIONS^({','.join(positive_codes)})^NS,RH")
            if negative_codes:
                lines.append(f"NEGATIVE MENTIONS^({','.join(negative_codes)})^NS,RH")
            lines.append(" ^ ^NS,RH")

            display_rows: list[dict[str, str]] = []
            display_by_key: dict[tuple[str, str, str], dict[str, str]] = {}
            for row in rows:
                display_key = (
                    self._wincross_label_key(row["label"]),
                    row["supercategory"].casefold(),
                    row["direction"],
                )
                existing_display = display_by_key.get(display_key)
                if existing_display is None:
                    display_row = dict(row)
                    display_by_key[display_key] = display_row
                    display_rows.append(display_row)
                else:
                    existing_display["code"] += f",{row['code']}"

            supercategory_order: list[str] = []
            grouped: dict[str, list[dict[str, str]]] = {}
            ungrouped: list[dict[str, str]] = []
            tail_rows: list[dict[str, str]] = []
            for row in display_rows:
                if self._wincross_label_key(row["label"]) in {"none", "dont know"}:
                    tail_rows.append(row)
                    continue
                supercategory = row["supercategory"]
                if not supercategory:
                    ungrouped.append(row)
                    continue
                key = supercategory.casefold()
                if key not in grouped:
                    grouped[key] = []
                    supercategory_order.append(key)
                grouped[key].append(row)

            other_keys = [
                key
                for key in supercategory_order
                if self._wincross_label_key(grouped[key][0]["supercategory"]) == "other"
            ]
            regular_group_keys = [key for key in supercategory_order if key not in other_keys]
            tail_rows.sort(
                key=lambda row: (
                    0 if self._wincross_label_key(row["label"]) == "none" else 1,
                    code_sort_key(row["code"].split(",", 1)[0]),
                )
            )

            if supercategory_order:
                lines.append(" ^ ^NS,DV")
                for key in regular_group_keys:
                    items = grouped[key]
                    supercategory = items[0]["supercategory"]
                    lines.append(f"{supercategory}^NET {len(items) + 1}^NS,SK")
                    for item in items:
                        lines.append(f"{item['label']}^({item['code']})^NS,SK")
                    lines.append(" ^ ^NS,DV")
                for item in ungrouped:
                    lines.append(f"{item['label']}^({item['code']})^NS,SK")
                if ungrouped:
                    lines.append(" ^ ^NS,DV")
                for key in other_keys:
                    items = grouped[key]
                    supercategory = items[0]["supercategory"]
                    lines.append(f"{supercategory}^NET {len(items) + 1}^NS,SK,RL")
                    for item in items:
                        flags = "NS,SK,RL" if self._wincross_label_key(item["label"]) == "other" else "NS,SK"
                        lines.append(f"{item['label']}^({item['code']})^{flags}")
                    lines.append(" ^ ^NS,RL,DV")
                for item in tail_rows:
                    lines.append(f"{item['label']}^({item['code']})^NS,SK,RL")
                if tail_rows:
                    lines.append(" ^ ^NS,RL,DV")
            else:
                for item in ungrouped:
                    lines.append(f"{item['label']}^({item['code']})^NS,SK")
                for item in tail_rows:
                    lines.append(f"{item['label']}^({item['code']})^NS,SK,RL")
                lines.append(" ^ ^NS,RL,DV" if tail_rows else " ^ ^NS,DV")

            rendered_code_count += len(rows)
            blocks.append("\n".join(lines).rstrip())

        text = "\n\n".join(blocks).rstrip()
        if text:
            text += "\n"
        return text, tuple(unlabeled), rendered_code_count

    def export_wincross_tables(self, output_path: Path) -> WinCrossExportSummary:
        text, unlabeled_codes, code_count = self.render_wincross_tables()
        if not text:
            raise ValueError("No open-end codes are available for WinCross export.")
        output_path = output_path.resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(text, encoding="utf-8", newline="\n")
        question_count = len([block for block in text.strip().split("\n\n") if block.strip()])
        with self.connection:
            self.connection.execute(
                "INSERT INTO audit_log(changed_at, action, detail) VALUES(?, 'export_wincross_tables', ?)",
                (
                    now_iso(),
                    json.dumps(
                        {
                            "path": str(output_path),
                            "questions": question_count,
                            "codes": code_count,
                            "unlabeled_codes": unlabeled_codes,
                        }
                    ),
                ),
            )
        return WinCrossExportSummary(
            output_path=output_path,
            question_count=question_count,
            code_count=code_count,
            unlabeled_codes=unlabeled_codes,
        )

    def export_audit_csv(self, output_path: Path) -> Path:
        rows = self.connection.execute("SELECT * FROM audit_log ORDER BY id").fetchall()
        headers = [
            "changed_at",
            "action",
            "question_name",
            "record_key",
            "old_value",
            "new_value",
            "detail",
        ]
        output_path = output_path.resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=headers)
            writer.writeheader()
            for row in rows:
                writer.writerow({header: row[header] for header in headers})
        return output_path


def suggest_latest_data_file(project_root: Path | None, work_dir: Path | None) -> Path | None:
    roots: list[Path] = []
    if project_root is not None:
        roots.extend(
            [
                project_root / "Data" / "Partial",
                project_root / "Data" / "Partial Data",
                project_root / "Data",
            ]
        )
    if work_dir is not None:
        roots.append(work_dir)
    excluded = ("quota", "codelist", "codebook", ".oe-coded", "audit", "mapping")
    candidates: list[Path] = []
    seen: set[Path] = set()
    for root in roots:
        if not root.exists() or not root.is_dir():
            continue
        for pattern in ("*.csv", "*.xlsx", "*.xlsm"):
            for path in root.glob(pattern):
                resolved = path.resolve()
                if resolved in seen or any(token in path.name.casefold() for token in excluded):
                    continue
                seen.add(resolved)
                candidates.append(resolved)
    if not candidates:
        return None
    return max(candidates, key=lambda path: path.stat().st_mtime)
