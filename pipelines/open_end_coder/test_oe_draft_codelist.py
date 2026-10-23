"""Tests for drafting a job's code lists into a review workbook (no model download needed)."""

import csv
import shutil
import sys
import unittest
import uuid
from pathlib import Path
from unittest import mock

THIS_DIR = Path(__file__).resolve().parent
if str(THIS_DIR) not in sys.path:
    sys.path.insert(0, str(THIS_DIR))

import oe_codelist_embed as embed_drafter  # noqa: E402
import oe_draft_codelist as job  # noqa: E402
from oe_codelist_eval import DraftCode  # noqa: E402
from test_oe_codelist_draft import fake_embed  # noqa: E402

TMP_ROOT = THIS_DIR / ".tmp_tests_draft_codelist"
TEXTS = (
    ["He is too old", "Too old to serve", "His age, he is old", "He is old and tired", "Old age"]
    + ["Gas prices", "Inflation and gas prices", "Cost of groceries and gas prices", "Gas prices are too high"]
    + ["Nothing", "Don't know", "Purple"]
)


class DraftJobTests(unittest.TestCase):
    def setUp(self):
        self.folder = TMP_ROOT / uuid.uuid4().hex
        self.folder.mkdir(parents=True)
        self.addCleanup(shutil.rmtree, self.folder, True)
        self.data = self.folder / "data.csv"
        with self.data.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(["survnum", "phone_1", "lname", "q5", "q6"])
            for k, text in enumerate(TEXTS, 1):
                writer.writerow([k, "5550000", "Smith", "1" if k % 2 else "2", text])

    def test_only_needed_columns_are_read(self):
        rows = job.read_rows(self.data, ["q5", "q6"])
        self.assertEqual(set(rows[0]), {"survnum", "q5", "q6"})

    def test_workbook_has_numbered_codes_and_no_other_columns(self):
        from openpyxl import load_workbook

        spec = {
            "data": str(self.data),
            "out": str(self.folder / "draft.xlsx"),
            "questions": [{"column": "q6", "text": "Why less favorable?",
                           "lead": {"column": "q5", "type": "opinion_trend", "options": ["More favorable", "Less favorable"]}}],
        }
        with mock.patch.object(embed_drafter, "embed", fake_embed):
            job.run(spec)
        book = load_workbook(spec["out"], read_only=True)
        self.assertEqual(book.sheetnames, ["About", "Summary", "q6 code list", "q6 responses"])
        summary = " ".join(str(v) for row in book["Summary"].iter_rows(values_only=True) for v in row if v is not None)
        self.assertIn("Top codes", summary)
        self.assertIn("Main codes, lead-in Less favorable", summary)
        codes = [row[0] for row in list(book["q6 code list"].iter_rows(values_only=True))[1:]]
        self.assertIn(97, codes)
        self.assertIn(99, codes)
        self.assertIn(1, codes)
        responses = list(book["q6 responses"].iter_rows(values_only=True))
        self.assertEqual(responses[0][:3], ("survnum", "Verbatim", "Lead-in answer"))
        flat = " ".join(str(v) for row in responses for v in row if v is not None)
        self.assertNotIn("5550000", flat)
        self.assertNotIn("Smith", flat)
        self.assertEqual(len(responses) - 1, len(TEXTS))

    def test_numbering_puts_net_pairs_first_and_standard_codes_last(self):
        codes = [
            DraftCode("Alone", [0, 1, 2]),
            DraftCode("Bad on police", [3, 4], net="Police", lead="negative"),
            DraftCode("Good on police", [5], net="Police", lead="positive"),
            DraftCode("Nothing", [6]),
        ]
        numbered = job.number_codes(codes)
        self.assertEqual([(n, c.label) for n, c in numbered],
                         [("1", "Good on police"), ("2", "Bad on police"), ("3", "Alone"), ("97", "Nothing")])


if __name__ == "__main__":
    unittest.main()
