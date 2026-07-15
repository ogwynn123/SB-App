#!/usr/bin/env python3
"""Build Phone/SMS questionnaire DOCX drafts from a final questionnaire.

The builder intentionally keeps the transformation layer JSON-driven:

* source DOCX run formatting is read through the existing WinCross
  ``survey_unformatter`` sidecar model;
* channel rules come from ``channel_rules.default.json`` or a supplied JSON file;
* each run writes a machine-readable transform report and output sidecars.

The default rules are conservative. They cover repeatable mechanical CATI to SMS
edits and highlight changed text in the generated DOCX for review.
"""

from __future__ import annotations

import argparse
import difflib
import json
import re
import sys
from copy import deepcopy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.text.paragraph import Paragraph
from docx.text.run import Run

THIS_DIR = Path(__file__).resolve().parent
WINCROSS_DIR = THIS_DIR.parent / "wincross_table_generator"
if str(WINCROSS_DIR) not in sys.path:
    sys.path.insert(0, str(WINCROSS_DIR))

from survey_unformatter import build_formatting_sidecar  # noqa: E402


DEFAULT_RULES_PATH = THIS_DIR / "channel_rules.default.json"


@dataclass
class ParagraphChange:
    kind: str
    docx_paragraph_index: int
    source_record_index: int
    original_text: str
    new_text: str
    changed_spans: list[dict[str, Any]]
    rules_applied: list[str] = field(default_factory=list)


@dataclass
class ChannelBuild:
    channel: str
    output_docx: Path
    output_formatting_json: Path | None
    changes: list[ParagraphChange]


def load_rules(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if int(payload.get("schema_version", 0)) != 1:
        raise SystemExit(f"Unsupported rule schema_version in {path}")
    channels = payload.get("channels")
    if not isinstance(channels, dict):
        raise SystemExit(f"Rules file has no channels object: {path}")
    return payload


def compile_regex(pattern: str) -> re.Pattern[str]:
    return re.compile(pattern)


def merge_spans(spans: Iterable[dict[str, int]]) -> list[dict[str, int]]:
    ordered = sorted(
        ({"start": int(s["start"]), "end": int(s["end"])} for s in spans if int(s["start"]) < int(s["end"])),
        key=lambda s: (s["start"], s["end"]),
    )
    if not ordered:
        return []
    merged = [dict(ordered[0])]
    for span in ordered[1:]:
        last = merged[-1]
        if span["start"] <= last["end"]:
            last["end"] = max(last["end"], span["end"])
        else:
            merged.append(dict(span))
    return merged


def deletion_anchor_span(new_text: str, offset: int) -> dict[str, int] | None:
    """Return a small nearby span to highlight when a diff only deleted text."""
    if not new_text:
        return None
    start = min(max(offset, 0), len(new_text) - 1)
    end = start
    while end < len(new_text) and not new_text[end].isspace() and new_text[end] != "\t":
        end += 1
    if end == start:
        while start > 0 and not new_text[start - 1].isspace() and new_text[start - 1] != "\t":
            start -= 1
        end = min(max(offset, 0) + 1, len(new_text))
    if start < end:
        return {"start": start, "end": end}
    return None


def diff_highlight_spans(original: str, new: str) -> list[dict[str, Any]]:
    matcher = difflib.SequenceMatcher(a=original, b=new, autojunk=False)
    spans: list[dict[str, int]] = []
    for tag, _i1, _i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            continue
        if j1 < j2:
            spans.append({"start": j1, "end": j2})
        else:
            anchor = deletion_anchor_span(new, j1)
            if anchor:
                spans.append(anchor)
    return [
        {
            "start": span["start"],
            "end": span["end"],
            "text": new[span["start"] : span["end"]],
            "highlight": True,
            "highlight_color": "yellow",
            "source": "channel_builder_diff",
        }
        for span in merge_spans(spans)
    ]


def rule_is_active_for_text(rule: dict[str, Any], text: str) -> bool:
    if rule.get("enabled", True) is False:
        return False
    when = rule.get("when")
    if when and not compile_regex(str(when)).search(text):
        return False
    unless = rule.get("unless")
    if unless and compile_regex(str(unless)).search(text):
        return False
    return True


def apply_regex_rule(text: str, rule: dict[str, Any]) -> tuple[str, bool]:
    pattern = compile_regex(str(rule["match"]))
    replacement = str(rule.get("replacement", ""))
    count = 1 if rule.get("limit") == "first" else 0
    new_text, replacements = pattern.subn(replacement, text, count=count)
    return new_text, replacements > 0


def apply_channel_rules(
    source_sidecar: dict[str, Any],
    channel_rules: list[dict[str, Any]],
) -> list[ParagraphChange]:
    changes: list[ParagraphChange] = []
    first_matching_applied: set[str] = set()
    paragraphs = source_sidecar.get("paragraphs", [])
    if not isinstance(paragraphs, list):
        return changes

    for record in paragraphs:
        if not isinstance(record, dict):
            continue
        original = str(record.get("text", ""))
        if not original:
            continue
        current = original
        applied: list[str] = []
        for rule in channel_rules:
            rule_id = str(rule.get("id", "unnamed_rule"))
            if str(rule.get("scope", "all_paragraphs")) == "first_matching_paragraph" and rule_id in first_matching_applied:
                continue
            if "match" not in rule:
                continue
            if not rule_is_active_for_text(rule, current):
                continue
            rule_type = str(rule.get("type", "regex_replace"))
            if rule_type == "regex_replace":
                updated, did_change = apply_regex_rule(current, rule)
                if did_change and updated != current:
                    current = updated
                    applied.append(rule_id)
                    if str(rule.get("scope", "all_paragraphs")) == "first_matching_paragraph":
                        first_matching_applied.add(rule_id)
            elif rule_type in {"insert_before", "insert_after"}:
                if compile_regex(str(rule["match"])).search(current):
                    insert_text = str(rule.get("insert_text", ""))
                    if insert_text:
                        changes.append(
                            ParagraphChange(
                                kind=rule_type,
                                docx_paragraph_index=int(record.get("docx_paragraph_index", record.get("index", 0))),
                                source_record_index=int(record.get("index", 0)),
                                original_text=current,
                                new_text=insert_text,
                                changed_spans=[
                                    {
                                        "start": 0,
                                        "end": len(insert_text),
                                        "text": insert_text,
                                        "highlight": True,
                                        "highlight_color": "yellow",
                                        "source": "channel_builder_insert",
                                    }
                                ],
                                rules_applied=[rule_id],
                            )
                        )
                    if str(rule.get("scope", "all_paragraphs")) == "first_matching_paragraph":
                        first_matching_applied.add(rule_id)
        if current != original:
            changes.append(
                ParagraphChange(
                    kind="replace",
                    docx_paragraph_index=int(record.get("docx_paragraph_index", record.get("index", 0))),
                    source_record_index=int(record.get("index", 0)),
                    original_text=original,
                    new_text=current,
                    changed_spans=diff_highlight_spans(original, current),
                    rules_applied=applied,
                )
            )
    return changes


def output_path_for_channel(source_docx: Path, out_dir: Path, channel_config: dict[str, Any], channel: str) -> Path:
    suffix = str(channel_config.get("output_suffix") or f" - {channel.title()}")
    stem = source_docx.stem
    for existing in (" - Phone", " - PHONE", " - phone", " - SMS", " - Sms", " - sms"):
        if stem.endswith(existing):
            stem = stem[: -len(existing)]
            break
    return out_dir / f"{stem}{suffix}{source_docx.suffix}"


def copy_run_properties(source_run: Run | None, target_run: Run) -> None:
    if source_run is None or source_run._r.rPr is None:
        return
    target_rpr = target_run._r.get_or_add_rPr()
    target_run._r.remove(target_rpr)
    target_run._r.insert(0, deepcopy(source_run._r.rPr))


def clear_paragraph_runs(paragraph: Paragraph) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def add_text_with_highlights(paragraph: Paragraph, text: str, spans: list[dict[str, Any]], template_run: Run | None) -> None:
    merged = merge_spans(spans)
    cursor = 0
    for span in merged:
        start = max(0, min(int(span["start"]), len(text)))
        end = max(start, min(int(span["end"]), len(text)))
        if cursor < start:
            run = paragraph.add_run(text[cursor:start])
            copy_run_properties(template_run, run)
        if start < end:
            run = paragraph.add_run(text[start:end])
            copy_run_properties(template_run, run)
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        cursor = end
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        copy_run_properties(template_run, run)
    if not text:
        run = paragraph.add_run("")
        copy_run_properties(template_run, run)


def apply_changes_to_docx(
    source_docx: Path,
    output_docx: Path,
    changes: list[ParagraphChange],
    *,
    overwrite: bool,
) -> None:
    if output_docx.exists() and not overwrite:
        raise SystemExit(f"Output exists; pass --overwrite to replace: {output_docx}")
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(str(source_docx))
    replace_map = {change.docx_paragraph_index: change for change in changes if change.kind == "replace"}
    insert_before_map: dict[int, list[ParagraphChange]] = {}
    insert_after_map: dict[int, list[ParagraphChange]] = {}
    for change in changes:
        if change.kind == "insert_before":
            insert_before_map.setdefault(change.docx_paragraph_index, []).append(change)
        elif change.kind == "insert_after":
            insert_after_map.setdefault(change.docx_paragraph_index, []).append(change)

    original_paragraphs = list(doc.paragraphs)
    for idx, paragraph in enumerate(original_paragraphs):
        template_run = paragraph.runs[0] if paragraph.runs else None
        for change in insert_before_map.get(idx, []):
            inserted = paragraph.insert_paragraph_before()
            inserted.style = paragraph.style
            add_text_with_highlights(inserted, change.new_text, change.changed_spans, template_run)
        change = replace_map.get(idx)
        if change is not None:
            clear_paragraph_runs(paragraph)
            add_text_with_highlights(paragraph, change.new_text, change.changed_spans, template_run)
        for change in insert_after_map.get(idx, []):
            inserted = paragraph.insert_paragraph_before()
            inserted.style = paragraph.style
            add_text_with_highlights(inserted, change.new_text, change.changed_spans, template_run)
            paragraph._p.addnext(inserted._p)
    doc.save(str(output_docx))


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def build_channel(
    *,
    source_docx: Path,
    out_dir: Path,
    channel: str,
    channel_config: dict[str, Any],
    source_sidecar: dict[str, Any],
    overwrite: bool,
    write_formatting_sidecars: bool,
) -> ChannelBuild:
    output_docx = output_path_for_channel(source_docx, out_dir, channel_config, channel)
    rules = channel_config.get("rules", [])
    if not isinstance(rules, list):
        raise SystemExit(f"Channel '{channel}' rules must be a list.")
    changes = apply_channel_rules(source_sidecar, rules)
    apply_changes_to_docx(source_docx, output_docx, changes, overwrite=overwrite)

    formatting_path: Path | None = None
    if write_formatting_sidecars:
        formatting_path = output_docx.with_name(f"{output_docx.stem}.formatting.json")
        generated_sidecar = build_formatting_sidecar(str(output_docx), raw_order_source="docx")
        generated_sidecar["channel_builder"] = {
            "channel": channel,
            "source_docx": str(source_docx),
            "change_count": len(changes),
        }
        write_json(formatting_path, generated_sidecar)

    return ChannelBuild(
        channel=channel,
        output_docx=output_docx,
        output_formatting_json=formatting_path,
        changes=changes,
    )


def report_payload(
    *,
    source_docx: Path,
    rules_path: Path,
    source_sidecar: dict[str, Any],
    builds: list[ChannelBuild],
) -> dict[str, Any]:
    return {
        "schema_version": 1,
        "source_docx": str(source_docx),
        "rules_path": str(rules_path),
        "source_summary": source_sidecar.get("summary", {}),
        "channels": [
            {
                "channel": build.channel,
                "output_docx": str(build.output_docx),
                "output_formatting_json": str(build.output_formatting_json) if build.output_formatting_json else "",
                "change_count": len(build.changes),
                "changes": [
                    {
                        "source_record_index": change.source_record_index,
                        "kind": change.kind,
                        "docx_paragraph_index": change.docx_paragraph_index,
                        "rules_applied": change.rules_applied,
                        "original_text": change.original_text,
                        "new_text": change.new_text,
                        "changed_spans": change.changed_spans,
                    }
                    for change in build.changes
                ],
            }
            for build in builds
        ],
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Create Phone/SMS questionnaire DOCX drafts from a final DOCX using "
            "JSON channel rules and DOCX formatting sidecars."
        )
    )
    parser.add_argument("--docx", required=True, help="Source final questionnaire DOCX.")
    parser.add_argument(
        "--channels",
        nargs="+",
        default=["phone", "sms"],
        help="Channels to build. Defaults to: phone sms",
    )
    parser.add_argument(
        "--rules",
        default=str(DEFAULT_RULES_PATH),
        help="JSON rules file. Defaults to channel_rules.default.json.",
    )
    parser.add_argument("--out-dir", help="Output directory. Defaults to the source DOCX folder.")
    parser.add_argument("--report", help="Transform report JSON path.")
    parser.add_argument("--overwrite", action="store_true", help="Replace existing output DOCX/JSON files.")
    parser.add_argument(
        "--no-formatting-sidecars",
        action="store_true",
        help="Skip output .formatting.json sidecars.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    source_docx = Path(args.docx)
    if not source_docx.exists():
        raise SystemExit(f"Source DOCX not found: {source_docx}")
    rules_path = Path(args.rules)
    rules_payload = load_rules(rules_path)
    out_dir = Path(args.out_dir) if args.out_dir else source_docx.parent
    report_path = Path(args.report) if args.report else out_dir / f"{source_docx.stem}.channel_builder.report.json"

    source_sidecar = build_formatting_sidecar(str(source_docx), raw_order_source="docx")
    channels_config = rules_payload["channels"]
    builds: list[ChannelBuild] = []
    for channel in args.channels:
        if channel not in channels_config:
            available = ", ".join(sorted(channels_config))
            raise SystemExit(f"Channel '{channel}' is not in {rules_path}. Available: {available}")
        build = build_channel(
            source_docx=source_docx,
            out_dir=out_dir,
            channel=channel,
            channel_config=channels_config[channel],
            source_sidecar=source_sidecar,
            overwrite=args.overwrite,
            write_formatting_sidecars=not args.no_formatting_sidecars,
        )
        builds.append(build)

    write_json(
        report_path,
        report_payload(
            source_docx=source_docx,
            rules_path=rules_path,
            source_sidecar=source_sidecar,
            builds=builds,
        ),
    )

    print("[channel-builder] Complete.")
    for build in builds:
        print(f"[channel-builder] {build.channel}: {build.output_docx} ({len(build.changes)} changed paragraph(s))")
        if build.output_formatting_json:
            print(f"[channel-builder] {build.channel} formatting: {build.output_formatting_json}")
    print(f"[channel-builder] report: {report_path}")


if __name__ == "__main__":
    main()
