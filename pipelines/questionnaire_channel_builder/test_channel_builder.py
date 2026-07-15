import json
import sys
import unittest
from pathlib import Path

from docx import Document

THIS_DIR = Path(__file__).resolve().parent
if str(THIS_DIR) not in sys.path:
    sys.path.insert(0, str(THIS_DIR))

import channel_builder as builder  # noqa: E402


class ChannelBuilderTests(unittest.TestCase):
    def test_diff_highlight_spans_anchors_deletions_to_remaining_text(self) -> None:
        spans = builder.diff_highlight_spans("(Don't know)\t3", "Don't know\t3")

        self.assertEqual(spans[0]["text"], "Don't")
        self.assertEqual(spans[0]["highlight_color"], "yellow")

    def test_sms_rules_apply_expected_json_driven_changes(self) -> None:
        rules = json.loads(builder.DEFAULT_RULES_PATH.read_text(encoding="utf-8"))
        sidecar = {
            "paragraphs": [
                {
                    "index": 0,
                    "docx_paragraph_index": 0,
                    "text": "Hello, my name is ___ and I'm calling from The Poll. Your telephone number was selected.",
                },
                {
                    "index": 1,
                    "docx_paragraph_index": 1,
                    "text": "(Don't know)\t3",
                },
                {
                    "index": 2,
                    "docx_paragraph_index": 2,
                    "text": "Now, I'm going to read you a list of names and I'd like you to tell me whether you like them.",
                },
            ]
        }

        changes = builder.apply_channel_rules(sidecar, rules["channels"]["sms"]["rules"])

        self.assertEqual(len(changes), 3)
        self.assertEqual(changes[0].new_text, "STANDARD SMS INTRO")
        self.assertEqual(changes[1].new_text, "Don\u2019t know\t3")
        self.assertIn("you will read a list of names and select whether", changes[2].new_text)

    def test_apply_changes_to_docx_highlights_changed_text(self) -> None:
        scratch = THIS_DIR.parents[1] / ".tmp_tests" / "questionnaire_channel_builder"
        scratch.mkdir(parents=True, exist_ok=True)
        source = scratch / "source.docx"
        output = scratch / "output.docx"

        doc = Document()
        doc.add_paragraph("(Don't know)\t3")
        doc.save(str(source))

        change = builder.ParagraphChange(
            kind="replace",
            docx_paragraph_index=0,
            source_record_index=0,
            original_text="(Don't know)\t3",
            new_text="Don't know\t3",
            changed_spans=builder.diff_highlight_spans("(Don't know)\t3", "Don't know\t3"),
            rules_applied=["sms_remove_parentheses_from_response_options"],
        )
        builder.apply_changes_to_docx(source, output, [change], overwrite=True)

        result = Document(str(output))
        self.assertEqual(result.paragraphs[0].text, "Don't know\t3")
        highlighted_runs = [run for run in result.paragraphs[0].runs if run.font.highlight_color is not None]
        self.assertTrue(highlighted_runs)

    def test_insert_before_rule_adds_highlighted_paragraph(self) -> None:
        scratch = THIS_DIR.parents[1] / ".tmp_tests" / "questionnaire_channel_builder"
        scratch.mkdir(parents=True, exist_ok=True)
        source = scratch / "insert_source.docx"
        output = scratch / "insert_output.docx"

        doc = Document()
        doc.add_paragraph("Last question, choose one.")
        doc.save(str(source))

        change = builder.ParagraphChange(
            kind="insert_before",
            docx_paragraph_index=0,
            source_record_index=0,
            original_text="Last question, choose one.",
            new_text="[MAKE QUESTION SKIPPABLE]",
            changed_spans=[
                {
                    "start": 0,
                    "end": len("[MAKE QUESTION SKIPPABLE]"),
                    "text": "[MAKE QUESTION SKIPPABLE]",
                    "highlight": True,
                    "highlight_color": "yellow",
                }
            ],
            rules_applied=["sms_make_last_question_skippable"],
        )
        builder.apply_changes_to_docx(source, output, [change], overwrite=True)

        result = Document(str(output))
        self.assertEqual(result.paragraphs[0].text, "[MAKE QUESTION SKIPPABLE]")
        self.assertEqual(result.paragraphs[1].text, "Last question, choose one.")
        self.assertTrue(any(run.font.highlight_color is not None for run in result.paragraphs[0].runs))


if __name__ == "__main__":
    unittest.main()
