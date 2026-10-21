from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable, Optional
import argparse
import json
import math
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


@dataclass(frozen=True)
class ExactQuota:
    section: str
    category: str
    label: str
    code: str
    split: str
    gender: str
    target: int
    source: str


@dataclass(frozen=True)
class LimitQuota:
    section: str
    category: str
    label: str
    minimum: Optional[int]
    maximum: Optional[int]
    minimum_share: Optional[float]
    maximum_share: Optional[float]
    source: str


@dataclass(frozen=True)
class ScrapedQuotaWorkbook:
    path: str
    title: str
    sheet: str
    exact_quotas: list[ExactQuota]
    limit_quotas: list[LimitQuota]
    notes: list[str]
    messages: list[str]

    def section_totals(self) -> dict[str, int]:
        totals: dict[str, int] = {}
        for quota in self.exact_quotas:
            totals[quota.section] = totals.get(quota.section, 0) + quota.target
        return totals

    def to_dict(self) -> dict:
        return {
            "path": self.path,
            "title": self.title,
            "sheet": self.sheet,
            "section_totals": self.section_totals(),
            "exact_quotas": [asdict(quota) for quota in self.exact_quotas],
            "limit_quotas": [asdict(quota) for quota in self.limit_quotas],
            "notes": list(self.notes),
            "messages": list(self.messages),
        }


def _text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _key(value: object) -> str:
    return _text(value).casefold()


def _number(value: object) -> Optional[float]:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    text = _text(value).replace(",", "")
    if not text or text.startswith("="):
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return number if math.isfinite(number) else None


def _integer(value: object) -> Optional[int]:
    number = _number(value)
    if number is None:
        return None
    return int(math.floor(number + 0.5))


def _code(value: object) -> str:
    number = _number(value)
    if number is not None and number.is_integer():
        return str(int(number))
    return _text(value)


def _section_name(value: object) -> str:
    text = _text(value)
    key = text.casefold()
    if "oversample" in key:
        return text
    return "Base Sample"


def _is_section_banner(value: object) -> bool:
    key = _key(value)
    if any(word in key for word in ("please", "cap ", "must ")):
        return False
    return "oversample" in key or ("base" in key and ("sample" in key or "complete" in key))


def _split_name(value: object) -> str:
    text = _text(value)
    if not text:
        return ""
    match = re.search(r"\bsplit\s+(.+)$", text, flags=re.IGNORECASE)
    return _text(match.group(1) if match else text).upper()


class _WorkbookView:
    def __init__(self, path: Path) -> None:
        try:
            self.formula_book = load_workbook(path, data_only=False, read_only=False)
            self.value_book = load_workbook(path, data_only=True, read_only=False)
        except PermissionError as exc:
            raise PermissionError(
                f"Could not read quota workbook. Close it in Excel/OneDrive and try again: {path}"
            ) from exc
        sheet_name = "Quotas" if "Quotas" in self.formula_book.sheetnames else self.formula_book.sheetnames[0]
        self.formula_sheet = self.formula_book[sheet_name]
        self.value_sheet = self.value_book[sheet_name]

    @property
    def sheet_name(self) -> str:
        return self.formula_sheet.title

    def value(self, row: int, column: int) -> object:
        calculated = self.value_sheet.cell(row, column).value
        if calculated is not None:
            return calculated
        raw = self.formula_sheet.cell(row, column).value
        if isinstance(raw, str) and raw.startswith("="):
            return None
        return raw

    def raw(self, row: int, column: int) -> object:
        return self.formula_sheet.cell(row, column).value

    def merged_value(self, row: int, column: int) -> object:
        value = self.value(row, column)
        if value is not None:
            return value
        coordinate = self.formula_sheet.cell(row, column).coordinate
        for merged in self.formula_sheet.merged_cells.ranges:
            if coordinate in merged:
                return self.value(merged.min_row, merged.min_col)
        return None


def _row_texts(view: _WorkbookView, row: int) -> list[tuple[int, str]]:
    values: list[tuple[int, str]] = []
    for column in range(1, view.formula_sheet.max_column + 1):
        text = _text(view.value(row, column))
        if text:
            values.append((column, text))
    return values


def _find_title(view: _WorkbookView) -> str:
    for row in range(1, min(view.formula_sheet.max_row, 10) + 1):
        for _column, text in _row_texts(view, row):
            return text
    return Path(view.formula_book.properties.title or "").name


def _section_by_row(view: _WorkbookView) -> dict[int, str]:
    sections: dict[int, str] = {}
    current = "Base Sample"
    for row in range(1, view.formula_sheet.max_row + 1):
        row_values = _row_texts(view, row)
        banner = (
            row_values[0][1]
            if row > 1 and len(row_values) == 1 and _is_section_banner(row_values[0][1])
            else ""
        )
        if banner:
            current = _section_name(banner)
        sections[row] = current
    return sections


def _region_header(column_text: str) -> bool:
    key = _key(column_text)
    return "region" in key and "code" not in key and "punch" not in key


def _code_header(column_text: str) -> bool:
    key = _key(column_text)
    return key == "code" or "region code" in key or "region punch" in key


def _gender_header(column_text: str) -> bool:
    return _key(column_text) in {"men", "women"}


def _exact_header_kind(view: _WorkbookView, row: int, section: str) -> tuple[str, int, int, list[int]] | None:
    headers = {column: text for column, text in _row_texts(view, row)}
    gender_columns = [column for column, text in headers.items() if _gender_header(text)]
    total_columns = [column for column, text in headers.items() if _key(text) == "total"]
    if total_columns:
        first_total = min(total_columns)
        gender_columns = [column for column in gender_columns if column < first_total]
    if not gender_columns or not any(_key(headers[column]) == "men" for column in gender_columns):
        return None
    if not any(_key(headers[column]) == "women" for column in gender_columns):
        return None

    region_columns = [column for column, text in headers.items() if _region_header(text)]
    code_columns = [column for column, text in headers.items() if _code_header(text)]
    if region_columns and code_columns:
        return "region", region_columns[0], code_columns[0], gender_columns
    if "oversample" in section.casefold():
        label_column = max(1, min(gender_columns) - 1)
        return "oversample", label_column, 0, gender_columns
    return None


def _parse_exact_quotas(
    view: _WorkbookView,
    sections: dict[int, str],
) -> tuple[list[ExactQuota], set[int]]:
    quotas: list[ExactQuota] = []
    consumed_rows: set[int] = set()
    ws = view.formula_sheet
    for header_row in range(1, ws.max_row + 1):
        section = sections[header_row]
        header = _exact_header_kind(view, header_row, section)
        if header is None:
            continue
        kind, label_column, code_column, gender_columns = header
        consumed_rows.add(header_row)

        split_by_column: dict[int, str] = {}
        if header_row > 1:
            for column in gender_columns:
                above = view.merged_value(header_row - 1, column)
                split_by_column[column] = _split_name(above) if "split" in _key(above) else ""
            if any(split_by_column.values()):
                consumed_rows.add(header_row - 1)

        for row in range(header_row + 1, ws.max_row + 1):
            if row != header_row + 1 and any(_is_section_banner(text) for _column, text in _row_texts(view, row)):
                break
            label = _text(view.value(row, label_column))
            code = _code(view.value(row, code_column)) if code_column else ""
            targets = {column: _integer(view.value(row, column)) for column in gender_columns}
            has_targets = any(target is not None for target in targets.values())
            if _key(label) == "total":
                consumed_rows.add(row)
                break
            if not label and not code and not has_targets:
                break
            if kind == "region" and (not label or not code):
                if has_targets:
                    consumed_rows.add(row)
                break
            if not has_targets:
                break

            consumed_rows.add(row)
            header_label = _text(view.value(header_row, label_column))
            if kind == "region":
                category = "Region"
            elif header_label:
                category = header_label
            elif _key(label) == "sex":
                category = "Sex"
            else:
                category = "Oversample"
            for column in gender_columns:
                target = targets[column]
                if target is None:
                    continue
                quotas.append(
                    ExactQuota(
                        section=section,
                        category=category,
                        label=label,
                        code=code,
                        split=split_by_column.get(column, ""),
                        gender=_text(view.value(header_row, column)).title(),
                        target=target,
                        source=f"{view.sheet_name}!{get_column_letter(column)}{row}",
                    )
                )
    return quotas, consumed_rows


def _limit_pairs(view: _WorkbookView, row: int) -> list[tuple[int, int]]:
    pairs: list[tuple[int, int]] = []
    for column in range(1, view.formula_sheet.max_column):
        if _key(view.value(row, column)) == "minimum" and _key(view.value(row, column + 1)) == "maximum":
            pairs.append((column, column + 1))
    return pairs


def _share(value: object) -> Optional[float]:
    number = _number(value)
    if number is None or number < 0 or number > 1:
        return None
    return number


def _parse_limit_quotas(
    view: _WorkbookView,
    sections: dict[int, str],
    section_totals: dict[str, int],
) -> tuple[list[LimitQuota], set[int]]:
    quotas: list[LimitQuota] = []
    consumed_rows: set[int] = set()
    ws = view.formula_sheet
    for header_row in range(1, ws.max_row + 1):
        pairs = _limit_pairs(view, header_row)
        if not pairs:
            continue
        first_min, first_max = pairs[0]
        label_column = first_min - 1
        category = _text(view.value(header_row, label_column))
        if not category or _key(category) in {"minimum", "maximum"}:
            continue
        section = sections[header_row]
        total = section_totals.get(section, 0)
        consumed_rows.add(header_row)

        for row in range(header_row + 1, ws.max_row + 1):
            label = _text(view.value(row, label_column))
            if not label:
                break
            if _key(label) in {"total", "minimum", "maximum"}:
                break

            minimum = _integer(view.value(row, first_min))
            maximum = _integer(view.value(row, first_max))
            minimum_share: Optional[float] = None
            maximum_share: Optional[float] = None
            for share_min, share_max in pairs[1:]:
                candidate_min = _share(view.value(row, share_min))
                candidate_max = _share(view.value(row, share_max))
                if candidate_min is not None:
                    minimum_share = candidate_min
                if candidate_max is not None:
                    maximum_share = candidate_max

            if minimum is None and minimum_share is not None and total:
                minimum = _integer(minimum_share * total)
            if maximum is None and maximum_share is not None and total:
                maximum = _integer(maximum_share * total)
            if minimum is None and maximum is None:
                break

            consumed_rows.add(row)
            cells = [
                f"{get_column_letter(column)}{row}"
                for column in (first_min, first_max)
                if view.raw(row, column) is not None
            ]
            quotas.append(
                LimitQuota(
                    section=section,
                    category=category,
                    label=label,
                    minimum=minimum,
                    maximum=maximum,
                    minimum_share=minimum_share,
                    maximum_share=maximum_share,
                    source=f"{view.sheet_name}!{','.join(cells)}",
                )
            )
    return quotas, consumed_rows


def _collect_notes(
    view: _WorkbookView,
    title: str,
    consumed_rows: set[int],
) -> list[str]:
    notes: list[str] = []
    for row in range(1, view.formula_sheet.max_row + 1):
        values = _row_texts(view, row)
        if row in consumed_rows or len(values) != 1:
            continue
        _column, text = values[0]
        if text == title or _is_section_banner(text):
            continue
        key = text.casefold()
        if "please" in key or "cap" in key or "must" in key:
            notes.append(f"{view.sheet_name}: {text}")

    for ws in view.value_book.worksheets:
        if ws.title == view.sheet_name:
            continue
        for row in ws.iter_rows():
            for cell in row:
                text = _text(cell.value)
                if text:
                    notes.append(f"{ws.title}: {text}")
    return list(dict.fromkeys(notes))


def scrape_quota_workbook(path: Path | str) -> ScrapedQuotaWorkbook:
    quota_path = Path(path)
    view = _WorkbookView(quota_path)
    title = _find_title(view)
    sections = _section_by_row(view)
    exact_quotas, exact_rows = _parse_exact_quotas(view, sections)
    section_totals: dict[str, int] = {}
    for quota in exact_quotas:
        section_totals[quota.section] = section_totals.get(quota.section, 0) + quota.target
    limit_quotas, limit_rows = _parse_limit_quotas(view, sections, section_totals)
    notes = _collect_notes(view, title, exact_rows | limit_rows)

    messages: list[str] = []
    if not exact_quotas:
        messages.append("No exact quota grid was detected.")
    if not limit_quotas:
        messages.append("No minimum/maximum quota blocks were detected.")
    return ScrapedQuotaWorkbook(
        path=str(quota_path),
        title=title,
        sheet=view.sheet_name,
        exact_quotas=exact_quotas,
        limit_quotas=limit_quotas,
        notes=notes,
        messages=messages,
    )


def scrape_quota_folder(path: Path | str) -> list[ScrapedQuotaWorkbook]:
    root = Path(path)
    workbooks = sorted(
        [*root.glob("*.xlsx"), *root.glob("*.xlsm")],
        key=lambda item: item.name.casefold(),
    )
    return [scrape_quota_workbook(workbook) for workbook in workbooks]


def _limit_text(value: Optional[int], share: Optional[float]) -> str:
    if value is None:
        return ""
    if share is None:
        return str(value)
    return f"{value} ({share:.0%})"


def render_markdown_report(results: Iterable[ScrapedQuotaWorkbook]) -> str:
    lines = [
        "# Quota workbook scrape verification",
        "",
        "This is a read-only extraction preview. Exact cells, min/max limits, oversamples, and workbook notes are listed for verification.",
    ]
    for result in results:
        lines.extend(
            [
                "",
                f"## {Path(result.path).name}",
                "",
                f"- Title: {result.title}",
                f"- Parsed sheet: {result.sheet}",
            ]
        )
        totals = result.section_totals()
        for section, total in totals.items():
            exact_count = sum(1 for quota in result.exact_quotas if quota.section == section)
            limit_count = sum(1 for quota in result.limit_quotas if quota.section == section)
            lines.append(f"- {section}: target {total}; {exact_count} exact cell(s); {limit_count} min/max quota(s)")

        sections = list(dict.fromkeys(
            [quota.section for quota in result.exact_quotas]
            + [quota.section for quota in result.limit_quotas]
        ))
        for section in sections:
            exact = [quota for quota in result.exact_quotas if quota.section == section]
            limits = [quota for quota in result.limit_quotas if quota.section == section]
            if exact:
                lines.extend(
                    [
                        "",
                        f"### {section} — exact cells",
                        "",
                        "| Category | Label | Code | Split | Gender | Target | Source |",
                        "|---|---|---:|---|---|---:|---|",
                    ]
                )
                for quota in exact:
                    lines.append(
                        f"| {quota.category} | {quota.label} | {quota.code} | {quota.split} | "
                        f"{quota.gender} | {quota.target} | {quota.source} |"
                    )
            if limits:
                lines.extend(
                    [
                        "",
                        f"### {section} — minimum/maximum quotas",
                        "",
                        "| Category | Label | Minimum | Maximum | Source |",
                        "|---|---|---:|---:|---|",
                    ]
                )
                for quota in limits:
                    lines.append(
                        f"| {quota.category} | {quota.label} | "
                        f"{_limit_text(quota.minimum, quota.minimum_share)} | "
                        f"{_limit_text(quota.maximum, quota.maximum_share)} | {quota.source} |"
                    )

        if result.notes:
            lines.extend(["", "### Notes and instructions", ""])
            lines.extend(f"- {note}" for note in result.notes)
        if result.messages:
            lines.extend(["", "### Parser messages", ""])
            lines.extend(f"- {message}" for message in result.messages)
    lines.append("")
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Extract quota definitions from related quota workbooks.")
    parser.add_argument("path", type=Path, help="Quota workbook or folder containing quota workbooks")
    parser.add_argument("--markdown", type=Path, help="Write a human-readable verification report")
    parser.add_argument("--json", type=Path, dest="json_path", help="Write the extracted structure as JSON")
    args = parser.parse_args(argv)

    results = scrape_quota_folder(args.path) if args.path.is_dir() else [scrape_quota_workbook(args.path)]
    if args.markdown:
        args.markdown.parent.mkdir(parents=True, exist_ok=True)
        args.markdown.write_text(render_markdown_report(results), encoding="utf-8")
    if args.json_path:
        args.json_path.parent.mkdir(parents=True, exist_ok=True)
        args.json_path.write_text(
            json.dumps([result.to_dict() for result in results], indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
    if not args.markdown and not args.json_path:
        print(render_markdown_report(results))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
