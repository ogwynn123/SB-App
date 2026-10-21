from __future__ import annotations

import sys
import unittest
from pathlib import Path

from openpyxl import Workbook


THIS_DIR = Path(__file__).resolve().parent
if str(THIS_DIR) not in sys.path:
    sys.path.insert(0, str(THIS_DIR))

from quota_workbook_scraper import scrape_quota_workbook  # noqa: E402


class QuotaWorkbookScraperTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp_dir = THIS_DIR / ".tmp_tests" / "scraper"
        self.tmp_dir.mkdir(parents=True, exist_ok=True)

    def test_scrapes_unsplit_base_and_all_limit_categories(self) -> None:
        path = self.tmp_dir / "unsplit.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Quotas"
        ws.append(["Test quotas"])
        ws.append([])
        ws.append(["IN-01 Region", "Code", "Men", "Women", "Total"])
        ws.append(["North", 1, 18, 22, "=SUM(C4:D4)"])
        ws.append(["Total", "", "=SUM(C4:C4)", "=SUM(D4:D4)", "=SUM(E4:E4)"])
        ws.append([])
        ws.append(["Age", "Minimum", "Maximum", "", "Age", "Minimum", "Maximum"])
        ws.append(["18-34", "=F8*$E$5", "", "", "18-34", 0.25, ""])
        ws.append([])
        ws.append(["Race", "Minimum", "Maximum", "", "Race", "Minimum", "Maximum"])
        ws.append(["White", "", "=G11*$E$5", "", "White", "", 0.8])
        ws.append([])
        ws.append(["Education", "Minimum", "Maximum", "", "Education", "Minimum", "Maximum"])
        ws.append(["College grad", "", 16, "", "College grad", "", 0.4])
        wb.save(path)

        result = scrape_quota_workbook(path)
        self.assertEqual(result.section_totals(), {"Base Sample": 40})
        self.assertEqual(len(result.exact_quotas), 2)
        limits = {(quota.category, quota.label): quota for quota in result.limit_quotas}
        self.assertEqual(limits[("Age", "18-34")].minimum, 10)
        self.assertEqual(limits[("Race", "White")].maximum, 32)
        self.assertEqual(limits[("Education", "College grad")].maximum, 16)

    def test_scrapes_merged_splits_oversample_and_notes(self) -> None:
        path = self.tmp_dir / "split_oversample.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "Split quotas"
        ws["A3"] = "PLEASE CAP SMS COMPLETES AT 40%"
        ws["A5"] = "BASE SAMPLE"
        ws["C7"] = "Split A"
        ws.merge_cells("C7:D7")
        ws["E7"] = "Split B"
        ws.merge_cells("E7:F7")
        for column, value in enumerate(
            ["Region Name", "Region Punch", "Men", "Women", "Men", "Women", "Total"], start=1
        ):
            ws.cell(8, column).value = value
        ws.append(["North", 1, 10, 11, 10, 9, "=SUM(C9:F9)"])
        ws["A12"] = "BLACK ON FILE OVERSAMPLE"
        ws["B14"] = "SPLIT A"
        ws.merge_cells("B14:C14")
        ws["D14"] = "SPLIT B"
        ws.merge_cells("D14:E14")
        for column, value in enumerate(["Race", "Men", "Women", "Men", "Women", "Total"], start=1):
            ws.cell(15, column).value = value
        for column, value in enumerate(["Black on File", 5, 6, 7, 8, "=SUM(B16:E16)"], start=1):
            ws.cell(16, column).value = value
        notes = wb.create_sheet("Regional Definitions")
        notes["A1"] = "Use the region field in sample"
        wb.save(path)

        result = scrape_quota_workbook(path)
        self.assertEqual(result.section_totals()["Base Sample"], 40)
        self.assertEqual(result.section_totals()["BLACK ON FILE OVERSAMPLE"], 26)
        base_splits = [quota.split for quota in result.exact_quotas if quota.section == "Base Sample"]
        self.assertEqual(base_splits, ["A", "A", "B", "B"])
        self.assertTrue(any("CAP SMS" in note for note in result.notes))
        self.assertTrue(any("region field" in note for note in result.notes))


if __name__ == "__main__":
    unittest.main()
