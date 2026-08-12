from __future__ import annotations

import sys
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook


THIS_DIR = Path(__file__).resolve().parent
if str(THIS_DIR) not in sys.path:
    sys.path.insert(0, str(THIS_DIR))

from sample_quota_checker_core import (  # noqa: E402
    BaseQuota,
    QuotaSpec,
    RangeQuota,
    check_sample_against_quota,
    load_quota_spec,
    quota_spec_from_dict,
    quota_spec_to_dict,
    suggest_column_mapping,
)
from sample_quota_checker_settings import (  # noqa: E402
    clear_quota_override,
    load_quota_override,
    load_saved_mapping,
    load_saved_quota_for_sample,
    save_mapping,
    save_quota_override,
)
from sample_quota_checker_gui import _dialog_initial_directory  # noqa: E402


class SampleQuotaCheckerTests(unittest.TestCase):
    def test_launcher_work_directory_is_used_as_dialog_fallback(self) -> None:
        fallback = THIS_DIR / ".tmp_tests"
        fallback.mkdir(exist_ok=True)
        self.assertEqual(_dialog_initial_directory("", fallback), str(fallback.resolve()))

    def test_parse_quota_workbook_and_check_sample(self) -> None:
        tmp_dir = THIS_DIR / ".tmp_tests"
        tmp_dir.mkdir(exist_ok=True)
        quota_path = tmp_dir / "quota.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "Quotas"
        ws["A5"] = ""
        ws["C5"] = "Split A/C"
        ws["E5"] = "Split A/D"
        ws["G5"] = "Split B/C"
        ws["I5"] = "Split B/D"
        headers = ["Region", "Region Code", "Men", "Women", "Men", "Women", "Men", "Women", "Men", "Women", "Total"]
        for col, value in enumerate(headers, start=1):
            ws.cell(6, col).value = value
        ws.append(["Cincinnati", 1, 2, 3, 4, 5, 6, 7, 8, 9, 44])
        ws.append(["The Neck", 2, 1, 1, 1, 1, 1, 1, 1, 1, 8])
        ws.append(["Rest of CD1", 3, 1, 1, 1, 1, 1, 1, 1, 1, 8])
        ws["A13"] = "Race"
        ws["B13"] = "Minimum"
        ws["C13"] = "Maximum"
        ws["E13"] = "Race"
        ws["F13"] = "Minimum"
        ws["G13"] = "Maximum"
        ws["A14"] = "White"
        ws["C14"] = 52
        ws["E14"] = "White"
        ws["G14"] = 0.87
        ws["A15"] = "Black"
        ws["B15"] = 4
        ws["E15"] = "Black"
        ws["F15"] = 0.07
        wb.save(quota_path)

        spec = load_quota_spec(quota_path)
        self.assertEqual(len(spec.base_quotas), 24)
        self.assertEqual(spec.target_total, 60)
        self.assertTrue(any(q.label == "White" and q.maximum == 52 for q in spec.range_quotas))
        self.assertTrue(any(q.label == "Black" and q.minimum == 4 for q in spec.range_quotas))

        sample = pd.DataFrame(
            [
                {"regioncode": 1, "gendergrp": 1, "num_ab": 1, "num_cd": 1, "race": "W", "stype": 1},
                {"regioncode": 1, "gendergrp": 1, "num_ab": 1, "num_cd": 1, "race": "W", "stype": 1},
                {"regioncode": 1, "gendergrp": 1, "num_ab": 1, "num_cd": 1, "race": "B", "stype": 1},
                {"regioncode": 1, "gendergrp": 1, "num_ab": 1, "num_cd": 1, "race": "B", "stype": 2},
                {"regioncode": 1, "gendergrp": 2, "num_ab": 1, "num_cd": 1, "race": "B", "stype": 1},
            ]
        )
        result = check_sample_against_quota(sample, spec)
        over = [row for row in result.rows if row.group == "Cincinnati | Split A/C | Men"][0]
        self.assertEqual(over.current, 3)
        self.assertEqual(over.status, "Over")
        oversample = [row for row in result.rows if row.section == "Sample Type" and row.group == "Oversample"][0]
        self.assertEqual(oversample.current, 1)
        black = [row for row in result.rows if row.section == "Race" and row.group == "Black"][0]
        self.assertEqual(black.status, "Under Min")

    def test_custom_column_mapping_and_saved_mapping(self) -> None:
        spec = QuotaSpec(
            base_quotas=[BaseQuota("Cincinnati", "1", "A/C", "Men", 1)],
            range_quotas=[RangeQuota("Race", "Black", 1, None)],
            target_total=1,
        )
        sample = pd.DataFrame(
            [
                {
                    "region_custom": 1,
                    "sex_custom": "M",
                    "ab_custom": "A",
                    "cd_custom": "C",
                    "race_custom": "B",
                    "sample_kind": 1,
                },
                {
                    "region_custom": 1,
                    "sex_custom": "M",
                    "ab_custom": "A",
                    "cd_custom": "C",
                    "race_custom": "B",
                    "sample_kind": 2,
                },
            ]
        )
        mapping = {
            "region": "region_custom",
            "gender": "sex_custom",
            "split_ab": "ab_custom",
            "split_cd": "cd_custom",
            "race": "race_custom",
            "stype": "sample_kind",
        }
        result = check_sample_against_quota(sample, spec, mapping=mapping)
        cell = [row for row in result.rows if row.section == "Base Sample"][0]
        self.assertEqual(cell.current, 1)
        self.assertEqual(cell.status, "Full")
        oversample = [row for row in result.rows if row.section == "Sample Type" and row.group == "Oversample"][0]
        self.assertEqual(oversample.current, 1)

        tmp_dir = THIS_DIR / ".tmp_tests"
        tmp_dir.mkdir(exist_ok=True)
        store_path = tmp_dir / "column_mappings.json"
        quota_path = tmp_dir / "quota_custom.xlsx"
        sample_path = tmp_dir / "sample_custom.csv"
        save_mapping(quota_path, sample_path, mapping, store_path=store_path)
        self.assertEqual(load_saved_mapping(quota_path, sample_path, store_path=store_path), mapping)
        self.assertEqual(
            load_saved_quota_for_sample(sample_path, store_path=store_path),
            str(quota_path.resolve(strict=False)),
        )

    def test_suggest_column_mapping(self) -> None:
        sample = pd.DataFrame(columns=["regnum", "sex", "s_ab", "s_cd", "qracegrp", "sample_type"])
        mapping = suggest_column_mapping(sample)
        self.assertEqual(mapping["region"], "regnum")
        self.assertEqual(mapping["gender"], "sex")
        self.assertEqual(mapping["split_ab"], "s_ab")
        self.assertEqual(mapping["split_cd"], "s_cd")
        self.assertEqual(mapping["race"], "qracegrp")
        self.assertEqual(mapping["stype"], "sample_type")

    def test_blank_base_columns_allowed_when_only_range_quotas_exist(self) -> None:
        spec = QuotaSpec(
            base_quotas=[],
            range_quotas=[RangeQuota("Race", "Black", 2, None)],
            target_total=0,
        )
        sample = pd.DataFrame(
            [
                {"race_only": "B"},
                {"race_only": "Black"},
                {"race_only": "W"},
            ]
        )
        result = check_sample_against_quota(
            sample,
            spec,
            mapping={
                "region": "",
                "gender": "",
                "split_ab": "",
                "split_cd": "",
                "race": "race_only",
                "stype": "",
            },
        )

        black = [row for row in result.rows if row.section == "Race" and row.group == "Black"][0]
        self.assertEqual(black.current, 2)
        self.assertEqual(black.status, "OK")
        self.assertTrue(any("region/gender/split mappings are optional" in msg for msg in result.messages))

    def test_blank_race_column_allowed_when_no_range_quotas_exist(self) -> None:
        spec = QuotaSpec(
            base_quotas=[BaseQuota("Cincinnati", "1", "A/C", "Men", 1)],
            range_quotas=[],
            target_total=1,
        )
        sample = pd.DataFrame(
            [
                {"regioncode": 1, "gendergrp": 1, "num_ab": 1, "num_cd": 1},
            ]
        )
        result = check_sample_against_quota(
            sample,
            spec,
            mapping={
                "region": "regioncode",
                "gender": "gendergrp",
                "split_ab": "num_ab",
                "split_cd": "num_cd",
                "race": "",
                "stype": "",
            },
        )

        cell = [row for row in result.rows if row.section == "Base Sample"][0]
        self.assertEqual(cell.current, 1)
        self.assertEqual(cell.status, "Full")
        self.assertTrue(any("race mapping is optional" in msg for msg in result.messages))

    def test_custom_quota_can_use_sample_column_and_scope(self) -> None:
        spec = QuotaSpec(
            base_quotas=[],
            range_quotas=[
                RangeQuota("Age", "Under 50", 1, 1, "agegrp", "1, 2", "base"),
                RangeQuota("Age", "Under 50 - All", 2, 2, "agegrp", "1, 2", "all"),
            ],
            target_total=0,
        )
        sample = pd.DataFrame(
            [
                {"agegrp": 1, "stype": 1},
                {"agegrp": 2, "stype": 2},
                {"agegrp": 3, "stype": 1},
            ]
        )

        result = check_sample_against_quota(sample, spec)
        rows = {row.group: row for row in result.rows if row.section == "Age"}
        self.assertEqual(rows["Under 50"].current, 1)
        self.assertEqual(rows["Under 50"].status, "OK")
        self.assertIn("base sample rows", rows["Under 50"].detail)
        self.assertEqual(rows["Under 50 - All"].current, 2)
        self.assertEqual(rows["Under 50 - All"].status, "OK")
        self.assertIn("all sample rows", rows["Under 50 - All"].detail)

    def test_quota_definition_round_trip_and_override_storage(self) -> None:
        spec = QuotaSpec(
            base_quotas=[BaseQuota("North", "10", "A/C", "Women", 12)],
            range_quotas=[RangeQuota("Age", "18-34", 8, 20, "agegrp", "1", "all")],
            target_total=999,
        )
        definition = quota_spec_to_dict(spec)
        restored = quota_spec_from_dict(definition)
        self.assertEqual(restored.target_total, 12)
        self.assertEqual(restored.base_quotas, spec.base_quotas)
        self.assertEqual(restored.range_quotas, spec.range_quotas)

        tmp_dir = THIS_DIR / ".tmp_tests"
        tmp_dir.mkdir(exist_ok=True)
        store_path = tmp_dir / "quota_overrides.json"
        quota_path = tmp_dir / "reviewed_quota.xlsx"
        save_quota_override(quota_path, definition, store_path=store_path)
        self.assertEqual(load_quota_override(quota_path, store_path=store_path), definition)
        self.assertTrue(clear_quota_override(quota_path, store_path=store_path))
        self.assertEqual(load_quota_override(quota_path, store_path=store_path), {})


if __name__ == "__main__":
    unittest.main()
