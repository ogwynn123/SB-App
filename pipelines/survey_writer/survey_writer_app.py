#!/usr/bin/env python3
"""Lightweight survey draft builder.

This app is intentionally separate from the production workspace flow. It saves
small draft JSON files and can export a readable questionnaire to HTML or DOCX.
"""

from __future__ import annotations

import html
import json
import os
import re
import sqlite3
import sys
import time
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
import tkinter as tk


def _configure_tcl_env_early() -> None:
    if os.environ.get("TCL_LIBRARY") and os.environ.get("TK_LIBRARY"):
        return

    py_dir = Path(sys.executable).resolve().parent
    for root in (py_dir / "tcl", py_dir.parent / "tcl", Path.home() / "tcl"):
        tcl_dir = root / "tcl8.6"
        tk_dir = root / "tk8.6"
        if (tcl_dir / "init.tcl").exists() and (tk_dir / "tk.tcl").exists():
            os.environ.setdefault("TCL_LIBRARY", str(tcl_dir))
            os.environ.setdefault("TK_LIBRARY", str(tk_dir))
            return


_configure_tcl_env_early()


BASE_DIR = Path(__file__).resolve().parents[2]
DRAFT_ROOT = BASE_DIR / "drafts" / "survey_writer"
QUESTION_DB = BASE_DIR / "repository" / "questions.sqlite"

QUESTION_TEMPLATES = (
    ("Single choice", "single_choice"),
    ("Open end", "open_end"),
    ("Ballot with leans", "ballot_with_leans"),
    ("Favorability 0-10", "favorability_0_10"),
    ("Favorability 4pt battery", "favorability_4pt_battery"),
    ("Job approval battery", "job_approval_battery"),
    ("Describes well battery", "describes_well_battery"),
    ("More likely battery", "more_likely_battery"),
    ("Less likely battery", "less_likely_battery"),
    ("Q100+ demographics block", "demographics_block"),
)


@dataclass
class SearchResult:
    question_id: int
    survey_id: str
    qid: str
    qnum: str
    question_text: str
    item_text: str
    option_texts: list[str]
    question_type: str
    question_family: str
    question_group: str
    question_target: str
    state_code: str
    region: str
    client: str


def clean_text(value: object) -> str:
    text = str(value or "")
    text = text.replace("\r", " ").replace("\n", " ").replace("\t", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def safe_file_stem(value: str, fallback: str = "survey_draft") -> str:
    stem = re.sub(r"[^A-Za-z0-9._ -]+", "", value or "").strip()
    stem = re.sub(r"\s+", "_", stem)
    return stem or fallback


def new_draft() -> dict:
    return {
        "version": 1,
        "survey_id": "",
        "title": "Untitled Survey Draft",
        "created_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
        "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
        "questions": [],
    }


def normalize_option_lines(raw: str) -> list[dict[str, str]]:
    options: list[dict[str, str]] = []
    next_value = 1
    for line in str(raw or "").splitlines():
        line = line.strip()
        if not line:
            continue
        match = re.match(r"^([A-Za-z0-9]+)\s*[\.\)]\s+(.+)$", line)
        if match:
            value = match.group(1).strip()
            text = match.group(2).strip()
        else:
            value = str(next_value)
            text = line
        options.append({"value": value, "text": text})
        if value.isdigit():
            next_value = max(next_value, int(value) + 1)
        else:
            next_value += 1
    return options


def option_lines(options: list[dict[str, str]]) -> str:
    lines = []
    for i, opt in enumerate(options, start=1):
        value = str(opt.get("value") or i)
        text = str(opt.get("text") or "")
        lines.append(f"{value}. {text}")
    return "\n".join(lines)


def make_question(
    *,
    qnum: int | str,
    qid: str | None = None,
    question_type: str = "single_choice",
    text: str = "",
    options: list[str] | list[dict[str, str]] | None = None,
    instructions: list[str] | None = None,
    item_text: str = "",
    linked_group: str = "",
    source: dict | None = None,
) -> dict:
    qnum_text = str(qnum)
    if options and isinstance(options[0], str):  # type: ignore[index]
        normalized_options = [
            {"value": str(i), "text": str(label)}
            for i, label in enumerate(options, start=1)  # type: ignore[arg-type]
        ]
    else:
        normalized_options = [
            {"value": str(opt.get("value") or i), "text": str(opt.get("text") or "")}
            for i, opt in enumerate(options or [], start=1)  # type: ignore[union-attr]
        ]
    return {
        "qnum": qnum_text,
        "qid": qid or qnum_text,
        "question_type": question_type,
        "text": text,
        "instructions": instructions or [],
        "item_text": item_text,
        "response_options": normalized_options,
        "linked_group": linked_group,
        "source": source or {},
    }


def max_numeric_qnum(questions: list[dict]) -> int:
    max_value = 0
    for q in questions:
        qnum = str(q.get("qnum") or "")
        match = re.match(r"^(\d+)", qnum)
        if match:
            max_value = max(max_value, int(match.group(1)))
    return max_value


def next_qnum(questions: list[dict], minimum: int = 1) -> int:
    return max(minimum - 1, max_numeric_qnum(questions)) + 1


def linked_group_id(prefix: str) -> str:
    return f"{prefix}_{int(time.time() * 1000)}"


def letter(index: int) -> str:
    return chr(ord("a") + index)


def template_questions(kind: str, start_qnum: int) -> list[dict]:
    q = start_qnum
    if kind == "single_choice":
        return [
            make_question(
                qnum=q,
                question_type="single_choice",
                text="Question text goes here.",
                options=["Yes", "No", "(Don't know)"],
            )
        ]
    if kind == "open_end":
        return [
            make_question(
                qnum=q,
                question_type="open_end",
                text="What is the most important issue you would like to see addressed?",
            )
        ]
    if kind == "ballot_with_leans":
        return [
            make_question(
                qnum=q,
                question_type="head_to_head_ballot",
                text=(
                    "If the election were held today and the candidates were: Candidate One, "
                    "the Democrat, and Candidate Two, the Republican, for whom would you vote? "
                    "Well, toward whom do you lean?"
                ),
                options=[
                    "Candidate One",
                    "(Lean Candidate One)",
                    "Candidate Two",
                    "(Lean Candidate Two)",
                    "(Undecided)",
                ],
            )
        ]
    if kind == "favorability_0_10":
        group = linked_group_id("fav010")
        text = (
            "Now, I'm going to read you a list of public figures and I'd like you to tell me "
            "whether you have a favorable or unfavorable opinion of each person using a scale "
            "that runs from 0 to 10. A 0 means very low or very unfavorable and a 10 means "
            "very high or very favorable. If you have never heard of a name, or don't know "
            "enough to rate them, please just say so."
        )
        options = [str(i) for i in range(0, 11)] + ["Can't rate", "Don't know"]
        return [
            make_question(
                qnum=q,
                qid=f"{q}{letter(i)}",
                question_type="favorability_0_10",
                text=text,
                item_text=f"{letter(i)}. {name}",
                options=options,
                linked_group=group,
            )
            for i, name in enumerate(("Candidate One", "Candidate Two", "Donald Trump"))
        ]
    if kind == "favorability_4pt_battery":
        group = linked_group_id("fav4")
        text = (
            "Now, I'm going to read you the names of some public figures. For each one, "
            "please tell me whether you have a very favorable, somewhat favorable, somewhat "
            "unfavorable or very unfavorable opinion of that person."
        )
        options = [
            "Very favorable",
            "Somewhat favorable",
            "(Mixed)",
            "Somewhat unfavorable",
            "Very unfavorable",
            "Can't rate/Never heard",
        ]
        return [
            make_question(
                qnum=q,
                qid=f"{q}{letter(i)}",
                question_type="favorability_4pt",
                text=text,
                item_text=f"{letter(i)}. {name}",
                options=options,
                linked_group=group,
            )
            for i, name in enumerate(("Candidate One", "Candidate Two", "Donald Trump"))
        ]
    if kind == "job_approval_battery":
        group = linked_group_id("job")
        text = "Do you approve or disapprove of the job each of the following is doing?"
        options = [
            "Strongly approve",
            "Somewhat approve",
            "Somewhat disapprove",
            "Strongly disapprove",
            "(Don't know)",
        ]
        items = (
            "Donald Trump is doing as President",
            "The Governor is doing",
            "Your member of Congress is doing",
        )
        return [
            make_question(
                qnum=q,
                qid=f"{q}{letter(i)}",
                question_type="approve_disapprove_4pt",
                text=text,
                item_text=f"{letter(i)}. {item}",
                options=options,
                linked_group=group,
            )
            for i, item in enumerate(items)
        ]
    if kind == "describes_well_battery":
        group = linked_group_id("describes")
        text = "Thinking about Candidate One, please tell me whether each phrase describes them very well, somewhat well, not well, or don't know."
        options = ["Very well", "Somewhat well", "Not well", "(Don't know)"]
        items = (
            "Works hard for the people of the district",
            "Stays in touch with the people",
            "Understands the needs of working people",
        )
        return [
            make_question(
                qnum=q,
                qid=f"{q}{letter(i)}",
                question_type="describes_well",
                text=text,
                item_text=f"{letter(i)}. {item}",
                options=options,
                linked_group=group,
            )
            for i, item in enumerate(items)
        ]
    if kind in {"more_likely_battery", "less_likely_battery"}:
        more = kind == "more_likely_battery"
        group = linked_group_id("likely")
        axis = "more" if more else "less"
        text = (
            f"For each one, please tell me if that statement makes you much {axis} likely "
            "or somewhat {axis} likely to vote for Candidate One."
        ).replace("{axis}", axis)
        options = (
            ["Much more likely", "Somewhat more likely", "No difference", "Less likely", "(Don't know)"]
            if more
            else [
                "Much less likely",
                "Somewhat less likely",
                "No difference",
                "More likely",
                "Not true",
                "(Don't know)",
            ]
        )
        tags = ("[JOBS] Candidate One created local jobs", "[COSTS] Candidate One lowered costs", "[VALUES] Candidate One shares our values")
        return [
            make_question(
                qnum=q,
                qid=f"{q}{letter(i)}",
                question_type=f"{axis}_likely_battery",
                text=text,
                item_text=f"{letter(i)}. {item}",
                options=options,
                linked_group=group,
            )
            for i, item in enumerate(tags)
        ]
    if kind == "demographics_block":
        start = max(100, q)
        return [
            make_question(qnum=start, question_type="demographic_sex", text="Sex", options=["Male", "Female"]),
            make_question(
                qnum=start + 1,
                question_type="demographic_age",
                text="Age",
                options=["18-29", "30-34", "35-39", "40-44", "45-49", "50-54", "55-59", "60-64", "65-69", "70+", "(Refused)"],
            ),
            make_question(
                qnum=start + 2,
                question_type="demographic_education",
                text="Education",
                options=[
                    "Some grade school",
                    "Some high school",
                    "Graduated high school",
                    "Technical/Vocational",
                    "Some college but no degree",
                    "Associates degree",
                    "Bachelors degree",
                    "Graduate/Professional",
                    "(Don't know/Refused)",
                ],
            ),
            make_question(
                qnum=start + 3,
                question_type="demographic_race",
                text="Race",
                options=["White", "Black", "AAPI", "Hispanic or Latino", "(Other)", "(Don't know/Refused)"],
            ),
            make_question(
                qnum=start + 4,
                question_type="demographic_latino",
                text="Do you consider yourself to be of Latino or Hispanic descent?",
                options=["Yes", "No", "(Don't know/Refused)"],
            ),
            make_question(
                qnum=start + 5,
                question_type="politics_identity_importance",
                text="How important is politics to your personal identity, very important, somewhat important, not that important, or not important at all?",
                options=["Very important", "Somewhat important", "Not that important", "Not important at all", "(Don't know)"],
            ),
            make_question(
                qnum=start + 6,
                question_type="abortion_identity",
                text="When it comes to the issue of abortion, do you generally think of yourself as pro-choice or pro-life? And would you say you are strongly or do you lean?",
                options=["Strongly pro-choice", "Lean pro-choice", "Lean pro-life", "Strongly pro-life", "(Neither)", "(Refused)"],
            ),
            make_question(
                qnum=start + 7,
                question_type="past_vote_recall",
                text="In the 2024 race for President, did you vote for Donald Trump, the Republican or Kamala Harris, the Democrat?",
                instructions=[
                    "[IF RESPONDENT DECLINES TO SAY]",
                    "This survey is being done for research purposes only. Your willingness to say how you voted is really important for ensuring the accuracy of our research.",
                ],
                options=["Trump", "Harris", "(Other)", "(I did not vote in 2024)", "(Refused)"],
            ),
        ]
    raise ValueError(f"Unsupported template: {kind}")


def draft_question_label(q: dict) -> str:
    qid = str(q.get("qid") or q.get("qnum") or "?")
    qtype = str(q.get("question_type") or "question")
    item = clean_text(q.get("item_text") or "")
    text = clean_text(q.get("text") or "")
    label_text = item or text
    if len(label_text) > 72:
        label_text = label_text[:69] + "..."
    linked = " *" if q.get("linked_group") else ""
    return f"Q{qid}{linked}  [{qtype}]  {label_text}"


def render_plain_text(draft: dict) -> str:
    lines = []
    title = clean_text(draft.get("title") or "Untitled Survey Draft")
    survey_id = clean_text(draft.get("survey_id") or "")
    lines.append(title if not survey_id else f"{survey_id} {title}")
    lines.append("")
    for q in draft.get("questions", []):
        qid = clean_text(q.get("qid") or q.get("qnum") or "")
        lines.append(f"Q{qid}. {clean_text(q.get('text') or '')}")
        for inst in q.get("instructions") or []:
            lines.append(f"[{clean_text(inst).strip('[]')}]")
        item = clean_text(q.get("item_text") or "")
        if item:
            lines.append(item)
        for opt in q.get("response_options") or []:
            value = clean_text(opt.get("value") or "")
            text = clean_text(opt.get("text") or "")
            lines.append(f"{value}. {text}")
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def render_html(draft: dict) -> str:
    title = clean_text(draft.get("title") or "Untitled Survey Draft")
    survey_id = clean_text(draft.get("survey_id") or "")
    doc_title = f"{survey_id} {title}".strip()
    body = [
        "<!doctype html>",
        "<html>",
        "<head>",
        "<meta charset=\"utf-8\">",
        f"<title>{html.escape(doc_title)}</title>",
        "<style>",
        "body{font-family:Arial,Helvetica,sans-serif;line-height:1.35;margin:36px;max-width:960px;}",
        "h1{font-size:22px;margin:0 0 20px 0;} .q{margin:0 0 22px 0;}",
        ".qid{font-weight:700;} .instruction{font-weight:700;color:#555;margin:4px 0;}",
        ".item{margin:6px 0 4px 24px;} ol{margin:4px 0 0 42px;} li{margin:2px 0;}",
        ".meta{color:#666;font-size:12px;margin-bottom:20px;}",
        "</style>",
        "</head>",
        "<body>",
        f"<h1>{html.escape(doc_title)}</h1>",
        f"<div class=\"meta\">Generated by Survey Writer on {html.escape(time.strftime('%Y-%m-%d %H:%M'))}</div>",
    ]
    for q in draft.get("questions", []):
        qid = clean_text(q.get("qid") or q.get("qnum") or "")
        body.append("<section class=\"q\">")
        body.append(
            f"<div><span class=\"qid\">Q{html.escape(qid)}.</span> "
            f"{html.escape(clean_text(q.get('text') or ''))}</div>"
        )
        for inst in q.get("instructions") or []:
            body.append(f"<div class=\"instruction\">[{html.escape(clean_text(inst).strip('[]'))}]</div>")
        item = clean_text(q.get("item_text") or "")
        if item:
            body.append(f"<div class=\"item\">{html.escape(item)}</div>")
        options = q.get("response_options") or []
        if options:
            body.append("<ol>")
            for opt in options:
                body.append(f"<li value=\"{html.escape(clean_text(opt.get('value') or ''))}\">{html.escape(clean_text(opt.get('text') or ''))}</li>")
            body.append("</ol>")
        body.append("</section>")
    body.extend(["</body>", "</html>"])
    return "\n".join(body) + "\n"


def export_html(draft: dict, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(render_html(draft), encoding="utf-8")


def export_docx(draft: dict, path: Path) -> None:
    try:
        from docx import Document
    except Exception as exc:  # pragma: no cover - environment-dependent
        raise RuntimeError("python-docx is required for DOCX export.") from exc

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    title = clean_text(draft.get("title") or "Untitled Survey Draft")
    survey_id = clean_text(draft.get("survey_id") or "")
    doc.add_heading(f"{survey_id} {title}".strip(), level=1)
    for q in draft.get("questions", []):
        qid = clean_text(q.get("qid") or q.get("qnum") or "")
        p = doc.add_paragraph()
        p.add_run(f"Q{qid}. ").bold = True
        p.add_run(clean_text(q.get("text") or ""))
        for inst in q.get("instructions") or []:
            p = doc.add_paragraph(f"[{clean_text(inst).strip('[]')}]")
            p.runs[0].bold = True
        item = clean_text(q.get("item_text") or "")
        if item:
            doc.add_paragraph(item)
        for opt in q.get("response_options") or []:
            doc.add_paragraph(
                f"{clean_text(opt.get('value') or '')}. {clean_text(opt.get('text') or '')}",
                style=None,
            )
        doc.add_paragraph("")
    doc.save(path)


def ascii_fold_text(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.replace("\u2019", "'").replace("\u2018", "'")
    text = re.sub(r"[^a-zA-Z0-9']+", " ", text.casefold())
    return re.sub(r"\s+", " ", text).strip()


SPANISH_CUES = {
    "ahora",
    "leer",
    "leerle",
    "gustaria",
    "dijera",
    "opinion",
    "desfavorable",
    "figuras",
    "publicas",
    "escala",
    "significa",
    "calificar",
}


def likely_english_question(question_text: str, item_text: str, source_path: str = "") -> bool:
    source = ascii_fold_text(source_path)
    if re.search(r"\b(spanish|espanol|castellano)\b", source):
        return False
    blob = ascii_fold_text(f"{question_text} {item_text}")
    if not blob:
        return True
    tokens = set(blob.split())
    spanish_score = len(tokens & SPANISH_CUES)
    english_cues = sum(
        1
        for cue in (
            "please tell me",
            "favorable or unfavorable",
            "do you approve",
            "would you vote",
            "i'm going to read",
            "i am going to read",
        )
        if cue in blob
    )
    return not (spanish_score >= 5 and english_cues == 0)


def parse_options_json(raw: object) -> list[str]:
    try:
        payload = json.loads(str(raw or "[]"))
    except Exception:
        return []
    if not isinstance(payload, list):
        return []
    return [clean_text(x) for x in payload if clean_text(x)]


def query_tokens(query: str) -> list[str]:
    return [token for token in re.findall(r"[A-Za-z0-9']+", query.casefold()) if token.strip("'")]


def search_question_archive(
    *,
    query: str,
    client: str = "",
    question_type: str = "",
    limit: int = 50,
    db_path: Path = QUESTION_DB,
) -> list[SearchResult]:
    if not db_path.exists():
        return []
    clauses = ["1=1"]
    params: list[object] = []
    columns = ["question_text", "item_text", "option_texts_json", "question_target", "question_type", "question_group"]
    for token in query_tokens(query):
        token_clause = " OR ".join(f"LOWER({col}) LIKE ?" for col in columns)
        clauses.append(f"({token_clause})")
        params.extend([f"%{token}%"] * len(columns))
    if client.strip():
        clauses.append("LOWER(client) = LOWER(?)")
        params.append(client.strip())
    if question_type.strip():
        clauses.append("LOWER(question_type) LIKE ?")
        params.append(f"%{question_type.casefold().strip()}%")
    sql = f"""
        SELECT question_id, survey_id, qid, qnum, question_text, item_text,
               option_texts_json, question_type, question_family, question_group,
               question_target, state_code, region, client, source_json_path,
               survey_code, question_ord
        FROM repo_questions
        WHERE {' AND '.join(clauses)}
        ORDER BY CAST(survey_code AS INTEGER) DESC, question_ord ASC
        LIMIT ?
    """
    params.append(int(limit))
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        rows = conn.execute(sql, params).fetchall()
    finally:
        conn.close()
    results: list[SearchResult] = []
    for row in rows:
        if not likely_english_question(row["question_text"], row["item_text"], row["source_json_path"]):
            continue
        results.append(
            SearchResult(
                question_id=int(row["question_id"]),
                survey_id=str(row["survey_id"] or ""),
                qid=str(row["qid"] or ""),
                qnum=str(row["qnum"] or ""),
                question_text=str(row["question_text"] or ""),
                item_text=str(row["item_text"] or ""),
                option_texts=parse_options_json(row["option_texts_json"]),
                question_type=str(row["question_type"] or ""),
                question_family=str(row["question_family"] or ""),
                question_group=str(row["question_group"] or ""),
                question_target=str(row["question_target"] or ""),
                state_code=str(row["state_code"] or ""),
                region=str(row["region"] or ""),
                client=str(row["client"] or ""),
            )
        )
    return results


def distinct_clients(db_path: Path = QUESTION_DB) -> list[str]:
    if not db_path.exists():
        return []
    conn = sqlite3.connect(db_path)
    try:
        rows = conn.execute(
            "SELECT DISTINCT client FROM repo_questions WHERE TRIM(client) <> '' ORDER BY LOWER(client) LIMIT 500"
        ).fetchall()
    finally:
        conn.close()
    return [str(row[0]) for row in rows if str(row[0]).strip()]


class SurveyWriterApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("Survey Writer")
        self.root.geometry("1180x760")
        DRAFT_ROOT.mkdir(parents=True, exist_ok=True)

        self.draft = new_draft()
        self.draft_path: Path | None = None
        self.selected_index: int | None = None
        self.search_results: list[SearchResult] = []
        self._loading_editor = False

        self.survey_id_var = tk.StringVar()
        self.title_var = tk.StringVar(value=self.draft["title"])
        self.status_var = tk.StringVar(value=f"Draft area: {DRAFT_ROOT}")
        self.template_var = tk.StringVar(value=QUESTION_TEMPLATES[0][0])

        self.qnum_var = tk.StringVar()
        self.qid_var = tk.StringVar()
        self.qtype_var = tk.StringVar()
        self.linked_var = tk.StringVar()

        self.search_query_var = tk.StringVar()
        self.search_client_var = tk.StringVar()
        self.search_type_var = tk.StringVar()
        self.search_status_var = tk.StringVar(value=f"Question DB: {QUESTION_DB}")

        self._build_ui()
        self._refresh_question_list()

    def _build_ui(self) -> None:
        top = ttk.Frame(self.root, padding=(10, 8))
        top.pack(fill=tk.X)
        ttk.Label(top, text="Survey ID").grid(row=0, column=0, sticky="w")
        ttk.Entry(top, textvariable=self.survey_id_var, width=18).grid(row=0, column=1, padx=(6, 16), sticky="w")
        ttk.Label(top, text="Title").grid(row=0, column=2, sticky="w")
        ttk.Entry(top, textvariable=self.title_var, width=72).grid(row=0, column=3, padx=(6, 16), sticky="we")
        top.columnconfigure(3, weight=1)
        ttk.Button(top, text="New", command=self.new_draft).grid(row=0, column=4, padx=2)
        ttk.Button(top, text="Open", command=self.open_draft).grid(row=0, column=5, padx=2)
        ttk.Button(top, text="Save", command=self.save_draft).grid(row=0, column=6, padx=2)
        ttk.Button(top, text="Export HTML", command=self.export_html_dialog).grid(row=0, column=7, padx=2)
        ttk.Button(top, text="Export DOCX", command=self.export_docx_dialog).grid(row=0, column=8, padx=2)

        main = ttk.PanedWindow(self.root, orient=tk.HORIZONTAL)
        main.pack(fill=tk.BOTH, expand=True, padx=10, pady=(0, 8))

        left = ttk.Frame(main, padding=(0, 0, 8, 0))
        main.add(left, weight=1)
        ttk.Label(left, text="Questions").pack(anchor="w")
        self.question_list = tk.Listbox(left, exportselection=False, width=52)
        self.question_list.pack(fill=tk.BOTH, expand=True, pady=(4, 6))
        self.question_list.bind("<<ListboxSelect>>", self.on_question_select)

        template_row = ttk.Frame(left)
        template_row.pack(fill=tk.X)
        ttk.Combobox(
            template_row,
            textvariable=self.template_var,
            values=[label for label, _kind in QUESTION_TEMPLATES],
            state="readonly",
            width=32,
        ).pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Button(template_row, text="Add", command=self.add_template).pack(side=tk.LEFT, padx=(6, 0))

        move_row = ttk.Frame(left)
        move_row.pack(fill=tk.X, pady=(6, 0))
        for text, command in (
            ("Duplicate", self.duplicate_question),
            ("Delete", self.delete_question),
            ("Up", lambda: self.move_question(-1)),
            ("Down", lambda: self.move_question(1)),
        ):
            ttk.Button(move_row, text=text, command=command).pack(side=tk.LEFT, padx=(0, 5))

        right = ttk.Notebook(main)
        main.add(right, weight=3)

        editor = ttk.Frame(right, padding=8)
        right.add(editor, text="Editor")
        self._build_editor(editor)

        search = ttk.Frame(right, padding=8)
        right.add(search, text="Archive Search")
        self._build_search(search)

        preview = ttk.Frame(right, padding=8)
        right.add(preview, text="Preview")
        self.preview_text = tk.Text(preview, wrap=tk.WORD)
        self.preview_text.pack(fill=tk.BOTH, expand=True)
        ttk.Button(preview, text="Refresh Preview", command=self.refresh_preview).pack(anchor="e", pady=(6, 0))

        status = ttk.Label(self.root, textvariable=self.status_var, anchor="w", padding=(10, 0))
        status.pack(fill=tk.X, side=tk.BOTTOM)

    def _build_editor(self, parent: ttk.Frame) -> None:
        meta = ttk.Frame(parent)
        meta.pack(fill=tk.X)
        for col, (label, var, width) in enumerate(
            (
                ("Q number", self.qnum_var, 10),
                ("QID", self.qid_var, 12),
                ("Type", self.qtype_var, 30),
                ("Linked group", self.linked_var, 26),
            )
        ):
            ttk.Label(meta, text=label).grid(row=0, column=col * 2, sticky="w", padx=(0 if col == 0 else 12, 4))
            ttk.Entry(meta, textvariable=var, width=width).grid(row=0, column=col * 2 + 1, sticky="w")

        ttk.Label(parent, text="Question text").pack(anchor="w", pady=(12, 2))
        self.question_text = tk.Text(parent, height=5, wrap=tk.WORD)
        self.question_text.pack(fill=tk.X)

        ttk.Label(parent, text="Instructions, one per line").pack(anchor="w", pady=(10, 2))
        self.instructions_text = tk.Text(parent, height=4, wrap=tk.WORD)
        self.instructions_text.pack(fill=tk.X)

        ttk.Label(parent, text="Item text").pack(anchor="w", pady=(10, 2))
        self.item_text = tk.Text(parent, height=3, wrap=tk.WORD)
        self.item_text.pack(fill=tk.X)

        ttk.Label(parent, text="Answer choices, one per line. Use '1. Choice' or just 'Choice'.").pack(anchor="w", pady=(10, 2))
        self.options_text = tk.Text(parent, height=10, wrap=tk.WORD)
        self.options_text.pack(fill=tk.BOTH, expand=True)

        buttons = ttk.Frame(parent)
        buttons.pack(fill=tk.X, pady=(8, 0))
        ttk.Button(buttons, text="Apply Question", command=self.apply_editor).pack(side=tk.RIGHT)

    def _build_search(self, parent: ttk.Frame) -> None:
        controls = ttk.Frame(parent)
        controls.pack(fill=tk.X)
        ttk.Label(controls, text="Query").grid(row=0, column=0, sticky="w")
        ttk.Entry(controls, textvariable=self.search_query_var, width=34).grid(row=0, column=1, sticky="we", padx=(6, 12))
        ttk.Label(controls, text="Client").grid(row=0, column=2, sticky="w")
        ttk.Combobox(controls, textvariable=self.search_client_var, values=distinct_clients(), width=26).grid(row=0, column=3, sticky="we", padx=(6, 12))
        ttk.Label(controls, text="Type").grid(row=0, column=4, sticky="w")
        ttk.Entry(controls, textvariable=self.search_type_var, width=22).grid(row=0, column=5, sticky="we", padx=(6, 12))
        ttk.Button(controls, text="Search", command=self.run_search).grid(row=0, column=6)
        controls.columnconfigure(1, weight=1)
        controls.columnconfigure(3, weight=1)

        cols = ("survey", "qid", "type", "client", "text")
        self.search_tree = ttk.Treeview(parent, columns=cols, show="headings", height=14)
        for col, label, width in (
            ("survey", "Survey", 90),
            ("qid", "QID", 70),
            ("type", "Type", 150),
            ("client", "Client", 150),
            ("text", "Question / Item", 560),
        ):
            self.search_tree.heading(col, text=label)
            self.search_tree.column(col, width=width, anchor="w")
        self.search_tree.pack(fill=tk.BOTH, expand=True, pady=(8, 6))
        self.search_tree.bind("<<TreeviewSelect>>", self.preview_search_result)
        self.search_tree.bind("<Double-1>", lambda _event: self.add_selected_search_result())

        self.search_preview = tk.Text(parent, height=9, wrap=tk.WORD)
        self.search_preview.pack(fill=tk.X)
        bottom = ttk.Frame(parent)
        bottom.pack(fill=tk.X, pady=(6, 0))
        ttk.Label(bottom, textvariable=self.search_status_var).pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Button(bottom, text="Add Selected to Draft", command=self.add_selected_search_result).pack(side=tk.RIGHT)

    def _save_metadata_from_controls(self) -> None:
        self.draft["survey_id"] = self.survey_id_var.get().strip()
        self.draft["title"] = self.title_var.get().strip() or "Untitled Survey Draft"
        self.draft["updated_at"] = time.strftime("%Y-%m-%dT%H:%M:%S")

    def _refresh_question_list(self, select_index: int | None = None) -> None:
        self.question_list.delete(0, tk.END)
        for q in self.draft.get("questions", []):
            self.question_list.insert(tk.END, draft_question_label(q))
        if select_index is not None and self.draft.get("questions"):
            select_index = max(0, min(select_index, len(self.draft["questions"]) - 1))
            self.question_list.selection_set(select_index)
            self.question_list.see(select_index)
            self.load_question(select_index)
        self.refresh_preview()

    def on_question_select(self, _event: object | None = None) -> None:
        if self._loading_editor:
            return
        self.apply_editor(silent=True)
        selection = self.question_list.curselection()
        if not selection:
            return
        self.load_question(int(selection[0]))

    def clear_editor(self) -> None:
        self._loading_editor = True
        try:
            for var in (self.qnum_var, self.qid_var, self.qtype_var, self.linked_var):
                var.set("")
            for widget in (self.question_text, self.instructions_text, self.item_text, self.options_text):
                widget.delete("1.0", tk.END)
        finally:
            self._loading_editor = False

    def load_question(self, index: int) -> None:
        questions = self.draft.get("questions", [])
        if not (0 <= index < len(questions)):
            self.selected_index = None
            self.clear_editor()
            return
        q = questions[index]
        self.selected_index = index
        self._loading_editor = True
        try:
            self.qnum_var.set(str(q.get("qnum") or ""))
            self.qid_var.set(str(q.get("qid") or ""))
            self.qtype_var.set(str(q.get("question_type") or ""))
            self.linked_var.set(str(q.get("linked_group") or ""))
            self.question_text.delete("1.0", tk.END)
            self.question_text.insert("1.0", str(q.get("text") or ""))
            self.instructions_text.delete("1.0", tk.END)
            self.instructions_text.insert("1.0", "\n".join(q.get("instructions") or []))
            self.item_text.delete("1.0", tk.END)
            self.item_text.insert("1.0", str(q.get("item_text") or ""))
            self.options_text.delete("1.0", tk.END)
            self.options_text.insert("1.0", option_lines(q.get("response_options") or []))
        finally:
            self._loading_editor = False

    def apply_editor(self, silent: bool = False) -> None:
        if self._loading_editor or self.selected_index is None:
            return
        questions = self.draft.get("questions", [])
        if not (0 <= self.selected_index < len(questions)):
            return
        q = questions[self.selected_index]
        q["qnum"] = self.qnum_var.get().strip() or q.get("qnum") or ""
        q["qid"] = self.qid_var.get().strip() or q.get("qid") or q["qnum"]
        q["question_type"] = self.qtype_var.get().strip() or "single_choice"
        q["linked_group"] = self.linked_var.get().strip()
        q["text"] = self.question_text.get("1.0", tk.END).strip()
        q["instructions"] = [line.strip() for line in self.instructions_text.get("1.0", tk.END).splitlines() if line.strip()]
        q["item_text"] = self.item_text.get("1.0", tk.END).strip()
        q["response_options"] = normalize_option_lines(self.options_text.get("1.0", tk.END))
        self._save_metadata_from_controls()
        self._refresh_question_list(self.selected_index)
        if not silent:
            self.status_var.set("Question applied.")

    def add_template(self) -> None:
        self.apply_editor(silent=True)
        label = self.template_var.get()
        kind = next((kind for template_label, kind in QUESTION_TEMPLATES if template_label == label), "single_choice")
        start = 100 if kind == "demographics_block" else next_qnum(self.draft.get("questions", []))
        created = template_questions(kind, start)
        self.draft.setdefault("questions", []).extend(created)
        self._save_metadata_from_controls()
        self._refresh_question_list(len(self.draft["questions"]) - len(created))
        self.status_var.set(f"Added {len(created)} question(s) from template: {label}")

    def duplicate_question(self) -> None:
        self.apply_editor(silent=True)
        if self.selected_index is None:
            return
        q = json.loads(json.dumps(self.draft["questions"][self.selected_index]))
        q["qnum"] = str(next_qnum(self.draft.get("questions", [])))
        q["qid"] = q["qnum"]
        q["linked_group"] = ""
        self.draft["questions"].insert(self.selected_index + 1, q)
        self._refresh_question_list(self.selected_index + 1)

    def delete_question(self) -> None:
        if self.selected_index is None:
            return
        del self.draft["questions"][self.selected_index]
        self._refresh_question_list(min(self.selected_index, len(self.draft["questions"]) - 1) if self.draft["questions"] else None)

    def move_question(self, delta: int) -> None:
        self.apply_editor(silent=True)
        if self.selected_index is None:
            return
        new_index = self.selected_index + delta
        questions = self.draft["questions"]
        if not (0 <= new_index < len(questions)):
            return
        questions[self.selected_index], questions[new_index] = questions[new_index], questions[self.selected_index]
        self._refresh_question_list(new_index)

    def new_draft(self) -> None:
        self.draft = new_draft()
        self.draft_path = None
        self.selected_index = None
        self.survey_id_var.set("")
        self.title_var.set(self.draft["title"])
        self.clear_editor()
        self._refresh_question_list()
        self.status_var.set("New draft started.")

    def open_draft(self) -> None:
        path = filedialog.askopenfilename(
            initialdir=DRAFT_ROOT,
            title="Open survey draft",
            filetypes=[("Survey draft JSON", "*.json"), ("All files", "*.*")],
        )
        if not path:
            return
        with open(path, "r", encoding="utf-8") as f:
            self.draft = json.load(f)
        self.draft_path = Path(path)
        self.selected_index = None
        self.survey_id_var.set(str(self.draft.get("survey_id") or ""))
        self.title_var.set(str(self.draft.get("title") or "Untitled Survey Draft"))
        self.clear_editor()
        self._refresh_question_list(0 if self.draft.get("questions") else None)
        self.status_var.set(f"Opened {self.draft_path}")

    def save_draft(self) -> None:
        self.apply_editor(silent=True)
        self._save_metadata_from_controls()
        if self.draft_path is None:
            default_name = safe_file_stem(self.draft.get("survey_id") or self.draft.get("title") or "survey_draft") + ".json"
            path = filedialog.asksaveasfilename(
                initialdir=DRAFT_ROOT,
                initialfile=default_name,
                title="Save survey draft",
                defaultextension=".json",
                filetypes=[("Survey draft JSON", "*.json")],
            )
            if not path:
                return
            self.draft_path = Path(path)
        self.draft_path.parent.mkdir(parents=True, exist_ok=True)
        self.draft_path.write_text(json.dumps(self.draft, indent=2), encoding="utf-8")
        self.status_var.set(f"Saved {self.draft_path}")

    def default_export_path(self, suffix: str) -> Path:
        stem = safe_file_stem(self.draft.get("survey_id") or self.draft.get("title") or "survey_draft")
        return DRAFT_ROOT / f"{stem}{suffix}"

    def export_html_dialog(self) -> None:
        self.apply_editor(silent=True)
        self._save_metadata_from_controls()
        default = self.default_export_path(".html")
        path = filedialog.asksaveasfilename(
            initialdir=default.parent,
            initialfile=default.name,
            title="Export HTML questionnaire",
            defaultextension=".html",
            filetypes=[("HTML", "*.html")],
        )
        if not path:
            return
        export_html(self.draft, Path(path))
        self.status_var.set(f"Exported HTML: {path}")

    def export_docx_dialog(self) -> None:
        self.apply_editor(silent=True)
        self._save_metadata_from_controls()
        default = self.default_export_path(".docx")
        path = filedialog.asksaveasfilename(
            initialdir=default.parent,
            initialfile=default.name,
            title="Export DOCX questionnaire",
            defaultextension=".docx",
            filetypes=[("Word document", "*.docx")],
        )
        if not path:
            return
        try:
            export_docx(self.draft, Path(path))
        except Exception as exc:
            messagebox.showerror("DOCX export", str(exc))
            return
        self.status_var.set(f"Exported DOCX: {path}")

    def refresh_preview(self) -> None:
        if not hasattr(self, "preview_text"):
            return
        self._save_metadata_from_controls()
        self.preview_text.configure(state=tk.NORMAL)
        self.preview_text.delete("1.0", tk.END)
        self.preview_text.insert("1.0", render_plain_text(self.draft))
        self.preview_text.configure(state=tk.DISABLED)

    def run_search(self) -> None:
        query = self.search_query_var.get().strip()
        if not query:
            messagebox.showinfo("Archive Search", "Enter a search query first.")
            return
        self.search_results = search_question_archive(
            query=query,
            client=self.search_client_var.get().strip(),
            question_type=self.search_type_var.get().strip(),
            limit=60,
        )
        for row in self.search_tree.get_children():
            self.search_tree.delete(row)
        for i, result in enumerate(self.search_results):
            text = clean_text(result.item_text or result.question_text)
            if len(text) > 120:
                text = text[:117] + "..."
            self.search_tree.insert(
                "",
                tk.END,
                iid=str(i),
                values=(result.survey_id, result.qid, result.question_type, result.client, text),
            )
        self.search_status_var.set(f"Found {len(self.search_results)} result(s). Double-click a row to add it.")

    def selected_search_result(self) -> SearchResult | None:
        selection = self.search_tree.selection()
        if not selection:
            return None
        index = int(selection[0])
        if not (0 <= index < len(self.search_results)):
            return None
        return self.search_results[index]

    def preview_search_result(self, _event: object | None = None) -> None:
        result = self.selected_search_result()
        self.search_preview.configure(state=tk.NORMAL)
        self.search_preview.delete("1.0", tk.END)
        if result is not None:
            lines = [
                f"Survey: {result.survey_id}  QID: {result.qid}",
                f"Client: {result.client}  Region: {result.region}  State: {result.state_code}",
                f"Type: {result.question_type}  Family: {result.question_family}  Group: {result.question_group}",
                "",
                result.question_text,
            ]
            if result.item_text:
                lines.extend(["", result.item_text])
            if result.option_texts:
                lines.append("")
                lines.extend(f"{i}. {text}" for i, text in enumerate(result.option_texts, start=1))
            self.search_preview.insert("1.0", "\n".join(lines))
        self.search_preview.configure(state=tk.DISABLED)

    def add_selected_search_result(self) -> None:
        self.apply_editor(silent=True)
        result = self.selected_search_result()
        if result is None:
            return
        qnum = next_qnum(self.draft.get("questions", []))
        q = make_question(
            qnum=qnum,
            question_type=result.question_type or "single_choice",
            text=result.question_text,
            item_text=result.item_text,
            options=result.option_texts,
            source={
                "survey_id": result.survey_id,
                "qid": result.qid,
                "client": result.client,
                "region": result.region,
                "state_code": result.state_code,
            },
        )
        self.draft.setdefault("questions", []).append(q)
        self._refresh_question_list(len(self.draft["questions"]) - 1)
        self.status_var.set(f"Added archive question from {result.survey_id} Q{result.qid}.")


def main() -> None:
    root = tk.Tk()
    app = SurveyWriterApp(root)
    root.protocol("WM_DELETE_WINDOW", root.destroy)
    root.mainloop()


if __name__ == "__main__":
    main()
