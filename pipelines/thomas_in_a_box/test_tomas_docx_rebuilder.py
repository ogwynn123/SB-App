import sys
import unittest
from pathlib import Path

from docx import Document


THIS_DIR = Path(__file__).resolve().parent
if str(THIS_DIR) not in sys.path:
    sys.path.insert(0, str(THIS_DIR))

import tomas_docx_rebuilder as rebuilder


class TomasDocxRebuilderTests(unittest.TestCase):
    def test_semantic_json_pairing_ignores_extra_spanish_directive(self) -> None:
        raw_en = "\n".join(
            [
                "STANDARD SMS INTRO",
                "How are you registered to vote, as a Democrat, a Republican, an Independent, or with another party?",
                "Bird",
            ]
        )
        raw_es = "\n".join(
            [
                "INTRODUCCIÓN ESTÁNDAR POR SMS",
                "¿Cómo está registrado para votar: como demócrata, republicano, independiente o con otro partido?",
                "Bird",
            ]
        )
        json_en = """
        {
          "survey_level_metadata": {
            "notes": [
              "Output follows PDF line order/numbering for parser compatibility."
            ]
          },
          "elements": [
            {
              "type": "directive",
              "text": "STANDARD SMS INTRO"
            },
            {
              "type": "question",
              "qid": "1",
              "text": "How are you registered to vote, as a Democrat, a Republican, an Independent, or with another party?",
              "response_options": [
                {"value": "1", "text": "Bird"}
              ]
            }
          ]
        }
        """
        json_es = """
        {
          "survey_level_metadata": {
            "notes": []
          },
          "elements": [
            {
              "type": "directive",
              "text": "Notas:"
            },
            {
              "type": "directive",
              "text": "INTRODUCCIÓN ESTÁNDAR POR SMS"
            },
            {
              "type": "question",
              "qid": "1",
              "text": "¿Cómo está registrado para votar: como demócrata, republicano, independiente o con otro partido?",
              "response_options": [
                {"value": "1", "text": "Bird"}
              ]
            }
          ]
        }
        """

        mapping = rebuilder.build_mapping(
            raw_en,
            raw_es,
            json_en_text=json_en,
            json_es_text=json_es,
        )

        self.assertEqual(mapping.exact.get("STANDARD SMS INTRO"), "INTRODUCCIÓN ESTÁNDAR POR SMS")
        self.assertEqual(
            mapping.exact.get(
                "How are you registered to vote, as a Democrat, a Republican, an Independent, or with another party?"
            ),
            "¿Cómo está registrado para votar: como demócrata, republicano, independiente o con otro partido?",
        )
        self.assertEqual(mapping.exact.get("Bird"), "Bird")
        self.assertNotEqual(mapping.exact.get("STANDARD SMS INTRO"), "Notas:")

    def test_merge_tags_from_source_preserves_q7_tag_positions(self) -> None:
        source = (
            "[IF Q7=1-6] "
            "And who would be your second choice in the Democratic primary? "
            "[DO NOT READ CHOICE FROM Q7, RANDOMIZE]"
        )
        translated = "¿Y quién sería su segunda opción en la primaria demócrata?"

        merged = rebuilder._merge_tags_from_source(source, translated)

        self.assertEqual(
            merged,
            "[IF Q7=1-6] ¿Y quién sería su segunda opción en la primaria demócrata? "
            "[DO NOT READ CHOICE FROM Q7, RANDOMIZE]",
        )

    def test_set_paragraph_text_preserves_bold_tag_groups(self) -> None:
        doc = Document()
        paragraph = doc.add_paragraph()

        r1 = paragraph.add_run("[IF ")
        r1.bold = True
        r2 = paragraph.add_run("Q7=1-6] ")
        r2.bold = True
        r3 = paragraph.add_run("And who would be your ")
        r3.bold = False
        r4 = paragraph.add_run("second choice in the Democratic primary? ")
        r4.bold = False
        r5 = paragraph.add_run("[DO NOT READ CHOICE ")
        r5.bold = True
        r6 = paragraph.add_run("FROM Q7, RANDOMIZE]")
        r6.bold = True

        new_text = (
            "[IF Q7=1-6] ¿Y quién sería su segunda opción en la primaria demócrata? "
            "[DO NOT READ CHOICE FROM Q7, RANDOMIZE]"
        )
        rebuilder._set_paragraph_text(paragraph, new_text)

        self.assertEqual("".join(run.text for run in paragraph.runs), new_text)
        bold_text = "".join(run.text for run in paragraph.runs if run.bold)
        non_bold_text = "".join(run.text for run in paragraph.runs if not run.bold)
        self.assertEqual(
            bold_text,
            "[IF Q7=1-6] [DO NOT READ CHOICE FROM Q7, RANDOMIZE]",
        )
        self.assertEqual(
            non_bold_text,
            "¿Y quién sería su segunda opción en la primaria demócrata? ",
        )


if __name__ == "__main__":
    unittest.main()
