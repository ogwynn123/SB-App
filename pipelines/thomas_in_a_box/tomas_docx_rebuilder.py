#!/usr/bin/env python3
"""Rebuild a formatted Spanish DOCX using aligned TOMAS RAW.en/RAW.es outputs."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from docx import Document
from docx.document import Document as DocumentType
from docx.section import Section
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run


OPS_LINE_RE = re.compile(
    r"^\s*("
    r"If\b.*\b(skip|go to|goto|terminate)\b.*|"
    r"Skip\s+to\b.*|"
    r"Go\s+to\b.*|"
    r"TERMINATE\b.*|"
    r"Programmer\s+note:.*|"
    r"Interviewer\s+note:.*|"
    r"Routing\s+note:.*|"
    r"Logic\s+note:.*|"
    r"Quota\s+note:.*|"
    r"Soft\s+check:.*|"
    r"Hard\s+check:.*"
    r")\s*$",
    re.IGNORECASE,
)
ONLY_BRACKET_TAGS_RE = re.compile(r"^\s*(?:\[[^\]]*\]\s*)+$")
WS_RE = re.compile(r"\s+")
LEADING_CODE_RE = re.compile(r"^\s*(?:\d+|[A-Za-z])(?:[\)\].:/\-]|/)\s+")
TRAILING_CODE_RE = re.compile(r"\s+(?:TERMINATE|\d+)(?:\s+(?:TERMINATE|\d+))*\s*$", re.IGNORECASE)
DOT_LEADER_SUFFIX_RE = re.compile(r"(?:\s+\.(?:\s*\.){2,}|\.{5,})\s*$")
BULLET_TOKEN_RE = re.compile(r"(?:\u2022|\u00b7|\uf0b7|\u00ef\u201a\u00b7)")
LEADING_BULLET_RE = re.compile(r"^(\s*)(?:\u2022|\u00b7|\uf0b7|\u00ef\u201a\u00b7)\s+")
LEADING_LEGACY_BULLET_RE = re.compile(r"^(\s*)(?:\uf0b7|\u00ef\u201a\u00b7)\s+")
OPTION_PAIR_RE = re.compile(r"\s*(.+?)\s*\.{2,}\s*(TERMINATE|\d+)\b", re.IGNORECASE)
BRACKET_TAG_RE = re.compile(r"\[[^\]]*\]")
CURLY_TOKEN_RE = re.compile(r"\{[^}]*\}")
ANGLE_TOKEN_RE = re.compile(r"<[^>]*>")
QID_RE = re.compile(r"\bQ\d+[A-Z]?\b", re.IGNORECASE)
MAP_LINE_RE = re.compile(r"^\s*(TERMINATE|\d+)\s*=\s*(.+)$", re.IGNORECASE)
SENTENCE_SPLIT_RE = re.compile(r"(?<=[\.\?\!])\s+")
ITEM_MARKER_RE = re.compile(r"^\s*([A-Za-z]|\d+)(?:[\)\].:/\-]|/)\s+")
BRACKET_MARKER_RE = re.compile(r"^\s*\[([^\]]+)\]")


def _is_ops_line(line: str) -> bool:
    return bool(OPS_LINE_RE.match(line) or ONLY_BRACKET_TAGS_RE.match(line))


def _normalize_text(text: str) -> str:
    s = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    s = s.replace("\u2018", "'").replace("\u2019", "'")
    s = s.replace("\u201c", '"').replace("\u201d", '"')
    s = s.replace("\u2013", "-").replace("\u2014", "-")
    s = s.replace("\u2026", "...")
    s = s.replace("\uf0b7", " ").replace("\u00ef\u201a\u00b7", " ")
    s = WS_RE.sub(" ", s.strip())
    return s


def _normalize_loose(text: str) -> str:
    s = _normalize_text(text)
    s = LEADING_CODE_RE.sub("", s)
    s = TRAILING_CODE_RE.sub("", s)
    s = DOT_LEADER_SUFFIX_RE.sub("", s)
    s = BULLET_TOKEN_RE.sub(" ", s)
    if s.startswith("(") and s.endswith(")") and len(s) >= 2:
        s = s[1:-1].strip()
    s = WS_RE.sub(" ", s.strip())
    return s


def _dominant_variant(variants: Counter[str]) -> str:
    ranked = sorted(variants.items(), key=lambda x: (-x[1], x[0]))
    return ranked[0][0]


def _is_code_token(text: str) -> bool:
    return bool(re.fullmatch(r"(?:TERMINATE|\d+)", _normalize_loose(text), re.IGNORECASE))


def _distribute_words(text: str, parts: int) -> list[str]:
    if parts <= 0:
        return []
    if parts == 1:
        return [text.strip()]
    words = text.split()
    if not words:
        return [""] * parts
    if len(words) <= parts:
        out = words + [""] * (parts - len(words))
        return out
    base, extra = divmod(len(words), parts)
    out: list[str] = []
    idx = 0
    for i in range(parts):
        take = base + (1 if i < extra else 0)
        out.append(" ".join(words[idx : idx + take]))
        idx += take
    return out


def _apply_tab_layout_from_source(source_core: str, translated_core: str) -> str:
    if "\t" not in source_core or "\t" in translated_core:
        return translated_core

    src_cells = source_core.split("\t")
    nonempty_idx = [i for i, cell in enumerate(src_cells) if cell.strip()]
    if not nonempty_idx:
        return translated_core

    filled = [""] * len(src_cells)
    content = translated_core.strip()

    code_slot: int | None = None
    if len(nonempty_idx) >= 2 and _is_code_token(src_cells[nonempty_idx[-1]]):
        code_slot = nonempty_idx[-1]

    if code_slot is not None:
        m = re.search(r"(?:^|\s)(TERMINATE|\d+)\s*$", content, re.IGNORECASE)
        code_val = _normalize_loose(src_cells[code_slot])
        if m:
            code_val = m.group(1)
            content = content[: m.start(1)].strip()
        filled[code_slot] = code_val
        text_slots = [idx for idx in nonempty_idx if idx != code_slot]
    else:
        text_slots = nonempty_idx

    chunks = _distribute_words(content, len(text_slots))
    for idx, slot in enumerate(text_slots):
        filled[slot] = chunks[idx] if idx < len(chunks) else ""

    return "\t".join(filled)


def _apply_surface_layout_from_source(source: str, translated: str) -> str:
    lead_match = re.match(r"^\s*", source)
    trail_match = re.search(r"\s*$", source)
    lead = lead_match.group(0) if lead_match else ""
    trail = trail_match.group(0) if trail_match else ""
    core_start = len(lead)
    core_end = len(source) - len(trail) if trail else len(source)
    source_core = source[core_start:core_end]
    translated_core = translated.strip()
    translated_core = _apply_tab_layout_from_source(source_core, translated_core)
    return f"{lead}{translated_core}{trail}"


def _looks_optionish(line: str) -> bool:
    s = _normalize_text(line)
    if OPTION_PAIR_RE.search(s):
        return True
    if DOT_LEADER_SUFFIX_RE.search(s):
        return True
    # Dense numeric code sequences are typically answer rows (e.g., scales).
    if re.search(r"\b(?:TERMINATE|\d+)(?:\s+(?:TERMINATE|\d+)){2,}\s*$", s, re.IGNORECASE):
        return True
    return False


def _split_bullet_fragments(line: str) -> list[str]:
    if not BULLET_TOKEN_RE.search(line):
        return []
    canonical = BULLET_TOKEN_RE.sub("\u2022", line)
    parts = [p.strip(" \t-") for p in canonical.split("\u2022")]
    return [p for p in parts if p]


def _option_pairs(line: str) -> list[tuple[str, str]]:
    pairs: list[tuple[str, str]] = []
    for m in OPTION_PAIR_RE.finditer(line):
        label = m.group(1).strip()
        code = m.group(2).strip()
        if label:
            pairs.append((label, code))
    return pairs


def _append_source_suffix_if_any(source: str, translated_loose: str) -> str:
    src_norm = _normalize_text(source)
    m = TRAILING_CODE_RE.search(src_norm)
    if not m:
        return translated_loose
    suffix = m.group(0).strip()
    if not suffix:
        return translated_loose
    if translated_loose.endswith(suffix):
        return translated_loose
    joiner = "\t" if "\t" in source else " "
    return f"{translated_loose}{joiner}{suffix}"


def _has_leading_bullet(line: str) -> bool:
    return bool(LEADING_BULLET_RE.match(line))


def _strip_leading_legacy_bullet(line: str) -> str:
    out = LEADING_LEGACY_BULLET_RE.sub(r"\1", line, count=1)
    # Rare fallback if mojibake literal leaks through matching.
    if out != line:
        return out
    if line.lstrip().startswith("â€¢ "):
        prefix_ws = line[: len(line) - len(line.lstrip())]
        return prefix_ws + line.lstrip()[len("â€¢ ") :]
    return line


def _sanitize_translated_segment(source: str, translated: str) -> str:
    """
    Remove legacy text-bullet glyphs when the source segment does not include
    a leading bullet marker. This prevents duplicate/boxed bullets in DOCX rows
    that already carry list formatting.
    """
    src_parts = re.split(r"(\r\n|\n|\r)", source)
    out_parts = re.split(r"(\r\n|\n|\r)", translated)
    if len(src_parts) != len(out_parts):
        return translated

    cleaned: list[str] = []
    for src_part, out_part in zip(src_parts, out_parts):
        # Keep original line separators untouched.
        if out_part in {"\r\n", "\n", "\r"}:
            cleaned.append(out_part)
            continue
        if src_part in {"\r\n", "\n", "\r"}:
            cleaned.append(out_part)
            continue

        if _has_leading_bullet(src_part):
            # Preserve explicit bullet text when source includes it.
            cleaned.append(out_part)
            continue

        cleaned.append(_strip_leading_legacy_bullet(out_part))
    return "".join(cleaned)


def _strip_bracket_tags(text: str) -> str:
    return BRACKET_TAG_RE.sub(" ", text)


def _split_bracket_segments(text: str) -> list[tuple[bool, str]]:
    segments: list[tuple[bool, str]] = []
    pos = 0
    for match in BRACKET_TAG_RE.finditer(text):
        if pos < match.start():
            non_tag = text[pos : match.start()]
            if non_tag.strip():
                segments.append((False, non_tag))
        segments.append((True, match.group(0)))
        pos = match.end()
    if pos < len(text):
        tail = text[pos:]
        if tail.strip():
            segments.append((False, tail))
    return segments


def _apply_spacing_template(template: str, translated: str) -> str:
    lead_match = re.match(r"^\s*", template)
    trail_match = re.search(r"\s*$", template)
    leading_ws = lead_match.group(0) if lead_match else ""
    trailing_ws = trail_match.group(0) if trail_match else ""
    return f"{leading_ws}{translated.strip()}{trailing_ws}"


def _merge_tags_from_source(source: str, translated: str) -> str:
    if not BRACKET_TAG_RE.search(source):
        return translated
    if BRACKET_TAG_RE.search(translated):
        return translated

    source_segments = _split_bracket_segments(source)
    tags = [segment for is_tag, segment in source_segments if is_tag]
    if not tags or not source_segments:
        return translated

    body_segments = [segment for is_tag, segment in source_segments if not is_tag]
    if len(body_segments) == 1:
        rebuilt: list[str] = []
        for is_tag, segment in source_segments:
            if is_tag:
                rebuilt.append(segment)
            else:
                rebuilt.append(_apply_spacing_template(segment, translated))
        return "".join(rebuilt).strip()

    stripped_source = source.strip()
    remaining_tags = list(tags)
    leading_tags = []
    while remaining_tags and stripped_source.startswith(remaining_tags[0]):
        leading_tags.append(remaining_tags.pop(0))
        stripped_source = stripped_source[len(leading_tags[-1]) :].lstrip()

    trailing_tags = []
    stripped_source_end = source.strip()
    while remaining_tags and stripped_source_end.endswith(remaining_tags[-1]):
        trailing_tags.insert(0, remaining_tags.pop())
        stripped_source_end = stripped_source_end[: -len(trailing_tags[0])].rstrip()

    body = translated.strip()
    out = body
    if leading_tags:
        out = (" ".join(leading_tags) + " " + body).strip()
    if not remaining_tags and not trailing_tags:
        return out

    if trailing_tags:
        out = (out + " " + " ".join(trailing_tags)).strip()
    if not remaining_tags:
        return out

    if leading_tags or trailing_tags:
        # Keep unresolved in-line tags near the translated body without reordering
        # the clearly leading/trailing stage directions.
        return (out + " " + " ".join(remaining_tags)).strip()

    # Keep a single in-line tag near the original word position when possible.
    if len(remaining_tags) == 1:
        pieces = BRACKET_TAG_RE.split(stripped_source, maxsplit=1)
        if len(pieces) == 2:
            left_words = len(_normalize_loose(pieces[0]).split())
            out_words = out.split()
            if out_words:
                idx = max(0, min(len(out_words), left_words))
                out = " ".join(out_words[:idx] + [remaining_tags[0]] + out_words[idx:])
                return out.strip()

    # Fallback for multiple/complex in-line tags.
    return (" ".join(remaining_tags) + " " + out).strip()


def _token_tuple(rx: re.Pattern[str], text: str, *, upper: bool = False) -> tuple[str, ...]:
    items = rx.findall(text)
    if upper:
        return tuple(item.upper() for item in items)
    return tuple(items)


def _trailing_code_key(text: str) -> str:
    norm = _normalize_text(text)
    m = TRAILING_CODE_RE.search(norm)
    if not m:
        return ""
    return WS_RE.sub(" ", m.group(0).strip().upper())


@dataclass(frozen=True)
class LineFeatures:
    line: str
    stripped: str
    is_blank: bool
    is_ops: bool
    bracket_tags: tuple[str, ...]
    tokens: tuple[str, ...]
    qids: tuple[str, ...]
    leading_prefix: str
    trailing_codes: str
    map_code: str
    looks_optionish: bool
    has_anchor: bool


def _extract_line_features(line: str) -> LineFeatures:
    stripped = line.strip()
    bracket_tags = _token_tuple(BRACKET_TAG_RE, line)
    tokens = _token_tuple(CURLY_TOKEN_RE, line) + _token_tuple(ANGLE_TOKEN_RE, line)
    qids = _token_tuple(QID_RE, line, upper=True)
    lead_match = re.match(r"^\s*(?:\d+|[A-Za-z])(?:[\)\].:/\-]|/)\s+", line)
    leading_prefix = lead_match.group(0) if lead_match else ""
    trailing_codes = _trailing_code_key(line)
    map_match = MAP_LINE_RE.match(line)
    map_code = map_match.group(1).upper() if map_match else ""
    looks_optionish = _looks_optionish(line)
    has_anchor = bool(bracket_tags or tokens or qids or leading_prefix or trailing_codes or map_code)
    return LineFeatures(
        line=line,
        stripped=stripped,
        is_blank=not stripped,
        is_ops=_is_ops_line(line),
        bracket_tags=bracket_tags,
        tokens=tokens,
        qids=qids,
        leading_prefix=leading_prefix,
        trailing_codes=trailing_codes,
        map_code=map_code,
        looks_optionish=looks_optionish,
        has_anchor=has_anchor,
    )


def _anchors_compatible(src: LineFeatures, out: LineFeatures) -> bool:
    if src.bracket_tags and src.bracket_tags != out.bracket_tags:
        return False
    if src.tokens and src.tokens != out.tokens:
        return False
    if src.qids and src.qids != out.qids:
        return False
    if src.leading_prefix and not out.line.startswith(src.leading_prefix):
        return False
    if src.trailing_codes and src.trailing_codes != out.trailing_codes:
        return False
    if src.map_code and src.map_code != out.map_code:
        return False
    return True


def _pair_score(src: LineFeatures, out: LineFeatures) -> int:
    # Strongly preserve blank/ops structure so alignment does not drift.
    if src.is_blank:
        return 2 if out.is_blank else -3
    if src.is_ops:
        return 12 if src.line == out.line else -10

    score = 0
    anchors = 0
    contradictions = 0

    if src.bracket_tags:
        anchors += 1
        if src.bracket_tags == out.bracket_tags:
            score += 8
        else:
            score -= 10
            contradictions += 1
    if src.tokens:
        anchors += 1
        if src.tokens == out.tokens:
            score += 8
        else:
            score -= 10
            contradictions += 1
    if src.qids:
        anchors += 1
        if src.qids == out.qids:
            score += 8
        else:
            score -= 10
            contradictions += 1
    if src.leading_prefix:
        anchors += 1
        if out.line.startswith(src.leading_prefix):
            score += 5
        else:
            score -= 6
            contradictions += 1
    if src.trailing_codes:
        anchors += 1
        if src.trailing_codes == out.trailing_codes:
            score += 7
        else:
            score -= 8
            contradictions += 1
    if src.map_code:
        anchors += 1
        if src.map_code == out.map_code:
            score += 7
        else:
            score -= 8
            contradictions += 1

    if src.looks_optionish and not out.trailing_codes:
        score -= 4
    if out.is_ops:
        score -= 6

    if anchors:
        score += 1
    if contradictions >= 2:
        score -= 4
    return score


def _align_pairs(src_feats: list[LineFeatures], out_feats: list[LineFeatures]) -> list[tuple[int | None, int | None]]:
    n = len(src_feats)
    m = len(out_feats)
    gap_penalty = -2

    scores = [[0] * (m + 1) for _ in range(n + 1)]
    back = [[""] * (m + 1) for _ in range(n + 1)]

    for i in range(1, n + 1):
        scores[i][0] = scores[i - 1][0] + gap_penalty
        back[i][0] = "U"
    for j in range(1, m + 1):
        scores[0][j] = scores[0][j - 1] + gap_penalty
        back[0][j] = "L"

    for i in range(1, n + 1):
        src = src_feats[i - 1]
        for j in range(1, m + 1):
            out = out_feats[j - 1]
            diag = scores[i - 1][j - 1] + _pair_score(src, out)
            up = scores[i - 1][j] + gap_penalty
            left = scores[i][j - 1] + gap_penalty

            best = diag
            step = "D"
            if up > best:
                best = up
                step = "U"
            if left > best:
                best = left
                step = "L"
            scores[i][j] = best
            back[i][j] = step

    aligned: list[tuple[int | None, int | None]] = []
    i = n
    j = m
    while i > 0 or j > 0:
        if i > 0 and j > 0 and back[i][j] == "D":
            aligned.append((i - 1, j - 1))
            i -= 1
            j -= 1
            continue
        if i > 0 and (j == 0 or back[i][j] == "U"):
            aligned.append((i - 1, None))
            i -= 1
            continue
        aligned.append((None, j - 1))
        j -= 1

    aligned.reverse()
    return aligned


def _semantic_item_marker(text: Any) -> str:
    if not isinstance(text, str):
        return ""
    m = ITEM_MARKER_RE.match(text)
    if m:
        return m.group(1).upper()
    m = BRACKET_MARKER_RE.match(text)
    if m:
        marker = _normalize_loose(m.group(1)).upper()
        return marker
    return ""


def _semantic_list_item_key(parent_key: str, item: Any) -> str | None:
    if not isinstance(item, dict):
        return None

    if parent_key == "elements":
        item_type = str(item.get("type") or "").strip().lower()
        if item_type == "directive":
            # Directive ordering can legitimately differ between the English and
            # Spanish JSONs, which makes positional list-path pairing unsafe.
            return None

        parts: list[str] = [item_type or "item"]
        qid = str(item.get("qid") or item.get("qnum") or "").strip()
        if qid:
            parts.append(f"qid={qid.upper()}")

        marker = _semantic_item_marker(item.get("item_text")) or _semantic_item_marker(item.get("text"))
        if marker:
            parts.append(f"marker={marker}")

        response_options = item.get("response_options")
        if isinstance(response_options, list):
            values = [
                str(opt.get("value") or "").strip().upper()
                for opt in response_options
                if isinstance(opt, dict) and str(opt.get("value") or "").strip()
            ]
            if values:
                parts.append(f"vals={','.join(values[:10])}")

        if len(parts) == 1:
            return None
        return "|".join(parts)

    if parent_key == "response_options":
        value = str(item.get("value") or "").strip()
        if value:
            return f"value={value.upper()}"
        marker = _semantic_item_marker(item.get("text"))
        if marker:
            return f"marker={marker}"
        return None

    if parent_key == "notes":
        return None

    return None


def _extract_json_text_fields(obj: Any, *, path: str = "", parent_key: str = "") -> dict[str, str]:
    out: dict[str, str] = {}
    if isinstance(obj, dict):
        for key, val in obj.items():
            if key == "notes":
                continue
            child = f"{path}.{key}" if path else key
            if key in {"text", "item_text"} and isinstance(val, str):
                out[child] = val
            nested = _extract_json_text_fields(val, path=child, parent_key=key)
            if nested:
                out.update(nested)
        return out
    if isinstance(obj, list):
        seen_semantic_keys: dict[str, int] = defaultdict(int)
        for idx, val in enumerate(obj):
            semantic_key = _semantic_list_item_key(parent_key, val)
            if parent_key in {"elements", "notes"} and semantic_key is None:
                continue
            if semantic_key is None:
                child = f"{path}[{idx}]"
            else:
                seen_semantic_keys[semantic_key] += 1
                ordinal = seen_semantic_keys[semantic_key]
                suffix = f"#{ordinal}" if ordinal > 1 else ""
                child = f"{path}[{semantic_key}{suffix}]"
            nested = _extract_json_text_fields(val, path=child)
            if nested:
                out.update(nested)
        return out
    return out


def _sentence_chunks(text: str) -> list[str]:
    s = text.strip()
    if not s:
        return []
    parts = [chunk.strip() for chunk in SENTENCE_SPLIT_RE.split(s) if chunk.strip()]
    if len(parts) <= 1:
        return []
    return parts


def _json_text_pairs(json_en_text: str, json_es_text: str) -> list[tuple[str, str]]:
    try:
        en_obj = json.loads(json_en_text)
        es_obj = json.loads(json_es_text)
    except Exception:
        return []

    en_fields = _extract_json_text_fields(en_obj)
    es_fields = _extract_json_text_fields(es_obj)
    pairs: list[tuple[str, str]] = []

    for path in sorted(set(en_fields).intersection(es_fields)):
        src = en_fields[path]
        out = es_fields[path]
        if not src.strip() or not out.strip():
            continue
        if _is_ops_line(src) and _is_ops_line(out):
            continue
        pairs.append((src, out))

        src_lines = src.splitlines()
        out_lines = out.splitlines()
        if len(src_lines) >= 2 and len(src_lines) == len(out_lines):
            for s_line, o_line in zip(src_lines, out_lines):
                if not s_line.strip() or not o_line.strip():
                    continue
                if _is_ops_line(s_line) and _is_ops_line(o_line):
                    continue
                pairs.append((s_line, o_line))

        src_sents = _sentence_chunks(src)
        out_sents = _sentence_chunks(out)
        if src_sents and len(src_sents) == len(out_sents):
            for s_sent, o_sent in zip(src_sents, out_sents):
                if not s_sent.strip() or not o_sent.strip():
                    continue
                if _is_ops_line(s_sent) and _is_ops_line(o_sent):
                    continue
                pairs.append((s_sent, o_sent))

    return pairs


@dataclass(frozen=True)
class MatchResult:
    translated: str
    match_type: str


@dataclass
class MappingBundle:
    exact: Dict[str, str]
    normalized: Dict[str, str]
    normalized_ambiguous: set[str]
    loose: Dict[str, str]
    loose_ambiguous: set[str]
    source_pairs: int
    skipped_ops_or_blank: int
    json_pairs_used: int


@dataclass
class ReplaceRecord:
    location: str
    match_type: str
    source: str
    target: str


@dataclass
class RebuildResult:
    matched_segments: int
    unmatched_segments: int
    changed_paragraphs: int
    total_paragraphs: int
    report: List[ReplaceRecord]
    mapping: MappingBundle


def build_mapping(
    raw_en_text: str,
    raw_es_text: str,
    *,
    json_en_text: str | None = None,
    json_es_text: str | None = None,
) -> MappingBundle:
    src_raw_lines = raw_en_text.splitlines(keepends=True)
    out_raw_lines = raw_es_text.splitlines(keepends=True)
    src_lines = [line.rstrip("\r\n") for line in src_raw_lines]
    out_lines = [line.rstrip("\r\n") for line in out_raw_lines]

    src_feats = [_extract_line_features(line) for line in src_lines]
    out_feats = [_extract_line_features(line) for line in out_lines]
    aligned = _align_pairs(src_feats, out_feats)
    alignment_risky = len(src_lines) != len(out_lines)

    exact_counts: dict[str, Counter[str]] = defaultdict(Counter)
    norm_counts: dict[str, Counter[str]] = defaultdict(Counter)
    loose_counts: dict[str, Counter[str]] = defaultdict(Counter)
    aligned_join_rows: list[tuple[int, int, str, str]] = []

    def add_pair(
        src: str,
        out: str,
        *,
        weight: int = 1,
        include_for_join: bool = False,
        src_idx: int | None = None,
        out_idx: int | None = None,
    ) -> None:
        if not src.strip() or _is_ops_line(src):
            return
        exact_counts[src][out] += weight
        norm_counts[_normalize_text(src)][out] += weight
        loose_src = _normalize_loose(src)
        loose_out = _normalize_loose(out)
        if loose_src:
            loose_counts[loose_src][loose_out] += weight

        # Instruction-flex matching: allow JSON text (without inline tags) to match
        # DOCX lines that include the same text with bracket tags in-line.
        if BRACKET_TAG_RE.search(src):
            src_tagless = _strip_bracket_tags(src)
            out_tagless = _strip_bracket_tags(out)
            norm_src_tagless = _normalize_text(src_tagless)
            if norm_src_tagless:
                norm_counts[norm_src_tagless][out_tagless.strip() or out] += weight
            loose_src_tagless = _normalize_loose(src_tagless)
            loose_out_tagless = _normalize_loose(out_tagless)
            if loose_src_tagless:
                loose_counts[loose_src_tagless][loose_out_tagless or loose_out] += weight

        src_bullets = _split_bullet_fragments(src)
        out_bullets = _split_bullet_fragments(out)
        if src_bullets and len(src_bullets) == len(out_bullets):
            for s_part, o_part in zip(src_bullets, out_bullets):
                key = _normalize_loose(s_part)
                if key:
                    loose_counts[key][_normalize_loose(o_part)] += weight

        src_pairs = _option_pairs(src)
        out_pairs = _option_pairs(out)
        if src_pairs and len(src_pairs) == len(out_pairs):
            for (s_label, s_code), (o_label, o_code) in zip(src_pairs, out_pairs):
                s_label_code = f"{s_label} {s_code}"
                o_label_code = f"{o_label} {o_code}"
                norm_counts[_normalize_text(s_label_code)][o_label_code] += weight
                key = _normalize_loose(s_label_code)
                if key:
                    loose_counts[key][_normalize_loose(o_label_code)] += weight
                label_key = _normalize_loose(s_label)
                if label_key:
                    loose_counts[label_key][_normalize_loose(o_label)] += weight

        if include_for_join and src_idx is not None and out_idx is not None:
            aligned_join_rows.append((src_idx, out_idx, src, out))

    json_pairs_used = 0
    if json_en_text is not None and json_es_text is not None:
        for src, out in _json_text_pairs(json_en_text, json_es_text):
            add_pair(src, out, weight=4, include_for_join=False)
            json_pairs_used += 1

    skipped = sum(1 for feat in src_feats if feat.is_blank or feat.is_ops)

    paired_rows: list[dict[str, object]] = []
    for src_idx, out_idx in aligned:
        if src_idx is None or out_idx is None:
            continue
        src_feat = src_feats[src_idx]
        if src_feat.is_blank or src_feat.is_ops:
            continue
        out_feat = out_feats[out_idx]
        score = _pair_score(src_feat, out_feat)
        anchors_ok = _anchors_compatible(src_feat, out_feat)
        paired_rows.append(
            {
                "src_idx": src_idx,
                "out_idx": out_idx,
                "delta": out_idx - src_idx,
                "src": src_feat.line,
                "out": out_feat.line,
                "score": score,
                "has_anchor": src_feat.has_anchor,
                "anchor_trusted": bool(src_feat.has_anchor and anchors_ok and score >= 0),
                "anchors_ok": anchors_ok,
            }
        )

    prev_anchor: list[int | None] = [None] * len(paired_rows)
    last_anchor: int | None = None
    for idx, row in enumerate(paired_rows):
        if bool(row["anchor_trusted"]):
            last_anchor = idx
        prev_anchor[idx] = last_anchor

    next_anchor: list[int | None] = [None] * len(paired_rows)
    last_anchor = None
    for idx in range(len(paired_rows) - 1, -1, -1):
        row = paired_rows[idx]
        if bool(row["anchor_trusted"]):
            last_anchor = idx
        next_anchor[idx] = last_anchor

    for idx, row in enumerate(paired_rows):
        score = int(row["score"])
        if score < 0:
            continue
        has_anchor = bool(row["has_anchor"])
        if has_anchor:
            if not bool(row["anchors_ok"]):
                continue
        elif alignment_risky:
            prev_idx = prev_anchor[idx]
            next_idx = next_anchor[idx]
            current_delta = int(row["delta"])
            if prev_idx is None or next_idx is None:
                # Allow top preamble rows before first trusted anchor when local
                # offset is stable and source position is near file start.
                if prev_idx is None and next_idx is not None:
                    next_delta = int(paired_rows[next_idx]["delta"])
                    src_pos = int(row["src_idx"])
                    if current_delta != next_delta or src_pos > 40:
                        continue
                else:
                    continue
            else:
                prev_delta = int(paired_rows[prev_idx]["delta"])
                next_delta = int(paired_rows[next_idx]["delta"])
                if current_delta != prev_delta or current_delta != next_delta:
                    continue
                src_span = int(paired_rows[next_idx]["src_idx"]) - int(paired_rows[prev_idx]["src_idx"])
                if src_span > 24:
                    continue

        src = str(row["src"])
        out = str(row["out"])
        add_pair(
            src,
            out,
            weight=1,
            include_for_join=True,
            src_idx=int(row["src_idx"]),
            out_idx=int(row["out_idx"]),
        )

    # Add joined wrapped-line variants from locally consecutive, aligned rows.
    for span in (2, 3, 4):
        for i in range(0, max(0, len(aligned_join_rows) - span + 1)):
            chunk = aligned_join_rows[i : i + span]
            consecutive = True
            for j in range(len(chunk) - 1):
                s0, o0, _, _ = chunk[j]
                s1, o1, _, _ = chunk[j + 1]
                if s1 != s0 + 1 or o1 != o0 + 1:
                    consecutive = False
                    break
            if not consecutive:
                continue
            src_chunk = [s for _, _, s, _ in chunk]
            out_chunk = [o for _, _, _, o in chunk]
            joined_src = _normalize_loose(" ".join(src_chunk))
            joined_out = _normalize_loose(" ".join(out_chunk))
            if joined_src:
                loose_counts[joined_src][joined_out] += 1

    exact = {src: _dominant_variant(cnt) for src, cnt in exact_counts.items()}
    normalized: dict[str, str] = {}
    normalized_ambiguous: set[str] = set()
    for key, cnt in norm_counts.items():
        if len(cnt) == 1:
            normalized[key] = _dominant_variant(cnt)
        else:
            normalized_ambiguous.add(key)

    loose: dict[str, str] = {}
    loose_ambiguous: set[str] = set()
    for key, cnt in loose_counts.items():
        if not key:
            continue
        if len(cnt) == 1:
            loose[key] = _dominant_variant(cnt)
        else:
            loose_ambiguous.add(key)

    return MappingBundle(
        exact=exact,
        normalized=normalized,
        normalized_ambiguous=normalized_ambiguous,
        loose=loose,
        loose_ambiguous=loose_ambiguous,
        source_pairs=min(len(src_lines), len(out_lines)),
        skipped_ops_or_blank=skipped,
        json_pairs_used=json_pairs_used,
    )


def match_segment(segment: str, mapping: MappingBundle) -> Optional[MatchResult]:
    if segment in mapping.exact:
        translated = _apply_surface_layout_from_source(segment, mapping.exact[segment])
        translated = _sanitize_translated_segment(segment, translated)
        return MatchResult(translated, "exact")

    normalized_key = _normalize_text(segment)
    if not normalized_key:
        return None
    if normalized_key in mapping.normalized_ambiguous:
        return None
    if normalized_key in mapping.normalized:
        translated = _apply_surface_layout_from_source(segment, mapping.normalized[normalized_key])
        translated = _sanitize_translated_segment(segment, translated)
        return MatchResult(translated, "normalized")

    if BRACKET_TAG_RE.search(segment):
        normalized_tagless = _normalize_text(_strip_bracket_tags(segment))
        if normalized_tagless and normalized_tagless not in mapping.normalized_ambiguous:
            if normalized_tagless in mapping.normalized:
                translated = _merge_tags_from_source(segment, mapping.normalized[normalized_tagless])
                translated = _apply_surface_layout_from_source(segment, translated)
                translated = _sanitize_translated_segment(segment, translated)
                return MatchResult(translated, "normalized_tagless")

    loose_key = _normalize_loose(segment)
    if not loose_key:
        return None
    if loose_key in mapping.loose_ambiguous:
        return None
    if loose_key in mapping.loose:
        translated = _append_source_suffix_if_any(segment, mapping.loose[loose_key])
        translated = _apply_surface_layout_from_source(segment, translated)
        translated = _sanitize_translated_segment(segment, translated)
        return MatchResult(translated, "loose")

    if BRACKET_TAG_RE.search(segment):
        loose_tagless = _normalize_loose(_strip_bracket_tags(segment))
        if loose_tagless and loose_tagless not in mapping.loose_ambiguous:
            if loose_tagless in mapping.loose:
                translated = _append_source_suffix_if_any(segment, mapping.loose[loose_tagless])
                translated = _merge_tags_from_source(segment, translated)
                translated = _apply_surface_layout_from_source(segment, translated)
                translated = _sanitize_translated_segment(segment, translated)
                return MatchResult(translated, "loose_tagless")
    return None


def translate_multiline_text(text: str, mapping: MappingBundle) -> tuple[str, int, int, list[tuple[str, str, str]]]:
    parts = re.split(r"(\r\n|\n|\r)", text)
    out_parts: list[str] = []
    matched = 0
    unmatched = 0
    details: list[tuple[str, str, str]] = []

    for idx, part in enumerate(parts):
        # separators are odd indexes because split keeps delimiters
        if idx % 2 == 1:
            out_parts.append(part)
            continue
        if not part.strip():
            out_parts.append(part)
            continue
        if _is_ops_line(part):
            # Keep ops/routing tags as-is and do not count as unmatched.
            out_parts.append(part)
            continue
        m = match_segment(part, mapping)
        if m is None:
            unmatched += 1
            out_parts.append(part)
            continue
        matched += 1
        out_parts.append(m.translated)
        details.append((m.match_type, part, m.translated))
    return ("".join(out_parts), matched, unmatched, details)


def _paragraph_text(paragraph: Paragraph) -> str:
    if paragraph.runs:
        return "".join(run.text for run in paragraph.runs)
    return paragraph.text


def _redistribute_text_across_runs(runs: List[Run], new_text: str) -> None:
    if not runs:
        return
    if len(runs) == 1:
        runs[0].text = new_text
        return

    lengths = [len(r.text) for r in runs]
    total = sum(lengths)
    if total <= 0:
        runs[0].text = new_text
        for r in runs[1:]:
            r.text = ""
        return

    pos = 0
    cumulative = 0
    for i in range(len(runs) - 1):
        cumulative += lengths[i]
        end = round((cumulative / total) * len(new_text))
        if end < pos:
            end = pos
        runs[i].text = new_text[pos:end]
        pos = end
    runs[-1].text = new_text[pos:]


def _segment_spans_with_tags(text: str) -> list[tuple[bool, int, int, str]]:
    spans: list[tuple[bool, int, int, str]] = []
    pos = 0
    for match in BRACKET_TAG_RE.finditer(text):
        if pos < match.start():
            part = text[pos : match.start()]
            if part.strip():
                spans.append((False, pos, match.start(), part))
        tag_end = match.end()
        while tag_end < len(text) and text[tag_end].isspace():
            tag_end += 1
        spans.append((True, match.start(), tag_end, text[match.start() : tag_end]))
        pos = tag_end
    if pos < len(text):
        tail = text[pos:]
        if tail.strip():
            spans.append((False, pos, len(text), tail))
    return spans


def _set_paragraph_text_preserving_tag_runs(paragraph: Paragraph, new_text: str) -> bool:
    runs = list(paragraph.runs)
    if not runs:
        return False

    source_text = _paragraph_text(paragraph)
    source_tags = BRACKET_TAG_RE.findall(source_text)
    new_tags = BRACKET_TAG_RE.findall(new_text)
    if not source_tags or source_tags != new_tags:
        return False

    source_segments = _segment_spans_with_tags(source_text)
    new_segments = _segment_spans_with_tags(new_text)
    if len(source_segments) != len(new_segments):
        return False
    if [is_tag for is_tag, *_ in source_segments] != [is_tag for is_tag, *_ in new_segments]:
        return False

    run_spans: list[tuple[Run, int, int]] = []
    cursor = 0
    for run in runs:
        next_cursor = cursor + len(run.text)
        run_spans.append((run, cursor, next_cursor))
        cursor = next_cursor

    grouped_runs: list[list[Run]] = []
    used_run_ids: set[int] = set()
    for _is_tag, seg_start, seg_end, _seg_text in source_segments:
        segment_runs: list[Run] = []
        for run, run_start, run_end in run_spans:
            if run_end <= seg_start or run_start >= seg_end:
                continue
            if run_start < seg_start or run_end > seg_end:
                return False
            segment_runs.append(run)
            used_run_ids.add(id(run))
        if not segment_runs:
            return False
        grouped_runs.append(segment_runs)

    for run in runs:
        if id(run) not in used_run_ids:
            run.text = ""

    for segment_runs, (_, _, _, segment_text) in zip(grouped_runs, new_segments):
        _redistribute_text_across_runs(segment_runs, segment_text)
    return True


def _set_paragraph_text(paragraph: Paragraph, new_text: str) -> None:
    if paragraph.runs:
        if _set_paragraph_text_preserving_tag_runs(paragraph, new_text):
            return
        _redistribute_text_across_runs(list(paragraph.runs), new_text)
    else:
        paragraph.add_run(new_text)


def _iter_table_paragraphs(
    table: Table,
    *,
    prefix: str,
    seen_cells: set[int],
    seen_paragraphs: set[int],
) -> Iterable[tuple[Paragraph, str]]:
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell_obj: _Cell = cell
            cell_key = id(cell_obj._tc)
            if cell_key in seen_cells:
                continue
            seen_cells.add(cell_key)

            for p_idx, paragraph in enumerate(cell_obj.paragraphs):
                p_key = id(paragraph._p)
                if p_key in seen_paragraphs:
                    continue
                seen_paragraphs.add(p_key)
                loc = f"{prefix}r{r_idx}:c{c_idx}:p{p_idx}"
                yield paragraph, loc

            for t_idx, nested in enumerate(cell_obj.tables):
                nested_prefix = f"{prefix}r{r_idx}:c{c_idx}:t{t_idx}:"
                yield from _iter_table_paragraphs(
                    nested,
                    prefix=nested_prefix,
                    seen_cells=seen_cells,
                    seen_paragraphs=seen_paragraphs,
                )


def _iter_story_paragraphs(
    paragraphs: List[Paragraph],
    tables: List[Table],
    *,
    prefix: str,
    seen_cells: set[int],
    seen_paragraphs: set[int],
) -> Iterable[tuple[Paragraph, str]]:
    for p_idx, paragraph in enumerate(paragraphs):
        p_key = id(paragraph._p)
        if p_key in seen_paragraphs:
            continue
        seen_paragraphs.add(p_key)
        yield paragraph, f"{prefix}p{p_idx}"

    for t_idx, table in enumerate(tables):
        yield from _iter_table_paragraphs(
            table,
            prefix=f"{prefix}t{t_idx}:",
            seen_cells=seen_cells,
            seen_paragraphs=seen_paragraphs,
        )


def _iter_all_paragraphs(document: DocumentType) -> Iterable[tuple[Paragraph, str]]:
    seen_cells: set[int] = set()
    seen_paragraphs: set[int] = set()

    yield from _iter_story_paragraphs(
        document.paragraphs,
        document.tables,
        prefix="doc:",
        seen_cells=seen_cells,
        seen_paragraphs=seen_paragraphs,
    )

    seen_parts: set[int] = set()
    for s_idx, section in enumerate(document.sections):
        sec: Section = section
        part_defs = [
            ("header", sec.header),
            ("footer", sec.footer),
            ("first_header", sec.first_page_header),
            ("first_footer", sec.first_page_footer),
            ("even_header", sec.even_page_header),
            ("even_footer", sec.even_page_footer),
        ]
        for label, part in part_defs:
            part_key = id(part._element)
            if part_key in seen_parts:
                continue
            seen_parts.add(part_key)
            yield from _iter_story_paragraphs(
                part.paragraphs,
                part.tables,
                prefix=f"sec{s_idx}:{label}:",
                seen_cells=seen_cells,
                seen_paragraphs=seen_paragraphs,
            )


def rebuild_docx(
    *,
    docx_in: Path,
    raw_en: Path,
    raw_es: Path,
    json_en: Path | None = None,
    json_es: Path | None = None,
    docx_out: Path,
) -> RebuildResult:
    raw_en_text = raw_en.read_text(encoding="utf-8", errors="replace")
    raw_es_text = raw_es.read_text(encoding="utf-8", errors="replace")
    json_en_text: str | None = None
    json_es_text: str | None = None
    if json_en is not None and json_es is not None and json_en.exists() and json_es.exists():
        json_en_text = json_en.read_text(encoding="utf-8", errors="replace")
        json_es_text = json_es.read_text(encoding="utf-8", errors="replace")

    mapping = build_mapping(
        raw_en_text,
        raw_es_text,
        json_en_text=json_en_text,
        json_es_text=json_es_text,
    )

    doc = Document(str(docx_in))
    report: list[ReplaceRecord] = []
    matched = 0
    unmatched = 0
    changed_paragraphs = 0
    total_paragraphs = 0

    for paragraph, location in _iter_all_paragraphs(doc):
        total_paragraphs += 1
        original = _paragraph_text(paragraph)
        if not original.strip():
            continue
        replaced, m_count, u_count, details = translate_multiline_text(original, mapping)
        matched += m_count
        unmatched += u_count
        if replaced == original:
            continue
        _set_paragraph_text(paragraph, replaced)
        changed_paragraphs += 1
        for m_type, source, target in details:
            report.append(
                ReplaceRecord(
                    location=location,
                    match_type=m_type,
                    source=source,
                    target=target,
                )
            )

    if docx_out.parent and not docx_out.parent.exists():
        docx_out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_out))

    return RebuildResult(
        matched_segments=matched,
        unmatched_segments=unmatched,
        changed_paragraphs=changed_paragraphs,
        total_paragraphs=total_paragraphs,
        report=report,
        mapping=mapping,
    )


def write_report(path: Path, result: RebuildResult) -> None:
    payload = {
        "summary": {
            "matched_segments": result.matched_segments,
            "unmatched_segments": result.unmatched_segments,
            "changed_paragraphs": result.changed_paragraphs,
            "total_paragraphs": result.total_paragraphs,
            "mapping_exact_keys": len(result.mapping.exact),
            "mapping_normalized_keys": len(result.mapping.normalized),
            "mapping_ambiguous_normalized_keys": len(result.mapping.normalized_ambiguous),
            "mapping_loose_keys": len(result.mapping.loose),
            "mapping_ambiguous_loose_keys": len(result.mapping.loose_ambiguous),
            "mapping_source_pairs_used": result.mapping.source_pairs,
            "mapping_skipped_ops_or_blank": result.mapping.skipped_ops_or_blank,
            "mapping_json_pairs_used": result.mapping.json_pairs_used,
        },
        "replacements": [
            {
                "location": r.location,
                "match_type": r.match_type,
                "source": r.source,
                "target": r.target,
            }
            for r in result.report
        ],
    }
    if path.parent and not path.parent.exists():
        path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Create a formatted Spanish DOCX by replacing English strings in a source DOCX "
            "using aligned TOMAS RAW.en/RAW.es outputs."
        )
    )
    parser.add_argument("--docx", required=True, help="Source English DOCX template")
    parser.add_argument("--raw-en", required=True, help="Aligned English RAW text")
    parser.add_argument("--raw-es", required=True, help="Aligned Spanish RAW text")
    parser.add_argument("--json-en", default="", help="Optional English JSON from RAW.en (preferred mapping source)")
    parser.add_argument("--json-es", default="", help="Optional Spanish JSON from RAW.es (preferred mapping source)")
    parser.add_argument("-o", "--out", required=True, help="Output Spanish DOCX path")
    parser.add_argument(
        "--report",
        default="",
        help="Optional JSON replacement report path (default: <out>.replace_report.json)",
    )
    return parser


def main(argv: Optional[list[str]] = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)

    docx_in = Path(args.docx)
    raw_en = Path(args.raw_en)
    raw_es = Path(args.raw_es)
    json_en = Path(args.json_en) if args.json_en else None
    json_es = Path(args.json_es) if args.json_es else None
    out = Path(args.out)
    report_path = Path(args.report) if args.report else out.with_suffix(out.suffix + ".replace_report.json")

    missing = [p for p in (docx_in, raw_en, raw_es) if not p.exists()]
    if json_en is not None and json_es is None:
        raise SystemExit("--json-en requires --json-es.")
    if json_es is not None and json_en is None:
        raise SystemExit("--json-es requires --json-en.")
    if json_en is not None and not json_en.exists():
        missing.append(json_en)
    if json_es is not None and not json_es.exists():
        missing.append(json_es)
    if missing:
        joined = "\n".join(f"- {p}" for p in missing)
        raise SystemExit(f"Missing required input file(s):\n{joined}")

    result = rebuild_docx(
        docx_in=docx_in,
        raw_en=raw_en,
        raw_es=raw_es,
        json_en=json_en,
        json_es=json_es,
        docx_out=out,
    )
    write_report(report_path, result)

    print("[tomas-docx] Complete.")
    print(f"[tomas-docx] Output DOCX: {out}")
    print(f"[tomas-docx] Report: {report_path}")
    print(
        "[tomas-docx] Summary: "
        f"changed_paragraphs={result.changed_paragraphs}/{result.total_paragraphs}, "
        f"matched_segments={result.matched_segments}, "
        f"unmatched_segments={result.unmatched_segments}"
    )
    if result.mapping.json_pairs_used:
        print(f"[tomas-docx] JSON-driven mapping pairs used: {result.mapping.json_pairs_used}")
    if result.mapping.normalized_ambiguous:
        print(
            "[tomas-docx] Note: "
            f"{len(result.mapping.normalized_ambiguous)} normalized source keys were ambiguous and skipped."
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
