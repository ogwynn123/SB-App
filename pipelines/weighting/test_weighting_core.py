from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill

THIS_DIR = Path(__file__).resolve().parent
if str(THIS_DIR) not in sys.path:
    sys.path.insert(0, str(THIS_DIR))

from weighting_core import (  # noqa: E402
    BlockAxes,
    Variable,
    WeightingConfig,
    assign_labels,
    build_dimensions,
    draft_config,
    parse_code_spec,
    rake,
    read_target_blocks,
    run_weighting,
    write_outputs,
)

YELLOW = PatternFill(fill_type="solid", fgColor="FFFFFF00")


def _write_block(ws, top: int, left: int, title: str, rows: list[str], cols: list[str], values) -> None:
    ws.cell(top - 2, 1, title)
    for c, label in enumerate(cols):
        ws.cell(top - 1, left + c, label)
    for r, label in enumerate(rows):
        ws.cell(top + r, left - 1, label)
        for c in range(len(cols)):
            cell = ws.cell(top + r, left + c, values[r][c])
            cell.fill = YELLOW


def _workbook(path: Path, region_sex, age_sex) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Weighting"
    _write_block(ws, 5, 3, "Sex within Region", ["Region 1", "Region 2"], ["Men", "Women"], region_sex)
    _write_block(ws, 12, 3, "Sex within Age", ["Young", "Old"], ["Men", "Women"], age_sex)
    ws.cell(20, 3).fill = YELLOW  # stray empty yellow cell must be ignored
    wb.save(path)


def _sample(seed: int = 7, n: int = 600) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    return pd.DataFrame(
        {
            "sex": rng.choice(["1", "2"], n, p=[0.35, 0.65]),
            "region": rng.choice(["1", "2"], n, p=[0.7, 0.3]),
            "age": rng.choice([str(a) for a in range(1, 11)], n),
        }
    )


def _config() -> WeightingConfig:
    return WeightingConfig(
        variables={
            "Sex": Variable("Sex", "sex", {"Men": "1", "Women": "2"}),
            "Region": Variable("Region", "region", {"Region 1": "1", "Region 2": "2"}),
            "Age": Variable("Age", "age", {"Young": "1-5", "Old": "*"}),
        },
        blocks=[
            BlockAxes("Sex within Region [C5:D6]", "Region", "Sex"),
            BlockAxes("Sex within Age [C12:D13]", "Age", "Sex"),
        ],
        cap=4.0,
    )


class CodeSpecTests(unittest.TestCase):
    def test_ranges_lists_blank_and_rest(self):
        self.assertEqual(parse_code_spec("1-3,7"), ({"1", "2", "3", "7"}, False))
        self.assertEqual(parse_code_spec("U,N,blank"), ({"U", "N", ""}, False))
        self.assertEqual(parse_code_spec("*"), (set(), True))

    def test_rest_label_takes_unclaimed_codes(self):
        var = Variable("Party", "party", {"Dem": "Democrat", "GOP": "Republican", "Ind": "*"})
        out = assign_labels(pd.Series(["Democrat", "Libertarian", "Republican", ""]), var)
        self.assertEqual(list(out), ["Dem", "Ind", "GOP", "Ind"])

    def test_float_codes_from_excel_match_integer_specs(self):
        var = Variable("Sex", "q100", {"Men": "1", "Women": "2"})
        self.assertEqual(list(assign_labels(pd.Series(["1.0", "2"]), var)), ["Men", "Women"])


class WorkbookTests(unittest.TestCase):
    def test_reads_yellow_blocks_with_labels(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "targets.xlsx"
            _workbook(path, [[0.2, 0.3], [0.25, 0.25]], [[0.25, 0.25], [0.25, 0.25]])
            blocks = read_target_blocks(path)
        self.assertEqual([b.key for b in blocks], ["Sex within Region [C5:D6]", "Sex within Age [C12:D13]"])
        self.assertEqual(blocks[0].row_labels, ["Region 1", "Region 2"])
        self.assertEqual(blocks[0].col_labels, ["Men", "Women"])
        self.assertAlmostEqual(blocks[0].total, 1.0)

    def test_draft_config_names_shared_axis_from_titles(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "targets.xlsx"
            _workbook(path, [[0.2, 0.3], [0.25, 0.25]], [[0.25, 0.25], [0.25, 0.25]])
            config = draft_config(read_target_blocks(path))
        self.assertEqual(set(config.variables), {"Sex", "Region", "Age"})
        self.assertEqual((config.blocks[0].row_variable, config.blocks[0].col_variable), ("Region", "Sex"))


class RakeTests(unittest.TestCase):
    def _run(self, region_sex, age_sex, cap=4.0, df=None):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "targets.xlsx"
            _workbook(path, region_sex, age_sex)
            blocks = read_target_blocks(path)
        config = _config()
        config.cap = cap
        return run_weighting(df if df is not None else _sample(), blocks, config), blocks, config

    def test_consistent_targets_are_hit_exactly(self):
        result, _blocks, _config_ = self._run([[0.2, 0.3], [0.25, 0.25]], [[0.22, 0.28], [0.23, 0.27]])
        self.assertTrue(result.converged)
        self.assertLess(result.max_error, 1e-4)
        self.assertAlmostEqual(result.weights.mean(), 1.0, places=6)

    def test_weights_respect_cap_and_never_go_negative(self):
        result, _blocks, _config_ = self._run([[0.05, 0.05], [0.45, 0.45]], [[0.25, 0.25], [0.25, 0.25]], cap=2.0)
        self.assertLessEqual(result.weights.max(), 2.0 + 1e-9)
        self.assertGreaterEqual(result.weights.min(), 0.0)
        self.assertFalse(result.converged)
        self.assertTrue(any("closest fit" in w for w in result.warnings))

    def test_empty_cell_target_moves_within_shared_axis(self):
        df = _sample()
        df.loc[(df.region == "2") & (df.sex == "1"), "region"] = "1"  # nobody is Region 2 | Men
        result, _blocks, _config_ = self._run([[0.2, 0.3], [0.25, 0.25]], [[0.22, 0.23], [0.23, 0.32]], df=df)
        self.assertTrue(any("share moved to the rest of 'Men'" in w for w in result.warnings))
        men = sum(c.weighted_share for c in result.cells if c.dimension == "Sex within Region" and c.cell.endswith("Men"))
        self.assertAlmostEqual(men, 0.45, places=3)

    def test_conflicting_targets_split_the_miss(self):
        # Region block says Men 50%, age block says Men 40%: no weights satisfy both.
        result, _blocks, _config_ = self._run([[0.25, 0.25], [0.25, 0.25]], [[0.2, 0.3], [0.2, 0.3]])
        self.assertTrue(any("least-squares" in w for w in result.warnings))
        men_region = sum(c.weighted_share for c in result.cells if c.dimension == "Sex within Region" and "Men" in c.cell)
        self.assertAlmostEqual(men_region, 0.45, delta=0.01)

    def test_unmapped_codes_are_reported(self):
        df = _sample()
        df.loc[0, "sex"] = "9"
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "targets.xlsx"
            _workbook(path, [[0.2, 0.3], [0.25, 0.25]], [[0.25, 0.25], [0.25, 0.25]])
            blocks = read_target_blocks(path)
        with self.assertRaisesRegex(ValueError, r"codes with no target label: 9 x1"):
            build_dimensions(df, blocks, _config())

    def test_zero_target_with_respondents_gets_minimum_weight(self):
        result, _blocks, _config_ = self._run([[0.5, 0.5], [0.0, 0.0]], [[0.25, 0.25], [0.25, 0.25]])
        region2 = _sample().region.to_numpy() == "2"
        np.testing.assert_allclose(result.weights[region2], 0.001, atol=1e-9)
        self.assertGreater(result.weights.min(), 0)
        self.assertAlmostEqual(result.weights.sum(), 600, places=6)

    def test_outputs_written(self):
        result, _blocks, config = self._run([[0.2, 0.3], [0.25, 0.25]], [[0.22, 0.28], [0.23, 0.27]])
        with tempfile.TemporaryDirectory() as tmp:
            data_out, report = Path(tmp) / "out.csv", Path(tmp) / "report.xlsx"
            write_outputs(_sample(), result, config, data_out, report)
            written = pd.read_csv(data_out, dtype={"sex": str})
            sheets = pd.read_excel(report, sheet_name=None)
        self.assertAlmostEqual(written["weight"].sum(), 600, places=9)
        self.assertAlmostEqual(result.summary()["weighted_n"], 600, places=6)
        self.assertEqual(set(sheets), {"Summary", "Targets vs achieved", "Weight spread", "Warnings"})


class DirectRakeTests(unittest.TestCase):
    def test_single_dimension_matches_targets(self):
        from weighting_core import Dimension

        index = np.array([0] * 30 + [1] * 70)
        dim = Dimension("Sex", ["Men", "Women"], np.array([0.5, 0.5]), index)
        result = rake([dim])
        self.assertTrue(result.converged)
        self.assertAlmostEqual(result.weights[0], 50 / 30, places=6)


if __name__ == "__main__":
    unittest.main()
