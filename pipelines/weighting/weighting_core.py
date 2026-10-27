"""Rake survey data toward target contributions read from a weighting workbook.

The workbook layout is the house "Weighting combos" sheet: each target block is
a rectangle of yellow cells holding the desired share of the total sample, with
column labels above it and row labels to its left (e.g. Sex within Region).
Each block is one raking dimension; the row and column axes are each tied to a
data column through a label -> codes mapping.
"""
from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

YELLOW_FILLS = {"FFFFFF00", "00FFFF00"}
DEFAULT_CAP = 4.0
DEFAULT_FLOOR = 0.001  # no respondent is zeroed out or negative
DEFAULT_TOLERANCE = 0.0001  # 0.01 percentage points
DEFAULT_MAX_ITERATIONS = 500
BLANK_TOKEN = "blank"
REST_TOKEN = "*"


# --------------------------------------------------------------------------- targets


@dataclass
class TargetBlock:
    title: str
    cell_range: str
    row_labels: list[str]
    col_labels: list[str]
    values: list[list[float]]
    enabled: bool = True

    @property
    def total(self) -> float:
        return float(sum(sum(row) for row in self.values))

    @property
    def key(self) -> str:
        return f"{self.title} [{self.cell_range}]"


def _is_yellow(cell) -> bool:
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return False
    rgb = getattr(fill.fgColor, "rgb", None)
    return isinstance(rgb, str) and rgb.upper() in YELLOW_FILLS


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _col_letter(col: int) -> str:
    letters = ""
    while col:
        col, rem = divmod(col - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def read_target_blocks(path: str | Path, sheet: str | None = None, max_row: int | None = None) -> list[TargetBlock]:
    """Find every yellow target rectangle on the sheet (rows up to max_row)."""
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else _default_sheet(wb)
    last_row = min(ws.max_row, max_row) if max_row else ws.max_row

    yellow: set[tuple[int, int]] = set()
    for row in ws.iter_rows(min_row=1, max_row=last_row):
        for cell in row:
            if cell.value is not None and _is_yellow(cell):
                yellow.add((cell.row, cell.column))

    blocks: list[TargetBlock] = []
    seen: set[tuple[int, int]] = set()
    for start in sorted(yellow):
        if start in seen:
            continue
        component = _connected(start, yellow)
        seen |= component
        r0 = min(r for r, _ in component)
        r1 = max(r for r, _ in component)
        c0 = min(c for _, c in component)
        c1 = max(c for _, c in component)
        values = [[_number(ws.cell(r, c).value) for c in range(c0, c1 + 1)] for r in range(r0, r1 + 1)]
        row_labels = [_label_left(ws, r, c0) or f"Row {r}" for r in range(r0, r1 + 1)]
        col_labels = [_label_above(ws, r0, c) or _col_letter(c) for c in range(c0, c1 + 1)]
        blocks.append(
            TargetBlock(
                title=_block_title(ws, r0) or f"Block {_col_letter(c0)}{r0}",
                cell_range=f"{_col_letter(c0)}{r0}:{_col_letter(c1)}{r1}",
                row_labels=row_labels,
                col_labels=col_labels,
                values=values,
            )
        )

    # A second block over the same labels is usually a scratch alternative; keep the first.
    signatures: set[frozenset] = set()
    for block in blocks:
        signature = frozenset({tuple(sorted(block.row_labels)), tuple(sorted(block.col_labels))})
        block.enabled = signature not in signatures and block.total > 0
        signatures.add(signature)
    return blocks


def _default_sheet(wb):
    for name in wb.sheetnames:
        if name.strip().lower() == "weighting":
            return wb[name]
    return wb.worksheets[0]


def _connected(start, cells):
    stack, found = [start], {start}
    while stack:
        r, c = stack.pop()
        for nxt in ((r + 1, c), (r - 1, c), (r, c + 1), (r, c - 1)):
            if nxt in cells and nxt not in found:
                found.add(nxt)
                stack.append(nxt)
    return found


def _number(value) -> float:
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _label_left(ws, row: int, col: int) -> str:
    for c in range(col - 1, max(0, col - 4), -1):
        value = ws.cell(row, c).value
        if isinstance(value, str) and value.strip():
            return value.strip()
    return ""


def _label_above(ws, row: int, col: int) -> str:
    for r in range(row - 1, max(0, row - 5), -1):
        value = ws.cell(r, col).value
        if isinstance(value, str) and value.strip():
            return value.strip()
    return ""


def _block_title(ws, row: int) -> str:
    for r in range(row, max(0, row - 8), -1):
        value = _text(ws.cell(r, 1).value)
        if "within" in value.lower():
            return value
    return ""


# --------------------------------------------------------------------------- mapping


@dataclass
class Variable:
    """A data column plus the codes behind each target label."""

    name: str
    column: str
    labels: dict[str, str] = field(default_factory=dict)  # label -> code spec ("1-4", "W", "*", "blank")


@dataclass
class BlockAxes:
    block_key: str
    row_variable: str
    col_variable: str
    enabled: bool = True


@dataclass
class WeightingConfig:
    variables: dict[str, Variable] = field(default_factory=dict)
    blocks: list[BlockAxes] = field(default_factory=list)
    cap: float = DEFAULT_CAP
    floor: float = DEFAULT_FLOOR
    weight_column: str = "weight"
    base_weight_column: str = ""


def config_to_dict(config: WeightingConfig) -> dict:
    return {
        "cap": config.cap,
        "floor": config.floor,
        "weight_column": config.weight_column,
        "base_weight_column": config.base_weight_column,
        "variables": {
            name: {"column": var.column, "labels": dict(var.labels)} for name, var in config.variables.items()
        },
        "blocks": [
            {"block": b.block_key, "rows": b.row_variable, "cols": b.col_variable, "enabled": b.enabled}
            for b in config.blocks
        ],
    }


def config_from_dict(data: dict) -> WeightingConfig:
    variables = {
        name: Variable(name=name, column=str(spec.get("column", "")), labels=dict(spec.get("labels", {})))
        for name, spec in (data.get("variables") or {}).items()
    }
    blocks = [
        BlockAxes(
            block_key=str(item["block"]),
            row_variable=str(item.get("rows", "")),
            col_variable=str(item.get("cols", "")),
            enabled=bool(item.get("enabled", True)),
        )
        for item in data.get("blocks") or []
    ]
    return WeightingConfig(
        variables=variables,
        blocks=blocks,
        cap=float(data.get("cap", DEFAULT_CAP)),
        floor=float(data.get("floor", DEFAULT_FLOOR)),
        weight_column=str(data.get("weight_column", "weight")),
        base_weight_column=str(data.get("base_weight_column", "")),
    )


def load_config(path: str | Path) -> WeightingConfig:
    return config_from_dict(json.loads(Path(path).read_text(encoding="utf-8")))


def save_config(config: WeightingConfig, path: str | Path) -> None:
    Path(path).write_text(json.dumps(config_to_dict(config), indent=2), encoding="utf-8")


def draft_config(blocks: list[TargetBlock], df: pd.DataFrame | None = None) -> WeightingConfig:
    """Group block axes that share labels into variables and name them from block titles.

    "Sex within Region" and "Sex within Age" share Men/Women, so that axis
    becomes "Sex"; the other axis of each block takes the remaining title word.
    """
    axes: list[tuple[int, str, list[str]]] = []
    for i, block in enumerate(blocks):
        axes.append((i, "rows", block.row_labels))
        axes.append((i, "cols", block.col_labels))

    groups: list[dict] = []
    for i, side, labels in axes:
        label_set = set(labels)
        for group in groups:
            overlap = len(label_set & group["labels"]) / max(1, min(len(label_set), len(group["labels"])))
            if overlap >= 0.5:
                group["labels"] |= label_set
                group["order"] += [l for l in labels if l not in group["order"]]
                group["members"].append((i, side))
                break
        else:
            groups.append({"labels": set(label_set), "order": list(labels), "members": [(i, side)]})

    def title_parts(block: TargetBlock) -> list[str]:
        parts = re.split(r"\s+within\s+", block.title, flags=re.IGNORECASE)
        return [p.strip() for p in parts if p.strip()]

    names: dict[int, str] = {}
    for g, group in enumerate(groups):
        if len(group["members"]) > 1:
            common = set.intersection(*(set(title_parts(blocks[i])) for i, _ in group["members"]))
            if len(common) == 1:
                names[g] = common.pop()
    for g, group in enumerate(groups):
        if g in names:
            continue
        for i, side in group["members"]:
            other = next(o for o, grp in enumerate(groups) if (i, "cols" if side == "rows" else "rows") in grp["members"])
            leftover = [p for p in title_parts(blocks[i]) if p != names.get(other)]
            if len(leftover) == 1:
                names[g] = leftover[0]
                break
        names.setdefault(g, " / ".join(group["order"][:2]))

    used: set[str] = set()
    variables: dict[str, Variable] = {}
    member_name: dict[tuple[int, str], str] = {}
    for g, group in enumerate(groups):
        name, k = names[g], 2
        while name in used:
            name, k = f"{names[g]} {k}", k + 1
        used.add(name)
        column = guess_column(name, group["order"], df) if df is not None else ""
        labels = suggest_labels(group["order"], df[column]) if column else {l: "" for l in group["order"]}
        variables[name] = Variable(name=name, column=column, labels=labels)
        for member in group["members"]:
            member_name[member] = name

    block_axes = [
        BlockAxes(block.key, member_name[(i, "rows")], member_name[(i, "cols")], block.enabled)
        for i, block in enumerate(blocks)
    ]
    return WeightingConfig(variables=variables, blocks=block_axes)


def guess_column(name: str, labels: list[str], df: pd.DataFrame) -> str:
    """Pick the column whose values best match the labels, preferring a name match."""
    best, best_score = "", 0.0
    low = name.lower()
    for column in df.columns:
        values = df[column]
        if values.nunique() > 40:
            continue
        matched = sum(1 for spec in suggest_labels(labels, values, positional=False).values() if spec)
        if matched / len(labels) <= 0.6:
            continue
        score = matched / len(labels) + (0.5 if low and low in column.lower() else 0)
        if score > best_score:
            best, best_score = column, score
    return best


def parse_code_spec(spec: str) -> tuple[set[str], bool]:
    """Return (codes, is_rest). "1-4,7" -> {1,2,3,4,7}; "*" -> everything unclaimed."""
    codes: set[str] = set()
    is_rest = False
    for token in str(spec).split(","):
        token = token.strip()
        if not token:
            continue
        if token == REST_TOKEN:
            is_rest = True
        elif token.lower() == BLANK_TOKEN:
            codes.add("")
        elif re.fullmatch(r"-?\d+\s*-\s*-?\d+", token):
            lo, hi = (int(x) for x in re.split(r"(?<=\d)\s*-\s*", token, maxsplit=1))
            codes.update(str(v) for v in range(min(lo, hi), max(lo, hi) + 1))
        else:
            codes.add(token)
    return codes, is_rest


def normalize_code(value) -> str:
    text = _text(value)
    if re.fullmatch(r"-?\d+\.0+", text):
        text = text.split(".")[0]
    return text


def assign_labels(series: pd.Series, variable: Variable) -> pd.Series:
    """Map each respondent's code to a target label; unmatched codes become NaN."""
    codes = series.map(normalize_code)
    result = pd.Series(np.nan, index=series.index, dtype=object)
    rest_labels = []
    claimed: set[str] = set()
    for label, spec in variable.labels.items():
        label_codes, is_rest = parse_code_spec(spec)
        if is_rest:
            rest_labels.append(label)
        claimed |= label_codes
        result[codes.isin(label_codes) & result.isna()] = label
    if rest_labels:
        result[~codes.isin(claimed) & result.isna()] = rest_labels[0]
    return result


def suggest_labels(labels: list[str], series: pd.Series, positional: bool = True) -> dict[str, str]:
    """Best-guess code spec for each label from the values present in the column."""
    present = sorted({normalize_code(v) for v in series}, key=_code_sort_key)
    suggestion: dict[str, str] = {}
    used: set[str] = set()
    for label in labels:
        low = label.lower()
        matches = [
            code
            for code in present
            if code and code not in used and (low.startswith(code.lower()) or code.lower().startswith(low))
        ]
        number = re.search(r"(\d+)\s*$", label)
        if len(matches) == 1:
            suggestion[label] = matches[0]
        elif number and number.group(1) in present and not re.search(r"\d+\s*-\s*\d+", label):
            suggestion[label] = number.group(1)
        else:
            continue
        used.add(suggestion[label])
    if positional and not suggestion and len(labels) == len(present):
        suggestion = {label: code or BLANK_TOKEN for label, code in zip(labels, present)}
    return {label: suggestion.get(label, "") for label in labels}


def _code_sort_key(code: str):
    return (0, int(code), "") if re.fullmatch(r"-?\d+", code) else (1, 0, code)


# --------------------------------------------------------------------------- raking


@dataclass
class Dimension:
    name: str
    cells: list[str]
    targets: np.ndarray  # shares summing to 1
    index: np.ndarray  # cell index per respondent


@dataclass
class CellResult:
    dimension: str
    cell: str
    n: int
    unweighted_share: float
    target_share: float
    weighted_share: float

    @property
    def diff(self) -> float:
        return self.weighted_share - self.target_share


@dataclass
class RakeResult:
    weights: np.ndarray
    iterations: int
    converged: bool
    max_error: float
    cells: list[CellResult]
    warnings: list[str]
    cap: float
    floor: float

    @property
    def n(self) -> int:
        return int(len(self.weights))

    @property
    def design_effect(self) -> float:
        w = self.weights
        return float(len(w) * np.sum(w**2) / np.sum(w) ** 2) if np.sum(w) else float("nan")

    @property
    def effective_n(self) -> float:
        return self.n / self.design_effect

    @property
    def efficiency(self) -> float:
        return 1.0 / self.design_effect

    def summary(self) -> dict[str, float | int | bool]:
        w = self.weights
        return {
            "respondents": self.n,
            "weighted_n": round(float(w.sum()), 4),
            "converged": self.converged,
            "iterations": self.iterations,
            "max_error_pts": round(self.max_error * 100, 3),
            "min_weight": round(float(w.min()), 4),
            "max_weight": round(float(w.max()), 4),
            "at_cap": int(np.sum(w >= self.cap - 1e-9)),
            "at_floor": int(np.sum(w <= self.floor + 1e-9)),
            "design_effect": round(self.design_effect, 3),
            "efficiency_pct": round(self.efficiency * 100, 1),
            "effective_n": round(self.effective_n, 1),
        }


def build_dimensions(df: pd.DataFrame, blocks: list[TargetBlock], config: WeightingConfig) -> tuple[list[Dimension], list[str]]:
    """Turn enabled target blocks into raking dimensions; raise on unmapped respondents."""
    by_key = {block.key: block for block in blocks}
    dims: list[Dimension] = []
    warnings: list[str] = []
    problems: list[str] = []
    usage: dict[str, int] = {}
    for axes in config.blocks:
        if axes.enabled:
            for name in {axes.row_variable, axes.col_variable}:
                usage[name] = usage.get(name, 0) + 1
    for axes in config.blocks:
        if not axes.enabled:
            continue
        block = by_key.get(axes.block_key)
        if block is None:
            problems.append(f"Target block not found in workbook: {axes.block_key}")
            continue
        row_var = config.variables.get(axes.row_variable)
        col_var = config.variables.get(axes.col_variable)
        if row_var is None or col_var is None:
            problems.append(f"{block.title}: choose a variable for both rows and columns.")
            continue
        missing = [v.column for v in (row_var, col_var) if v.column not in df.columns]
        if missing:
            problems.append(f"{block.title}: column(s) not in data: {', '.join(missing)}")
            continue

        row_labels = assign_labels(df[row_var.column], row_var)
        col_labels = assign_labels(df[col_var.column], col_var)
        for var, assigned in ((row_var, row_labels), (col_var, col_labels)):
            unmapped = df.loc[assigned.isna(), var.column].map(normalize_code).value_counts()
            if len(unmapped):
                shown = ", ".join(f"{code or '(blank)'} x{count}" for code, count in unmapped.items())
                problems.append(f"{block.title}: {var.name} ({var.column}) codes with no target label: {shown}")

        cells = [f"{rl} | {cl}" for rl in block.row_labels for cl in block.col_labels]
        cell_of = {name: i for i, name in enumerate(cells)}
        index = np.array(
            [
                cell_of.get(f"{rl} | {cl}", -1) if isinstance(rl, str) and isinstance(cl, str) else -1
                for rl, cl in zip(row_labels, col_labels)
            ]
        )
        grid = np.array(block.values, dtype=float)
        counts = np.bincount(index[index >= 0], minlength=len(cells)).reshape(grid.shape)
        keep_columns = usage.get(axes.col_variable, 0) >= usage.get(axes.row_variable, 0)
        grid, moved = _reallocate_empty_cells(grid, counts, block, keep_columns)
        warnings.extend(f"{block.title}: {note}" for note in moved)
        target_arr = grid.ravel()
        total = target_arr.sum()
        if total <= 0:
            problems.append(f"{block.title}: targets sum to zero.")
            continue
        if abs(total - 1.0) > 0.005:
            warnings.append(f"{block.title}: targets sum to {total:.2%}; rescaled to 100%.")
        dims.append(Dimension(name=block.title, cells=cells, targets=target_arr / total, index=index))
    if problems:
        raise ValueError("\n".join(problems))
    if not dims:
        raise ValueError("No target blocks are enabled.")
    return dims, warnings


def _reallocate_empty_cells(
    grid: np.ndarray, counts: np.ndarray, block: TargetBlock, keep_columns: bool
) -> tuple[np.ndarray, list[str]]:
    """Move the target of a cell nobody falls into onto populated cells along one axis.

    Raking cannot put weight in an empty cell, and leaving its target in place
    makes every other block chase an impossible total. The share goes to the
    populated cells of the same column (or row), in proportion to their own
    targets, so that axis keeps its total - e.g. Men stay at 47%.
    """
    grid = grid.copy()
    notes: list[str] = []
    for r, c in zip(*np.nonzero((counts == 0) & (grid > 0))):
        share = grid[r, c]
        for along_column in (keep_columns, not keep_columns):
            mask = np.zeros_like(grid, dtype=bool)
            if along_column:
                mask[:, c] = counts[:, c] > 0
            else:
                mask[r, :] = counts[r, :] > 0
            if not mask.any():
                continue
            basis = np.where(mask, grid, 0.0)
            if basis.sum() <= 0:
                basis = np.where(mask, counts.astype(float), 0.0)
            grid += share * basis / basis.sum()
            grid[r, c] = 0.0
            kept = block.col_labels[c] if along_column else block.row_labels[r]
            notes.append(
                f"no respondents in '{block.row_labels[r]} | {block.col_labels[c]}' "
                f"(target {share:.2%}); share moved to the rest of '{kept}'."
            )
            break
    return grid, notes


def rake(
    dims: list[Dimension],
    base_weights: np.ndarray | None = None,
    cap: float = DEFAULT_CAP,
    floor: float = DEFAULT_FLOOR,
    tolerance: float = DEFAULT_TOLERANCE,
    max_iterations: int = DEFAULT_MAX_ITERATIONS,
) -> RakeResult:
    """Iterative proportional fitting with weights held to [floor, cap] (mean weight = 1).

    When the cap stops a target from being reached the loop keeps the closest
    solution seen (smallest worst-cell error) rather than the last one.
    """
    n = len(dims[0].index)
    w = np.ones(n) if base_weights is None else np.asarray(base_weights, dtype=float).copy()
    w = w * n / w.sum()
    warnings: list[str] = []

    for dim in dims:
        counts = np.bincount(dim.index, minlength=len(dim.cells))
        for i, cell in enumerate(dim.cells):
            if counts[i] == 0 and dim.targets[i] > 0:
                warnings.append(f"{dim.name}: no respondents in '{cell}' (target {dim.targets[i]:.2%}) - cannot be hit.")
            elif counts[i] > 0 and dim.targets[i] == 0:
                warnings.append(f"{dim.name}: {counts[i]} respondent(s) in '{cell}' with a 0% target are weighted down toward the minimum ({floor:g}).")

    best_w, best_error, iterations, converged = w.copy(), float("inf"), 0, False
    for iterations in range(1, max_iterations + 1):
        for dim in dims:
            current = np.bincount(dim.index, weights=w, minlength=len(dim.cells))
            wanted = dim.targets * w.sum()
            factor = np.divide(wanted, current, out=np.ones_like(current), where=current > 0)
            w = w * factor[dim.index]
        w = w * n / w.sum()
        # Clip then renormalise a few times so the cap holds at mean 1.
        for _ in range(20):
            clipped = np.clip(w, floor, cap)
            if np.allclose(clipped, w):
                break
            w = clipped * n / clipped.sum()
        w = np.clip(w, floor, cap)

        error = _max_error(dims, w)
        if error < best_error - 1e-12:
            best_w, best_error = w.copy(), error
        if error < tolerance:
            converged = True
            break

    polished = False
    if not converged:
        candidate = polish_least_squares(dims, best_w, cap=cap, floor=floor)
        if _squared_error(dims, candidate) < _squared_error(dims, best_w) - 1e-12:
            best_w, best_error, polished = candidate, _max_error(dims, candidate), True
            converged = best_error < tolerance
    if polished:
        warnings.append("Targets conflict, so raking alone could not meet all of them; a least-squares pass spread the misses across cells.")
    if not converged:
        warnings.append(
            f"Did not reach every target within {tolerance * 100:.2f} pts after {iterations} passes; "
            f"closest fit is off by at most {best_error * 100:.2f} pts (weight cap {cap:g})."
        )
    return RakeResult(
        weights=best_w,
        iterations=iterations,
        converged=converged,
        max_error=best_error,
        cells=_cell_results(dims, best_w),
        warnings=warnings,
        cap=cap,
        floor=floor,
    )


def polish_least_squares(
    dims: list[Dimension],
    start: np.ndarray,
    cap: float = DEFAULT_CAP,
    floor: float = DEFAULT_FLOOR,
    ridge: float = 1e-5,
    iterations: int = 5000,
) -> np.ndarray:
    """Minimise total squared share error with floor <= w <= cap and mean weight 1.

    Raking resolves conflicting targets by letting the last block win; this
    spreads the unavoidable misses across cells instead. The small ridge keeps
    weights near the raked solution so it does not trade efficiency for
    hundredths of a point. Accelerated projected gradient on a quadratic.
    """
    n = len(start)
    w0 = start.copy()
    max_count = sum(int(np.bincount(d.index).max()) for d in dims)
    step = 1.0 / (2.0 * max_count / n**2 + 2.0 * ridge / n)

    def gradient(w):
        g = 2.0 * ridge / n * (w - w0)
        for d in dims:
            residual = np.bincount(d.index, weights=w, minlength=len(d.cells)) / n - d.targets
            g += 2.0 / n * residual[d.index]
        return g

    w, y, t = w0.copy(), w0.copy(), 1.0
    for _ in range(iterations):
        w_next = _project(y - step * gradient(y), n, floor, cap)
        t_next = (1 + np.sqrt(1 + 4 * t * t)) / 2
        y = w_next + (t - 1) / t_next * (w_next - w)
        if np.max(np.abs(w_next - w)) < 1e-10:
            w = w_next
            break
        w, t = w_next, t_next
    return w


def _project(v: np.ndarray, total: float, floor: float, cap: float) -> np.ndarray:
    """Euclidean projection onto {floor <= w <= cap, sum(w) = total} by bisection on a shift."""
    lo, hi = np.min(v) - cap, np.max(v) - floor
    for _ in range(100):
        mid = (lo + hi) / 2
        if np.clip(v - mid, floor, cap).sum() > total:
            lo = mid
        else:
            hi = mid
    return np.clip(v - (lo + hi) / 2, floor, cap)


def _squared_error(dims: list[Dimension], w: np.ndarray) -> float:
    total = w.sum()
    return float(
        sum(np.sum((np.bincount(d.index, weights=w, minlength=len(d.cells)) / total - d.targets) ** 2) for d in dims)
    )


def _max_error(dims: list[Dimension], w: np.ndarray) -> float:
    total = w.sum()
    worst = 0.0
    for dim in dims:
        shares = np.bincount(dim.index, weights=w, minlength=len(dim.cells)) / total
        worst = max(worst, float(np.max(np.abs(shares - dim.targets))))
    return worst


def _cell_results(dims: list[Dimension], w: np.ndarray) -> list[CellResult]:
    rows: list[CellResult] = []
    n, total = len(w), w.sum()
    for dim in dims:
        counts = np.bincount(dim.index, minlength=len(dim.cells))
        weighted = np.bincount(dim.index, weights=w, minlength=len(dim.cells)) / total
        for i, cell in enumerate(dim.cells):
            if counts[i] == 0 and dim.targets[i] == 0:
                continue
            rows.append(
                CellResult(
                    dimension=dim.name,
                    cell=cell,
                    n=int(counts[i]),
                    unweighted_share=counts[i] / n,
                    target_share=float(dim.targets[i]),
                    weighted_share=float(weighted[i]),
                )
            )
    return rows


# --------------------------------------------------------------------------- io


def read_data(path: str | Path) -> pd.DataFrame:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm", ".xls"):
        return pd.read_excel(path, dtype=str, keep_default_na=False)
    if suffix == ".sav":
        try:
            import pyreadstat
        except ImportError as exc:  # pragma: no cover - optional dependency
            raise RuntimeError("Reading .sav files needs pyreadstat (pip install pyreadstat).") from exc
        df, _meta = pyreadstat.read_sav(str(path))
        return df.astype(object).where(df.notna(), "").astype(str)
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return pd.read_csv(path, dtype=str, keep_default_na=False, encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Could not decode {path.name} as UTF-8 or Windows-1252.")


def run_weighting(
    df: pd.DataFrame,
    blocks: list[TargetBlock],
    config: WeightingConfig,
    tolerance: float = DEFAULT_TOLERANCE,
    max_iterations: int = DEFAULT_MAX_ITERATIONS,
) -> RakeResult:
    dims, warnings = build_dimensions(df, blocks, config)
    base = None
    if config.base_weight_column:
        base = pd.to_numeric(df[config.base_weight_column], errors="coerce").fillna(0).to_numpy()
    result = rake(dims, base, cap=config.cap, floor=config.floor, tolerance=tolerance, max_iterations=max_iterations)
    result.warnings[:0] = warnings
    return result


def rounded_weights(weights: np.ndarray, cap: float, decimals: int = 6) -> np.ndarray:
    """Round weights so they still sum to exactly the number of respondents.

    The rounding remainder goes onto the largest weight that has room under the cap.
    """
    n = len(weights)
    out = np.round(weights * n / weights.sum(), decimals)
    remainder = round(n - out.sum(), decimals)
    if remainder:
        room = np.where(out + remainder <= cap, out, -np.inf)
        out[int(np.argmax(room))] += remainder
    return np.round(out, decimals)


def write_outputs(df: pd.DataFrame, result: RakeResult, config: WeightingConfig, data_out: Path, report_out: Path) -> None:
    out = df.copy()
    out[config.weight_column] = rounded_weights(result.weights, result.cap)
    if data_out.suffix.lower() in (".xlsx", ".xlsm"):
        out.to_excel(data_out, index=False)
    else:
        out.to_csv(data_out, index=False, encoding="cp1252", errors="replace")

    cells = pd.DataFrame(
        [
            {
                "Block": c.dimension,
                "Cell": c.cell,
                "N": c.n,
                "Unweighted %": round(c.unweighted_share * 100, 2),
                "Target %": round(c.target_share * 100, 2),
                "Weighted %": round(c.weighted_share * 100, 2),
                "Diff (pts)": round(c.diff * 100, 2),
            }
            for c in result.cells
        ]
    )
    summary = pd.DataFrame(list(result.summary().items()), columns=["Measure", "Value"])
    warnings = pd.DataFrame({"Warning": result.warnings or ["None"]})
    weights = pd.Series(result.weights)
    bins = [0, 0.25, 0.5, 0.75, 1, 1.5, 2, 3, result.cap + 1e-9]
    spread = pd.cut(weights, bins=bins, include_lowest=True).value_counts().sort_index()
    distribution = pd.DataFrame({"Weight range": [str(b) for b in spread.index], "Respondents": spread.values})
    with pd.ExcelWriter(report_out) as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        cells.to_excel(writer, sheet_name="Targets vs achieved", index=False)
        distribution.to_excel(writer, sheet_name="Weight spread", index=False)
        warnings.to_excel(writer, sheet_name="Warnings", index=False)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Rake survey data to the yellow targets in a weighting workbook.")
    parser.add_argument("--data", required=True)
    parser.add_argument("--targets", required=True, help="Weighting combos workbook")
    parser.add_argument("--config", required=True, help="JSON mapping of target labels to data codes")
    parser.add_argument("--sheet", default=None)
    parser.add_argument("--max-row", type=int, default=None)
    parser.add_argument("--cap", type=float, default=None)
    parser.add_argument("--out-data", default=None)
    parser.add_argument("--out-report", default=None)
    args = parser.parse_args(argv)

    config = load_config(args.config)
    if args.cap is not None:
        config.cap = args.cap
    df = read_data(args.data)
    blocks = read_target_blocks(args.targets, args.sheet, args.max_row)
    result = run_weighting(df, blocks, config)
    data_path = Path(args.data)
    out_data = Path(args.out_data) if args.out_data else data_path.with_name(f"{data_path.stem}_weighted{data_path.suffix}")
    out_report = Path(args.out_report) if args.out_report else data_path.with_name(f"{data_path.stem}_weight_report.xlsx")
    write_outputs(df, result, config, out_data, out_report)
    for key, value in result.summary().items():
        print(f"{key:>16}: {value}")
    for warning in result.warnings:
        print(f"WARNING: {warning}")
    print(f"Weighted data: {out_data}\nReport: {out_report}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
