from __future__ import annotations

import os
from pathlib import Path
import sys
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook

from weighting_core import (
    BLANK_TOKEN,
    REST_TOKEN,
    RakeResult,
    TargetBlock,
    WeightingConfig,
    assign_labels,
    draft_config,
    load_config,
    normalize_code,
    read_data,
    read_target_blocks,
    run_weighting,
    save_config,
    suggest_labels,
    write_outputs,
)


def _launch_work_directory() -> Path | None:
    raw = os.environ.get("SURVEY_PLATFORM_WORK_DIR", "").strip()
    if not raw:
        return None
    path = Path(raw).expanduser().resolve(strict=False)
    return path if path.is_dir() else None


def _config_path_for(data_path: str) -> Path:
    path = Path(data_path)
    return path.with_name(f"{path.stem}_weighting.json")


class WeightingApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("Sample Weighting")
        self.root.geometry("1280x860")
        self.launch_work_dir = _launch_work_directory()

        self.data_var = tk.StringVar()
        self.targets_var = tk.StringVar()
        self.sheet_var = tk.StringVar()
        self.max_row_var = tk.StringVar()
        self.cap_var = tk.StringVar(value="4")
        self.floor_var = tk.StringVar(value="0.001")
        self.weight_col_var = tk.StringVar(value="weight")
        self.rows_var = tk.StringVar()
        self.cols_var = tk.StringVar()
        self.column_var = tk.StringVar()
        self.summary_var = tk.StringVar(value="Choose the data file and the weighting workbook.")

        self.df = None
        self.blocks: list[TargetBlock] = []
        self.config = WeightingConfig()
        self.result: RakeResult | None = None
        self.current_variable = ""
        self.code_vars: dict[str, tk.StringVar] = {}

        self._build_ui()

    # ------------------------------------------------------------------ layout

    def _build_ui(self) -> None:
        outer = ttk.Frame(self.root, padding=10)
        outer.pack(fill=tk.BOTH, expand=True)

        files = ttk.LabelFrame(outer, text="Files", padding=8)
        files.pack(fill=tk.X)
        files.columnconfigure(1, weight=1)
        self._file_row(files, 0, "Data CSV/XLSX", self.data_var, self._browse_data)
        self._file_row(files, 1, "Weighting workbook", self.targets_var, self._browse_targets)
        sheet_row = ttk.Frame(files)
        sheet_row.grid(row=2, column=0, columnspan=3, sticky="ew", pady=(5, 0))
        ttk.Label(sheet_row, text="Sheet").pack(side=tk.LEFT)
        self.sheet_combo = ttk.Combobox(sheet_row, textvariable=self.sheet_var, state="readonly", width=28)
        self.sheet_combo.pack(side=tk.LEFT, padx=(6, 12))
        ttk.Label(sheet_row, text="Read targets through row").pack(side=tk.LEFT)
        ttk.Entry(sheet_row, textvariable=self.max_row_var, width=6).pack(side=tk.LEFT, padx=(6, 12))
        ttk.Button(sheet_row, text="Load Targets", command=self.load_targets).pack(side=tk.LEFT)
        ttk.Button(sheet_row, text="Open Setup...", command=self.open_setup).pack(side=tk.LEFT, padx=(18, 0))
        ttk.Button(sheet_row, text="Save Setup...", command=self.save_setup).pack(side=tk.LEFT, padx=(8, 0))

        setup = ttk.PanedWindow(outer, orient=tk.HORIZONTAL)
        setup.pack(fill=tk.BOTH, expand=True, pady=(8, 0))

        blocks_frame = ttk.LabelFrame(setup, text="Target blocks (double-click to turn on/off)", padding=6)
        self.block_tree = ttk.Treeview(
            blocks_frame, columns=("use", "block", "rows", "cols", "total"), show="headings", height=8
        )
        for col, text, width in (
            ("use", "Use", 45),
            ("block", "Block", 260),
            ("rows", "Rows", 110),
            ("cols", "Columns", 110),
            ("total", "Total", 70),
        ):
            self.block_tree.heading(col, text=text)
            self.block_tree.column(col, width=width, anchor=tk.W)
        self.block_tree.pack(fill=tk.BOTH, expand=True)
        self.block_tree.bind("<Double-1>", self._toggle_block)
        self.block_tree.bind("<<TreeviewSelect>>", self._block_selected)
        axis_row = ttk.Frame(blocks_frame)
        axis_row.pack(fill=tk.X, pady=(6, 0))
        ttk.Label(axis_row, text="Rows are").pack(side=tk.LEFT)
        self.rows_combo = ttk.Combobox(axis_row, textvariable=self.rows_var, state="readonly", width=16)
        self.rows_combo.pack(side=tk.LEFT, padx=(4, 10))
        ttk.Label(axis_row, text="Columns are").pack(side=tk.LEFT)
        self.cols_combo = ttk.Combobox(axis_row, textvariable=self.cols_var, state="readonly", width=16)
        self.cols_combo.pack(side=tk.LEFT, padx=(4, 0))
        self.rows_combo.bind("<<ComboboxSelected>>", self._axis_changed)
        self.cols_combo.bind("<<ComboboxSelected>>", self._axis_changed)
        setup.add(blocks_frame, weight=1)

        vars_frame = ttk.LabelFrame(setup, text="Variables: which data codes fall under each target label", padding=6)
        top = ttk.Frame(vars_frame)
        top.pack(fill=tk.X)
        self.var_list = tk.Listbox(top, height=7, exportselection=False, width=18)
        self.var_list.pack(side=tk.LEFT, fill=tk.Y)
        self.var_list.bind("<<ListboxSelect>>", self._variable_selected)
        right = ttk.Frame(top)
        right.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(8, 0))
        col_row = ttk.Frame(right)
        col_row.pack(fill=tk.X)
        ttk.Label(col_row, text="Data column").pack(side=tk.LEFT)
        self.column_combo = ttk.Combobox(col_row, textvariable=self.column_var, state="readonly", width=24)
        self.column_combo.pack(side=tk.LEFT, padx=(6, 8))
        self.column_combo.bind("<<ComboboxSelected>>", self._column_changed)
        ttk.Button(col_row, text="Suggest Codes", command=self._suggest_codes).pack(side=tk.LEFT)
        self.values_label = ttk.Label(right, text="", foreground="#555555", wraplength=520, justify=tk.LEFT)
        self.values_label.pack(fill=tk.X, pady=(4, 0))
        ttk.Label(
            right,
            text=f"Codes: 1-4,7  |  W  |  Democrat,D  |  {BLANK_TOKEN} = empty  |  {REST_TOKEN} = every code not listed elsewhere",
            foreground="#555555",
        ).pack(fill=tk.X)
        self.codes_frame = ttk.Frame(vars_frame)
        self.codes_frame.pack(fill=tk.BOTH, expand=True, pady=(6, 0))
        setup.add(vars_frame, weight=1)

        actions = ttk.Frame(outer, padding=(0, 8, 0, 8))
        actions.pack(fill=tk.X)
        ttk.Label(actions, text="Max weight").pack(side=tk.LEFT)
        ttk.Entry(actions, textvariable=self.cap_var, width=6).pack(side=tk.LEFT, padx=(4, 10))
        ttk.Label(actions, text="Min weight").pack(side=tk.LEFT)
        ttk.Entry(actions, textvariable=self.floor_var, width=6).pack(side=tk.LEFT, padx=(4, 10))
        ttk.Label(actions, text="Weight column").pack(side=tk.LEFT)
        ttk.Entry(actions, textvariable=self.weight_col_var, width=14).pack(side=tk.LEFT, padx=(4, 14))
        ttk.Button(actions, text="Run Weighting", command=self.run).pack(side=tk.LEFT)
        ttk.Button(actions, text="Save Weighted Data + Report", command=self.save_outputs).pack(side=tk.LEFT, padx=(8, 0))
        ttk.Label(actions, textvariable=self.summary_var).pack(side=tk.LEFT, padx=(16, 0))

        results = ttk.PanedWindow(outer, orient=tk.VERTICAL)
        results.pack(fill=tk.BOTH, expand=True)
        table_frame = ttk.Frame(results)
        self.result_tree = ttk.Treeview(
            table_frame,
            columns=("block", "cell", "n", "unweighted", "target", "weighted", "diff"),
            show="headings",
            height=14,
        )
        for col, text, width in (
            ("block", "Block", 200),
            ("cell", "Cell", 220),
            ("n", "N", 50),
            ("unweighted", "Unweighted %", 100),
            ("target", "Target %", 80),
            ("weighted", "Weighted %", 90),
            ("diff", "Diff (pts)", 80),
        ):
            self.result_tree.heading(col, text=text)
            self.result_tree.column(col, width=width, anchor=tk.W)
        self.result_tree.tag_configure("ok", background="#eaf7ea")
        self.result_tree.tag_configure("near", background="#fff8dc")
        self.result_tree.tag_configure("off", background="#fdeaea")
        scroll = ttk.Scrollbar(table_frame, orient=tk.VERTICAL, command=self.result_tree.yview)
        self.result_tree.configure(yscrollcommand=scroll.set)
        self.result_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scroll.pack(side=tk.RIGHT, fill=tk.Y)
        results.add(table_frame, weight=3)
        self.notes = tk.Text(results, height=8, wrap=tk.WORD)
        results.add(self.notes, weight=1)

    def _file_row(self, parent, row: int, label: str, var: tk.StringVar, command) -> None:
        ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", padx=(0, 8), pady=2)
        ttk.Entry(parent, textvariable=var).grid(row=row, column=1, sticky="ew", pady=2)
        ttk.Button(parent, text="Browse...", command=command).grid(row=row, column=2, padx=(8, 0), pady=2)

    def _initial_dir(self, current: str) -> str | None:
        if current:
            parent = Path(current).expanduser().parent
            if parent.is_dir():
                return str(parent)
        return str(self.launch_work_dir) if self.launch_work_dir else None

    # ------------------------------------------------------------------ files

    def _browse_data(self) -> None:
        path = filedialog.askopenfilename(
            title="Choose survey data",
            initialdir=self._initial_dir(self.data_var.get() or self.targets_var.get()),
            filetypes=[("Data files", "*.csv *.xlsx *.sav"), ("All files", "*.*")],
        )
        if not path:
            return
        self.data_var.set(path)
        try:
            self.df = read_data(path)
        except Exception as exc:  # noqa: BLE001 - surface any read failure to the user
            messagebox.showerror("Data", f"Could not read data:\n{exc}")
            return
        self.column_combo["values"] = list(self.df.columns)
        self.summary_var.set(f"Loaded {len(self.df):,} respondents.")
        saved = _config_path_for(path)
        if saved.exists() and self.blocks:
            self._apply_config(load_config(saved))
        elif self.blocks:
            self._apply_config(draft_config(self.blocks, self.df))

    def _browse_targets(self) -> None:
        path = filedialog.askopenfilename(
            title="Choose weighting workbook",
            initialdir=self._initial_dir(self.targets_var.get() or self.data_var.get()),
            filetypes=[("Excel", "*.xlsx *.xlsm"), ("All files", "*.*")],
        )
        if not path:
            return
        self.targets_var.set(path)
        try:
            sheets = load_workbook(path, read_only=True).sheetnames
        except PermissionError:
            messagebox.showerror("Workbook", "The workbook is locked. Close it in Excel, or save a copy, and try again.")
            return
        self.sheet_combo["values"] = sheets
        default = next((s for s in sheets if s.strip().lower() == "weighting"), sheets[0])
        self.sheet_var.set(default)
        self.load_targets()

    def load_targets(self) -> None:
        path = self.targets_var.get().strip()
        if not path:
            return
        try:
            max_row = int(self.max_row_var.get()) if self.max_row_var.get().strip() else None
            self.blocks = read_target_blocks(path, self.sheet_var.get() or None, max_row)
        except PermissionError:
            messagebox.showerror("Workbook", "The workbook is locked. Close it in Excel, or save a copy, and try again.")
            return
        except Exception as exc:  # noqa: BLE001
            messagebox.showerror("Workbook", f"Could not read targets:\n{exc}")
            return
        if not self.blocks:
            messagebox.showwarning("Workbook", "No yellow target cells found on that sheet.")
            return
        saved = _config_path_for(self.data_var.get()) if self.data_var.get() else None
        if saved and saved.exists():
            self._apply_config(load_config(saved))
        else:
            self._apply_config(draft_config(self.blocks, self.df))
        self.summary_var.set(f"Found {len(self.blocks)} target blocks.")

    def open_setup(self) -> None:
        path = filedialog.askopenfilename(
            title="Open weighting setup",
            initialdir=self._initial_dir(self.data_var.get()),
            filetypes=[("Weighting setup", "*.json")],
        )
        if path:
            self._apply_config(load_config(path))

    def save_setup(self) -> None:
        self._store_codes()
        self._read_settings()
        default = _config_path_for(self.data_var.get()) if self.data_var.get() else Path("weighting.json")
        path = filedialog.asksaveasfilename(
            title="Save weighting setup",
            initialdir=str(default.parent),
            initialfile=default.name,
            defaultextension=".json",
            filetypes=[("Weighting setup", "*.json")],
        )
        if path:
            save_config(self.config, path)
            self.summary_var.set(f"Saved setup: {Path(path).name}")

    # ------------------------------------------------------------------ setup editing

    def _apply_config(self, config: WeightingConfig) -> None:
        known = {b.block_key for b in config.blocks}
        for block in self.blocks:
            if block.key not in known:
                drafted = draft_config([block], self.df)
                config.blocks.extend(drafted.blocks)
                for name, var in drafted.variables.items():
                    config.variables.setdefault(name, var)
        self.config = config
        self.cap_var.set(f"{config.cap:g}")
        self.floor_var.set(f"{config.floor:g}")
        self.weight_col_var.set(config.weight_column)
        self.current_variable = ""
        names = list(config.variables)
        self.rows_combo["values"] = names
        self.cols_combo["values"] = names
        self.var_list.delete(0, tk.END)
        for name in names:
            self.var_list.insert(tk.END, name)
        self._refresh_blocks()
        if names:
            self.var_list.selection_set(0)
            self._variable_selected()

    def _refresh_blocks(self) -> None:
        totals = {b.key: b.total for b in self.blocks}
        self.block_tree.delete(*self.block_tree.get_children())
        for i, axes in enumerate(self.config.blocks):
            if axes.block_key not in totals:
                continue
            self.block_tree.insert(
                "",
                tk.END,
                iid=str(i),
                values=(
                    "Yes" if axes.enabled else "-",
                    axes.block_key,
                    axes.row_variable,
                    axes.col_variable,
                    f"{totals[axes.block_key]:.2%}",
                ),
            )

    def _selected_block(self):
        selection = self.block_tree.selection()
        return self.config.blocks[int(selection[0])] if selection else None

    def _toggle_block(self, _event=None) -> None:
        axes = self._selected_block()
        if axes:
            axes.enabled = not axes.enabled
            self._refresh_blocks()

    def _block_selected(self, _event=None) -> None:
        axes = self._selected_block()
        if axes:
            self.rows_var.set(axes.row_variable)
            self.cols_var.set(axes.col_variable)

    def _axis_changed(self, _event=None) -> None:
        axes = self._selected_block()
        if axes:
            selected = self.block_tree.selection()
            axes.row_variable = self.rows_var.get()
            axes.col_variable = self.cols_var.get()
            self._refresh_blocks()
            self.block_tree.selection_set(selected)

    def _variable_selected(self, _event=None) -> None:
        self._store_codes()
        selection = self.var_list.curselection()
        if not selection:
            return
        self.current_variable = self.var_list.get(selection[0])
        var = self.config.variables[self.current_variable]
        self.column_var.set(var.column)
        self._render_codes()

    def _column_changed(self, _event=None) -> None:
        var = self.config.variables.get(self.current_variable)
        if var is None:
            return
        self._store_codes()
        var.column = self.column_var.get()
        if not any(spec.strip() for spec in var.labels.values()):
            self._suggest_codes()
        else:
            self._render_codes()

    def _suggest_codes(self) -> None:
        var = self.config.variables.get(self.current_variable)
        if var is None or self.df is None or var.column not in self.df.columns:
            return
        suggested = suggest_labels(list(var.labels), self.df[var.column])
        var.labels = {label: suggested.get(label) or var.labels.get(label, "") for label in var.labels}
        self._render_codes()

    def _render_codes(self) -> None:
        for child in self.codes_frame.winfo_children():
            child.destroy()
        self.code_vars = {}
        var = self.config.variables.get(self.current_variable)
        if var is None:
            return
        counts = {}
        if self.df is not None and var.column in self.df.columns:
            values = self.df[var.column].map(normalize_code).value_counts()
            shown = ", ".join(f"{code or '(blank)'}: {n}" for code, n in values.head(25).items())
            self.values_label.configure(text=f"Values in {var.column}: {shown}")
            counts = assign_labels(self.df[var.column], var).value_counts().to_dict()
            unmatched = int(assign_labels(self.df[var.column], var).isna().sum())
        else:
            self.values_label.configure(text="Pick the data column for this variable.")
            unmatched = 0
        ttk.Label(self.codes_frame, text="Target label", font=("Segoe UI", 9, "bold")).grid(row=0, column=0, sticky="w")
        ttk.Label(self.codes_frame, text="Codes", font=("Segoe UI", 9, "bold")).grid(row=0, column=1, sticky="w")
        ttk.Label(self.codes_frame, text="N", font=("Segoe UI", 9, "bold")).grid(row=0, column=2, sticky="w")
        for r, (label, spec) in enumerate(var.labels.items(), start=1):
            ttk.Label(self.codes_frame, text=label).grid(row=r, column=0, sticky="w", padx=(0, 10), pady=1)
            code_var = tk.StringVar(value=spec)
            entry = ttk.Entry(self.codes_frame, textvariable=code_var, width=22)
            entry.grid(row=r, column=1, sticky="w", pady=1)
            entry.bind("<FocusOut>", lambda _e: self._codes_edited())
            entry.bind("<Return>", lambda _e: self._codes_edited())
            self.code_vars[label] = code_var
            ttk.Label(self.codes_frame, text=str(counts.get(label, 0))).grid(row=r, column=2, sticky="w", padx=(8, 0))
        if unmatched:
            ttk.Label(
                self.codes_frame,
                text=f"{unmatched} respondent(s) have a code with no label - add it to a label or use {REST_TOKEN}.",
                foreground="#b00020",
            ).grid(row=len(var.labels) + 1, column=0, columnspan=3, sticky="w", pady=(6, 0))

    def _codes_edited(self) -> None:
        self._store_codes()
        self._render_codes()

    def _store_codes(self) -> None:
        var = self.config.variables.get(self.current_variable)
        if var is None:
            return
        for label, code_var in self.code_vars.items():
            if label in var.labels:
                var.labels[label] = code_var.get().strip()

    def _read_settings(self) -> None:
        self.config.cap = float(self.cap_var.get())
        self.config.floor = float(self.floor_var.get() or 0.001)
        self.config.weight_column = self.weight_col_var.get().strip() or "weight"

    # ------------------------------------------------------------------ run

    def run(self) -> None:
        if self.df is None or not self.blocks:
            messagebox.showinfo("Weighting", "Load the data file and the weighting workbook first.")
            return
        self._store_codes()
        try:
            self._read_settings()
        except ValueError:
            messagebox.showerror("Weighting", "Max and min weight must be numbers.")
            return
        if self.config.cap <= 1 or self.config.floor <= 0 or self.config.floor >= 1:
            messagebox.showerror("Weighting", "Max weight must be above 1 and min weight above 0 and below 1.")
            return
        try:
            self.result = run_weighting(self.df, self.blocks, self.config)
        except ValueError as exc:
            messagebox.showerror("Weighting", f"Fix the setup first:\n\n{exc}")
            return
        save_config(self.config, _config_path_for(self.data_var.get()))
        self._show_result(self.result)

    def _show_result(self, result: RakeResult) -> None:
        self.result_tree.delete(*self.result_tree.get_children())
        for cell in result.cells:
            gap = abs(cell.diff) * 100
            tag = "ok" if gap < 0.1 else "near" if gap < 0.5 else "off"
            self.result_tree.insert(
                "",
                tk.END,
                values=(
                    cell.dimension,
                    cell.cell,
                    cell.n,
                    f"{cell.unweighted_share:.2%}",
                    f"{cell.target_share:.2%}",
                    f"{cell.weighted_share:.2%}",
                    f"{cell.diff * 100:+.2f}",
                ),
                tags=(tag,),
            )
        s = result.summary()
        status = "Converged" if s["converged"] else f"Closest fit, worst cell off {s['max_error_pts']:.2f} pts"
        self.summary_var.set(
            f"{status} | weights {s['min_weight']:.2f}-{s['max_weight']:.2f} | "
            f"weighted N {s['weighted_n']:g} of {s['respondents']} | "
            f"efficiency {s['efficiency_pct']:.1f}% (effective n {s['effective_n']:.0f})"
        )
        self.notes.delete("1.0", tk.END)
        lines = [f"{k.replace('_', ' ')}: {v}" for k, v in s.items()]
        lines += [""] + (result.warnings or ["No warnings."])
        self.notes.insert("1.0", "\n".join(lines))

    def save_outputs(self) -> None:
        if self.result is None or self.df is None:
            messagebox.showinfo("Weighting", "Run the weighting first.")
            return
        data_path = Path(self.data_var.get())
        out = filedialog.asksaveasfilename(
            title="Save weighted data",
            initialdir=str(data_path.parent),
            initialfile=f"{data_path.stem}_weighted{data_path.suffix if data_path.suffix != '.sav' else '.csv'}",
            filetypes=[("CSV", "*.csv"), ("Excel", "*.xlsx")],
        )
        if not out:
            return
        out_path = Path(out)
        report = out_path.with_name(f"{data_path.stem}_weight_report.xlsx")
        try:
            write_outputs(self.df, self.result, self.config, out_path, report)
        except PermissionError:
            messagebox.showerror("Weighting", "A file is open in Excel. Close it and save again.")
            return
        self.summary_var.set(f"Saved {out_path.name} and {report.name}")


def main() -> None:
    root = tk.Tk()
    WeightingApp(root)
    root.mainloop()


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--help":
        print("Sample Weighting: rake survey data to the yellow targets in a weighting workbook.")
    else:
        main()
