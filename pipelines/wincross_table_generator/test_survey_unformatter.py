import sys
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX


THIS_DIR = Path(__file__).resolve().parent
if str(THIS_DIR) not in sys.path:
    sys.path.insert(0, str(THIS_DIR))

from survey_unformatter import build_formatting_sidecar, extract_docx_paragraph_records  # noqa: E402


class SurveyUnformatterFormattingTests(unittest.TestCase):
    def _scratch_path(self, name: str) -> Path:
        base = THIS_DIR.parents[1] / ".tmp_tests" / "wincross_formatting"
        base.mkdir(parents=True, exist_ok=True)
        return base / name

    def test_extract_docx_paragraph_records_preserves_all_bold_spans(self) -> None:
        path = self._scratch_path("formatting_spans.docx")
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("1. Do you support ")
        bracket = paragraph.add_run("[PLAN]")
        bracket.bold = True
        paragraph.add_run(" and ")
        plain = paragraph.add_run("health care")
        plain.bold = True
        paragraph.add_run("?")
        doc.save(str(path))

        records = extract_docx_paragraph_records(str(path))

        self.assertEqual(len(records), 1)
        spans = records[0]["bold_spans"]
        self.assertEqual([span["text"] for span in spans], ["[PLAN]", "health care"])
        self.assertTrue(spans[0]["bracketed"])
        self.assertEqual(spans[0]["bold_coverage"], "all")
        self.assertFalse(spans[1]["bracketed"])

    def test_extract_docx_paragraph_records_preserves_yellow_highlight_spans(self) -> None:
        path = self._scratch_path("formatting_highlight_spans.docx")
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("1. Read ")
        bracket = paragraph.add_run("[CHOICE]")
        bracket.font.highlight_color = WD_COLOR_INDEX.YELLOW
        paragraph.add_run(" and ")
        plain = paragraph.add_run("this section")
        plain.font.highlight_color = WD_COLOR_INDEX.YELLOW
        paragraph.add_run(".")
        doc.save(str(path))

        records = extract_docx_paragraph_records(str(path))

        self.assertEqual(len(records), 1)
        spans = records[0]["highlight_spans"]
        self.assertEqual([span["text"] for span in spans], ["[CHOICE]", "this section"])
        self.assertTrue(spans[0]["bracketed"])
        self.assertEqual(spans[0]["highlight_color"], "yellow")
        self.assertEqual(spans[0]["highlight_coverage"], "all")
        self.assertFalse(spans[1]["bracketed"])

    def test_extract_docx_paragraph_records_preserves_strikethrough_spans(self) -> None:
        path = self._scratch_path("formatting_strikethrough_spans.docx")
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("1. Remove ")
        bracket = paragraph.add_run("[OLD]")
        bracket.font.strike = True
        paragraph.add_run(" and ")
        plain = paragraph.add_run("this phrase")
        plain.font.strike = True
        doc.save(str(path))

        records = extract_docx_paragraph_records(str(path))

        self.assertEqual(len(records), 1)
        spans = records[0]["strikethrough_spans"]
        self.assertEqual([span["text"] for span in spans], ["[OLD]", "this phrase"])
        self.assertTrue(spans[0]["bracketed"])
        self.assertEqual(spans[0]["strikethrough_style"], "single")
        self.assertEqual(spans[0]["strikethrough_coverage"], "all")
        self.assertFalse(spans[1]["bracketed"])

    def test_build_formatting_sidecar_counts_non_bracketed_bold(self) -> None:
        path = self._scratch_path("formatting_summary.docx")
        doc = Document()
        paragraph = doc.add_paragraph()
        bold = paragraph.add_run("Plain bold")
        bold.bold = True
        paragraph.add_run(" ")
        highlighted = paragraph.add_run("Plain highlighted")
        highlighted.font.highlight_color = WD_COLOR_INDEX.YELLOW
        paragraph.add_run(" ")
        struck = paragraph.add_run("Plain struck")
        struck.font.strike = True
        doc.save(str(path))

        payload = build_formatting_sidecar(str(path), raw_order_source="pdf")

        self.assertEqual(payload["raw_order_source"], "pdf")
        self.assertEqual(payload["summary"]["bold_span_count"], 1)
        self.assertEqual(payload["summary"]["non_bracketed_bold_span_count"], 1)
        self.assertEqual(payload["summary"]["highlight_span_count"], 1)
        self.assertEqual(payload["summary"]["non_bracketed_highlight_span_count"], 1)
        self.assertEqual(payload["summary"]["strikethrough_span_count"], 1)
        self.assertEqual(payload["summary"]["non_bracketed_strikethrough_span_count"], 1)


if __name__ == "__main__":
    unittest.main()
