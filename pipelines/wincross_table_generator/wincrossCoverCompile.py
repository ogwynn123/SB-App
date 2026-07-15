from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple
import json
import re
import sys

from docx import Document
from docx.shared import Pt

from wincrossTableCompile import (
    Question,
    ResponseOption,
    assign_standard_vote_headers,
    get_leans_pairs,
    is_dont_know_label,
    is_other_label,
    is_undecided_label,
    norm_spaces,
    parse_survey,
    standard_vote_header_label,
    strip_parens,
    vote_candidate_signature,
    vote_office_label,
    vote_primary_prefix,
)


@dataclass(frozen=True)
class BannerPoint:
    banner_name: str
    index: int
    expression: str
    width: int


@dataclass(frozen=True)
class CoverEntry:
    section: str
    label: str
    definition: str
    source: str


SECTION_ORDER = [
    "DEF VTRS",
    "VOTE HIST",
    "REGION",
    "RACE",
    "EDUCATION",
    "SEX/EDUCATION",
    "WHITE/EDU",
    "IDEOLOGY",
    "TYPE OF AG",
    "VOTE FOR ATTY GENERAL",
    "SPLIT",
    "TARGETS",
    "OTHER BANNER POINTS",
]


def parse_banner_code(text: str) -> List[BannerPoint]:
    points: List[BannerPoint] = []
    banner_name = ""
    index = 0
    for raw in text.splitlines():
        line = raw.rstrip()
        banner_match = re.match(r"^\*Banner\s*(\d+)", line, flags=re.IGNORECASE)
        if banner_match:
            banner_name = f"Banner{banner_match.group(1)}"
            index = 0
            continue
        point_match = re.match(r"^\s*(?P<expr>.+?)\^W(?P<width>\d+)\s*$", line)
        if not point_match or line.lstrip().upper().startswith("ST:"):
            continue
        index += 1
        points.append(
            BannerPoint(
                banner_name=banner_name,
                index=index,
                expression=norm_spaces(point_match.group("expr")),
                width=int(point_match.group("width")),
            )
        )
    return points


def qid_key(qid: str) -> str:
    text = str(qid).strip().upper()
    return text[1:] if text.startswith("Q") else text


def qref(q: Question) -> str:
    return f"Q{qid_key(q.qid)}"


def question_lookup(questions: Sequence[Question]) -> Dict[str, Question]:
    return {qid_key(q.qid): q for q in questions}


def qids_in_expression(expression: str) -> List[str]:
    seen: List[str] = []
    for match in re.finditer(r"\bq(?P<qid>\d+[a-z]?)\b", expression, flags=re.IGNORECASE):
        key = qid_key(match.group("qid"))
        if key not in seen:
            seen.append(key)
    return seen


def first_question(questions: Sequence[Question], qtype: str) -> Optional[Question]:
    return next((q for q in questions if str(q.question_type).casefold() == qtype.casefold()), None)


def selector_for_q(expression: str, q: Question) -> Optional[str]:
    pattern = re.compile(rf"\bq{re.escape(qid_key(q.qid))}\s*\(([^)]*)\)", flags=re.IGNORECASE)
    match = pattern.search(expression)
    return norm_spaces(match.group(1)) if match else None


def expression_has_q(expression: str, q: Optional[Question]) -> bool:
    if q is None:
        return False
    return qid_key(q.qid) in qids_in_expression(expression)


def selector_values(selector: str) -> List[str]:
    values: List[str] = []
    for part in re.split(r"\s*,\s*", selector.strip()):
        if not part:
            continue
        range_match = re.fullmatch(r"(-?\d+)\s*-\s*(-?\d+)", part)
        if range_match:
            start = int(range_match.group(1))
            end = int(range_match.group(2))
            step = 1 if end >= start else -1
            values.extend(str(n) for n in range(start, end + step, step))
        else:
            values.append(part.strip())
    return values


def selected_options(q: Question, selector: str) -> List[ResponseOption]:
    values = set(selector_values(selector))
    return [opt for opt in q.options if str(opt.value).strip() in values]


def option_phrase(q: Question, selector: str) -> str:
    texts = [norm_spaces(strip_parens(opt.text)) for opt in selected_options(q, selector)]
    texts = [text for text in texts if text]
    if not texts:
        return selector
    if len(texts) == 1:
        return f'"{texts[0]}"'
    if len(texts) == 2:
        return f'"{texts[0]}" or "{texts[1]}"'
    return ", ".join(f'"{text}"' for text in texts[:-1]) + f', or "{texts[-1]}"'


def compact_selector_label(selector: str) -> str:
    return selector.replace(",", "/").replace(" ", "")


def education_label(selector: str) -> str:
    values = set(selector_values(selector))
    if values and values.issubset({"1", "2", "3", "4", "5", "6"}):
        return "No Deg"
    if values and values.issubset({"7", "8", "9", "10"}):
        return "Coll Grad"
    return compact_selector_label(selector)


def education_definition(q: Question, selector: str) -> str:
    return f"Voters whose last grade completed in school was {option_phrase(q, selector)} ({qref(q)})."


def gender_label(q: Question, selector: str) -> str:
    text = option_phrase(q, selector).replace('"', "").casefold()
    if "female" in text or "woman" in text:
        return "Women"
    return "Men"


def gender_definition_word(q: Question, selector: str) -> str:
    text = option_phrase(q, selector).replace('"', "").casefold()
    return "Female" if "female" in text or "woman" in text else "Male"


def is_simple_q_expression(expression: str, q: Optional[Question]) -> bool:
    if q is None:
        return False
    pattern = re.compile(rf"^\s*q{re.escape(qid_key(q.qid))}\s*\([^)]*\)\s*$", re.IGNORECASE)
    return bool(pattern.match(expression))


def is_total_expression(expression: str) -> bool:
    return expression.strip().upper() == "TN"


def race_label_and_definition(expression: str) -> Tuple[str, str]:
    expr = expression.casefold()
    if "not" in expr and "race(w)" in expr.replace(" ", ""):
        return "VOC", "Voters who are not listed as White on the voter file."
    if "race(h)" in expr.replace(" ", ""):
        return "Hispanic", "Voters who are listed as Hispanic on the voter file."
    if "race(w)" in expr.replace(" ", ""):
        return "Wht", "Voters who are listed as White on the voter file."
    return expression, f"Voters matching {expression}."


def ideology_label(selector: str, expression: str) -> str:
    values = set(selector_values(selector))
    expr = expression.casefold()
    if "race(w)" in expr.replace(" ", ""):
        return "WHT COLL LIB"
    if values == {"1"}:
        return "Very Lib"
    if values == {"2"}:
        return "Smwt Lib"
    if values == {"1", "2"}:
        return "Totl Lib"
    if values.issuperset({"3", "4", "5"}) or values.issuperset({"3", "4", "5", "6"}):
        return "Mod+Con"
    return compact_selector_label(selector)


def ideology_definition(q: Question, selector: str, expression: str, edu_q: Optional[Question]) -> str:
    if "race(w)" in expression.casefold().replace(" ", "") and edu_q is not None:
        return (
            f"Voters who are listed as White on the voter file, are college graduates ({qref(edu_q)}), "
            f"and say they are {option_phrase(q, selector)} ({qref(q)})."
        )
    return f"Voters who say they are {option_phrase(q, selector)} ({qref(q)})."


def question_text(q: Question) -> str:
    return norm_spaces(" ".join(part for part in [q.title or "", q.item_text or ""] if part))


def is_type_of_candidate_question(q: Question) -> bool:
    text = question_text(q).casefold()
    return "type of candidate" in text or "type of candidates" in text


def office_short_label(q: Question) -> str:
    office = vote_office_label(q)
    if office == "Attorney General":
        return "AG"
    return office.upper() if office != "Office" else "CANDIDATE"


def type_of_candidate_label(text: str) -> str:
    lowered = text.casefold()
    if "trump" in lowered and ("overreach" in lowered or "administration" in lowered):
        return "Stop DJT"
    if "prosecut" in lowered:
        return "Prosecute"
    if "abortion" in lowered or "civil rights" in lowered:
        return "Def Rght"
    if "managing" in lowered or "management" in lowered or "agency" in lowered:
        return "Mgmt Exp"
    words = re.findall(r"[A-Za-z0-9]+", text)
    return " ".join(words[:2]) if words else text


def candidate_full_name(q: Question, opt: ResponseOption) -> str:
    label = norm_spaces(strip_parens(opt.text))
    if not label:
        return label
    text = question_text(q)
    pattern = re.compile(rf"\b([A-Z][a-zA-Z]+)\s+({re.escape(label)})(?:\s+(Junior|Jr\.?))?\b")
    match = pattern.search(text)
    if match:
        return f"{match.group(1)} {match.group(2)}"
    return label


def candidate_initials(q: Question, opt: ResponseOption) -> str:
    full = candidate_full_name(q, opt)
    parts = re.findall(r"[A-Za-z]+", full)
    if len(parts) >= 2:
        return f"{parts[0][0]}{parts[1][0]}".upper()
    if parts:
        word = parts[0].upper()
        return word[:4] if len(word) <= 4 else word[:3]
    return ""


def pair_for_selector(q: Question, selector: str) -> Optional[Tuple[ResponseOption, ResponseOption]]:
    values = set(selector_values(selector))
    for pair in get_leans_pairs(q) or []:
        if {str(pair[0].value), str(pair[1].value)} == values:
            return pair
    return None


def vote_stage_questions(questions: Sequence[Question]) -> List[Question]:
    groups: Dict[Tuple[str, ...], List[Question]] = {}
    for q in questions:
        if str(q.question_type).casefold() != "head_to_head_ballot":
            continue
        sig = vote_candidate_signature(q)
        if sig:
            groups.setdefault(sig, []).append(q)
    return max(groups.values(), key=len, default=[])


def vote_stage_description(q: Question, stage_number: int) -> str:
    standard = standard_vote_header_label(q) or getattr(q, "_standard_vote_header", None) or "Vote"
    standard = re.sub(r"\s+-\s+(?:Three|Four|Five)\s+Way\b", "", standard, flags=re.IGNORECASE)
    standard = re.sub(r"\s+(?:II|III|IV|V|VI|VII|VIII|IX|X|\d+)$", "", standard, flags=re.IGNORECASE)
    if standard.startswith("Democratic Primary Vote for "):
        base = "Democratic primary vote for " + standard.removeprefix("Democratic Primary Vote for ")
    elif standard.startswith("Republican Primary Vote for "):
        base = "Republican primary vote for " + standard.removeprefix("Republican Primary Vote for ")
    else:
        base = standard.replace("Vote for", "vote for", 1)
    if stage_number == 1:
        prefix = "the initial "
    elif stage_number == 2:
        prefix = "the informed "
    elif "one last time" in question_text(q).casefold() or stage_number >= 4:
        prefix = "the final "
    else:
        prefix = f"the stage {stage_number} "
    return f"{prefix}{base}"


def split_sections(questions: Sequence[Question]) -> Dict[str, Tuple[str, str, str]]:
    out: Dict[str, Tuple[str, str, str]] = {}
    tag_re = re.compile(r"\[(?:(?P<label>[^]\-]+?)\s*-\s*)?SPLIT\s+(?P<marker>[A-Z0-9]+)\]", re.IGNORECASE)
    for q in questions:
        text = question_text(q)
        matches = list(tag_re.finditer(text))
        for idx, match in enumerate(matches):
            marker = match.group("marker").upper()
            label = norm_spaces(match.group("label") or "")
            start = match.end()
            end = matches[idx + 1].start() if idx + 1 < len(matches) else len(text)
            section_text = norm_spaces(text[start:end])
            section_text = re.split(r"\bHaving heard\b|\bAfter hearing\b|\bNow\b", section_text, maxsplit=1)[0].strip()
            if not label:
                lowered = section_text.casefold()
                if any(token in lowered for token in ("liberal", "progressive", "abortion", "l-g-b-t-q", "lgbtq", "civil rights")):
                    label = "Liberal"
                elif "prosecut" in lowered:
                    label = "Prosecute"
                else:
                    label = f"Split {marker}"
            out[marker] = (label, section_text, qref(q))
    return out


def marker_for_split_expression(expression: str) -> Optional[str]:
    match = re.match(r"\s*num_([a-z0-9]+)\((\d+)\)\s*$", expression, flags=re.IGNORECASE)
    if not match:
        return None
    markers = match.group(1).upper()
    idx = int(match.group(2)) - 1
    if 0 <= idx < len(markers):
        return markers[idx]
    return None


def classify_target_expression(
    point: BannerPoint,
    lookup: Dict[str, Question],
    stages: Sequence[Question],
) -> Optional[CoverEntry]:
    expression = point.expression
    if re.fullmatch(r"PERSUADE\([^)]*\)", expression, flags=re.IGNORECASE):
        qrefs = ", ".join(qref(q) for q in stages) if stages else "the vote questions"
        return CoverEntry(
            "TARGETS",
            "Persuade",
            f"Voters who change their support across {qrefs}.",
            expression,
        )

    match = re.match(
        r"\s*NOT\s+Q(?P<from>\d+[a-z]?)\s*\((?P<from_sel>[^)]*)\)\s+AND\s+Q(?P<to>\d+[a-z]?)\s*\((?P<to_sel>[^)]*)\)\s*$",
        expression,
        flags=re.IGNORECASE,
    )
    if not match:
        return None
    from_q = lookup.get(qid_key(match.group("from")))
    to_q = lookup.get(qid_key(match.group("to")))
    if from_q is None or to_q is None:
        return None
    pair = pair_for_selector(to_q, match.group("to_sel"))
    if pair is None:
        return None
    candidate = candidate_full_name(to_q, pair[0])
    candidate_short = strip_parens(pair[0].text).split()[0]
    stage_by_qid = {qid_key(q.qid): idx for idx, q in enumerate(stages, start=1)}
    from_stage = stage_by_qid.get(qid_key(from_q.qid), 0)
    to_stage = stage_by_qid.get(qid_key(to_q.qid), 0)
    stage_label = f"{from_stage}-{to_stage}" if from_stage and to_stage else f"{qref(from_q)}-{qref(to_q)}"
    return CoverEntry(
        "TARGETS",
        f"To {candidate_short} {stage_label}",
        (
            f"Voters who do not support {candidate} in {vote_stage_description(from_q, from_stage or 1)} "
            f"({qref(from_q)}) but do support {candidate} in {vote_stage_description(to_q, to_stage or 2)} ({qref(to_q)})."
        ),
        expression,
    )


def classify_point(
    point: BannerPoint,
    questions: Sequence[Question],
    lookup: Dict[str, Question],
    split_map: Dict[str, Tuple[str, str, str]],
    stages: Sequence[Question],
) -> Optional[CoverEntry]:
    expression = point.expression
    gender_q = first_question(questions, "gender_binary")
    age_q = first_question(questions, "demographics_age")
    edu_q = first_question(questions, "demographics_education")
    ideology_q = first_question(questions, "ideology_5pt")
    likelihood_q = first_question(questions, "likelihood_vote")
    qids = qids_in_expression(expression)
    qset = set(qids)

    if is_total_expression(expression):
        return None
    if is_simple_q_expression(expression, gender_q) or is_simple_q_expression(expression, age_q):
        return None
    if expression_has_q(expression, gender_q) and expression_has_q(expression, age_q):
        return None

    if likelihood_q is not None and is_simple_q_expression(expression, likelihood_q):
        selector = selector_for_q(expression, likelihood_q) or ""
        return CoverEntry(
            "DEF VTRS",
            "DEF VTRS",
            f"Voters who say they will {option_phrase(likelihood_q, selector).lower()} ({qref(likelihood_q)}).",
            expression,
        )

    if re.search(r"\bPVOTEHIST\b", expression, flags=re.IGNORECASE):
        selector = re.search(r"\(([^)]*)\)", expression)
        sel = selector.group(1) if selector else ""
        label = "3+" if sel in {"3,4", "3-4"} else "2 or Less" if sel in {"0-2", "0,1,2"} else compact_selector_label(sel)
        return CoverEntry(
            "VOTE HIST",
            label,
            f"Voters in vote-history group {sel}. [Fill in election years and vote-history definition.]",
            expression,
        )

    if re.search(r"\bREGION", expression, flags=re.IGNORECASE):
        selector = re.search(r"\(([^)]*)\)", expression)
        sel = selector.group(1) if selector else ""
        return CoverEntry(
            "REGION",
            f"Region {sel}",
            f"Voters in {expression}. [Fill in geography for this region group.]",
            expression,
        )

    compact_expr = expression.casefold().replace(" ", "")
    if "race(" in compact_expr and not expression_has_q(expression, edu_q) and not expression_has_q(expression, ideology_q):
        label, definition = race_label_and_definition(expression)
        return CoverEntry("RACE", label, definition, expression)

    if edu_q is not None and is_simple_q_expression(expression, edu_q):
        selector = selector_for_q(expression, edu_q) or ""
        return CoverEntry("EDUCATION", education_label(selector), education_definition(edu_q, selector), expression)

    if gender_q is not None and edu_q is not None and expression_has_q(expression, gender_q) and expression_has_q(expression, edu_q):
        gender_selector = selector_for_q(expression, gender_q) or ""
        edu_selector = selector_for_q(expression, edu_q) or ""
        label = f"{gender_label(gender_q, gender_selector)} {'Ndeg' if education_label(edu_selector) == 'No Deg' else 'Coll'}"
        definition = (
            f"{gender_definition_word(gender_q, gender_selector)} voters ({qref(gender_q)}) "
            f"whose last grade completed in school was {option_phrase(edu_q, edu_selector)} ({qref(edu_q)})."
        )
        return CoverEntry("SEX/EDUCATION", label, definition, expression)

    if (
        edu_q is not None
        and expression_has_q(expression, edu_q)
        and "race(w)" in compact_expr
        and not expression_has_q(expression, ideology_q)
    ):
        selector = selector_for_q(expression, edu_q) or ""
        return CoverEntry(
            "WHITE/EDU",
            education_label(selector),
            f"Voters who are listed as White on the voter file whose last grade completed in school was {option_phrase(edu_q, selector)} ({qref(edu_q)}).",
            expression,
        )

    if ideology_q is not None and expression_has_q(expression, ideology_q):
        selector = selector_for_q(expression, ideology_q) or ""
        return CoverEntry(
            "IDEOLOGY",
            ideology_label(selector, expression),
            ideology_definition(ideology_q, selector, expression, edu_q),
            expression,
        )

    marker = marker_for_split_expression(expression)
    if marker:
        label, text, qref_text = split_map.get(marker, (f"Split {marker}", "", ""))
        definition = f"Voters assigned to {label} split."
        if text:
            definition = f"Voters who heard: {text} ({qref_text})."
        return CoverEntry("SPLIT", label, definition, expression)

    target = classify_target_expression(point, lookup, stages)
    if target:
        return target

    if len(qids) == 1 and qids[0] in lookup:
        q = lookup[qids[0]]
        selector = selector_for_q(expression, q) or ""
        if is_type_of_candidate_question(q):
            opts = selected_options(q, selector)
            text = norm_spaces(strip_parens(opts[0].text)) if opts else selector
            return CoverEntry(
                f"TYPE OF {office_short_label(q)}",
                type_of_candidate_label(text),
                f"Voters who are most likely to support an Attorney General candidate \"{text}\" ({qref(q)}).",
                expression,
            )
        pair = pair_for_selector(q, selector)
        if pair is not None:
            section = (getattr(q, "_standard_vote_header", None) or standard_vote_header_label(q) or f"Vote {qref(q)}").upper()
            section = section.replace("DEMOCRATIC PRIMARY ", "").replace("VOTE FOR ATTORNEY GENERAL", "VOTE FOR ATTY GENERAL")
            section = re.sub(r"\s+-\s+(?:THREE|FOUR|FIVE)\s+WAY\b", "", section)
            section = re.sub(r"\s+(?:II|III|IV|V|VI|VII|VIII|IX|X|\d+)$", "", section)
            candidate = candidate_full_name(q, pair[0])
            stage_by_qid = {qid_key(stage.qid): idx for idx, stage in enumerate(stages, start=1)}
            stage_number = stage_by_qid.get(qid_key(q.qid), 1)
            return CoverEntry(
                section,
                candidate_initials(q, pair[0]),
                f"Voters who support {candidate} in {vote_stage_description(q, stage_number)} ({qref(q)}).",
                expression,
            )
        if selector and any(is_undecided_label(strip_parens(opt.text)) for opt in selected_options(q, selector)):
            section = (getattr(q, "_standard_vote_header", None) or standard_vote_header_label(q) or f"Vote {qref(q)}").upper()
            section = section.replace("DEMOCRATIC PRIMARY ", "").replace("VOTE FOR ATTORNEY GENERAL", "VOTE FOR ATTY GENERAL")
            section = re.sub(r"\s+-\s+(?:THREE|FOUR|FIVE)\s+WAY\b", "", section)
            section = re.sub(r"\s+(?:II|III|IV|V|VI|VII|VIII|IX|X|\d+)$", "", section)
            return CoverEntry(section, "Und", f"Voters who are unsure of their support ({qref(q)}).", expression)

    return CoverEntry(
        "OTHER BANNER POINTS",
        expression,
        f"Voters matching {expression}. [Review and fill in definition.]",
        expression,
    )


def build_cover_entries(questions: Sequence[Question], banner_points: Sequence[BannerPoint]) -> List[CoverEntry]:
    questions_list = list(questions)
    assign_standard_vote_headers(questions_list)
    lookup = question_lookup(questions_list)
    splits = split_sections(questions_list)
    stages = vote_stage_questions(questions_list)
    entries: List[CoverEntry] = []
    seen: set[Tuple[str, str, str]] = set()
    for point in banner_points:
        entry = classify_point(point, questions_list, lookup, splits, stages)
        if entry is None:
            continue
        key = (entry.section, entry.label, entry.definition)
        if key in seen:
            continue
        seen.add(key)
        entries.append(entry)
    return entries


def grouped_entries(entries: Sequence[CoverEntry]) -> List[Tuple[str, List[CoverEntry]]]:
    by_section: Dict[str, List[CoverEntry]] = {}
    for entry in entries:
        by_section.setdefault(entry.section, []).append(entry)
    ordered_sections = [section for section in SECTION_ORDER if section in by_section]
    ordered_sections.extend(section for section in by_section if section not in ordered_sections)
    return [(section, by_section[section]) for section in ordered_sections]


def write_subgroup_docx(entries: Sequence[CoverEntry], outfile: Path) -> None:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    doc.add_heading("SUBGROUP DEFINITIONS", level=1)
    intro = (
        "Throughout this report, references are made to various subgroups within the sample which may require "
        "further definition. These groups are defined below:"
    )
    doc.add_paragraph(intro)

    for section, section_entries in grouped_entries(entries):
        doc.add_heading(section, level=2)
        for entry in section_entries:
            p = doc.add_paragraph()
            label = p.add_run(entry.label)
            label.bold = True
            p.add_run("\t")
            p.add_run(f"{entry.definition}")

    outfile.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(outfile))


def build_cover_docx(json_path: Path, banner_path: Path, outfile: Path) -> None:
    survey = json.loads(json_path.read_text(encoding="utf-8"))
    questions = parse_survey(survey)
    banner_points = parse_banner_code(banner_path.read_text(encoding="utf-8", errors="replace"))
    entries = build_cover_entries(questions, banner_points)
    write_subgroup_docx(entries, outfile)


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = list(argv if argv is not None else sys.argv[1:])
    if len(args) != 3:
        raise SystemExit("Usage: python wincrossCoverCompile.py survey.json banners.txt cover.docx")
    build_cover_docx(Path(args[0]), Path(args[1]), Path(args[2]))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
