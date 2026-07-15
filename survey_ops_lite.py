#!/usr/bin/env python3
"""Small daily-operations shell for active survey jobs.

This app intentionally stays thin. It keeps one project JSON per survey and
delegates heavy work to the existing WinCross, merge, and OE tools.
"""

from __future__ import annotations

import csv
import json
import os
import shutil
import subprocess
import sys
import threading
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

import tkinter as tk
from tkinter import filedialog, messagebox, ttk


APP_ROOT = Path(__file__).resolve().parent
LITE_WORKSPACE = APP_ROOT / "lite_workspace"
PROJECT_FILENAME = "project.json"


def utc_now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def clean_project_id(value: str) -> str:
    cleaned = "".join(ch for ch in value.strip() if ch.isalnum() or ch in {"_", "-"})
    return cleaned.upper()


def require_project_id(value: str) -> str:
    cleaned = clean_project_id(value)
    if not cleaned:
        raise ValueError("Project ID is required.")
    return cleaned


def project_dir(project_id: str, workspace_root: Path = LITE_WORKSPACE) -> Path:
    return workspace_root / require_project_id(project_id)


def project_path(project_id: str, workspace_root: Path = LITE_WORKSPACE) -> Path:
    return project_dir(project_id, workspace_root) / PROJECT_FILENAME


def rel_project_path(project_root: Path, *parts: str) -> Path:
    return project_root.joinpath(*parts)


def ensure_project_dirs(project_root: Path) -> None:
    for relative in (
        ("inputs", "questionnaire"),
        ("inputs", "data", "current"),
        ("inputs", "data", "timeseries"),
        ("inputs", "data", "codelists"),
        ("outputs", "wincross", "phone"),
        ("outputs", "wincross", "sms"),
        ("outputs", "merge"),
        ("outputs", "oe"),
        ("runs",),
    ):
        rel_project_path(project_root, *relative).mkdir(parents=True, exist_ok=True)


def default_project(project_id: str) -> dict[str, Any]:
    cleaned = require_project_id(project_id)
    return {
        "schema_version": 1,
        "project_id": cleaned,
        "current_day": 1,
        "created_at": utc_now(),
        "updated_at": utc_now(),
        "paths": {
            "phone_docx": "",
            "phone_pdf": "",
            "sms_docx": "",
            "sms_pdf": "",
            "current_data": "",
            "timeseries": "",
            "daily_codelist": "",
            "last_wincross_phone_tables": "",
            "last_wincross_sms_tables": "",
            "last_merge_output": "",
            "last_oe_output": "",
        },
        "data": {
            "current_header": [],
            "current_source": "",
            "current_tracked_at": "",
            "daily_codelists": {},
        },
        "runs": [],
    }


def load_project(project_id: str, workspace_root: Path = LITE_WORKSPACE) -> dict[str, Any]:
    path = project_path(project_id, workspace_root)
    if not path.exists():
        project = default_project(project_id)
        save_project(project, workspace_root)
        return project
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError(f"Invalid project file: {path}")
    data.setdefault("paths", {})
    data.setdefault("data", {})
    data.setdefault("runs", [])
    return data


def save_project(project: dict[str, Any], workspace_root: Path = LITE_WORKSPACE) -> Path:
    project_id = require_project_id(str(project.get("project_id", "")))
    project["project_id"] = project_id
    project["updated_at"] = utc_now()
    root = project_dir(project_id, workspace_root)
    ensure_project_dirs(root)
    path = root / PROJECT_FILENAME
    path.write_text(json.dumps(project, indent=2), encoding="utf-8")
    return path


def _read_csv_header(path: Path) -> list[str]:
    encodings = ("utf-8-sig", "utf-8", "latin1")
    last_error: Exception | None = None
    for enc in encodings:
        try:
            with path.open("r", encoding=enc, newline="") as handle:
                reader = csv.reader(handle)
                row = next(reader, [])
                return [str(cell).strip() for cell in row]
        except Exception as exc:
            last_error = exc
    if last_error:
        raise last_error
    return []


def read_table_header(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv_header(path)
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except Exception as exc:  # pragma: no cover - dependency error path
            raise RuntimeError("openpyxl is required to inspect XLSX headers.") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            return ["" if cell is None else str(cell).strip() for cell in row]
        finally:
            workbook.close()
    raise ValueError(f"Unsupported data file type: {path.suffix}")


def headers_match(left: list[str], right: list[str]) -> bool:
    return [x.strip() for x in left] == [x.strip() for x in right]


def copy_into_project(source: Path, dest: Path) -> Path:
    if not source.exists():
        raise FileNotFoundError(source)
    dest.parent.mkdir(parents=True, exist_ok=True)
    try:
        if source.resolve() == dest.resolve():
            return dest
    except OSError:
        pass
    shutil.copy2(source, dest)
    return dest


def require_source_path(value: str, label: str) -> Path:
    raw = value.strip()
    if not raw:
        raise ValueError(f"{label} path is required.")
    return Path(raw)


def channel_label(channel: str) -> str:
    ch = channel.lower()
    if ch == "sms":
        return "SMS"
    return ch.title()


def track_current_data(project: dict[str, Any], source: Path, workspace_root: Path = LITE_WORKSPACE) -> Path:
    project_id = require_project_id(str(project.get("project_id", "")))
    header = read_table_header(source)
    if not header:
        raise ValueError(f"Could not read a header row from {source}")
    existing_header = project.get("data", {}).get("current_header") or []
    if existing_header and not headers_match(existing_header, header):
        raise ValueError("Current data headers do not match the tracked delivery.")
    root = project_dir(project_id, workspace_root)
    dest = rel_project_path(root, "inputs", "data", "current", f"{project_id}_current_data{source.suffix.lower()}")
    copy_into_project(source, dest)
    project.setdefault("paths", {})["current_data"] = str(dest)
    project.setdefault("data", {})["current_header"] = header
    project["data"]["current_source"] = str(source)
    project["data"]["current_tracked_at"] = utc_now()
    save_project(project, workspace_root)
    return dest


def track_timeseries(project: dict[str, Any], source: Path, workspace_root: Path = LITE_WORKSPACE) -> Path:
    project_id = require_project_id(str(project.get("project_id", "")))
    root = project_dir(project_id, workspace_root)
    dest = rel_project_path(root, "inputs", "data", "timeseries", f"{project_id}_timeseries{source.suffix.lower()}")
    copy_into_project(source, dest)
    project.setdefault("paths", {})["timeseries"] = str(dest)
    save_project(project, workspace_root)
    return dest


def track_daily_codelist(
    project: dict[str, Any],
    source: Path,
    day: int,
    workspace_root: Path = LITE_WORKSPACE,
) -> Path:
    project_id = require_project_id(str(project.get("project_id", "")))
    root = project_dir(project_id, workspace_root)
    dest = rel_project_path(root, "inputs", "data", "codelists", f"day_{day}", source.name)
    copy_into_project(source, dest)
    project.setdefault("paths", {})["daily_codelist"] = str(dest)
    codelists = project.setdefault("data", {}).setdefault("daily_codelists", {})
    codelists[str(day)] = {"path": str(dest), "source": str(source), "tracked_at": utc_now()}
    save_project(project, workspace_root)
    return dest


def wincross_output_paths(project: dict[str, Any], channel: str, workspace_root: Path = LITE_WORKSPACE) -> dict[str, Path]:
    project_id = require_project_id(str(project.get("project_id", "")))
    ch = channel.lower()
    root = project_dir(project_id, workspace_root)
    out_dir = rel_project_path(root, "outputs", "wincross", ch)
    base = f"{project_id}_{ch}"
    return {
        "raw": out_dir / f"{base}RAW.txt",
        "json": out_dir / f"{base}JSON.json",
        "tables": out_dir / f"{base}TABLES.autogen.txt",
        "banners": out_dir / f"{base}BANNERS.autogen.txt",
        "cover": out_dir / f"{base}COVER.autogen.docx",
    }


def build_wincross_command(
    project: dict[str, Any],
    channel: str,
    *,
    py_exe: str = sys.executable,
    workspace_root: Path = LITE_WORKSPACE,
) -> tuple[list[str], dict[str, Path]]:
    ch = channel.lower()
    paths = project.setdefault("paths", {})
    docx_raw = str(paths.get(f"{ch}_docx") or "").strip()
    pdf_raw = str(paths.get(f"{ch}_pdf") or "").strip()
    if not docx_raw:
        raise FileNotFoundError(f"{channel_label(channel)} DOCX path is required.")
    if not pdf_raw:
        raise FileNotFoundError(f"{channel_label(channel)} PDF path is required.")
    docx = Path(docx_raw)
    pdf = Path(pdf_raw)
    if not docx.exists():
        raise FileNotFoundError(f"{channel_label(channel)} DOCX not found: {docx}")
    if not pdf.exists():
        raise FileNotFoundError(f"{channel_label(channel)} PDF not found: {pdf}")
    out = wincross_output_paths(project, ch, workspace_root)
    script = APP_ROOT / "pipelines" / "wincross_table_generator" / "survey_pipeline.py"
    cmd = [
        py_exe,
        str(script),
        "--docx",
        str(docx),
        "--raw-out",
        str(out["raw"]),
        "--json-out",
        str(out["json"]),
        "--tables-out",
        str(out["tables"]),
        "--banners-out",
        str(out["banners"]),
        "--cover-out",
        str(out["cover"]),
        "--pdf",
        str(pdf),
    ]
    return cmd, out


def record_run(project: dict[str, Any], stage: str, status: str, detail: str = "") -> None:
    project.setdefault("runs", []).append(
        {
            "stage": stage,
            "status": status,
            "detail": detail,
            "at": utc_now(),
        }
    )


def open_path(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    if sys.platform.startswith("win"):
        os.startfile(str(path))  # type: ignore[attr-defined]
    elif sys.platform == "darwin":
        subprocess.Popen(["open", str(path)])
    else:
        subprocess.Popen(["xdg-open", str(path)])


def launcher_path(relative: str) -> Path:
    return APP_ROOT / relative


def launch_file(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    if path.suffix.lower() in {".bat", ".cmd"} and sys.platform.startswith("win"):
        subprocess.Popen(["cmd", "/c", str(path)], cwd=str(path.parent))
    elif path.suffix.lower() == ".ps1" and sys.platform.startswith("win"):
        subprocess.Popen(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(path)], cwd=str(path.parent))
    else:
        subprocess.Popen([sys.executable, str(path)], cwd=str(path.parent))


@dataclass
class PathField:
    key: str
    label: str
    patterns: list[tuple[str, str]]
    var: tk.StringVar = field(init=False)

    def bind(self, master: tk.Misc) -> None:
        self.var = tk.StringVar(master=master)


class SurveyOpsLiteApp:
    def __init__(self, root: tk.Tk, workspace_root: Path = LITE_WORKSPACE) -> None:
        self.root = root
        self.workspace_root = workspace_root
        self.project: dict[str, Any] | None = None
        self.project_id_var = tk.StringVar(value="")
        self.day_var = tk.StringVar(value="1")
        self.status_var = tk.StringVar(value="No project loaded")
        self.path_fields = [
            PathField("phone_docx", "Phone DOCX", [("Word documents", "*.docx")]),
            PathField("phone_pdf", "Phone PDF", [("PDF files", "*.pdf")]),
            PathField("sms_docx", "SMS DOCX", [("Word documents", "*.docx")]),
            PathField("sms_pdf", "SMS PDF", [("PDF files", "*.pdf")]),
            PathField("current_data", "Current data", [("Data files", "*.csv *.xlsx")]),
            PathField("timeseries", "Timeseries", [("Data files", "*.csv *.xlsx")]),
            PathField("daily_codelist", "Daily codelist", [("Data files", "*.csv *.xlsx")]),
        ]
        for field in self.path_fields:
            field.bind(root)
        self._build()

    def _build(self) -> None:
        self.root.title("Survey Ops Lite")
        self.root.geometry("980x680")
        self.root.minsize(820, 560)

        shell = ttk.Frame(self.root, padding=10)
        shell.pack(fill=tk.BOTH, expand=True)

        top = ttk.Frame(shell)
        top.pack(fill=tk.X)
        ttk.Label(top, text="Survey").pack(side=tk.LEFT)
        ttk.Entry(top, textvariable=self.project_id_var, width=18).pack(side=tk.LEFT, padx=6)
        ttk.Label(top, text="Day").pack(side=tk.LEFT, padx=(10, 0))
        ttk.Spinbox(top, from_=1, to=99, textvariable=self.day_var, width=5).pack(side=tk.LEFT, padx=6)
        ttk.Button(top, text="Load", command=self.load_project_click).pack(side=tk.LEFT, padx=4)
        ttk.Button(top, text="Save", command=self.save_project_click).pack(side=tk.LEFT, padx=4)
        ttk.Button(top, text="Open Folder", command=self.open_project_folder).pack(side=tk.LEFT, padx=4)
        ttk.Label(top, textvariable=self.status_var).pack(side=tk.RIGHT)

        notebook = ttk.Notebook(shell)
        notebook.pack(fill=tk.BOTH, expand=True, pady=(10, 0))

        setup = ttk.Frame(notebook, padding=10)
        daily = ttk.Frame(notebook, padding=10)
        actions = ttk.Frame(notebook, padding=10)
        notebook.add(setup, text="Project")
        notebook.add(daily, text="Daily Data")
        notebook.add(actions, text="Actions")

        self._build_project_tab(setup)
        self._build_daily_tab(daily)
        self._build_actions_tab(actions)

        log_frame = ttk.LabelFrame(shell, text="Log", padding=6)
        log_frame.pack(fill=tk.BOTH, expand=False, pady=(10, 0))
        self.log_text = tk.Text(log_frame, height=10, wrap=tk.WORD)
        self.log_text.pack(fill=tk.BOTH, expand=True)

    def _build_project_tab(self, parent: ttk.Frame) -> None:
        for row, field in enumerate(self.path_fields[:4]):
            self._path_row(parent, row, field)
        ttk.Separator(parent).grid(row=5, column=0, columnspan=3, sticky="ew", pady=12)
        ttk.Button(parent, text="Run Phone WinCross", command=lambda: self.run_wincross("phone")).grid(row=6, column=0, sticky="w", pady=4)
        ttk.Button(parent, text="Run SMS WinCross", command=lambda: self.run_wincross("sms")).grid(row=6, column=1, sticky="w", pady=4)
        ttk.Button(parent, text="Open WinCross GUI", command=self.open_wincross_gui).grid(row=6, column=2, sticky="w", pady=4)
        parent.columnconfigure(1, weight=1)

    def _build_daily_tab(self, parent: ttk.Frame) -> None:
        for idx, field in enumerate(self.path_fields[4:], start=0):
            self._path_row(parent, idx, field)
        ttk.Separator(parent).grid(row=4, column=0, columnspan=3, sticky="ew", pady=12)
        ttk.Button(parent, text="Track Current Data", command=self.track_current_data_click).grid(row=5, column=0, sticky="w", pady=4)
        ttk.Button(parent, text="Track Timeseries", command=self.track_timeseries_click).grid(row=5, column=1, sticky="w", pady=4)
        ttk.Button(parent, text="Track Codelist", command=self.track_codelist_click).grid(row=5, column=2, sticky="w", pady=4)
        parent.columnconfigure(1, weight=1)

    def _build_actions_tab(self, parent: ttk.Frame) -> None:
        ttk.Button(parent, text="Open Merge GUI", command=self.open_merge_gui).grid(row=0, column=0, sticky="w", pady=4)
        ttk.Button(parent, text="Open OE Coder", command=self.open_oe_gui).grid(row=0, column=1, sticky="w", pady=4)
        ttk.Button(parent, text="Open Outputs", command=self.open_outputs_folder).grid(row=0, column=2, sticky="w", pady=4)
        ttk.Button(parent, text="Open Current Data", command=lambda: self.open_project_path("current_data")).grid(row=1, column=0, sticky="w", pady=4)
        ttk.Button(parent, text="Open Timeseries", command=lambda: self.open_project_path("timeseries")).grid(row=1, column=1, sticky="w", pady=4)
        ttk.Button(parent, text="Open Last Tables", command=self.open_last_tables).grid(row=1, column=2, sticky="w", pady=4)
        parent.columnconfigure(3, weight=1)

    def _path_row(self, parent: ttk.Frame, row: int, field: PathField) -> None:
        ttk.Label(parent, text=field.label).grid(row=row, column=0, sticky="w", padx=(0, 8), pady=4)
        ttk.Entry(parent, textvariable=field.var).grid(row=row, column=1, sticky="ew", pady=4)
        ttk.Button(parent, text="Browse", command=lambda f=field: self.browse_path(f)).grid(row=row, column=2, sticky="e", padx=(8, 0), pady=4)

    def log(self, message: str) -> None:
        timestamp = datetime.now().strftime("%H:%M:%S")
        self.log_text.insert(tk.END, f"[{timestamp}] {message}\n")
        self.log_text.see(tk.END)
        self.root.update_idletasks()

    def current_day(self) -> int:
        try:
            value = int(self.day_var.get())
        except Exception as exc:
            raise ValueError("Day must be a number.") from exc
        if value < 1:
            raise ValueError("Day must be 1 or greater.")
        return value

    def ensure_project(self) -> dict[str, Any]:
        if self.project is None:
            self.load_project_click()
        assert self.project is not None
        return self.project

    def load_project_click(self) -> None:
        try:
            project_id = require_project_id(self.project_id_var.get())
            self.project = load_project(project_id, self.workspace_root)
            self.project_id_var.set(project_id)
            self.day_var.set(str(self.project.get("current_day", 1) or 1))
            paths = self.project.setdefault("paths", {})
            for field in self.path_fields:
                field.var.set(str(paths.get(field.key, "")))
            self.status_var.set(f"Loaded {project_id}")
            self.log(f"Loaded project {project_id}")
        except Exception as exc:
            messagebox.showerror("Load project", str(exc))

    def save_project_click(self) -> None:
        try:
            project = self.project or default_project(self.project_id_var.get())
            project["current_day"] = self.current_day()
            paths = project.setdefault("paths", {})
            for field in self.path_fields:
                paths[field.key] = field.var.get().strip()
            save_project(project, self.workspace_root)
            self.project = project
            self.status_var.set(f"Saved {project['project_id']}")
            self.log(f"Saved project {project['project_id']}")
        except Exception as exc:
            messagebox.showerror("Save project", str(exc))

    def browse_path(self, field: PathField) -> None:
        picked = filedialog.askopenfilename(title=field.label, filetypes=[*field.patterns, ("All files", "*.*")])
        if picked:
            field.var.set(picked)
            self.save_project_click()

    def _refresh_project_paths_from_fields(self) -> dict[str, Any]:
        project = self.ensure_project()
        project["current_day"] = self.current_day()
        paths = project.setdefault("paths", {})
        for field in self.path_fields:
            if field.var.get().strip():
                paths[field.key] = field.var.get().strip()
        return project

    def track_current_data_click(self) -> None:
        try:
            project = self._refresh_project_paths_from_fields()
            source = require_source_path(str(project["paths"].get("current_data") or ""), "Current data")
            dest = track_current_data(project, source, self.workspace_root)
            self.path_fields[4].var.set(str(dest))
            self.status_var.set("Current data tracked")
            self.log(f"Current data tracked: {dest}")
        except Exception as exc:
            messagebox.showerror("Track current data", str(exc))

    def track_timeseries_click(self) -> None:
        try:
            project = self._refresh_project_paths_from_fields()
            source = require_source_path(str(project["paths"].get("timeseries") or ""), "Timeseries")
            dest = track_timeseries(project, source, self.workspace_root)
            self.path_fields[5].var.set(str(dest))
            self.status_var.set("Timeseries tracked")
            self.log(f"Timeseries tracked: {dest}")
        except Exception as exc:
            messagebox.showerror("Track timeseries", str(exc))

    def track_codelist_click(self) -> None:
        try:
            project = self._refresh_project_paths_from_fields()
            source = require_source_path(str(project["paths"].get("daily_codelist") or ""), "Daily codelist")
            dest = track_daily_codelist(project, source, self.current_day(), self.workspace_root)
            self.path_fields[6].var.set(str(dest))
            self.status_var.set("Codelist tracked")
            self.log(f"Codelist tracked: {dest}")
        except Exception as exc:
            messagebox.showerror("Track codelist", str(exc))

    def run_wincross(self, channel: str) -> None:
        try:
            project = self._refresh_project_paths_from_fields()
            cmd, outputs = build_wincross_command(project, channel, workspace_root=self.workspace_root)
        except Exception as exc:
            messagebox.showerror("WinCross", str(exc))
            return

        def worker() -> None:
            try:
                self.log(f"Running {channel} WinCross")
                proc = subprocess.Popen(
                    cmd,
                    cwd=str(APP_ROOT),
                    stdout=subprocess.PIPE,
                    stderr=subprocess.STDOUT,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                )
                assert proc.stdout is not None
                for line in proc.stdout:
                    line = line.rstrip()
                    if line:
                        self.root.after(0, self.log, line)
                code = proc.wait()
                if code != 0:
                    raise subprocess.CalledProcessError(code, cmd)
                project = self.ensure_project()
                paths = project.setdefault("paths", {})
                paths[f"last_wincross_{channel.lower()}_tables"] = str(outputs["tables"])
                paths[f"last_wincross_{channel.lower()}_banners"] = str(outputs["banners"])
                paths[f"last_wincross_{channel.lower()}_cover"] = str(outputs["cover"])
                record_run(project, f"wincross_{channel.lower()}", "ok", str(outputs["tables"]))
                save_project(project, self.workspace_root)
                label = channel_label(channel)
                self.root.after(0, self.status_var.set, f"{label} WinCross complete")
                self.root.after(0, self.log, f"{label} TABLES: {outputs['tables']}")
                self.root.after(0, self.log, f"{label} BANNERS: {outputs['banners']}")
                self.root.after(0, self.log, f"{label} COVER: {outputs['cover']}")
            except Exception as exc:
                project = self.project
                if project is not None:
                    record_run(project, f"wincross_{channel.lower()}", "error", str(exc))
                    save_project(project, self.workspace_root)
                self.root.after(0, messagebox.showerror, "WinCross", str(exc))

        threading.Thread(target=worker, daemon=True).start()

    def project_root(self) -> Path:
        project = self.ensure_project()
        return project_dir(str(project["project_id"]), self.workspace_root)

    def open_project_folder(self) -> None:
        try:
            open_path(self.project_root())
        except Exception as exc:
            messagebox.showerror("Open folder", str(exc))

    def open_outputs_folder(self) -> None:
        try:
            open_path(self.project_root() / "outputs")
        except Exception as exc:
            messagebox.showerror("Open outputs", str(exc))

    def open_project_path(self, key: str) -> None:
        try:
            project = self.ensure_project()
            path = Path(str(project.setdefault("paths", {}).get(key) or ""))
            open_path(path)
        except Exception as exc:
            messagebox.showerror("Open file", str(exc))

    def open_last_tables(self) -> None:
        try:
            project = self.ensure_project()
            paths = project.setdefault("paths", {})
            candidate = paths.get("last_wincross_sms_tables") or paths.get("last_wincross_phone_tables")
            if not candidate:
                raise FileNotFoundError("No WinCross TABLES output recorded yet.")
            open_path(Path(str(candidate)))
        except Exception as exc:
            messagebox.showerror("Open tables", str(exc))

    def open_wincross_gui(self) -> None:
        try:
            launch_file(launcher_path("pipelines/wincross_table_generator/open_survey_pipeline_gui.bat"))
            self.log("Opened WinCross GUI")
        except Exception as exc:
            messagebox.showerror("WinCross GUI", str(exc))

    def open_merge_gui(self) -> None:
        try:
            launch_file(launcher_path("pipelines/table_merger/open_sheet_merge_gui.bat"))
            self.log("Opened merge GUI")
        except Exception as exc:
            messagebox.showerror("Merge GUI", str(exc))

    def open_oe_gui(self) -> None:
        try:
            launch_file(launcher_path("pipelines/open_end_coder/Launch_OE_Codebook_GUI.bat"))
            self.log("Opened OE coder")
        except Exception as exc:
            messagebox.showerror("OE coder", str(exc))


def main() -> None:
    root = tk.Tk()
    app = SurveyOpsLiteApp(root)
    last = ""
    cfg_path = APP_ROOT / "survey_tool_config.json"
    try:
        cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
        last = str((cfg.get("dashboard") or {}).get("last_main_survey_id") or "")
    except Exception:
        last = ""
    if last:
        app.project_id_var.set(last)
    root.mainloop()


if __name__ == "__main__":
    main()
